"""올린 견적서로 교체한 뒤 이어서 고치기 — 고객 지적 재현(사용자 2026-10-15).

흐름: 견적 발행(01) → 받은 엑셀을 고침 → 같은 번호로 올려 교체 → 작성 상태·미리보기가 고친 파일 내용인지 → 이어서 고쳐 다시 발행.

    MSYS_NO_PATHCONV=1 docker compose run --rm --no-deps -e PYTHONPATH=/app -e PYTHONIOENCODING=utf-8 \
      -v ./backend/app:/app/app -v ./test_cases:/tc backend python /tc/runners/api/quote_sync.py
"""
from __future__ import annotations

import io
import sys

import openpyxl
import pymupdf

sys.path.insert(0, "/tc/runners")
import tc  # noqa: E402

Q = "/v1/product-recommend/quotes"
LINES = [{"name": "MG10", "commodity": "Magbot Gripper", "qty": 3, "unit_price": 1250000},
         {"name": "TCV1", "qty": 2, "unit_price": 480000}]


def edited(xlsx: bytes, *, qty: int, remark: str, comments: str) -> bytes:
    """엑셀에서 사람이 고친 것처럼 — 2번째 품목 수량, 1번째 품목 비고, Comments."""
    wb = openpyxl.load_workbook(io.BytesIO(xlsx))
    ws = wb.active
    head = next(r for r in range(1, ws.max_row + 1) if ws.cell(r, 1).value == "Commodity")
    ws.cell(head + 2, 4).value = qty
    ws.cell(head + 1, 7).value = remark
    ws.cell(next(r for r in range(1, head) if ws.cell(r, 4).value == "Comments"), 5).value = comments
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def pdf_text(data: bytes) -> str:
    return "".join(p.get_text() for p in pymupdf.open(stream=data, filetype="pdf"))


async def body(s: tc.Suite) -> None:
    r = await s.call("POST", f"{Q}/manual", "lucas", json={"form": {}})
    qid = r.json()["id"]
    await s.call("POST", f"{Q}/{qid}/edit", "lucas", json={"form": {"customer": f"{tc.MARK}싱크고객", "to": "김철수 대표님"},
                                                           "lines": LINES, "delivery": {"case": 1}})
    iss = (await s.call("POST", f"{Q}/{qid}/issue", "lucas", json={})).json()
    no = (iss.get("quote") or iss)["quote_no"]
    xlsx = (await s.call("GET", f"{Q}/{qid}/file", "lucas")).content
    state: dict = {}

    async def up(data: bytes):
        return await s.call("POST", f"{Q}/{qid}/upload", "lucas",
                            files={"file": ("고친견적서.xlsx", data, "application/octet-stream")})

    async def qs079():
        r = await up(edited(xlsx, qty=5, remark="엑셀에서 고친 비고", comments="www.magbot.kr\n- 엑셀에서 고친 공급 건"))
        j = r.json()
        state["up"] = j
        ok = r.status_code == 200 and j["read"].get("template") and j["read"].get("doc_no") == no and j.get("sync", {}).get("applies")
        return ok, f"{r.status_code} 회사 양식={j['read'].get('template')} 읽은 번호={j['read'].get('doc_no')} (발행 {no}) " \
                   f"바뀔 점={j.get('sync', {}).get('changes')}"
    await s.case("QS-079", "고친 회사 양식 견적서 올리기 → 칸 그대로 읽고 바뀔 점 미리 보여 줌", qs079)

    async def qs080():
        j = state["up"]
        body_ = {"upload": j["upload"], "items": j["read"]["items"], "initials": "LU", "doc_no": no}
        r1 = await s.call("POST", f"{Q}/{qid}/upload/confirm", "lucas", json=body_)
        r2 = await s.call("POST", f"{Q}/{qid}/upload/confirm", "lucas", json={**body_, "overwrite": True})
        return r1.status_code == 409 and r2.status_code == 200, f"같은 번호 {no} → {r1.status_code}(덮어쓸지 확인) → 덮어쓰기 {r2.status_code}"
    await s.case("QS-080", "같은 번호로 파일 교체", qs080)

    async def qs081():
        g = (await s.call("GET", f"{Q}/{qid}", "lucas")).json()
        ln = g["state"]["lines"]
        ok = ln[1]["qty"] == 5 and ln[0]["remark"] == "엑셀에서 고친 비고" and "엑셀에서 고친 공급 건" in g["state"]["form"]["comments"] \
            and g["quote_no"] == no
        return ok, f"작성 상태 수량 {ln[1]['qty']} 비고 {ln[0]['remark']!r} Comments {g['state']['form']['comments']!r} 번호 {g['quote_no']}"
    await s.case("QS-081", "교체 뒤 [견적서 고치기] 작성 상태 = 고친 파일 내용", qs081)

    async def qs082():
        p = await s.call("GET", f"{Q}/{qid}/preview", "lucas")
        t = pdf_text(p.content) if p.status_code == 200 else ""
        ok = "엑셀에서 고친 비고" in t and "엑셀에서 고친 공급 건" in t and "960,000" not in t and "2,400,000" in t
        return ok, f"미리보기 {p.status_code}: 고친 비고={'엑셀에서 고친 비고' in t}, 고친 Comments={'엑셀에서 고친 공급 건' in t}, " \
                   f"TCV1 5개 금액 2,400,000={'2,400,000' in t}, 예전 2개 금액 960,000={'960,000' in t}"
    await s.case("QS-082", "교체 뒤 미리보기 = 고친 파일 내용(예전 01 아님)", qs082)

    async def qs083():
        g = (await s.call("GET", f"{Q}/{qid}", "lucas")).json()
        lines = [{k: v for k, v in ln.items() if k in ("id", "name", "commodity", "qty", "unit_price", "remark")} for ln in g["state"]["lines"]]
        lines[0]["unit_price"] = 1200000
        e = await s.call("POST", f"{Q}/{qid}/edit", "lucas", json={"lines": lines})
        i = await s.call("POST", f"{Q}/{qid}/issue", "lucas", json={})
        f = await s.call("GET", f"{Q}/{qid}/file", "lucas")
        wb = openpyxl.load_workbook(io.BytesIO(f.content))
        ws = wb.active
        vals = [str(c.value) for row in ws.iter_rows() for c in row if c.value is not None]
        ok = e.status_code == 200 and i.status_code == 200 and "엑셀에서 고친 비고" in vals and 5 in [c.value for row in ws.iter_rows() for c in row]
        return ok, f"이어서 단가 수정 {e.status_code} → 다시 발행 {i.status_code}, 새 파일에 고친 비고 유지={'엑셀에서 고친 비고' in vals}"
    await s.case("QS-083", "교체한 내용에서 이어서 고쳐 다시 발행(고친 내용 유지)", qs083)

    async def qs084():
        cur = (await s.call("GET", f"{Q}/{qid}/file", "lucas")).content
        r = await up(edited(cur, qty=9, remark="끈 경우", comments="www.magbot.kr"))
        j = r.json()
        c = await s.call("POST", f"{Q}/{qid}/upload/confirm", "lucas",
                         json={"upload": j["upload"], "items": j["read"]["items"], "initials": "LU", "doc_no": "", "sync": False})
        g = (await s.call("GET", f"{Q}/{qid}", "lucas")).json()
        return c.status_code == 200 and g["state"]["lines"][1]["qty"] == 5, \
            f"[작성 상태도 바꾸기] 끄고 발행 {c.status_code} → 작성 상태 수량 {g['state']['lines'][1]['qty']}(5 그대로여야)"
    await s.case("QS-084", "작성 상태 맞추기를 끄면 예전처럼 발행 기록만", qs084)

    # ── 2단계: 엑셀에서 꾸민 서식(글꼴·색·채우기·병합·그림)이 이어서 고친 뒤에도 남는가 ──
    r = await s.call("POST", f"{Q}/manual", "lucas", json={"form": {}})
    q2 = r.json()["id"]
    await s.call("POST", f"{Q}/{q2}/edit", "lucas", json={"form": {"customer": f"{tc.MARK}서식고객", "to": "이영희 과장님"},
                                                          "lines": LINES, "delivery": {"case": 1}})
    prof = (await s.call("GET", "/v1/me/profile", "lucas")).json()
    if not all(prof.get(k) for k in ("initials", "mobile", "email")):
        for cid in ("QS-085", "QS-086", "QS-087"):
            s.skip(cid, "서식 유지 흐름", "lucas 내 정보(이니셜·휴대폰·이메일)가 비어 발행할 수 없음")
        return
    i2 = (await s.call("POST", f"{Q}/{q2}/issue", "lucas", json={})).json()
    no2 = (i2.get("quote") or i2)["quote_no"]
    raw = (await s.call("GET", f"{Q}/{q2}/file", "lucas")).content
    from openpyxl.styles import Font, PatternFill
    wb = openpyxl.load_workbook(io.BytesIO(raw))
    ws = wb.active
    head = next(r for r in range(1, ws.max_row + 1) if ws.cell(r, 1).value == "Commodity")
    ws.cell(head + 1, 7).value = "빨간 굵은 비고"
    ws.cell(head + 1, 7).font = Font(name="맑은 고딕", size=11, bold=True, color="FFFF0000")
    ws.cell(head + 2, 1).fill = PatternFill("solid", fgColor="FFCCE5FF")
    bottom = ws.max_row + 2
    ws.merge_cells(f"A{bottom}:C{bottom + 1}")
    ws.cell(bottom, 1).value = "직접 병합한 안내"
    buf = io.BytesIO()
    wb.save(buf)
    up2 = (await s.call("POST", f"{Q}/{q2}/upload", "lucas", files={"file": ("꾸민견적서.xlsx", buf.getvalue(), "application/octet-stream")})).json()
    await s.call("POST", f"{Q}/{q2}/upload/confirm", "lucas", json={"upload": up2["upload"], "items": up2["read"]["items"],
                                                                     "initials": prof["initials"], "doc_no": no2, "overwrite": True})

    async def qs085():
        g = (await s.call("GET", f"{Q}/{q2}", "lucas")).json()
        return g.get("on_file") == no2, f"화면 표시 on_file={g.get('on_file')} (기대 {no2})"
    await s.case("QS-085", "교체 뒤 '올린 파일 위에서 고치는 중' 표시", qs085)

    async def qs086():
        g = (await s.call("GET", f"{Q}/{q2}", "lucas")).json()
        lines = [{k: v for k, v in ln.items() if k in ("id", "name", "commodity", "qty", "unit_price", "remark")} for ln in g["state"]["lines"]]
        lines[1]["qty"] = 4
        lines.append({"name": "출장 교육", "commodity": "Service", "qty": 1, "unit_price": 300000, "remark": "1일"})
        await s.call("POST", f"{Q}/{q2}/edit", "lucas", json={"lines": lines})
        p = await s.call("GET", f"{Q}/{q2}/preview", "lucas")
        t = pdf_text(p.content) if p.status_code == 200 else ""
        want = f"{3 * 1250000 + 4 * 480000 + 300000:,}"
        merged = "직접병합한안내" in "".join(t.split())          # 병합 칸 폭에 따라 PDF 글이 줄바꿈된다
        return p.status_code == 200 and "출장 교육" in t and want in t and merged, \
            f"미리보기 {p.status_code}: 새 줄={'출장 교육' in t} 합계 {want}={want in t} 직접 병합 글={merged}"
    await s.case("QS-086", "이어서 고친 미리보기 — 고친 값 + 엑셀에서 넣은 내용 유지", qs086)

    async def qs087():
        i = await s.call("POST", f"{Q}/{q2}/issue", "lucas", json={})
        f = await s.call("GET", f"{Q}/{q2}/file", "lucas")
        w = openpyxl.load_workbook(io.BytesIO(f.content)).active
        rr = next((r for r in range(1, w.max_row + 1) if w.cell(r, 7).value == "빨간 굵은 비고"), None)
        red = rr is not None and w.cell(rr, 7).font.b and (w.cell(rr, 7).font.color.rgb if w.cell(rr, 7).font.color else "") == "FFFF0000"
        fill = rr is not None and w.cell(rr + 1, 1).fill.fgColor.rgb == "FFCCE5FF"
        merged = any(w.cell(m.min_row, 1).value == "직접 병합한 안내" for m in w.merged_cells.ranges)
        new_row = any(str(w.cell(r, 2).value or "") == "출장 교육" for r in range(1, w.max_row + 1))
        ok = i.status_code == 200 and red and fill and merged and new_row
        return ok, f"발행 {i.status_code}, 받은 파일: 빨간 굵은 글씨={red} 채우기={fill} 직접 병합={merged} 새 줄={new_row}"
    await s.case("QS-087", "다시 발행한 파일에 엑셀 서식 유지(글씨·채우기·병합) + 새 줄", qs087)


if __name__ == "__main__":
    tc.main("quote_sync", body)
