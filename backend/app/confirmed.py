"""확정 결과 저장·재활용 — 파인튜닝 대신 "확정본 검색 → few-shot" 으로 입맛을 반영한다.

흐름 (docs/design/proposal_automation_design.md §1 "확정은 영업 담당자", §4.3 proposals):
  1) 생성 직후 create_draft() — 요청 원문 + AI 결과(JSON)를 초안으로 남긴다.
  2) 담당자가 채팅의 확정 카드를 누르면 confirm() — 그대로 또는 수정본(.drawio/.pptx/.docx)과 함께.
     수정본이 있으면 사람이 고친 글자를 우선 저장한다(입맛은 고친 곳에 있다).
     확정본의 요청 원문 임베딩을 같은 행에 둔다. documents(일반 RAG)에 넣지 않는 이유:
     sales 채팅 검색에 다른 담당자의 요청 원문이 근거로 섞여 나오면 안 된다.
  3) 다음 생성 때 get_confirmed_examples(query) — 비슷한 요청의 확정본을 형식 참고 예시로 돌려준다.

왜 파인튜닝이 아닌가: 확정본 1건이 쌓이는 즉시 다음 생성에 반영되고, 잘못 확정한 건은 행 하나로
되돌릴 수 있다. 수십 건으로는 파인튜닝이 일반화되지 않는다(설계 문서 §5.1, §11).
"""
from __future__ import annotations

import io
import json
import logging
import re
import zipfile
from dataclasses import dataclass
from typing import Any
from xml.etree import ElementTree as ET

from sqlalchemy import text
from sqlalchemy.ext.asyncio import AsyncSession

from .rag import _vec_literal, embed_text
from .storage import get_object_stream

logger = logging.getLogger("confirmed")

KINDS = ("concept_map", "proposal_body")
KIND_LABEL = {"concept_map": "공정 개념도", "proposal_body": "제안서 본문"}
# 확정 예시가 이보다 덜 비슷하면 넣지 않는다 — 엉뚱한 공정 예시는 복사만 부른다.
MIN_SCORE = 0.55
MAX_EXAMPLE_CHARS = 3000

# 테이블은 migrations.py 의 _CREATE_PROPOSAL_RECORDS_TABLE 이 만든다.


class RecordError(ValueError):
    """확정 요청을 처리할 수 없는 사용자 오류(없음·권한·형식)."""


@dataclass
class Record:
    id: int
    user_id: int
    kind: str
    title: str
    request_text: str
    content: dict[str, Any]
    status: str
    revised_text: str | None


def confirm_link(record_id: int) -> str:
    """채팅 본문에 넣는 확정 카드 마커. frontend message.tsx 가 카드로 렌더한다."""
    return f"[이 결과를 확정](/api/proposal-records/{record_id}/confirm)"


async def create_draft(
    db: AsyncSession,
    *,
    user_id: int,
    kind: str,
    title: str,
    request_text: str,
    content: dict[str, Any],
    result_attachment_ids: list[int],
) -> int:
    if kind not in KINDS:
        raise ValueError(f"unknown kind: {kind}")
    res = await db.execute(
        text(
            "INSERT INTO proposal_records "
            "(user_id, kind, title, request_text, content, result_attachment_ids) "
            "VALUES (:u, :k, :t, :r, CAST(:c AS jsonb), CAST(:a AS jsonb)) RETURNING id"
        ),
        {"u": user_id, "k": kind, "t": title[:300], "r": request_text,
         "c": json.dumps(content, ensure_ascii=False), "a": json.dumps(result_attachment_ids)},
    )
    rid = int(res.scalar_one())
    await db.commit()
    return rid


async def get_record(db: AsyncSession, record_id: int) -> Record | None:
    row = (await db.execute(
        text("SELECT id, user_id, kind, title, request_text, content, status, revised_text "
             "FROM proposal_records WHERE id = :id"), {"id": record_id})).first()
    if row is None:
        return None
    content = row.content if isinstance(row.content, dict) else json.loads(row.content)
    return Record(row.id, row.user_id, row.kind, row.title, row.request_text, content,
                  row.status, row.revised_text)


# --- 수정본 글자 추출 --------------------------------------------------------

def _drawio_text(data: bytes) -> str:
    """draw.io 파일의 셀 라벨(HTML 태그 제거)을 위→아래, 왼→오른 순으로 잇는다."""
    root = ET.fromstring(data.decode("utf-8", errors="replace"))
    if root.find(".//mxGraphModel") is None:
        raise RecordError("압축된 draw.io 파일은 읽을 수 없습니다. "
                          "draw.io 에서 '파일 > 속성 > 압축' 을 끄고 저장해 주세요.")
    cells = []
    for c in root.iter("mxCell"):
        label = re.sub(r"<br\s*/?>", " / ", c.get("value") or "")
        label = re.sub(r"<[^>]+>", "", label)
        label = re.sub(r"&nbsp;", " ", label)
        label = re.sub(r"\s+", " ", label).strip()
        if not label:
            continue
        g = c.find("mxGeometry")
        y = float(g.get("y", 0)) if g is not None else 0.0
        x = float(g.get("x", 0)) if g is not None else 0.0
        cells.append((y, x, label))
    return "\n".join(lbl for _, _, lbl in sorted(cells))


def _docx_text(data: bytes) -> str:
    import docx  # python-docx

    try:
        d = docx.Document(io.BytesIO(data))
    except (zipfile.BadZipFile, KeyError, ValueError) as e:
        raise RecordError(f"Word 파일을 읽을 수 없습니다: {e}") from e
    lines = [p.text.strip() for p in d.paragraphs if p.text.strip()]
    for t in d.tables:
        for r in t.rows:
            lines.append(" | ".join(c.text.strip() for c in r.cells))
    return "\n".join(lines)


_A_NS = "{http://schemas.openxmlformats.org/drawingml/2006/main}"


def _pptx_text(data: bytes) -> str:
    """슬라이드 순서대로 문단(a:p) 글자를 모은다. 제안서 본문 산출물이 .pptx 다(proposal.pptx_render)."""
    try:
        z = zipfile.ZipFile(io.BytesIO(data))
    except zipfile.BadZipFile as e:
        raise RecordError(f"PowerPoint 파일을 읽을 수 없습니다: {e}") from e
    slides = sorted((n for n in z.namelist() if re.fullmatch(r"ppt/slides/slide\d+\.xml", n)),
                    key=lambda n: int(re.search(r"(\d+)\.xml$", n).group(1)))
    lines: list[str] = []
    for i, name in enumerate(slides, 1):
        try:
            root = ET.fromstring(z.read(name))
        except ET.ParseError as e:
            raise RecordError(f"PowerPoint 슬라이드를 읽을 수 없습니다: {e}") from e
        paras = ["".join(t.text or "" for t in p.iter(f"{_A_NS}t")).strip()
                 for p in root.iter(f"{_A_NS}p")]
        paras = [p for p in paras if p]
        if paras:
            lines.append(f"[슬라이드 {i}]")
            lines.extend(paras)
    return "\n".join(lines)


def extract_revised_text(filename: str, data: bytes) -> str:
    name = filename.lower()
    if name.endswith(".drawio") or name.endswith(".xml"):
        try:
            return _drawio_text(data)
        except ET.ParseError as e:
            raise RecordError(f"draw.io 파일 형식이 아닙니다: {e}") from e
    if name.endswith(".docx"):
        return _docx_text(data)
    if name.endswith(".pptx"):
        return _pptx_text(data)
    raise RecordError("수정본은 .drawio, .pptx 또는 .docx 파일만 올릴 수 있습니다.")


def _read_attachment(object_key: str) -> bytes:
    resp = get_object_stream(object_key)
    try:
        return b"".join(resp.stream(amt=256 * 1024))
    finally:
        resp.close()
        resp.release_conn()


# --- 확정 -------------------------------------------------------------------

async def confirm(
    db: AsyncSession,
    *,
    record_id: int,
    user_id: int,
    is_superadmin: bool = False,
    revised_attachment_id: int | None = None,
) -> Record:
    rec = await get_record(db, record_id)
    if rec is None or (rec.user_id != user_id and not is_superadmin):
        raise RecordError("확정할 결과를 찾을 수 없습니다.")

    revised_text = None
    if revised_attachment_id is not None:
        att = (await db.execute(
            text("SELECT user_id, object_key, original_filename FROM attachments WHERE id = :id"),
            {"id": revised_attachment_id})).first()
        if att is None or (att.user_id != user_id and not is_superadmin):
            raise RecordError("수정본 첨부를 찾을 수 없습니다.")
        revised_text = extract_revised_text(att.original_filename or "",
                                            _read_attachment(att.object_key)).strip()
        if not revised_text:
            raise RecordError("수정본에서 글자를 찾지 못했습니다.")

    # 검색 키는 요청 원문 — 다음에 "비슷한 요청" 이 오면 이 확정본을 찾는다.
    emb = _vec_literal(await embed_text(rec.request_text))
    await db.execute(
        text("UPDATE proposal_records SET status = 'confirmed', confirmed_at = NOW(), "
             "revised_attachment_id = COALESCE(:a, revised_attachment_id), "
             "revised_text = COALESCE(:t, revised_text), embedding = CAST(:e AS vector) "
             "WHERE id = :id"),
        {"a": revised_attachment_id, "t": revised_text, "e": emb, "id": record_id},
    )
    await db.commit()
    logger.info("[confirmed] record=%s kind=%s revised=%s", rec.id, rec.kind,
                revised_text is not None)
    return await get_record(db, record_id)  # type: ignore[return-value]


# --- 재활용 -----------------------------------------------------------------

def _example_body(rec: Record) -> str:
    """few-shot 에 넣을 확정 결과 글자. 수정본이 있으면 그것, 없으면 AI 결과 JSON."""
    if rec.revised_text:
        return f"(담당자가 수정한 최종본)\n{rec.revised_text}"
    return json.dumps(rec.content, ensure_ascii=False, separators=(",", ":"))


async def find_confirmed(
    db: AsyncSession, query: str, *, kind: str, top_k: int = 1, min_score: float = MIN_SCORE,
) -> list[tuple[Record, float]]:
    qvec = _vec_literal(await embed_text(query))
    rows = (await db.execute(
        text("SELECT id, 1 - (embedding <=> CAST(:q AS vector)) AS score "
             "FROM proposal_records WHERE kind = :kind AND status = 'confirmed' "
             "AND embedding IS NOT NULL ORDER BY embedding <=> CAST(:q AS vector) LIMIT :k"),
        {"q": qvec, "kind": kind, "k": top_k})).all()
    out = []
    for r in rows:
        if r.score < min_score:
            continue
        rec = await get_record(db, int(r.id))
        if rec:
            out.append((rec, float(r.score)))
    return out


async def get_confirmed_examples(
    db: AsyncSession, query: str, *, kind: str, top_k: int = 1,
) -> str:
    """비슷한 요청의 확정본 → 프롬프트용 참고 블록. 없으면 빈 문자열."""
    hits = await find_confirmed(db, query, kind=kind, top_k=top_k)
    if not hits:
        return ""
    parts = [
        f"[확정된 과거 {KIND_LABEL[kind]} 예시] 담당자가 확정한 결과다. "
        "항목 구성·표기 방식·설명 밀도만 참고한다. 이 예시의 공정·수치·고객 정보는 "
        "이번 요청과 무관하므로 옮겨 쓰지 않는다."
    ]
    for rec, score in hits:
        body = _example_body(rec)[:MAX_EXAMPLE_CHARS]
        parts.append(f"- 예시 요청: {rec.request_text[:400]}\n- 확정 결과:\n{body}")
    return "\n\n".join(parts)
