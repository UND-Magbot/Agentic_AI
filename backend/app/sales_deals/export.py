# -*- coding: utf-8 -*-
"""제품 영업 건 관리 리스트 → 엑셀(사용자 2026-10-08). 화면 리스트와 같은 순서(최근에 바뀐 건이 위), 금액은 공급가액.

시트 1 '영업 건' — 프로젝트 한 줄 + 바로 아래에 그 히스토리 줄들(엑셀 그룹 +/− 로 접고 펼침, 화면의 [히스토리 N]과 같게).
시트 2 '히스토리 전체' — 모든 프로젝트의 히스토리를 한 표로(필터·정렬용). 시트 3 '요약' — 단계별 건수·분기별 수주금액·매출액.
히스토리는 사용자 2026-10-08 요청(엑셀에 히스토리가 전부 들어가게).
"""
from __future__ import annotations

import datetime as dt
import io
from typing import Any
from zoneinfo import ZoneInfo

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from . import flow

# 프로젝트가 맨 앞(사용자 2026-10-08) — 번호는 프로젝트끼리 같을 수 있어 고객사·건명으로 구분
HEAD = ["고객사", "건명", "번호", "담당·입력", "견적일·바뀐 때", "공급가액", "단계", "수주일", "결제 조건", "확인한 입금(공급가액)",
        "다음 예정 입금일", "첫 출하일", "알림"]
WIDTH = [22, 30, 18, 10, 14, 14, 18, 12, 34, 16, 14, 12, 40]
HIST_HEAD = ["고객사", "건명", "프로젝트 번호", "히스토리 번호", "바뀐 때", "바뀐 내용", "단계", "공급가액", "입력", "비고"]
HIST_WIDTH = [22, 30, 14, 18, 17, 40, 16, 14, 10, 16]
KST = ZoneInfo("Asia/Seoul")
HIST_FILL = PatternFill("solid", fgColor="F6F7F9")
HIST_FONT = Font(color="555555", size=10)
# 화면(deal-format.ACTION_LABEL)과 같은 이름
ACTION_LABEL = {
    "create": "새 건 등록", "won": "수주 상태 변경", "invoice": "세금계산서 기록", "payment": "입금 확인", "paid_full": "완납 표시 변경",
    "shipment": "출하 기록", "remove_invoice": "계산서 기록 삭제", "remove_payment": "입금 기록 삭제", "remove_shipment": "출하 기록 삭제",
    "review_done": "확인 완료", "review_reopen": "확인 되돌리기", "po_upload": "발주서 첨부", "remove_po_file": "발주서 삭제",
    "po_compare": "발주서 품목 대조", "shipment_edit": "출하 기록 수정", "reset_to_quote": "견적 상태로 되돌림(테스트)",
    "statement": "거래명세서 발급", "pay_case": "결제 방식 선택", "pay_terms": "결제 조건 저장", "drop": "드랍", "undrop": "드랍 되살리기",
    "recommend": "제품 추천 확정", "quote": "견적서 발행", "quote_revision": "견적서 수정", "meeting": "미팅 정보 입력",
}
WON = "#,##0"
BOLD = Font(bold=True)
HEAD_FILL = PatternFill("solid", fgColor="E8EEF7")
LATE_FILL = PatternFill("solid", fgColor="FDE2E1")        # 예정 입금일 지남 — 화면 리스트 강조와 같게


def _row(r: dict[str, Any]) -> list[Any]:
    fl = r.get("flow") or {}
    nxt = next((t for t in fl.get("terms") or [] if not t.get("confirmed")), None)
    ships = sorted(s.get("date") for s in r.get("shipments") or [] if s.get("date"))
    return [r.get("customer") or "", r.get("title") or "", r.get("display_no") or r["deal_no"] or "번호 미부여", r.get("owner") or "",
            r.get("base_date") or "", flow.supply_total(r), r.get("stage") or "", r.get("won_date") or "",
            fl.get("pay_case_label") or "", fl.get("paid") or 0, (nxt or {}).get("expected_date") or "",
            ships[0] if ships else "", " / ".join(a["text"] for a in r.get("alerts") or [])]


def _when(iso: str | None) -> str:
    """이력 시각 → 한국 시각 '2026-10-08 14:03'."""
    try:
        return dt.datetime.fromisoformat(str(iso)).astimezone(KST).strftime("%Y-%m-%d %H:%M")
    except ValueError:
        return str(iso or "")[:16].replace("T", " ")


def _what(h: dict[str, Any]) -> str:
    """바뀐 내용 — 이름 + 짧은 덧붙임(금액·날짜 등)."""
    d = h.get("detail") or {}
    a = h.get("action")
    extra = ""
    if a == "won":
        extra = {True: "→ 수주", False: "→ 드랍"}.get(d.get("to"), "→ 미확인")
    elif a == "payment":
        extra = f"{d.get('note') or ''} {d.get('date') or ''} · {d.get('amount') or 0:,.0f}원".strip()
    elif a in ("shipment", "shipment_edit"):
        extra = f"{d.get('date') or ''} · {d.get('carrier') or ''}".strip(" ·")
    elif a == "statement":
        extra = f"{d.get('no') or ''} · 합계 {d.get('total') or 0:,.0f}원"
    elif a == "pay_terms":
        extra = ", ".join(d.get("terms") or [])
    elif a in ("quote", "quote_revision"):
        extra = f"V{d.get('revision')} · {d.get('amount') or 0:,.0f}원" if d.get("revision") else ""
    elif a == "drop":
        extra = str(d.get("reason") or "")
    elif a in ("po_upload", "remove_po_file"):
        extra = str(d.get("filename") or "")
    name = ACTION_LABEL.get(a, a or "")
    return f"{name} — {extra}" if extra else name


def _stage(h: dict[str, Any]) -> str:
    st = h.get("stage") or ""
    return f"{st} V{h['quote_rev']}" if st == "견적" and h.get("quote_rev") else st


def _hist_row(r: dict[str, Any], h: dict[str, Any]) -> list[Any]:
    return [r.get("customer") or "", r.get("title") or "", r.get("deal_no") or "", h.get("ver_no") or "", _when(h.get("created_at")),
            _what(h), _stage(h), h.get("amount"), h.get("user_name") or "시스템", "버전 업데이트됨" if h.get("superseded") else ""]


def build_xlsx(data: dict[str, Any]) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "영업 건"
    ws.append(HEAD)
    for c in ws[1]:
        c.font, c.fill = BOLD, HEAD_FILL
        c.alignment = Alignment(horizontal="center", vertical="center")
    ws.sheet_properties.outlinePr.summaryBelow = False       # 그룹 +/− 단추를 프로젝트 줄(위)에
    for r in data["rows"]:
        ws.append(_row(r))
        for c in ws[ws.max_row][:2]:
            c.font = BOLD
        if (r.get("flow") or {}).get("overdue"):
            for c in ws[ws.max_row]:
                c.fill = LATE_FILL
        # 프로젝트 아래에 히스토리(최신 먼저) — 번호·바뀐 때·바뀐 내용·단계·공급가액·입력, 그룹으로 접힘
        for h in r.get("history") or []:
            ws.append(["", f"  └ {_what(h)}", h.get("ver_no") or "", h.get("user_name") or "시스템", _when(h.get("created_at")),
                       h.get("amount"), _stage(h) + (" (버전 업데이트됨)" if h.get("superseded") else "")])
            ws.row_dimensions[ws.max_row].outlineLevel = 1
            for c in ws[ws.max_row]:
                c.fill, c.font = HIST_FILL, HIST_FONT
    for i, w in enumerate(WIDTH, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    for col in ("F", "J"):
        for c in ws[col][1:]:
            c.number_format = WON
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(HEAD))}{ws.max_row}"

    hs = wb.create_sheet("히스토리 전체")
    hs.append(HIST_HEAD)
    for c in hs[1]:
        c.font, c.fill = BOLD, HEAD_FILL
        c.alignment = Alignment(horizontal="center", vertical="center")
    for r in data["rows"]:
        for h in r.get("history") or []:
            hs.append(_hist_row(r, h))
    for i, w in enumerate(HIST_WIDTH, 1):
        hs.column_dimensions[get_column_letter(i)].width = w
    for c in hs["H"][1:]:
        c.number_format = WON
    hs.freeze_panes = "A2"
    hs.auto_filter.ref = f"A1:{get_column_letter(len(HIST_HEAD))}{max(hs.max_row, 1)}"

    sm = wb.create_sheet("요약")
    k = data.get("kpi") or {}
    sm.append(["기준일", data.get("today")])
    sm.append([])
    sm.append([f"{k.get('year')}년 수주금액(공급가액, 수주일 기준)"])
    for q, v in enumerate(k.get("won_by_quarter") or [], 1):
        sm.append([f"{q}분기", v])
    sm.append(["올해 합계", k.get("won_total")])
    if k.get("won_estimated"):
        sm.append([f"※ 수주일이 없는 옮긴 건 {k['won_estimated']}건은 첫 세금계산서 발행일(없으면 출하일·견적일)로 셈"])
    sm.append([])
    sm.append([f"{k.get('year')}년 매출액(공급가액, 출하일 기준)", k.get("sales")])
    sm.append([])
    sm.append(["단계", "건수"])
    for name, n in (data.get("summary") or {}).items():
        sm.append([name, n])
    sm.column_dimensions["A"].width = 40
    sm.column_dimensions["B"].width = 18
    for row in sm.iter_rows(min_col=2, max_col=2):
        for c in row:
            if isinstance(c.value, (int, float)):
                c.number_format = WON
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
