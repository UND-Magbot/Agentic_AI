# -*- coding: utf-8 -*-
"""거래명세서(납품표) 엑셀 — 회사 양식(docs/examples/statements/(주)유엔디로보틱스_거래명세서_대한TES_VT260825.pdf)을 따른다.

A4 한 장에 같은 명세서 두 장: 위 '공급자보관용', 아래 '공급받는자보관용'(사용자 2026-10-07: 수주 확인 뒤 먼저 발급).
머리(발행번호·발행일자·공급자·공급받는자) → 품목(품명·규격·수량·단가·금액·비고) → 인수자 칸 + 공급가액·부가세·합계.
공급자 칸은 양식 그대로(등록번호·상호·성명·주소, 직인), 공급받는자는 건마다 입력(고객사별로 기억 — service.statement_buyer).
품목 줄은 양식의 8줄보다 많으면 늘린다(한 장에 맞춰 인쇄 — 쪽 맞춤).
세트로 묶어서 발급할 수 있다(사용자 2026-10-08, 출하처럼 'TCC1 툴체인저(ATC) 1 SET') — 품목 한 줄 + unit 'SET',
원래 품목은 parts 로 같이 남긴다(양식에는 세트 한 줄만, 품목별로 되돌릴 때 쓴다).
"""
from __future__ import annotations

import io
import re
from datetime import date
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.drawing.spreadsheet_drawing import AnchorMarker, OneCellAnchor
from openpyxl.drawing.xdr import XDRPositiveSize2D
from openpyxl.utils.units import pixels_to_EMU
from openpyxl.styles import Alignment, Border, Font, Side

from ..company_knowledge.product_prices import COMPANY_PROFILE

# 양식의 공급자 칸(샘플 거래명세서에 적힌 값)
SUPPLIER = COMPANY_PROFILE["statement_supplier"]
MIN_ROWS = 8                       # 양식의 품목 줄 수
MAX_ITEMS = 30
SEAL = Path(__file__).resolve().parents[1] / "company_knowledge" / "data" / "quote" / "seal.png"
FONT = "맑은 고딕"
WON = '"₩"#,##0'
# 오른쪽 맞춤 숫자 — 끝에 빈칸 한 글자('_ ')를 둬서 오른쪽 선에 닿지 않게(LibreOffice 는 오른쪽 맞춤 들여쓰기를 무시한다)
WON_R = '"₩"#,##0_ '
QTY_R = '#,##0_ '
_THIN = Side(style="thin", color="000000")
_MED = Side(style="medium", color="000000")
_HAIR = Side(style="hair", color="000000")          # 양식의 안쪽 가로줄(점선처럼 가는 줄)

# ── 양식 실측(사용자 2026-10-07: 원본과 다르게 보이면 안 됨) ─────────────────────────────────
# 샘플 PDF 는 맑은 고딕 10pt 를 86% 로 인쇄한 것(종이 위 8.6pt). 아래 값은 PDF 에서 잰 '종이 위 pt' 이고,
# 시트에는 ÷ PRINT_SCALE 해서 넣는다. 가로는 좌우 여백을 맞춰 쪽 맞춤 배율이 정확히 86% 가 되게 한다
# (품목이 8줄보다 많으면 세로 쪽 맞춤이 더 줄여 한 장에 넣는다).
PRINT_SCALE = 0.86
# 열 경계(종이 위 pt): 40.7 | 60.0 | 90.2 | 110.1 | 204.8 | 240.8 | 289.4 | 314.3 | 349.6 | 428.8 | 454.4 | 550.8
#   머리: A 공급자 | B 등록번호·상호·주소 | C~F 값(상호값 C·D, 성명 E, 성명값 F) | G 공급받는자 | H 칸이름 | I~K 값(상호값 I, 성명 J, 성명값 K)
#   품목: A~C 품명 | D~E 규격 | F 수량 | G~H 단가 | I 금액 | J~K 비고 / 아래: A 인수자 | B 칸이름 | C~D 값 | E (인) | F~H 공급가액 | I 부가세 | J~K 합계
COL_PT = {"A": 19.3, "B": 30.2, "C": 19.9, "D": 94.7, "E": 36.0, "F": 48.6, "G": 24.9, "H": 35.3, "I": 79.2, "J": 25.6,
          "K": 96.4}
CONTENT_PT = sum(COL_PT.values())                  # 510.1
A4_W_PT = 595.3
SIDE_MARGIN_IN = (A4_W_PT - CONTENT_PT) / 2 / 72   # ≈ 0.59in — 이 여백이어야 가로 쪽 맞춤이 86%
# 행 높이(종이 위 pt) — 한 장
ROW_PT = {"title": 40.0, "no": 12.5, "date": 12.5, "gap1": 6.7, "reg": 27.5, "name": 17.3, "addr": 17.1, "gap2": 19.4,
          "head": 18.5, "item": 15.85, "gap3": 12.5, "sum": 13.6}
GAP_BETWEEN_PT = (12.0, 12.0, 14.0)                # 두 장 사이: 빈 줄 · 자르는 선 · 빈 줄
# 직인 — 샘플 실측(렌더 비교): 지름 약 40pt, 왼쪽 끝 x≈266(성명값 칸 안), 위쪽 y≈144(상호 줄 가운데 — 상호·주소 줄에 걸침)
SEAL_PT = 40.0
SEAL_LEFT_PT = 266.0 - 240.8                       # F 열 왼쪽 끝에서
SEAL_TOP_PT = 144.0 - 107.7                        # 등록번호 줄 위에서


def _sheet_pt(paper_pt: float) -> float:
    return paper_pt / PRINT_SCALE


def _col_width(paper_pt: float) -> float:
    """종이 위 pt → 열 너비(문자 수). 기본 글꼴(11pt) 숫자 너비 7px, 칸 여백 5px, 1px = 0.75pt."""
    return round((_sheet_pt(paper_pt) / 0.75 - 5) / 7, 2)


def _px(paper_pt: float) -> int:
    return round(_sheet_pt(paper_pt) / 0.75)


class StatementError(ValueError):
    """거래명세서를 만들 수 없는 이유(사람이 읽는 문장)."""


SET_UNIT = "SET"


def unit_of(it: dict[str, Any]) -> str:
    """수량 단위(예: SET) — 서식 문자열에 들어가므로 글자·숫자만."""
    return re.sub(r"[^0-9A-Za-z가-힣]", "", str(it.get("unit") or ""))[:6]


def bundle(items: list[dict[str, Any]], name: str, qty: float = 1) -> tuple[list[dict[str, Any]], list[dict[str, Any]]]:
    """품목들 → 세트 한 줄(단가 = 품목 합계 ÷ 세트 수량, 공급가액은 그대로). 반환 (새 품목, 구성 품목)."""
    parts = [{k: it.get(k) for k in ("name", "spec", "qty", "unit_price", "note")} for it in items if str(it.get("name") or "").strip()]
    supply = sum((it.get("qty") or 0) * (it.get("unit_price") or 0) for it in parts)
    qty = qty if qty and qty > 0 else 1
    qty = int(qty) if qty == int(qty) else qty            # 말에서 받은 1.0 → 1('1 SET')
    price = round(supply / qty)
    return [{"name": name.strip()[:80], "spec": "", "qty": qty, "unit": SET_UNIT, "unit_price": price, "note": ""}], parts


def tidy_name(name: Any) -> str:
    """품명 정리 — 단가표 품명의 '0.3m(meter)' 같은 단위 풀이는 뺀다(칸이 좁아 글자가 아주 작게 줄어들었다)."""
    return re.sub(r"(\d)\s*m\s*\(\s*meters?\s*\)", lambda m: m.group(1) + "m", str(name or ""), flags=re.I)


def split_name(name: str) -> tuple[str, str]:
    """견적 품목 글 → (품명, 규격). 견적서 품목은 'TCV1\\n- MASTER T.C (Payload 10Kg)' · 'PPM\\n-Pogo Pin Male' 꼴 —
    양식 예시처럼 'TCV1 MASTER T.C' / 'Payload 10Kg', 'PPM' / 'Pogo Pin Male' 로 나눈다. 꼴이 다르면 첫 줄 = 품명, 나머지 = 규격."""
    text = tidy_name(name)
    lines = [x.strip().lstrip("-").strip() for x in text.splitlines() if x.strip()]
    if not lines:
        return "", ""
    code, rest = lines[0], " ".join(lines[1:])
    if "(" in rest and rest.endswith(")"):
        head, spec = rest[:rest.index("(")].strip(), rest[rest.index("(") + 1:-1].strip()
        return (f"{code} {head}".strip(), spec)
    return code, rest


def totals(items: list[dict[str, Any]]) -> dict[str, int]:
    supply = int(round(sum((it.get("qty") or 0) * (it.get("unit_price") or 0) for it in items)))
    vat = int(round(supply * 0.1))
    return {"supply": supply, "vat": vat, "total": supply + vat}


def check(st: dict[str, Any]) -> None:
    """발급 전 검사 — 빠진 값이 있으면 StatementError(무엇이 빠졌는지)."""
    b = st.get("buyer") or {}
    miss = [lab for k, lab in (("reg_no", "공급받는자 등록번호"), ("name", "공급받는자 상호"), ("address", "공급받는자 주소"))
            if not str(b.get(k) or "").strip()]
    items = st.get("items") or []
    if not items:
        miss.append("품목")
    if any(not str(it.get("name") or "").strip() or not it.get("qty") or it.get("unit_price") is None for it in items):
        miss.append("품목의 품명·수량·단가")
    if miss:
        raise StatementError("거래명세서에 꼭 필요한 값이 비어 있습니다: " + ", ".join(miss))
    if len(items) > MAX_ITEMS:
        raise StatementError(f"품목은 {MAX_ITEMS}줄까지입니다.")


def _cell(ws, ref: str, value: Any = None, *, size: float = 10, bold: bool = False, h: str = "center",
          fmt: str | None = None, merge: str | None = None, underline: str | None = None, indent: int = 0,
          shrink: bool = False) -> None:
    """값·글꼴·정렬. 테두리는 칸 묶음(_frame)으로 따로 그린다 — 양식처럼 바깥은 굵게, 안쪽 가로줄은 가늘게."""
    if merge:
        ws.merge_cells(f"{ref}:{merge}")
    c = ws[ref]
    c.value = value
    c.font = Font(name=FONT, size=size, bold=bold, underline=underline)
    c.alignment = Alignment(horizontal=h, vertical="center", wrap_text=not shrink, shrink_to_fit=shrink, indent=indent)
    if fmt:
        c.number_format = fmt


def _frame(ws, top: int, left: str, bottom: int, right: str) -> None:
    """칸 묶음 테두리 — 바깥 굵은 선, 안쪽 세로 가는 실선, 안쪽 가로 점선(양식 그대로)."""
    cols = [chr(c) for c in range(ord(left), ord(right) + 1)]
    for r in range(top, bottom + 1):
        for col in cols:
            ws[f"{col}{r}"].border = Border(left=_MED if col == left else _THIN, right=_MED if col == right else _THIN,
                                            top=_MED if r == top else _HAIR, bottom=_MED if r == bottom else _HAIR)


def _height(ws, row: int, paper_pt: float) -> None:
    ws.row_dimensions[row].height = round(_sheet_pt(paper_pt), 2)


def _copy(ws, top: int, st: dict[str, Any], label: str, n_rows: int) -> int:
    """명세서 한 장을 top 행부터 그린다. 다음 장을 그릴 행을 돌려준다."""
    b, sup = st["buyer"], SUPPLIER
    day: date = date.fromisoformat(st["date"])
    r = top
    _cell(ws, f"A{r}", "거래명세서 (납 품 표)", size=22.5, bold=True, merge=f"K{r}", underline="single")
    _height(ws, r, ROW_PT["title"])
    r += 1
    _cell(ws, f"A{r}", "발행번호", h="left", merge=f"B{r}", indent=1)
    _cell(ws, f"C{r}", st.get("no") or "", h="left", merge=f"E{r}")
    _cell(ws, f"I{r}", label, h="right", merge=f"K{r}", indent=1)
    _height(ws, r, ROW_PT["no"])
    r += 1
    _cell(ws, f"A{r}", "발행일자", h="left", merge=f"B{r}", indent=1)
    _cell(ws, f"C{r}", f"{day.year}년 {day.month}월 {day.day}일", h="left", merge=f"E{r}")
    _height(ws, r, ROW_PT["date"])
    r += 1
    _height(ws, r, ROW_PT["gap1"])
    r += 1
    # 공급자 / 공급받는자 — 3줄(등록번호 · 상호/성명 · 주소)
    h0 = r
    _cell(ws, f"A{r}", "공\n급\n자", merge=f"A{r + 2}")
    _cell(ws, f"B{r}", "등록\n번호")                          # 양식은 '등록번/호'로 어색하게 꺾였다 — 두 줄로 깔끔하게
    _cell(ws, f"C{r}", sup["reg_no"], size=16, merge=f"F{r}")
    _cell(ws, f"G{r}", "공급\n받는\n자", size=9.2, merge=f"G{r + 2}")
    _cell(ws, f"H{r}", "등록\n번호")
    _cell(ws, f"I{r}", b.get("reg_no") or "", size=16, merge=f"K{r}")
    _height(ws, r, ROW_PT["reg"])
    r += 1
    _cell(ws, f"B{r}", "상호")
    _cell(ws, f"C{r}", sup["name"], merge=f"D{r}")
    _cell(ws, f"E{r}", "성명")
    _cell(ws, f"F{r}", sup["ceo"])
    _cell(ws, f"H{r}", "상호")
    _cell(ws, f"I{r}", b.get("name") or "", shrink=True)
    _cell(ws, f"J{r}", "성명")
    _cell(ws, f"K{r}", b.get("ceo") or "")
    _height(ws, r, ROW_PT["name"])
    r += 1
    _cell(ws, f"B{r}", "주소")
    _cell(ws, f"C{r}", sup["address"], size=9.2, merge=f"F{r}", shrink=True)
    _cell(ws, f"H{r}", "주소")
    _cell(ws, f"I{r}", b.get("address") or "", size=9.2, merge=f"K{r}", shrink=True)
    _height(ws, r, ROW_PT["addr"])
    _frame(ws, h0, "A", r, "F")
    _frame(ws, h0, "G", r, "K")
    if SEAL.exists():
        img = XLImage(str(SEAL))
        img.width = img.height = _px(SEAL_PT)
        mark = AnchorMarker(col=5, colOff=pixels_to_EMU(_px(SEAL_LEFT_PT)), row=h0 - 1,
                            rowOff=pixels_to_EMU(_px(SEAL_TOP_PT)))
        img.anchor = OneCellAnchor(_from=mark, ext=XDRPositiveSize2D(pixels_to_EMU(_px(SEAL_PT)), pixels_to_EMU(_px(SEAL_PT))))
        ws.add_image(img)
    r += 1
    _height(ws, r, ROW_PT["gap2"])
    r += 1
    # 품목
    head = r
    for col, (lab, end) in {"A": ("품명", "C"), "D": ("규격", "E"), "F": ("수량", None), "G": ("단가", "H"),
                            "I": ("금액", None), "J": ("비고", "K")}.items():
        _cell(ws, f"{col}{r}", lab, merge=f"{end}{r}" if end else None)
    _height(ws, r, ROW_PT["head"])
    r += 1
    items = st["items"]
    for i in range(n_rows):
        it = items[i] if i < len(items) else None
        _cell(ws, f"A{r}", tidy_name(it["name"]) if it else None, size=9.2, merge=f"C{r}", shrink=True)
        _cell(ws, f"D{r}", it.get("spec") if it else None, merge=f"E{r}", shrink=True)
        # 숫자 칸 — 오른쪽 선과 한 글자 띄우고(서식), 큰 금액이면 칸에 맞게 살짝 줄인다
        # 세트면 '1 SET' 처럼 단위를 붙여 보이게(값은 숫자 그대로 — 금액 = 수량 × 단가)
        unit = unit_of(it) if it else ""
        _cell(ws, f"F{r}", it["qty"] if it else None, h="right", fmt=f'#,##0" {unit}"_ ' if unit else QTY_R, shrink=True)
        _cell(ws, f"G{r}", it["unit_price"] if it else None, h="right", fmt=WON_R, merge=f"H{r}", shrink=True)
        _cell(ws, f"I{r}", f"=F{r}*G{r}" if it else None, h="right", fmt=WON_R, shrink=True)
        _cell(ws, f"J{r}", it.get("note") if it else None, size=9.2, merge=f"K{r}", shrink=True)
        _height(ws, r, ROW_PT["item"])
        r += 1
    first, last = r - n_rows, r - 1
    _frame(ws, head, "A", last, "K")
    _height(ws, r, ROW_PT["gap3"])
    r += 1
    # 인수자 + 합계
    _cell(ws, f"A{r}", "인\n수\n자", merge=f"A{r + 2}")
    for j, lab in enumerate(("TEL", "부서", "성명")):
        _cell(ws, f"B{r + j}", lab)
        _cell(ws, f"C{r + j}", None, merge=f"D{r + j}")
        _height(ws, r + j, ROW_PT["sum"])
    _cell(ws, f"E{r + 2}", "(인)")
    _frame(ws, r, "A", r + 2, "D")
    _cell(ws, f"F{r}", "공급가액", merge=f"H{r}")
    _cell(ws, f"I{r}", "부가세")
    _cell(ws, f"J{r}", "합계", merge=f"K{r}")
    _cell(ws, f"F{r + 1}", f"=SUM(I{first}:I{last})", fmt=WON, merge=f"H{r + 2}")
    _cell(ws, f"I{r + 1}", f"=ROUND(F{r + 1}*0.1,0)", fmt=WON, merge=f"I{r + 2}")
    _cell(ws, f"J{r + 1}", f"=F{r + 1}+I{r + 1}", fmt=WON, merge=f"K{r + 2}")
    _frame(ws, r, "F", r + 2, "K")
    return r + 3


def build_xlsx(st: dict[str, Any], *, strict: bool = True) -> bytes:
    """거래명세서 정보 → 엑셀 바이트(공급자보관용 + 공급받는자보관용). 빠진 값이 있으면 StatementError.
    strict=False 는 발급 전 미리보기 — 빈 칸은 빈 채로, 품명이 빈 줄은 뺀다."""
    if strict:
        check(st)
    else:
        st = {**st, "buyer": st.get("buyer") or {},
              "items": [{**it, "qty": it.get("qty") or 0, "unit_price": it.get("unit_price") or 0}
                        for it in st.get("items") or [] if str(it.get("name") or "").strip()][:MAX_ITEMS]}
    wb = Workbook()
    ws = wb.active
    ws.title = "거래명세서"
    ws.sheet_view.showGridLines = False
    for col, w in COL_PT.items():
        ws.column_dimensions[col].width = _col_width(w)
    n = max(MIN_ROWS, len(st["items"]))
    nxt = _copy(ws, 1, st, "공급자보관용", n)
    # 두 장 사이 — 빈 줄 · 자르는 선(굵은 가로줄) · 빈 줄
    for k, h in enumerate(GAP_BETWEEN_PT):
        _height(ws, nxt + k, h)
    for col in COL_PT:
        ws[f"{col}{nxt}"].border = Border(bottom=_MED)
    _copy(ws, nxt + len(GAP_BETWEEN_PT), st, "공급받는자보관용", n)
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.orientation = "portrait"
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 1
    ws.print_options.horizontalCentered = True
    ws.page_margins.left = ws.page_margins.right = round(SIDE_MARGIN_IN, 3)
    ws.page_margins.top, ws.page_margins.bottom = 0.5, 0.4
    ws.page_margins.header = ws.page_margins.footer = 0.2
    wb.calculation.fullCalcOnLoad = True
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def file_name(st: dict[str, Any]) -> str:
    """'(주)유엔디로보틱스_거래명세서_{고객}_{발행번호}' — 샘플 파일 이름 꼴."""
    import re

    customer = re.sub(r'[\\/:*?"<>|]+', " ", str(st["buyer"].get("name") or "")).strip() or "고객사"
    return f"(주)유엔디로보틱스_거래명세서_{customer}_{st.get('no') or st['date']}"
