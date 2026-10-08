# -*- coding: utf-8 -*-
"""기존 엑셀(견적·세금·수주 관리 시트 + 영업 출하시트) → 건 단위 이관 미리보기.

두 단계로 나눈다.
  1) `build_snapshot(관리시트, 출하시트)` — 엑셀을 읽어 원자료(건·품목·출하·사진 위치)를 dict 로.
     엑셀이 필요한 단계라 스크립트(scripts/build_sales_deals_preview.py)에서 한 번 돌려 JSON 으로 저장.
  2) `build_preview(snapshot, today)` — 원자료 → 리스트 행(단계·계산서·입금·출하·알림).
     오늘 날짜에 따라 '결과 미확인 N일' 같은 알림이 바뀌므로 API 요청마다 계산한다.

원칙
  - 원본 엑셀은 읽기만 한다.
  - 금액·상태는 칸에 적힌 값에서만 계산한다. 판단이 애매하면 지어내지 않고 '확인 필요' 알림으로 남긴다.
  - 같은 고객 견적이 차수인지 별건인지는 사람이 정한다(후보만 제안).
"""
from __future__ import annotations

import datetime as dt
import difflib
import re
import zipfile
from pathlib import Path
from typing import Any

import openpyxl

from ..finance.vendor_registry import normalize_vendor
from .status import STAGES, derive

BROKEN_CELL = "########"          # 열 너비가 좁아 날짜가 '#'로 저장된 칸(원본에 값이 남아 있지 않음)
REVISION_WINDOW_DAYS = 45         # 같은 고객 견적이 이 기간 안에 이어지면 차수 후보
SHIP_LOOKBACK_DAYS = 200          # 출하 ↔ 견적 연결 시 출하일 기준 이 기간 안의 견적만 후보
SHIP_LOOKAHEAD_DAYS = 7           # 견적일이 출하일보다 조금 늦게 적힌 경우 허용
PARTIAL_RATIO = 0.9               # 발행액/견적액이 이보다 작으면 일부 청구(그 위는 할인·옵션 차이)
VAT_RATE = 0.1

# 관리 시트 열 번호(1부터). 머리글: 2행.
M_NO, M_DATE, M_OWNER, M_CUST, M_CONTACT = 2, 3, 4, 5, 6
M_ITEM, M_PRICE, M_QTY, M_AMOUNT, M_NOTE = 7, 8, 9, 10, 11
M_SENT, M_RESULT, M_INV_FLAG, M_INV_DATE = 12, 13, 14, 15
M_SUPPLY, M_VAT, M_TOTAL, M_PAID = 16, 17, 18, 19

# 출하시트 열 번호. 머리글: 2행.
S_NO, S_CUST, S_DATE, S_ITEM, S_QTY, S_PRICE, S_AMOUNT = 1, 2, 3, 4, 5, 6, 7
S_PHOTO, S_CARRIER, S_TRACK, S_RECEIVER, S_ADDR, S_OWNER, S_PURPOSE = 8, 9, 10, 11, 12, 13, 14


# ── 셀 값 정리 ────────────────────────────────────────────────────────────────

def _text(v: Any) -> str:
    return "" if v is None else str(v).strip()


def _num(v: Any) -> float | None:
    if isinstance(v, bool):
        return None
    if isinstance(v, (int, float)):
        return float(v)
    s = _text(v).replace(",", "").replace("₩", "").replace("\\", "")
    return float(s) if re.fullmatch(r"-?\d+(\.\d+)?", s) else None


def _iso(v: Any) -> str | None:
    if isinstance(v, dt.datetime):
        return v.date().isoformat()
    if isinstance(v, dt.date):
        return v.isoformat()
    return None


def _cell_date(v: Any) -> tuple[str | None, str | None]:
    """(날짜 ISO, 날짜가 아닌 글) — 깨진 칸·'7월6일 예정' 같은 글은 두 번째로."""
    d = _iso(v)
    if d:
        return d, None
    s = _text(v)
    return None, (s or None)


def _won_amounts(text: str) -> list[int]:
    return [int(m.replace(",", "")) for m in re.findall(r"₩\s*([\d,]{4,})", text)]


def customer_key(name: str | None) -> str:
    """고객사 비교용 키 — 회사형태·괄호·공장/사업장 표기·기호 제거."""
    s = normalize_vendor(name)
    s = re.sub(r"대구공장|화성공장|공장|구미\s*2?\s*사업.*|건설부문", "", s)
    return re.sub(r"[\s\-·.,]", "", s)


def same_customer(a: str | None, b: str | None) -> bool:
    ka, kb = customer_key(a), customer_key(b)
    if not ka or not kb:
        return False
    if ka == kb or (min(len(ka), len(kb)) >= 3 and (ka in kb or kb in ka)):
        return True
    return difflib.SequenceMatcher(None, ka, kb).ratio() >= 0.75


# ── 결제조건 ─────────────────────────────────────────────────────────────────

_KW = [("선금", "발주 시"), ("발주", "발주 시"), ("advance", "발주 시"),
       ("납품 전", "납품 전"), ("납품전", "납품 전"), ("before delivery", "납품 전"),
       ("납품 후", "납품 후"), ("납품후", "납품 후"), ("잔금", None)]


def parse_pay_terms(text: str) -> list[dict]:
    """비고 글 → 결제 회차 [{label, pct, when, ship_gate}]. 못 읽으면 빈 목록(지어내지 않음)."""
    t = re.sub(r"\s+", " ", text or "")
    low = t.lower()
    gate = bool(re.search(r"입금\s*확인\s*후\s*(출고|출하)", t))
    out: list[dict] = []
    for m in re.finditer(r"(\d{1,3})\s*%", t):
        pct = int(m.group(1))
        before = low[max(0, m.start() - 30):m.start()]
        after = low[m.end():m.end() + 15]
        if "할인" in before[-8:] or "할인" in after[:6] or not 0 < pct <= 100:
            continue
        hits = [(before.rfind(k), when) for k, when in _KW if before.rfind(k) >= 0]
        if not hits:
            continue
        when = max(hits)[1]
        if when is None:  # '잔금' 만 있으면 앞뒤에서 납품 전/후를 찾는다
            around = after + " " + before
            when = "납품 후" if re.search(r"납품\s*후", around) else "납품 전"
        out.append({"label": "선금" if when == "발주 시" else "잔금", "pct": pct, "when": when,
                    "ship_gate": gate and when == "납품 전"})
    if not out and re.search(r"익월\s*(초|말)", t):
        when = "납품 후 익월" + ("말" if "익월말" in t.replace(" ", "") else "초")
        out.append({"label": "잔금", "pct": 100, "when": when, "ship_gate": False})
    if len(out) == 1 and out[0]["pct"] == 100:
        out[0]["label"] = "전액"
    return out


# ── 관리 시트 ────────────────────────────────────────────────────────────────

def parse_management(path: str | Path) -> list[dict]:
    """관리 시트 → 건 원자료. B열(번호)이 있는 행이 건의 시작, 아래 G열 행들이 품목."""
    ws = openpyxl.load_workbook(path, data_only=True).worksheets[0]
    deals: list[dict] = []
    cur: dict | None = None
    for r in range(3, ws.max_row + 1):
        v = lambda c: ws.cell(r, c).value  # noqa: E731
        if v(M_NO) is not None:
            no = _text(v(M_NO))
            quote_date, quote_date_text = _cell_date(v(M_DATE))
            sent, _ = _cell_date(v(M_SENT))
            inv_date, inv_text = _cell_date(v(M_INV_DATE))
            paid, paid_text = _cell_date(v(M_PAID))
            cur = {
                "legacy_no": no, "kind": "sales_only" if no.upper().startswith("T") else "quote",
                "row_start": r, "row_end": r,
                "quote_date": quote_date, "quote_date_text": quote_date_text,
                "owner": _text(v(M_OWNER)) or None, "customer": _text(v(M_CUST)) or None,
                "contact": _text(v(M_CONTACT)) or None, "note": _text(v(M_NOTE)),
                "sent_date": sent, "result": _text(v(M_RESULT)) or None,
                "inv_flag": _text(v(M_INV_FLAG)) or None,
                "inv_date": inv_date, "inv_date_text": inv_text,
                "supply": _num(v(M_SUPPLY)), "vat": _num(v(M_VAT)), "total": _num(v(M_TOTAL)),
                "paid_date": paid, "paid_text": paid_text, "items": [],
            }
            deals.append(cur)
        if cur is None:
            continue
        if v(M_ITEM) is not None or v(M_PRICE) is not None:
            cur["row_end"] = r
            cur["items"].append({"name": _text(v(M_ITEM)), "unit_price": _num(v(M_PRICE)),
                                 "qty": _num(v(M_QTY)), "qty_text": _text(v(M_QTY)) or None,
                                 "amount": _num(v(M_AMOUNT))})
    return deals


# ── 출하시트 ─────────────────────────────────────────────────────────────────

def _month_day(text: str) -> tuple[int, int] | None:
    m = re.search(r"(\d{1,2})\s*월\s*(\d{1,2})\s*일", text)
    return (int(m.group(1)), int(m.group(2))) if m else None


def parse_shipments(path: str | Path) -> list[dict]:
    """출하시트 → 출하 원자료. NO 나 출하일이 있는 행이 출하의 시작, 빈 NO·날짜 행은 앞 출하에 이어 붙인다."""
    ws = openpyxl.load_workbook(path, data_only=True).worksheets[0]
    photos = shipment_photo_rows(path)
    ships: list[dict] = []
    last: dt.date | None = None
    for r in range(3, ws.max_row + 1):
        v = lambda c: ws.cell(r, c).value  # noqa: E731
        starts = v(S_NO) is not None or v(S_DATE) is not None
        if not starts:
            if ships and any(v(c) for c in (S_ITEM, S_QTY, S_TRACK)):
                ships[-1]["row_end"] = r
            if ships and r in photos:
                ships[-1]["photo_refs"] += photos[r]
            continue
        date_iso = _iso(v(S_DATE))
        guessed = False
        if not date_iso and (md := _month_day(_text(v(S_DATE)))):
            # 연도 없는 '09월 11일' — 위 행 날짜의 연도로, 달이 크게 거꾸로 가면 다음 해로 본다
            year = last.year if last else dt.date.today().year
            d = dt.date(year, *md)
            if last and d < last - dt.timedelta(days=60):
                d = dt.date(year + 1, *md)
            date_iso, guessed = d.isoformat(), True
        if date_iso:
            last = dt.date.fromisoformat(date_iso)
        track = [t.strip(" /") for t in re.split(r"[\n|]", _text(v(S_TRACK))) if t.strip(" /")]
        ships.append({
            "row_start": r, "row_end": r, "legacy_no": _text(v(S_NO)) or None,
            "customer": re.sub(r"\s+", " ", _text(v(S_CUST))) or None,
            "date": date_iso, "date_guessed": guessed, "date_text": None if date_iso else (_text(v(S_DATE)) or None),
            "items": _text(v(S_ITEM)), "qty_text": _text(v(S_QTY)) or None,
            "unit_price": _num(v(S_PRICE)), "amount": _num(v(S_AMOUNT)), "amount_text": _text(v(S_AMOUNT)) or None,
            "carrier": _text(v(S_CARRIER)) or None,
            "tracking": [t for t in track if t.lower() != "x"],
            "receiver": _text(v(S_RECEIVER)) or None, "address": _text(v(S_ADDR)) or None,
            "owner": _text(v(S_OWNER)) or None, "purpose": _text(v(S_PURPOSE)) or None,
            "photo_refs": list(photos.get(r, [])),
        })
    _fill_blank_customers(ships)
    return ships


def _fill_blank_customers(ships: list[dict]) -> None:
    """고객사 칸이 빈 출하 — 같은 받는 사람(연락처)이 적힌 다른 출하의 고객사로 채우고 표시한다."""
    by_receiver = {s["receiver"]: s["customer"] for s in ships if s["customer"] and s["receiver"]}
    for s in ships:
        s["customer_guessed"] = False
        if not s["customer"] and s["receiver"] in by_receiver:
            s["customer"], s["customer_guessed"] = by_receiver[s["receiver"]], True


def shipment_photo_rows(path: str | Path) -> dict[int, list[str]]:
    """출하시트 사진 위치 → {엑셀 행: [xlsx 안 media 경로]}. 셀 안 사진(richData)과 떠 있는 그림 모두."""
    rows: dict[int, list[str]] = {}
    with zipfile.ZipFile(path) as z:
        names = set(z.namelist())
        sheet = z.read("xl/worksheets/sheet1.xml").decode("utf-8")
        # 셀 안 사진: <c r="H10" vm="1"> → rdrichvalue 의 (vm-1)번째 rv → 첫 값 = richValueRel 순번 → media
        if "xl/richData/rdrichvalue.xml" in names:
            rv = re.findall(r"<rv [^>]*><v>(\d+)</v>", z.read("xl/richData/rdrichvalue.xml").decode())
            rel_ids = re.findall(r'r:id="([^"]+)"', z.read("xl/richData/richValueRel.xml").decode())
            targets = dict(re.findall(r'Id="([^"]+)"[^>]*Target="\.\./media/([^"]+)"',
                                      z.read("xl/richData/_rels/richValueRel.xml.rels").decode()))
            for col_row, vm in re.findall(r'<c r="[A-Z]+(\d+)"[^>]*vm="(\d+)"', sheet):
                i = int(vm) - 1
                if i < len(rv) and int(rv[i]) < len(rel_ids) and rel_ids[int(rv[i])] in targets:
                    rows.setdefault(int(col_row), []).append("xl/media/" + targets[rel_ids[int(rv[i])]])
        # 떠 있는 그림: drawing 의 앵커 시작 행(0부터) → 엑셀 행
        if "xl/drawings/drawing1.xml" in names:
            drawing = z.read("xl/drawings/drawing1.xml").decode("utf-8")
            dtargets = dict(re.findall(r'Id="([^"]+)"[^>]*Target="\.\./media/([^"]+)"',
                                       z.read("xl/drawings/_rels/drawing1.xml.rels").decode()))
            for anchor in re.findall(r"<xdr:(?:twoCellAnchor|oneCellAnchor)[\s\S]*?</xdr:(?:twoCellAnchor|oneCellAnchor)>", drawing):
                row = re.search(r"<xdr:from>[\s\S]*?<xdr:row>(\d+)</xdr:row>", anchor)
                emb = re.search(r'r:embed="([^"]+)"', anchor)
                if row and emb and emb.group(1) in dtargets:
                    rows.setdefault(int(row.group(1)) + 1, []).append("xl/media/" + dtargets[emb.group(1)])
    return rows


def build_snapshot(management_path: str | Path, shipment_path: str | Path) -> dict:
    """엑셀 두 파일 → 원자료 스냅샷(JSON 으로 저장 가능)."""
    return {
        "sources": {"management": Path(management_path).name, "shipments": Path(shipment_path).name},
        "deals": parse_management(management_path),
        "shipments": parse_shipments(shipment_path),
    }


# ── 미리보기 계산 ────────────────────────────────────────────────────────────

def _line_amount(it: dict) -> float | None:
    if it["amount"] is not None:
        return it["amount"]
    if it["unit_price"] is not None:
        return it["unit_price"] * (it["qty"] if it["qty"] is not None else 1)
    return None


def _quote_amount(d: dict) -> float | None:
    vals = [a for a in (_line_amount(i) for i in d["items"]) if a is not None]
    return sum(vals) if vals else None


def _title(d: dict) -> str:
    first = next((i["name"] for i in d["items"] if i["name"]), "") or d.get("note", "")
    line = first.split("\n")[0]
    line = re.sub(r"^(모델명\s*:\s*)", "", line).strip(" -")
    extra = len([i for i in d["items"] if i["name"]]) - 1
    return (line[:40] + ("…" if len(line) > 40 else "")) + (f" 외 {extra}" if extra > 0 else "")


def _revision_groups(deals: list[dict]) -> dict[int, list[int]]:
    """같은 고객 견적이 REVISION_WINDOW_DAYS 안에 이어지면 한 묶음(차수 후보). {deal 순번: 묶음 순번들}."""
    quotes = sorted((i for i, d in enumerate(deals) if d["kind"] == "quote" and d["customer"] and d["quote_date"]),
                    key=lambda i: deals[i]["quote_date"])
    parent = {i: i for i in quotes}

    def find(i: int) -> int:
        while parent[i] != i:
            parent[i] = parent[parent[i]]
            i = parent[i]
        return i

    for a_pos, a in enumerate(quotes):
        da = dt.date.fromisoformat(deals[a]["quote_date"])
        for b in quotes[a_pos + 1:]:
            db = dt.date.fromisoformat(deals[b]["quote_date"])
            if (db - da).days > REVISION_WINDOW_DAYS:
                break
            if same_customer(deals[a]["customer"], deals[b]["customer"]):
                parent[find(b)] = find(a)
    groups: dict[int, list[int]] = {}
    for i in quotes:
        groups.setdefault(find(i), []).append(i)
    return {i: g for g in groups.values() if len(g) > 1 for i in g}


def _match_shipments(deals: list[dict], ships: list[dict]) -> dict[int, list[tuple[int, bool]]]:
    """출하 → 건 연결. {deal 순번: [(출하 순번, 금액으로 확인됐는지)]}. 못 찾은 출하는 연결하지 않는다."""
    linked: dict[int, list[tuple[int, bool]]] = {}
    for si, s in enumerate(ships):
        if not s["customer"] or not s["date"]:
            continue
        sd = dt.date.fromisoformat(s["date"])
        best: tuple[tuple, int, bool] | None = None
        for di, d in enumerate(deals):
            base = d["quote_date"] or d["inv_date"]
            if not base or not same_customer(s["customer"], d["customer"]):
                continue
            gap = (sd - dt.date.fromisoformat(base)).days
            if not -SHIP_LOOKAHEAD_DAYS <= gap <= SHIP_LOOKBACK_DAYS:
                continue
            amounts = {a for i in d["items"] for a in (i["unit_price"], _line_amount(i)) if a}
            amounts |= {_quote_amount(d)} | set(_won_amounts(d["note"]))
            by_amount = any(a and a in amounts for a in (s["amount"], s["unit_price"]))
            score = (by_amount, -abs(gap))
            if best is None or score > best[0]:
                best = (score, di, by_amount)
        if best:
            linked.setdefault(best[1], []).append((si, best[2]))
    return linked


def _alert(code: str, level: str, text: str) -> dict:
    return {"code": code, "level": level, "text": text}


def _won(n: float | None) -> str:
    return f"{n:,.0f}원" if n is not None else "-"


def build_preview(snapshot: dict, today: dt.date | None = None) -> dict:
    """스냅샷 → 리스트 화면용 데이터 {rows, summary, sources, today}."""
    today = today or dt.date.today()
    deals, ships = snapshot["deals"], snapshot["shipments"]
    rev = _revision_groups(deals)
    linked = _match_shipments(deals, ships)
    used_ships = {si for v in linked.values() for si, _ in v}
    no_count: dict[str, int] = {}
    for d in deals:
        no_count[d["legacy_no"]] = no_count.get(d["legacy_no"], 0) + 1

    rows: list[dict] = []
    prev_date: str | None = None
    for di, d in enumerate(deals):
        row = _deal_row(d, di, deals, ships, rev, linked.get(di, []), no_count, prev_date)
        prev_date = row["base_date"] or prev_date
        rows.append(row)
    for si, s in enumerate(ships):
        if si not in used_ships:
            rows.append(_shipment_only_row(s))

    rows.sort(key=lambda r: (r["base_date"] or "0000", r["source"]["row_start"]))
    seq: dict[str, int] = {}
    for r in rows:
        yy = (r["base_date"] or "0000")[2:4]
        seq[yy] = seq.get(yy, 0) + 1
        r["id"] = f"S{yy}-{seq[yy]:04d}"
    rows.reverse()  # 최근 건이 위로
    for r in rows:
        view = derive(r, today)
        r.update(stage=view["stage"], shipped=view["shipped"], invoice_status=view["invoice_status"],
                 alerts=r["review_alerts"] + view["alerts"])

    summary = {s: sum(1 for r in rows if r["stage"] == s) for s in STAGES}
    summary["확인 필요"] = sum(1 for r in rows if any(a["level"] == "review" for a in r["alerts"]))
    summary["알림"] = sum(1 for r in rows if any(a["level"] == "alert" for a in r["alerts"]))
    return {"today": today.isoformat(), "sources": snapshot["sources"], "summary": summary, "rows": rows,
            "counts": {"deals": len(deals), "shipments": len(ships), "shipments_linked": len(used_ships)}}


def _deal_row(d: dict, di: int, deals: list[dict], ships: list[dict], rev: dict[int, list[int]],
              links: list[tuple[int, bool]], no_count: dict[str, int], prev_date: str | None) -> dict:
    """관리 시트 한 건 → 저장할 건 기록(+ 이관 확인 알림). 단계·업무 알림은 status.derive 가 계산."""
    alerts: list[dict] = []
    note = d["note"]
    amount = _quote_amount(d)
    vat_included = bool(re.search(r"부가세\s*포함|VAT\s*포함", note, re.I))
    currency = "USD" if re.search(r"USD|US\$", note + " ".join(i["name"] for i in d["items"])) else "KRW"
    terms = parse_pay_terms(note)

    # 계산서(관리 시트에는 발행 1회만 적혀 있다)
    invoices: list[dict] = []
    if d["inv_date"] or d["total"] or d["inv_date_text"]:
        issuer = next((ln.strip("※ ") for ln in note.split("\n") if "명의" in ln), None)
        invoices.append({"date": d["inv_date"], "supply": d["supply"], "vat": d["vat"], "total": d["total"],
                         "issuer_note": issuer, "partial": False})
    expected = amount * (1 if vat_included else 1 + VAT_RATE) if amount else None
    inv_total = sum(i["total"] or 0 for i in invoices) or None
    inv_supply = sum(i["supply"] or 0 for i in invoices) or None
    # 견적 단가가 부가세 포함이었는데 비고에 안 적힌 경우(발행 합계 = 견적액)도 전액으로 본다
    inv_matches = bool(inv_total and amount and any(
        x and abs(x - y) <= y * 0.03 for x, y in ((inv_total, expected), (inv_total, amount), (inv_supply, amount))))
    inv_ratio = inv_total / expected if inv_total and expected else None
    partial_invoice = bool(inv_ratio and not inv_matches and inv_ratio < PARTIAL_RATIO)
    if invoices:
        invoices[0]["partial"] = partial_invoice

    # 입금
    payments: list[dict] = []
    paid_full = False
    if d["paid_date"]:
        payments.append({"date": d["paid_date"], "amount": inv_total, "note": None})
        paid_full = not partial_invoice
    elif d["paid_text"] == BROKEN_CELL:  # 입금일이 깨졌을 뿐 입금은 적혀 있었다
        payments.append({"date": None, "amount": inv_total, "note": "입금일 깨짐(########)"})
        paid_full = not partial_invoice
    elif d["paid_text"]:
        payments.append({"date": None, "amount": None, "note": d["paid_text"]})
        paid_full = not re.search(r"%|선금|예정", d["paid_text"])
    if "입금완료" in note.replace(" ", ""):
        if not payments:
            payments.append({"date": None, "amount": None, "note": "비고에 '입금완료'"})
        paid_full = paid_full or not partial_invoice

    shipments = [_shipment_view(ships[si], by_amount) for si, by_amount in links]
    shipped_note = not shipments and "출하완료" in note.replace(" ", "")
    # 수주 근거가 있으면 수주(True), 없으면 결과 미확인(None) — 엑셀에 실주 표시는 없어 False 는 만들지 않는다
    won = True if ((d["result"] or "").strip().upper() == "O" or invoices or payments or shipments or shipped_note) else None

    base_date = d["quote_date"] or d["inv_date"] or d["paid_date"]

    # ── 이관 확인(review) ──
    if d["kind"] == "quote" and not d["quote_date"]:
        why = "깨져서(########)" if d["quote_date_text"] == BROKEN_CELL else "비어 있어서"
        alerts.append(_alert("QUOTE_DATE", "review", f"견적일이 {why} 앞 건 날짜({prev_date or '-'}) 기준으로 정렬했습니다"))
    broken = [lbl for lbl, val in (("계산서 발행일", d["inv_date_text"]), ("입금일", d["paid_text"])) if val == BROKEN_CELL]
    if broken:
        alerts.append(_alert("BROKEN_DATE", "review", f"{'·'.join(broken)}이 ########로 깨져 있어 원본에 날짜가 없습니다 — 다시 입력 필요"))
    if d["paid_text"] and d["paid_text"] != BROKEN_CELL:
        alerts.append(_alert("PAID_TEXT", "review", f"입금일 칸이 글로 적혀 있음: \"{d['paid_text']}\""))
    if d["inv_flag"] and d["inv_flag"].upper() not in ("O", "X"):
        alerts.append(_alert("INV_FLAG_TEXT", "review", f"세금계산서 칸이 O·X가 아님: \"{d['inv_flag'][:40]}\""))
    if di in rev:
        others = [deals[i]["legacy_no"] for i in rev[di] if i != di]
        alerts.append(_alert("REVISION", "review", f"같은 고객 견적이 {REVISION_WINDOW_DAYS}일 안에 이어짐(옛 번호 {', '.join(others)}) — 같은 건의 차수인지 확인"))
    totals = _won_amounts(note)
    if totals and amount:
        # 비고의 ₩ 금액 중 하나가 품목 합계(부가세 포함 포함) 또는 '합계 − 비고의 할인 금액'이면 맞는 것으로 본다
        bases = [amount] + [amount - t for t in totals if t < amount]
        ok = any(abs(t - b * k) < 2 for t in totals for b in bases for k in (1, 1 + VAT_RATE))
        if not ok:
            alerts.append(_alert("TOTAL_MISMATCH", "review", f"품목 합계 {_won(amount)}와 비고 총액 {_won(max(totals))}이 다름(할인·절사·옵션 확인)"))
    if inv_ratio and not inv_matches and not any(abs(inv_ratio - t["pct"] / 100) < 0.03 for t in terms):
        why = "결제조건 비율과도 맞지 않음" if partial_invoice else "할인·옵션 확인"
        alerts.append(_alert("INVOICE_DIFF", "review", f"견적 {_won(expected)}(부가세 포함)과 발행 {_won(inv_total)}이 다름 — {why}"))
    if invoices and not d["inv_date"] and d["inv_date_text"] != BROKEN_CELL:
        alerts.append(_alert("INV_NO_DATE", "review", "세금계산서 금액은 있는데 발행일이 비어 있음"))
    if any(i["issuer_note"] for i in invoices):
        alerts.append(_alert("ISSUER", "review", "계산서 발행처가 고객사와 다름 — 발행처를 따로 지정해야 함"))
    if no_count.get(d["legacy_no"], 0) > 1:
        alerts.append(_alert("DUP_NO", "review", f"옛 번호 {d['legacy_no']}이 시트에 두 번 있음 — 새 건번호로 구분"))
    if any(i["qty"] is None and i["qty_text"] for i in d["items"]):
        q = next(i["qty_text"] for i in d["items"] if i["qty"] is None and i["qty_text"])
        alerts.append(_alert("QTY_TEXT", "review", f"수량이 숫자가 아님(\"{q}\") — 기간·세트 단위 확인"))
    if currency == "USD":
        alerts.append(_alert("USD", "review", "외화(USD) 견적 — 통화 확인"))
    if terms and sum(t["pct"] for t in terms) != 100:
        alerts.append(_alert("TERMS", "review", f"결제조건 비율 합이 {sum(t['pct'] for t in terms)}% — 비고 글 확인"))
    for s in shipments:
        if not s["linked_by_amount"]:
            alerts.append(_alert("SHIP_LINK", "review", f"출하({s['date'] or '-'})를 고객사·날짜로만 연결함 — 맞는 건인지 확인"))
            break

    return {
        "kind": d["kind"], "legacy_no": d["legacy_no"], "customer": d["customer"], "contact": d["contact"],
        "owner": d["owner"] or next((s["owner"] for s in shipments if s["owner"]), None),
        "title": _title(d), "base_date": base_date or prev_date,
        "quote": None if d["kind"] != "quote" else {
            "revision": 1, "date": d["quote_date"], "sent_date": d["sent_date"], "amount": amount,
            "vat_included": vat_included, "currency": currency, "items": d["items"]},
        "pay_terms": terms, "won": won, "invoices": invoices, "payments": payments,
        "paid_full": paid_full, "shipments": shipments, "shipped_note": shipped_note,
        "review_alerts": alerts, "note": note,
        "source": {"file": "management", "row_start": d["row_start"], "row_end": d["row_end"]},
    }


def _shipment_view(s: dict, by_amount: bool) -> dict:
    return {
        "date": s["date"], "date_guessed": s["date_guessed"], "date_text": s["date_text"],
        "customer": s["customer"], "items": s["items"], "qty_text": s["qty_text"],
        "unit_price": s["unit_price"], "amount": s["amount"], "carrier": s["carrier"], "tracking": s["tracking"],
        "receiver": s["receiver"], "address": s["address"], "owner": s["owner"], "purpose": s["purpose"],
        "photos": s.get("photos", []), "linked_by_amount": by_amount,
        "source": {"file": "shipments", "row_start": s["row_start"], "row_end": s["row_end"]},
    }


def _shipment_only_row(s: dict) -> dict:
    view = _shipment_view(s, False)
    alerts = []
    if (s["purpose"] or "") != "개발":
        alerts.append(_alert("SHIP_ONLY", "review", "연결할 견적을 찾지 못한 출하 — 견적 건에 붙이거나 '견적 없는 출하'로 확정"))
    if not s["customer"]:
        alerts.append(_alert("SHIP_NO_CUST", "review", "고객사 칸이 비어 있음"))
    elif s.get("customer_guessed"):
        alerts.append(_alert("SHIP_CUST_GUESS", "review", "고객사가 비어 있어 같은 받는 사람이 적힌 출하의 고객사로 채움"))
    if s["date_guessed"]:
        alerts.append(_alert("SHIP_YEAR", "review", "출하일에 연도가 없어 앞뒤 행으로 연도를 추정함"))
    if not s["date"]:
        alerts.append(_alert("SHIP_NO_DATE", "review", "출하일이 비어 있음"))
    first = (s["items"] or "").split("\n")[0]
    return {
        "kind": "shipment_only", "legacy_no": s["legacy_no"], "customer": s["customer"], "contact": s["receiver"],
        "owner": s["owner"], "title": first[:40] or "(품목 미기재)", "base_date": s["date"],
        "quote": None, "pay_terms": [], "won": None, "invoices": [], "payments": [], "paid_full": False,
        "shipments": [view], "shipped_note": False, "review_alerts": alerts, "note": "",
        "source": {"file": "shipments", "row_start": s["row_start"], "row_end": s["row_end"]},
    }
