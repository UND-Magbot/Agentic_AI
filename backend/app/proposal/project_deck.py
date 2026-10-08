"""제안서 프로젝트(문항별 값·상태·승인 자산) → 고객용 PPT(10쪽 이내).

UND 영업 AI 제안서 작성 가이드(2026-09-28) 표 7 "페이지 매핑"·표 9 "검수 사례" 기준.
  - 페이지 구성은 데이터(project["pages"] = 승인된 구성)로 받는다. 표 7 은 대안 3개 사례라 코드에 고정하지 않고,
    build_default_pages() 가 대안 수에 맞춘 기본안을 만든다(구성 승인 화면의 초기값).
  - 문구는 페이지마다 해당 문항 값만 근거로 로컬 모델이 짧게 쓴다(compose). 문항에 없는 수치·등급은 가리고,
    확인된 문항의 수치가 빠지면 한 번 다시 쓰게 한 뒤에도 빠지면 보고한다.
  - 견적·주요 항목 표는 quote_lines 에서 코드가 만든다: 공통 + 대안 1종 구조, 대안끼리 합산하지 않음,
    근거 있는 단가가 없으면 "별도 협의". 두 표가 같은 목록을 쓰므로 장비 수량이 어긋나지 않는다.
  - 공정도·라벨·표는 모두 PowerPoint 도형(편집 가능). 컨셉 이미지는 승인된 것만 크게 쓴다.
배치는 실제 고객 제출본(docs/examples/proposals/UND_삼성웰스토리_자동화실증_제안서_10P.pptx, 2026-09-30 사용자 비교 요청)을 기준으로 한다:
  제목 + 회색 한 줄 요지(KEY MESSAGE 띠·로고·장 번호 없음), 개요는 핵심 수치 줄, 공정은 실제 모듈 이름의 두 줄 흐름,
  컨셉은 그림을 쪽 전체로 + 요점 3개, 주요 항목은 대안 비교표 + 제품 사진 + 공통 모듈 수량, 견적은 묶음 행.
색·글자 맞춤 헬퍼는 pptx_premium 것을 쓴다(그 모듈의 원클릭 디자인판 틀은 건드리지 않는다).
완성 후 프롬프트로 수정(edit_deck): 쪽별 문구(DeckText)를 출력본에 저장해 두고, 요청대로 고친 문구로 다시 조판한다.

project 형식은 proposal_project.service.get_project() 의 반환값:
  items{code: {value, unit, status, ...}}, alternatives[{id, name, item_codes}],
  images[{id, attachment_id, alt_id, version, approved, labels[{text, x, y}]}] (x·y 는 그림 기준 0~1),
  pages[{page_no, type, title, item_codes, asset_ids}], quote_lines[{group, alt_id, item, qty, unit,
  unit_price, currency, basis, included}], output{max_pages}.
"""
from __future__ import annotations

import io
import json
import logging
import re
from dataclasses import dataclass, field
from datetime import datetime
from typing import Any

from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.dml import MSO_LINE_DASH_STYLE
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE, MSO_SHAPE_TYPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

from .. import proposal_llm
from ..company_context import COMPANY_NAME
from ..diagram.llm_spec import _extract_json
from . import pptx_premium as pp
from .body import (NUMBER_MASK, SPEC_MASK, Card, ProposalBody, _numbers_in, mask_unsupported_numbers,
                   mask_unsupported_specs)
from .pptx_render import SLIDE_H, SLIDE_W, RenderReport, _set_font

_log = logging.getLogger("proposal.project_deck")

PAGE_TYPES = ("cover", "overview", "flow", "common_concept", "alternative", "equipment", "poc", "quote")
# 가이드 표 7 의 문항 매핑을 페이지 유형별로 일반화(대안 페이지는 대안 고유 문항이 더해진다).
DEFAULT_CODES: dict[str, tuple[str, ...]] = {
    "cover": ("P01", "P02", "P03", "P13"),
    "overview": ("P01", "P02", "P03", "P04", "P05", "P06", "P07", "P08"),
    "flow": ("P09", "P10", "P11", "P12", "P17", "P18", "P19", "P29", "P30"),
    "common_concept": ("P25", "P26", "P27", "P28", "P29", "P30", "P35", "P36"),
    "alternative": ("P20", "P21", "P22", "P23", "P24", "P25", "P26", "P27", "P29"),
    "equipment": ("K04", "K05", "K06", "P21", "P22", "P23", "P24", "P38", "P41"),
    "poc": ("P31", "P32", "P33", "P34", "P35", "P36", "P37", "P38", "P41", "P42"),
    "quote": ("P23", "P28", "P43", "P44"),
}
DEFAULT_TITLES = {"cover": "표지", "overview": "프로젝트 개요", "flow": "전체 공정 흐름",
                  "common_concept": "공통 적용 컨셉", "equipment": "주요 항목", "poc": "실증과 현장 적용",
                  "quote": "견적 구성"}
_ENG = {"overview": "Overview", "flow": "Process Flow", "common_concept": "Common Concept",
        "alternative": "Concept", "equipment": "Key Items", "poc": "PoC & Site", "quote": "Quotation"}
_DROP_ORDER = ("poc", "common_concept")          # 쪽수 상한을 넘으면 먼저 빼는 페이지(대안·견적은 필수)
_ALT_CODES = ("P14", "P15", "P16")               # 기본 대안 A·B·C 문항(대안에 item_codes 가 없을 때)

USABLE = ("confirmed", "adopted", "assumed")     # 근거로 쓸 수 있는 상태
STATUS_KO = {"confirmed": "확인됨", "adopted": "채택 컨셉", "assumed": "가정", "unknown": "미확인",
             "conflict": "자료 충돌", "empty": "비어 있음"}
PRICE_TBD = "별도 협의"
# 견적 쪽 핵심 메시지는 고정 — 모델이 "투자 회수" 같은 말을 붙였다(실측, 가이드 P44: ROI 강제 금지).
QUOTE_HEADLINE = "기본 구성은 공통 설비에 로봇 대안 1종을 더한 구조이며, 금액은 근거 확인 후 확정합니다"
QUOTE_HEADLINE_SINGLE = "기본 구성은 공통 설비와 공정 컨셉 설비로 이루어지며, 금액은 근거 확인 후 확정합니다"
ALT_NOTE = "대안은 비교 대상입니다. 기본 구성 = 공통 설비 + 선택 대안 1종(대안끼리 합산하지 않음)."
SINGLE_LABEL = "공정 컨셉"       # 컨셉이 하나뿐일 때(기본) — "대안"·"택 1" 같은 비교 표현을 쓰지 않는다
CONCEPT_NOTE = pp.CONCEPT_NOTE
# "버퍼"는 빠짐 — 속도 조정 버퍼는 정상 흐름이다(가이드 P10).
_EXCEPTION_RE = re.compile(r"정지|미취출|놓침|수동\s*전환|예외|재시도")
_TAG_RE = re.compile(r"\[[^\]]{1,12}\]\s*")      # 가이드식 답변 머리표 "[확인]" "[채택 컨셉]"


# ── 기본 페이지 구성 ─────────────────────────────────────────────────────────

def approved_images(project: dict[str, Any], alt_id: str | None) -> list[dict[str, Any]]:
    """해당 대안(None=공통)의 승인 이미지 중 최신 버전 1장."""
    imgs = [i for i in project.get("images") or [] if i.get("approved") and (i.get("alt_id") or None) == alt_id]
    return sorted(imgs, key=lambda i: i.get("version") or 0)[-1:]


def build_default_pages(project: dict[str, Any]) -> list[dict[str, Any]]:
    """대안 수에 맞춘 기본 구성. 쪽수 상한(output.max_pages)을 넘으면 _DROP_ORDER 순으로 뺀다."""
    alts = project.get("alternatives") or []
    pages: list[dict[str, Any]] = []

    def add(ptype: str, title: str, codes: tuple[str, ...], assets: list[int], alt_id: str | None = None) -> None:
        pages.append({"type": ptype, "title": title, "item_codes": list(codes), "asset_ids": assets,
                      **({"alt_id": alt_id} if alt_id else {})})

    for t in ("cover", "overview", "flow"):
        add(t, DEFAULT_TITLES[t], DEFAULT_CODES[t], [])
    add("common_concept", DEFAULT_TITLES["common_concept"], DEFAULT_CODES["common_concept"],
        [i["id"] for i in approved_images(project, None)])
    for k, a in enumerate(alts):
        own = tuple(a.get("item_codes") or (_ALT_CODES[k:k + 1] if k < len(_ALT_CODES) else ()))
        add("alternative", a.get("name") or (f"대안 {a['id']}" if len(alts) > 1 else SINGLE_LABEL),
            own + DEFAULT_CODES["alternative"],
            [i["id"] for i in approved_images(project, a["id"])], a["id"])
    for t in ("equipment", "poc", "quote"):
        add(t, DEFAULT_TITLES[t], DEFAULT_CODES[t], [])
    limit = int((project.get("output") or {}).get("max_pages") or 10)
    for t in _DROP_ORDER:
        if len(pages) <= limit:
            break
        pages = [p for p in pages if p["type"] != t]
    for n, p in enumerate(pages, 1):
        p["page_no"] = n
    return pages


# ── 문구 작성(로컬 모델) ──────────────────────────────────────────────────────

@dataclass
class FlowStep:
    label: str
    note: str = ""
    kind: str = "main"          # main | branch | exception
    src: int | None = None      # branch/exception 이 갈라지는 main 단계 번호(0부터)


@dataclass
class PageText:
    headline: str = ""
    cards: list[Card] = field(default_factory=list)
    steps: list[FlowStep] = field(default_factory=list)
    kpis: list[dict[str, str]] = field(default_factory=list)       # 개요 핵심 수치 [{label, value, unit}]
    keywords: list[str] = field(default_factory=list)              # 컨셉 쪽 요점 3개
    checks: list[list[str]] = field(default_factory=list)          # 실증 쪽 [항목, 검토 방향]
    confirm: str = ""                                              # 실증 쪽 '확인 자료' 한 줄


@dataclass
class DeckText:
    title: str = ""
    subtitle: str = ""
    customer: str = ""
    pages: dict[int, PageText] = field(default_factory=dict)       # page_no → 문구


@dataclass
class DeckReport:
    pages: int = 0
    overflow: list[str] = field(default_factory=list)
    masked: list[str] = field(default_factory=list)                # 근거 없어 가린 수치·등급
    missing_facts: list[str] = field(default_factory=list)         # 확인된 문항 수치가 해당 페이지에 없음
    warnings: list[str] = field(default_factory=list)              # 쪽수 초과·이미지 없음·모델 실패 등
    softened: list[str] = field(default_factory=list)              # 성능 단정("대응 가능" 등)을 "검토"로 바꾼 곳


_SYSTEM = """너는 유엔디로보틱스의 수석 기술영업이다. 고객에게 제출할 로봇 자동화 제안서의 한 쪽 문구만 쓴다.
규칙:
- 아래 [문항] 값만 근거로 쓴다. 문항에 없는 수치·모델명·등급·가격·성능을 만들지 않는다.
- 상태별 표현: 확인됨=사실로, 채택 컨셉=제안 구성으로, 가정="~을 전제로 검토"로, 미확인=확정값처럼 쓰지 말고
  꼭 필요하면 "후속 확인"으로만, 자료 충돌=쓰지 않는다.
- 내부 질문 목록이나 미확인 항목 전체를 나열하지 않는다. 금액은 쓰지 않는다(견적 표는 따로 만든다).
- 컨셉 그림이 주인공이다. 문구는 짧게: headline 60자 이내, cards 2~4개, 카드마다 bullets 1~3개, bullet 45자 이내.
- 문항 원문을 길게 옮기지 말고 고객이 읽을 문장으로 요약한다.
- 성능·등급·환경 대응은 근거 자료로 검증되기 전까지 "대응 가능·보장·확보·달성" 같은 단정 대신 "검토"·"검증 예정"으로
  쓴다. 현장 조건(예: 클린룸 등급)이 사실이어도 그 조건을 충족한다는 뜻은 아니다. 보호 커버가 있어도 등급 확보가 아니다.
- "극대화·최고·완벽" 같은 과장 표현을 쓰지 않는다.
출력은 JSON 하나: {"headline": "...", "cards": [{"heading": "...", "bullets": ["..."]}]%s}
%s"""
# 쪽 유형별로 더 받는 칸(제출본 배치에 맞춤). (JSON 키, 규칙)
_EXTRA: dict[str, tuple[str, str]] = {
    "overview": (', "kpis": [{"label": "8자 이내", "value": "수치 포함 12자 이내", "unit": "6자 이내"}]',
                 "kpis 규칙: [문항] 중 확인됨 값에 적힌 핵심 수치(규모·수량·처리량 등)만 2~4개. '1종'·'1대' 같은 개수 1은 "
                 "넣지 않는다. 수치를 새로 계산하거나 "
                 "만들지 않는다. 없으면 빈 배열. cards 는 2개(대상과 범위 / 운영 조건처럼 나눠서)."),
    "flow": (', "steps": [{"label": "10자 이내", "note": "20자 이내", "kind": "main|branch|exception", "from": null}]',
             """steps 규칙:
- main = 대상물이 자동화 시작점부터 최종 인계(적재·반출)까지 거치는 정상 경로만, 순서대로 4~8개.
  label 은 "세척 전 비전"·"승강 적재"처럼 실제 장치·모듈 이름으로 쓴다("단계 1" 같은 말 금지).
  인식(비전)·취출(로봇)·적재처럼 대상물이 실제로 지나는 단계를 빠뜨리지 않는다.
- branch = 로봇 대상에서 제외해 다른 곳으로 보내는 흐름(예: 소형 분리).
- exception = 놓침 감지·정지·버퍼 수용·수동 전환 같은 예외 처리. main 에 넣지 않는다.
- branch·exception 의 "from" 에는 갈라지는 main 단계 번호(0부터)를 적는다.
- cards 는 1~2개(운영 연계처럼 흐름 밖에서 알아야 할 것)."""),
    "common_concept": (', "keywords": ["14자 이내"]',
                       "keywords 는 모든 컨셉에 공통으로 들어가는 모듈 3개. cards 는 그 모듈 3개를 같은 순서로 "
                       "(heading=모듈 이름, bullets 1~2개)."),
    "alternative": (', "keywords": ["14자 이내"]',
                    "keywords 는 이 컨셉을 한눈에 보여 주는 요점 3개(로봇 구성·배치·특징, 예: '8종 승강 적재')."),
    "poc": (', "checks": [{"item": "18자 이내", "direction": "40자 이내"}], "confirm": "60자 이내"',
            "cards 는 실증 단계 3개를 순서대로(heading=단계 이름, bullets=그 단계에서 확인할 것). "
            "checks 는 현장·조건별 검토 방향 0~4개. confirm 은 고객에게 받을 확인 자료 한 줄."),
}


def _system(ptype: str) -> str:
    keys, rules = _EXTRA.get(ptype, ("", ""))
    return _SYSTEM % (keys, rules)


_COVER = """고객 제안서 표지 문구를 [문항] 값만으로 JSON 하나로 써라:
{"title": "30자 이내 제안서 제목", "subtitle": "50자 이내 부제(목적·대안)", "customer": "고객사명(없으면 빈 문자열)"}
제목·부제의 목적은 P02(고객이 이번에 결정할 것)의 단계를 그대로 따른다. 실증·검토 단계면 "실증"·"검토" 제안으로 쓰고,
양산·투자를 확정하는 제안처럼 쓰지 않는다."""

# 성능 단정 — 근거 자료로 검증되기 전에는 "검토"로 낮춘다(실측: "Class 100 클린룸 환경 대응 가능",
# 문항은 환경 사실뿐이었음. 가이드 K06·P35: 확인된 항목만 확정값, 보호 커버 ≠ 등급 확보).
_CAN_RE = re.compile(r"(대응|적용|구현|처리|달성|충족|만족|운영|사용|구축)\s*(?:이\s*)?가능(?!성)(합니다|함|하며|하고|한)?")
_GUARANTEE_RE = re.compile(r"(?:보장|보증)(?:합니다|함|하는|된)?")
_ACHIEVE_RE = re.compile(r"(확보|달성|충족|만족)(합니다|함|하여|하는)?(?!\s*(?:검토|검증|여부|방안|목표))")
_PERF_RE = re.compile(r"[Cc]lass\s*\d|IP\s*\d|등급|인증|처리량|정밀도|사이클|C/?T|속도|성능|%|시간당|개/")
# 과장 표현(실측: "생산성 극대화") — 근거 없는 최상급을 평이한 말로.
_HYPE = (("극대화", "향상"), ("완벽한 ", ""), ("완벽하게 ", ""), ("최고의 ", ""))
_HEDGED_RE = re.compile(r"검토|검증|확인|목표|예정|여부")
_STAGE_BAD_RE = re.compile(r"양산[^,.·]{0,8}(?:투자|도입|적용|구축)")


def _item_lines(project: dict[str, Any], codes: list[str]) -> list[str]:
    items = project.get("items") or {}
    out = []
    for c in codes:
        it = items.get(c)
        if not it or it.get("status") in ("empty", "conflict") or not str(it.get("value") or "").strip():
            continue
        unit = f" ({it['unit']})" if it.get("unit") else ""
        out.append(f"{c} [{STATUS_KO.get(it.get('status'), it.get('status'))}] {str(it['value']).strip()}{unit}")
    return out


def _allowed(project: dict[str, Any]) -> tuple[set[str], str]:
    """근거로 인정하는 수치 키와 원문(등급·규격 대조용): 근거 상태 문항 값 + 고객 요청 원문."""
    parts = [str(project.get("request_text") or "")]
    for it in (project.get("items") or {}).values():
        if it.get("status") in USABLE:
            parts.append(f"{it.get('value') or ''} {it.get('unit') or ''}")
    src = "\n".join(parts)
    return _numbers_in(src), src


def _facts(project: dict[str, Any], codes: list[str]) -> dict[str, set[str]]:
    """확인됨 문항별 수치 키 — 해당 페이지 문구에 나와야 할 사실."""
    items = project.get("items") or {}
    out = {}
    for c in codes:
        it = items.get(c)
        if it and it.get("status") == "confirmed":
            # '1종'·'1대' 같은 개수 1은 핵심 수치에서 빼므로(_TRIVIAL_KPI_RE) 빠진 확인값으로 치지 않는다(QA 3회차 충돌).
            keys = {k for k in _numbers_in(_TAG_RE.sub("", f"{it.get('value') or ''} {it.get('unit') or ''}"))
                    if not k.endswith("|") and k.split("|")[0] not in ("0", "1")}
            if keys:
                out[c] = keys
    return out


def _clean(text: Any, limit: int) -> str:
    return re.sub(r"\s+", " ", str(text or "")).strip()[:limit]


def _parse_page(data: dict[str, Any], flow: bool) -> PageText:
    cards = []
    for c in (data.get("cards") or [])[:4]:
        if not isinstance(c, dict):
            continue
        bullets = [_clean(b, 80) for b in (c.get("bullets") or [])[:3] if _clean(b, 80)]
        if c.get("heading") or bullets:
            cards.append(Card(_clean(c.get("heading"), 30) or "-", bullets))
    steps = []
    if flow:
        mains = 0
        for s in (data.get("steps") or [])[:12]:
            if not isinstance(s, dict) or not _clean(s.get("label"), 20):
                continue
            kind = s.get("kind") if s.get("kind") in ("main", "branch", "exception") else "main"
            src = s.get("from")
            src = src if isinstance(src, int) and 0 <= src < max(mains, 1) else None
            if kind == "main":
                mains += 1
            elif src is None:
                src = max(mains - 1, 0)
            label = _clean(s["label"], 20)
            # 예외 처리는 정상 흐름이 아니다 — 모델이 main 에 넣어도 예외로 내린다(실측: 정지·버퍼가 main 에).
            if kind == "main" and mains > 1 and _EXCEPTION_RE.search(label):
                kind, src, mains = "exception", mains - 2, mains - 1
            steps.append(FlowStep(label, _clean(s.get("note"), 40), kind, src))
    kpis = []
    for k in (data.get("kpis") or [])[:4]:
        if isinstance(k, dict) and re.search(r"\d", _clean(k.get("value"), 16)):
            kpis.append({"label": _clean(k.get("label"), 14), "value": _clean(k.get("value"), 16),
                         "unit": _clean(k.get("unit"), 10)})
    keywords = [_clean(x, 20) for x in (data.get("keywords") or [])[:3] if _clean(x, 20)]
    checks = []
    for c in (data.get("checks") or [])[:4]:
        pair = [c.get("item"), c.get("direction")] if isinstance(c, dict) else (
            list(c[:2]) if isinstance(c, (list, tuple)) and len(c) >= 2 else [None, None])
        if _clean(pair[0], 24):
            checks.append([_clean(pair[0], 24), _clean(pair[1], 60)])
    return PageText(_clean(data.get("headline"), 90), cards, steps, kpis, keywords, checks,
                    _clean(data.get("confirm"), 90))


def _fallback(project: dict[str, Any], page: dict[str, Any]) -> PageText:
    """모델 실패 시: 근거 상태 문항 값을 그대로 짧게(지어내지 않는다)."""
    items = project.get("items") or {}
    cards = []
    for c in page.get("item_codes") or []:
        it = items.get(c)
        if it and it.get("status") in USABLE and it.get("value"):
            cards.append(Card(c, [_clean(_TAG_RE.sub("", str(it["value"])), 80)]))
    return PageText(page.get("title") or "", cards[:4])


def _page_text_all(t: PageText) -> str:
    return "\n".join([t.headline] + [c.heading + " " + " ".join(c.bullets) for c in t.cards]
                     + [f"{s.label} {s.note}" for s in t.steps]
                     + [f"{k['label']} {k['value']}{k['unit']}" for k in t.kpis] + t.keywords
                     + [f"{a} {b}" for a, b in t.checks] + [t.confirm])


def soften(text: str) -> tuple[str, list[str]]:
    """성능 단정을 검토 표현으로. 반환: (고친 문장, 바꾼 표현들)."""
    hits: list[str] = []

    def can(m: re.Match[str]) -> str:
        hits.append(m.group(0))
        end = m.group(2) or ""
        if end == "합니다":            # "대응 가능합니다" → "대응을 검토합니다"
            return f"{m.group(1)}{'를' if m.group(1) == '처리' else '을'} 검토합니다"
        return f"{m.group(1)} 검토" + (end if end in ("함", "하며", "하고") else "")

    def guarantee(m: re.Match[str]) -> str:
        hits.append(m.group(0))
        return "검토"

    def achieve(m: re.Match[str]) -> str:
        hits.append(m.group(0))
        return f"{m.group(1)} 검토" + ("합니다" if m.group(2) == "합니다" else "")

    text = _CAN_RE.sub(can, text)
    text = _GUARANTEE_RE.sub(guarantee, text)
    for word, plain in _HYPE:
        if word in text:
            hits.append(word.strip())
            text = text.replace(word, plain)
    # "확보·달성"은 공간 확보처럼 중립적으로도 쓰여, 성능 표현과 함께 단정할 때만 낮춘다.
    if _PERF_RE.search(text) and not _HEDGED_RE.search(text):
        text = _ACHIEVE_RE.sub(achieve, text)
    return text, hits


_TRIVIAL_KPI_RE = re.compile(r"(?:약\s*)?[01](?:\.0)?\s*(?:종|대|개|식|set|EA|ea)?")


def _mask_page(t: PageText, nums: set[str], src: str, report: DeckReport) -> None:
    def fix(s: str) -> str:
        s, m1 = mask_unsupported_numbers(s, nums, set())
        s, m2 = mask_unsupported_specs(s, src)
        s, m3 = soften(s)
        report.masked.extend(m1 + m2)
        report.softened.extend(m3)
        return s

    t.headline = fix(t.headline)
    for c in t.cards:
        c.heading = fix(c.heading)
        c.bullets = [fix(b) for b in c.bullets]
    for s in t.steps:
        s.label, s.note = fix(s.label), fix(s.note)
    kept = []
    for k in t.kpis:
        # 핵심 수치 칸은 수치가 주인공 — 근거가 없어 가려졌으면 칸째 뺀다(가림 표시만 크게 남지 않게).
        value, unit = fix(k["value"]), fix(k["unit"])
        joined = fix(f"{k['value']} {k['unit']}") if k["unit"] else value    # 단위까지 붙여 근거 확인
        if any(m in x for x in (value, unit, joined) for m in (NUMBER_MASK, SPEC_MASK)):
            continue
        if _TRIVIAL_KPI_RE.fullmatch(value.strip()):      # '1종'·'1대' 같은 개수 1은 핵심 수치가 아니다(QA 2회차)
            continue
        kept.append({"label": fix(k["label"]), "value": value, "unit": unit})
    t.kpis = kept
    t.keywords = [fix(x) for x in t.keywords]
    t.checks = [[fix(a), fix(b)] for a, b in t.checks]
    t.confirm = fix(t.confirm)


async def _ask(system: str, user: str) -> dict[str, Any]:
    raw = await proposal_llm.chat([{"role": "system", "content": system}, {"role": "user", "content": user}],
                                  num_predict=2000)
    data = _extract_json(raw)
    if not isinstance(data, dict):
        raise ValueError("JSON 객체가 아닙니다.")
    return data


def _stage_problem(project: dict[str, Any], title: str, subtitle: str) -> str:
    """표지 목적이 P02 단계와 어긋나면 사유(실측: P02 '실증 후 양산 투자 결정' → 부제 '양산 투자 제안')."""
    p02 = (project.get("items") or {}).get("P02") or {}
    if p02.get("status") not in USABLE or "실증" not in str(p02.get("value") or ""):
        return ""
    cover = f"{title} {subtitle}"
    if "실증" in cover:
        return ""
    if _STAGE_BAD_RE.search(cover):
        return "P02 는 실증 단계인데 양산·투자 제안처럼 썼습니다"
    return "P02 는 실증 단계인데 제목·부제에 실증 목적이 없습니다"


async def _compose_cover(project: dict[str, Any], cover: dict[str, Any], deck: DeckText, report: DeckReport) -> None:
    user = "\n".join(_item_lines(project, cover.get("item_codes") or [])) or "(문항 없음)"
    for attempt in range(2):
        try:
            d = await _ask(_COVER, user)
        except Exception as e:  # noqa: BLE001 — 표지는 프로젝트 제목으로 대신한다
            report.warnings.append(f"표지 문구 작성 실패 → 프로젝트 제목 사용({e.__class__.__name__})")
            break
        deck.title, deck.subtitle, deck.customer = (_clean(d.get("title"), 40), _clean(d.get("subtitle"), 70),
                                                     _clean(d.get("customer"), 30))
        problem = _stage_problem(project, deck.title, deck.subtitle)
        if not problem:
            break
        if attempt:
            # 다시 써도 어긋나면 부제를 문항 근거(실증 목적 + 대안 이름)로 바꾼다.
            alts = [a.get("name") for a in project.get("alternatives") or [] if a.get("name")]
            deck.subtitle = "자동화 실증 제안" + (f" · {' · '.join(alts)}" if alts else "")
            report.warnings.append(f"표지: {problem} → 부제를 실증 제안으로 고침")
            break
        user += f"\n[보완] {problem}. P02 단계에 맞게 다시 쓰세요."
    deck.title = deck.title or _clean(project.get("title"), 40) or "로봇 자동화 제안"


async def compose(project: dict[str, Any], pages: list[dict[str, Any]], report: DeckReport) -> DeckText:
    """페이지별 문구. 모델이 실패한 페이지는 문항 값 그대로(보고에 남김)."""
    nums, src = _allowed(project)
    deck = DeckText()
    cover = next((p for p in pages if p["type"] == "cover"), None)
    if cover:
        await _compose_cover(project, cover, deck, report)
        cleaned = []
        for s_ in (deck.title, deck.subtitle):
            s2, m1 = mask_unsupported_numbers(s_, nums, set())
            s2, m3 = soften(s2)
            report.masked.extend(m1)
            report.softened.extend(m3)
            cleaned.append(s2)
        deck.title, deck.subtitle = cleaned
    for p in pages:
        if p["type"] == "cover":
            continue
        flow = p["type"] == "flow"
        lines = _item_lines(project, p.get("item_codes") or [])
        user = (f"[페이지] {p.get('page_no')}쪽 {p.get('title')} ({p['type']})\n[문항]\n"
                + ("\n".join(lines) or "(근거 문항 없음 — 일반적인 설명만 짧게)"))
        facts = _facts(project, p.get("item_codes") or [])
        text: PageText | None = None
        for attempt in range(2):
            try:
                text = _parse_page(await _ask(_system(p["type"]), user), flow)
            except Exception as e:  # noqa: BLE001
                _log.warning("[project_deck] %s쪽 문구 실패: %s", p.get("page_no"), e)
                break
            body = _page_text_all(text)
            have = _numbers_in(body)
            missing = [c for c, keys in facts.items() if not keys & have]
            if not missing or attempt:
                report.missing_facts += [f"{p.get('page_no')}쪽 {c}" for c in missing]
                break
            user += ("\n[보완] 다음 확인된 문항의 수치가 빠졌습니다. 그대로 넣어 다시 쓰세요: "
                     + ", ".join(missing))
        if text is None or not (text.headline or text.cards):
            report.warnings.append(f"{p.get('page_no')}쪽 문구 작성 실패 → 문항 값을 그대로 사용")
            text = _fallback(project, p)
        _mask_page(text, nums, src, report)
        deck.pages[p["page_no"]] = text
    return deck


# ── 렌더링(제출본 배치) ───────────────────────────────────────────────────────

M = 0.6                                  # 좌우 여백 — 제출본은 여백이 좁고 그림·표가 넓다
CW = SLIDE_W - 2 * M
Y0, Y1 = 1.6, 6.8                        # 본문 위·아래(제목·요지 아래, 바닥글 위)
SKY = RGBColor(0x9F, 0xC3, 0xF5)         # 표지 고객사
SOFT = RGBColor(0xC9, 0xD3, 0xE3)
GREEN_TINT = RGBColor(0xE8, 0xF5, 0xEE)
ORANGE_TINT = RGBColor(0xFF, 0xF3, 0xE6)
ORANGE_DARK = RGBColor(0xC2, 0x5A, 0x0C)
_ROBOT_RE = re.compile(r"로봇|AMR|휴머노이드|델타|양팔|암\b|팔\b|협동|스카라|머니퓰레이터|ACR", re.I)
_TOOL_RE = re.compile(r"그리퍼|EOAT|툴|석션|흡착|핑거|mDPG|MG\b", re.I)


def _arrow(line) -> None:
    ln = line.line._get_or_add_ln()
    etree.SubElement(ln, qn("a:tailEnd"), type="triangle", w="med", len="med")


def _connector(slide, x1, y1, x2, y2, color, *, dashed=False, width=1.5, arrow=True):
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    c.line.color.rgb = color
    c.line.width = Pt(width)
    if dashed:
        c.line.dash_style = MSO_LINE_DASH_STYLE.DASH
    if arrow:
        _arrow(c)
    return c


def _box_text(shp, text: str, size: float, color, *, bold=False, sub: str = "", sub_color=None,
              wrap: bool = True) -> None:
    """도형 안 가운데 글(선택: 작은 둘째 줄)."""
    tf = shp.text_frame
    tf.word_wrap = wrap
    tf.margin_left = tf.margin_right = Inches(0.06)
    tf.margin_top = tf.margin_bottom = Inches(0.03)
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = text
    _set_font(r, size, color, bold=bold)
    if sub:
        p2 = tf.add_paragraph()
        p2.alignment = PP_ALIGN.CENTER
        r2 = p2.add_run()
        r2.text = sub
        _set_font(r2, max(size - 3, 8), sub_color or color)


class _Pages(pp._Deck):
    """pptx_premium 의 덱 상태(쪽 번호·넘침 보고)에 바닥글 문구만 제출본 식으로."""

    def __init__(self, prs, body: ProposalBody, short: str):
        super().__init__(prs, body)
        self.short = short


def _page(deck: _Pages, title: str, lead: str = "", tag: str = ""):
    """본문 쪽 틀: 큰 제목 + 회색 요지 한 줄 + 아주 옅은 바닥글(제출본 식 — 띠·로고·장 번호 없음)."""
    s = deck.slide(pp.WHITE)
    if tag:
        pp._label(s, M, 0.22, 3, 0.26, tag, 9.5, pp.ORANGE, bold=True, spacing=100)
    pp._fitted_text(s, deck.report, f"{title} 제목", M, 0.38, CW, 0.62, [title], max_pt=26, min_pt=18,
                    color=pp.NAVY, bold=True, anchor=MSO_ANCHOR.MIDDLE)
    if lead:
        pp._fitted_text(s, deck.report, f"{title} 요지", M, 1.0, CW, 0.42, [lead], max_pt=13, min_pt=10.5,
                        color=pp.GREY)
    pp._label(s, M, 7.03, 3, 0.25, "UND ROBOTICS", 7.5, pp.MUTED, bold=True, spacing=50)
    pp._label(s, SLIDE_W - M - 6.0, 7.03, 6.0, 0.25, f"{deck.short}      {deck.page:02d}", 7.5, pp.MUTED,
              align=PP_ALIGN.RIGHT)
    return s


def _cover(deck: _Pages, now: datetime, alts: list[str]) -> None:
    s = deck.slide(pp.NAVY)
    body = deck.body
    pp._logo(s, M + 0.08, 0.55, 0.5)
    if body.customer:
        pp._label(s, M, 1.75, 10, 0.5, body.customer, 20, SKY, bold=True)
    pp._fitted_text(s, deck.report, "표지 제목", M, 2.3, 9.2, 1.95, [body.title or "제안서"], max_pt=40, min_pt=26,
                    color=pp.WHITE, bold=True, anchor=MSO_ANCHOR.TOP)
    pp._shape(s, MSO_SHAPE.RECTANGLE, M, 4.38, 0.9, 0.05, pp.ORANGE)
    if body.subtitle:
        pp._fitted_text(s, deck.report, "표지 부제", M, 4.58, 10.5, 0.55, [body.subtitle], max_pt=16, min_pt=12,
                        color=SOFT)
    if len(alts) > 1:
        pp._label(s, M, 5.2, 11, 0.4, " / ".join(alts), 12, pp.MUTED)
    pp._label(s, M, 6.75, 4, 0.3, f"{now:%Y.%m.%d}", 10, pp.MUTED)
    pp._label(s, SLIDE_W - M - 5, 6.72, 5, 0.35, COMPANY_NAME, 12, pp.WHITE, bold=True, align=PP_ALIGN.RIGHT)


def _columns(deck: _Pages, s, where: str, cards: list[Card], y: float, h: float, *, max_pt: float = 12.5) -> None:
    """제출본 개요식 칸: 굵은 소제목 + 설명 줄(글머리표 없이). 칸마다 같은 글자 크기."""
    cards = cards[:3]
    if not cards or h < 0.8:        # 소제목 + 한 줄도 안 들어가는 자리엔 그리지 않는다(음수 높이 글상자 → 파일 손상)
        return
    gap = 0.55
    w = (CW - gap * (len(cards) - 1)) / len(cards)
    size, _ = pp._column_font([c.bullets for c in cards], w - 0.1, h - 0.5, max_pt, 9.5, bullet=False)
    for i, c in enumerate(cards):
        x = M + i * (w + gap)
        if c.heading not in ("", "-"):
            pp._fitted_text(s, deck.report, f"{where} 소제목", x, y, w, 0.4, [c.heading], max_pt=15, min_pt=11,
                            color=pp.NAVY, bold=True)
        pp._fitted_text(s, deck.report, f"{where} 설명", x, y + 0.48, w, h - 0.5, c.bullets, max_pt=max_pt,
                        min_pt=9.5, color=pp.INK, size=size)


def _overview(deck: _Pages, page: dict[str, Any], t: PageText) -> None:
    s = _page(deck, page.get("title") or "", t.headline)
    y = Y0 + 0.05
    if t.kpis:
        n = len(t.kpis)
        w = CW / n
        for i, k in enumerate(t.kpis):
            x = M + i * w
            pp._label(s, x, y, w - 0.2, 0.3, k["label"], 10.5, pp.GREY)
            pp._fitted_text(s, deck.report, f"{page['page_no']}쪽 핵심 수치", x, y + 0.3, w - 0.2, 0.72,
                            [k["value"]], max_pt=30, min_pt=18, color=pp.NAVY, bold=True, anchor=MSO_ANCHOR.MIDDLE)
            if k["unit"]:
                pp._label(s, x, y + 1.04, w - 0.2, 0.3, k["unit"], 10, pp.MUTED)
        pp._shape(s, MSO_SHAPE.RECTANGLE, M, y + 1.5, CW, 0.012, pp.HAIR)
        y += 1.8
    _columns(deck, s, f"{page['page_no']}쪽", t.cards or [Card("-", [])], y, Y1 - y, max_pt=14)


def _step_box(slide, report: RenderReport, where: str, x, y, w, h, st: FlowStep, no: int | None) -> None:
    """공정 단계 상자(제출본 식): 옅은 바탕 + '01 모듈 이름' + 회색 설명. 분기는 초록, 예외는 주황."""
    fill, line, lab = pp.TINT, None, pp.NAVY
    if st.kind == "branch":
        fill, lab = GREEN_TINT, pp.OK_GREEN
    elif st.kind == "exception":
        fill, line, lab = ORANGE_TINT, pp.ORANGE, ORANGE_DARK
    box = pp._shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h, fill, line=line, line_w=1, radius=0.08)
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(0.12)
    tf.margin_top = tf.margin_bottom = Inches(0.04)
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    head = (f"{no:02d}  " if no else "") + st.label
    size = 13.5 if h > 1 else 12.5
    while size > 9.5 and pp._lines_needed([head], w - 0.3, size) > 1:
        size -= 0.5
    note = 10.5 if h > 1 else 9.5
    need = (pp._lines_needed([head], w - 0.3, size) * size
            + (pp._lines_needed([st.note], w - 0.3, note) * note if st.note else 0)) * 1.25 / 72 + 0.1
    if need > h:
        report.overflow.append(where)
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.LEFT
    if no:
        r = p.add_run()
        r.text = f"{no:02d}  "
        _set_font(r, size - 1, pp.ORANGE, bold=True)
    r = p.add_run()
    r.text = st.label
    _set_font(r, size, lab, bold=True)
    if st.note:
        p2 = tf.add_paragraph()
        p2.alignment = PP_ALIGN.LEFT
        r2 = p2.add_run()
        r2.text = st.note
        _set_font(r2, note, pp.GREY)


def _flow(deck: _Pages, page: dict[str, Any], text: PageText) -> None:
    """편집 가능한 공정도. 정상 흐름 5단계 이하는 한 줄, 6~8단계는 두 줄(ㄹ자: 윗줄 →, 아랫줄 ←).
    분기·예외는 갈라지는 단계 바로 아래에(윗줄이면 두 줄 사이, 아랫줄이면 그 아래)."""
    s = _page(deck, page.get("title") or "", text.headline)
    where = f"{page['page_no']}쪽 공정도"
    mains = [st for st in text.steps if st.kind == "main"][:8]
    if not mains:
        deck.report.overflow.append(f"{page['page_no']}쪽 공정 단계 없음")
        mains = [FlowStep(c.heading) for c in text.cards] or [FlowStep("-")]
    side = [st for st in text.steps if st.kind != "main"]
    n = len(mains)
    rows = [list(range(n))] if n <= 5 else [list(range((n + 1) // 2)), list(range((n + 1) // 2, n))]
    per = len(rows[0])
    gap, h = (0.55, 1.1) if len(rows) == 1 else (0.55, 0.86)
    w = (CW - gap * (per - 1)) / per
    row_y = [Y0 + 0.45] if len(rows) == 1 else [Y0 + 0.1, Y0 + 2.05]
    pos: dict[int, tuple[float, float, int]] = {}
    for r, idxs in enumerate(rows):
        for j, i in enumerate(idxs):
            col = j if r == 0 else per - 1 - j
            pos[i] = (M + col * (w + gap), row_y[r], r)
            _step_box(s, deck.report, where, pos[i][0], pos[i][1], w, h, mains[i], i + 1)
    for i in range(n - 1):
        (x1, y1, r1), (x2, y2, r2) = pos[i], pos[i + 1]
        if r1 == r2:
            if x2 > x1:
                _connector(s, x1 + w + 0.03, y1 + h / 2, x2 - 0.03, y2 + h / 2, pp.NAVY_MID)
            else:
                _connector(s, x1 - 0.03, y1 + h / 2, x2 + w + 0.03, y2 + h / 2, pp.NAVY_MID)
        else:
            _connector(s, x1 + w / 2, y1 + h + 0.03, x2 + w / 2, y2 - 0.03, pp.NAVY_MID)
    by_src: dict[int, list[FlowStep]] = {}
    for st in side:
        by_src.setdefault(min(st.src or 0, n - 1), []).append(st)
    lowest = max(y for _, y, _ in pos.values()) + h
    bh = 0.8 if len(rows) == 1 else 0.7
    for src, group in by_src.items():
        sx, sy, r = pos[src]
        # 윗줄 단계에서 갈라지면 두 줄 사이, 아니면 그 단계 아래. 두 줄 사이에 끼면 아랫줄 상자와 겹치지 않게
        # 가로로 비켜 둔다(아랫줄은 같은 칸에 상자가 있으므로 칸 사이 틈 쪽으로).
        by = sy + h + (0.5 if len(rows) == 1 else 0.32)
        cw = max(1.6, min(w, 2.3))
        total = len(group) * cw + (len(group) - 1) * 0.15
        cx = sx + w / 2
        x0 = min(max(cx - total / 2, M), M + CW - total)
        for j, st in enumerate(group):
            bx = x0 + j * (cw + 0.15)
            _step_box(s, deck.report, where, bx, by, cw, bh, st, None)
            exc = st.kind == "exception"
            _connector(s, cx, sy + h + 0.02, bx + cw / 2, by - 0.02, pp.ORANGE if exc else pp.OK_GREEN, dashed=True,
                       width=1.25)
        lowest = max(lowest, by + bh)
    notes = [c for c in text.cards if c.bullets][:2]
    ny = lowest + 0.4
    if notes and Y1 - 0.4 - ny >= 0.9:
        pp._shape(s, MSO_SHAPE.RECTANGLE, M, ny - 0.2, CW, 0.012, pp.HAIR)
        room = Y1 - 0.4 - ny
        keep = 3 if room >= 1.5 else (2 if room >= 1.1 else 1)     # 두 줄 공정도면 자리가 좁다 — 줄 수를 맞춘다
        _columns(deck, s, f"{page['page_no']}쪽 공정 요점",
                 [Card(c.heading, c.bullets[:keep]) for c in notes], ny, room, max_pt=12.5)
    for k, (label, color, dashed) in enumerate((("정상 흐름", pp.NAVY_MID, False), ("분기", pp.OK_GREEN, True),
                                                ("예외·정지", pp.ORANGE, True))):
        lx = SLIDE_W - M - 4.8 + k * 1.6
        _connector(s, lx, 6.92, lx + 0.4, 6.92, color, dashed=dashed, width=1.5)
        pp._label(s, lx + 0.46, 6.79, 1.1, 0.26, label, 8.5, pp.GREY)


def _labels(slide, report: RenderReport, labels: list[dict[str, Any]], ix, iy, iw, ih) -> None:
    """그림 위 라벨(편집 가능한 도형): 점 + 말풍선 글상자. 좌표는 그림 기준 0~1."""
    for lab in labels[:10]:
        try:
            fx, fy = float(lab.get("x")), float(lab.get("y"))
        except (TypeError, ValueError):
            continue
        text = _clean(lab.get("text"), 24)
        if not text or not (0 <= fx <= 1 and 0 <= fy <= 1):
            continue
        px, py = ix + fx * iw, iy + fy * ih
        pp._shape(slide, MSO_SHAPE.OVAL, px - 0.07, py - 0.07, 0.14, 0.14, pp.ORANGE, line=pp.WHITE, line_w=1.5)
        tw = min(3.0, pp._text_width_pt(text, 10) * 1.1 / 72 + 0.4)     # 굵은 글씨 폭 여유(실측: 두 줄로 꺾임)
        tx = px + 0.15 if px + 0.15 + tw <= ix + iw else px - 0.15 - tw
        ty = min(max(py - 0.42, iy + 0.05), iy + ih - 0.36)
        pill = pp._shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, tx, ty, tw, 0.32, pp.NAVY, radius=0.5, alpha=0.88)
        _box_text(pill, text, 10, pp.WHITE, bold=True, wrap=False)


def _keyword_row(deck: _Pages, s, where: str, words: list[str], y: float) -> None:
    """그림 아래 요점 줄(제출본: 'RM AIDAL 양팔 구성   8종 승강 적재   외측 카세트 교체')."""
    words = [w for w in words if w][:3]
    if not words:
        return
    x = M
    for w_ in words:
        width = min(CW / len(words), pp._text_width_pt(w_, 12) * 1.08 / 72 + 0.45)
        pp._shape(s, MSO_SHAPE.RECTANGLE, x, y + 0.1, 0.09, 0.09, pp.ORANGE)
        pp._fitted_text(s, deck.report, where, x + 0.16, y, width - 0.16, 0.32, [w_], max_pt=12, min_pt=9.5,
                        color=pp.NAVY, bold=True, anchor=MSO_ANCHOR.MIDDLE)
        x += width + 0.35


def _concept_page(deck: _Pages, page: dict[str, Any], text: PageText, image: bytes | None,
                  labels: list[dict[str, Any]], report: DeckReport, tag: str) -> None:
    """컨셉 그림을 쪽 전체 너비로 + 아래 요점 3개(제출본 식). 요지는 제목 아래 한 줄."""
    from PIL import Image

    s = _page(deck, page.get("title") or "", text.headline, tag)
    top, bottom = Y0 - 0.1, Y1 - 0.5
    box_w, box_h = CW, bottom - top
    if image:
        w_px, h_px = Image.open(io.BytesIO(image)).size
        scale = min(box_w / w_px, box_h / h_px)
        iw, ih = w_px * scale, h_px * scale
        ix, iy = M + (box_w - iw) / 2, top + (box_h - ih) / 2
        s.shapes.add_picture(io.BytesIO(image), Inches(ix), Inches(iy), Inches(iw), Inches(ih))
        _labels(s, deck.report, labels, ix, iy, iw, ih)
    else:
        ph = pp._shape(s, MSO_SHAPE.RECTANGLE, M, top, box_w, box_h, pp.TINT, line=pp.MUTED)
        ph.line.dash_style = MSO_LINE_DASH_STYLE.DASH
        _box_text(ph, "승인된 컨셉 이미지가 없습니다", 14, pp.MUTED, bold=True)
        report.warnings.append(f"{page['page_no']}쪽 승인된 컨셉 이미지 없음")
    words = text.keywords or [c.heading for c in text.cards if c.heading not in ("", "-")]
    _keyword_row(deck, s, f"{page['page_no']}쪽 요점", words, Y1 - 0.38)
    pp._label(s, M, Y1 + 0.02, CW, 0.22, CONCEPT_NOTE, 8, pp.MUTED, italic=True, align=PP_ALIGN.RIGHT)


def _common_modules(deck: _Pages, page: dict[str, Any], text: PageText) -> None:
    """공통 컨셉 이미지가 없을 때: 모듈 3칸(번호·이름·설명) — 빈 자리 표시 대신(제출본의 공통 모듈 3칸 구성)."""
    s = _page(deck, page.get("title") or "", text.headline)
    cards = [c for c in text.cards if c.heading not in ("", "-") or c.bullets][:3]
    if not cards:
        cards = [Card(k, []) for k in text.keywords] or [Card("-", [])]
    extra = [c for c in text.cards if c not in cards]
    gap = 0.4
    w = (CW - gap * (len(cards) - 1)) / len(cards)
    avail = (Y1 - Y0) - (0.75 if extra else 0.1)
    size, need = pp._column_font([c.bullets for c in cards], w - 0.6, avail - 1.8, 14, 9.5, bullet=False)
    h = min(avail, max(3.0, 1.8 + need + 0.35))
    for i, c in enumerate(cards):
        x = M + i * (w + gap)
        pp._shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, x, Y0, w, h, pp.TINT, radius=0.04)
        pp._label(s, x + 0.3, Y0 + 0.25, 1.4, 0.6, f"{i + 1:02d}", 30, pp.ORANGE, bold=True)
        pp._fitted_text(s, deck.report, f"{page['page_no']}쪽 모듈 이름", x + 0.3, Y0 + 0.95, w - 0.6, 0.5,
                        [c.heading], max_pt=16, min_pt=11, color=pp.NAVY, bold=True)
        pp._shape(s, MSO_SHAPE.RECTANGLE, x + 0.3, Y0 + 1.5, 0.6, 0.03, pp.NAVY)
        pp._fitted_text(s, deck.report, f"{page['page_no']}쪽 모듈 설명", x + 0.3, Y0 + 1.65, w - 0.6, h - 1.8,
                        c.bullets, max_pt=13, min_pt=9.5, color=pp.INK, size=size)
    if extra:
        line = " / ".join(((c.heading + ": ") if c.heading not in ("", "-") else "") + " ".join(c.bullets[:1])
                          for c in extra)
        pp._fitted_text(s, deck.report, f"{page['page_no']}쪽 추가 검토", M, Y0 + h + 0.3, CW, 0.5, [line],
                        max_pt=11, min_pt=9, color=pp.GREY)


# 견적·주요 항목 ------------------------------------------------------------

_QTY_RE = re.compile(r"^\s*(\d+(?:\.\d+)?)")


def _group_kind(line: dict[str, Any]) -> str:
    g = str(line.get("group") or "")
    if line.get("alt_id") or g.startswith("대안"):
        return "alt"
    if g.startswith("옵션"):
        return "option"
    if g.startswith("현장"):
        return "site"
    return "common"


def _qty(line: dict[str, Any]) -> str:
    q = line.get("qty")
    return f"{q}{line.get('unit') or ''}" if q not in (None, "") else "-"


def _amount(line: dict[str, Any]) -> float | None:
    """근거가 있는 단가 × 수량. 근거·단가·수량 중 하나라도 없으면 None(= 별도 협의)."""
    price, basis = line.get("unit_price"), str(line.get("basis") or "").strip()
    m = _QTY_RE.match(str(line.get("qty") if line.get("qty") is not None else ""))
    if price in (None, "") or not basis or not m:
        return None
    try:
        return float(price) * float(m.group(1))
    except (TypeError, ValueError):
        return None


def _money(v: float | None, currency: str = "원") -> str:
    if v is None:
        return PRICE_TBD
    return f"{v:,.0f}{currency}" if currency == "원" else f"{currency} {v:,.0f}"


def comparing(project: dict[str, Any]) -> bool:
    """비교안이 2개 이상인가(기본은 공정 컨셉 1개 — 고객이 비교를 원할 때만 비교안을 더한다)."""
    alts = project.get("alternatives") or []
    ids = {a.get("id") for a in alts} or {l.get("alt_id") for l in project.get("quote_lines") or [] if l.get("alt_id")}
    return len(ids) > 1


def _alt_name(project: dict[str, Any], alt_id: str | None) -> str:
    if not comparing(project):
        name = next((a.get("name") for a in project.get("alternatives") or [] if a.get("id") == alt_id), "")
        return f"{SINGLE_LABEL} · {name}" if name else SINGLE_LABEL
    for a in project.get("alternatives") or []:
        if a.get("id") == alt_id:
            return f"{alt_id}안 · {a.get('name') or ''}".strip(" ·")
    return f"{alt_id}안" if alt_id else "대안"


def quote_rows(project: dict[str, Any]) -> tuple[list[list[str]], list[list[str]]]:
    """(견적 표 행, 대안별 기본 구성 합계 행). 대안끼리는 절대 더하지 않는다."""
    lines = project.get("quote_lines") or []
    cur = next((str(l.get("currency")) for l in lines if l.get("currency")), "원")
    rows: list[list[str]] = []
    for kind, label in (("common", "공통"), ("alt", None), ("option", "옵션"), ("site", "현장 적용")):
        for l in (x for x in lines if _group_kind(x) == kind):
            grp = _alt_name(project, l.get("alt_id")) if kind == "alt" else (str(l.get("group") or label))
            note = ("선택 1종" if comparing(project) else "포함") if kind == "alt" else (
                "별도(선택)" if kind in ("option", "site") else
                                                    ("포함" if l.get("included", True) else "별도"))
            rows.append([grp, _clean(l.get("item"), 40), _qty(l), _money(_amount(l), cur), note])
    common = [l for l in lines if _group_kind(l) == "common" and l.get("included", True)]
    totals: list[list[str]] = []
    alt_ids = [a["id"] for a in project.get("alternatives") or []] or sorted(
        {l["alt_id"] for l in lines if l.get("alt_id")})
    for aid in alt_ids:
        part = common + [l for l in lines if l.get("alt_id") == aid]
        amounts = [_amount(l) for l in part]
        total = None if (not part or any(a is None for a in amounts)) else sum(amounts)
        name = _alt_name(project, aid) if comparing(project) else SINGLE_LABEL
        totals.append([f"기본 구성 = 공통 + {name}", _money(total, cur)])
    return rows, totals


def equipment_rows(project: dict[str, Any]) -> list[list[str]]:
    """주요 항목 표 — 견적과 같은 목록(현장 적용비 제외)이라 수량이 서로 어긋나지 않는다."""
    out = []
    for kind in ("common", "alt", "option"):
        for l in (x for x in project.get("quote_lines") or [] if _group_kind(x) == kind):
            grp = _alt_name(project, l.get("alt_id")) if kind == "alt" else str(l.get("group") or
                                                                              ("옵션" if kind == "option" else "공통"))
            out.append([grp, _clean(l.get("item"), 40), _qty(l),
                        _clean(l.get("basis") if kind != "option" else "선택 검토", 40) or "-"])
    return out


_ROW_PAD = 0.1         # 표 행 위아래 여백(in) — 견적은 14행 안팎이 한 쪽에 들어가야 한다


def _grid_row_h(cells: list[str], widths: list[float], size: float) -> float:
    return max(pp._lines_needed([c], w - 0.14, size) for c, w in zip(cells, widths)) * size * 1.2 / 72 + _ROW_PAD



# 표 ------------------------------------------------------------------------

def _grid(slide, report: RenderReport, where: str, x, y, widths, rows: list[list[str]], avail_h: float,
          *, max_pt: float = 11.0, merge_first: bool = True) -> float:
    """표(첫 행 머리). 같은 구분이 이어지면 구분 칸은 첫 행에만. 넘치면 글자를 줄이고, 최소 크기로도 넘치면
    보고한다(쪽수 상한 때문에 다음 쪽으로 나누지 않는다)."""
    shown = [rows[0]] + [[("" if merge_first and i and r[0] == rows[1:][i - 1][0] else r[0])] + r[1:]
                         for i, r in enumerate(rows[1:])]
    size = max_pt
    while size > 8 and sum(_grid_row_h(r, widths, size) for r in shown) > avail_h:
        size -= 0.5
    total = sum(_grid_row_h(r, widths, size) for r in shown)
    if total > avail_h + 0.02:
        report.overflow.append(where)
    shp = slide.shapes.add_table(len(shown), len(widths), Inches(x), Inches(y), Inches(sum(widths)),
                                 Inches(total))
    tbl = shp.table
    for j, wv in enumerate(widths):
        tbl.columns[j].width = Inches(wv)
    for i, row in enumerate(shown):
        tbl.rows[i].height = Emu(Inches(_grid_row_h(row, widths, size)))
        for j, val in enumerate(row):
            cell = tbl.cell(i, j)
            cell.fill.solid()
            cell.fill.fore_color.rgb = pp.NAVY if i == 0 else (pp.TINT if i % 2 == 0 else pp.WHITE)
            cell.margin_left = cell.margin_right = Inches(0.09)
            cell.margin_top = cell.margin_bottom = Inches(0.02)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            cell.text_frame.word_wrap = True
            para = cell.text_frame.paragraphs[0]
            para.font.size = Pt(size)      # 빈 칸도 같은 크기 — 기본 18pt 로 행이 높아졌다(실측)
            r = para.add_run()
            r.text = val
            color = pp.WHITE if i == 0 else (pp.NAVY if j == 0 else (pp.MUTED if val == PRICE_TBD else pp.INK))
            _set_font(r, size, color, bold=(i == 0 or j == 0))
    return total


def _join_items(lines: list[dict[str, Any]]) -> str:
    """'후단 이송 컨베이어 연계 · 비전 2대 …' — 수량이 1이 아니면 붙인다."""
    out = []
    for l in lines:
        q = str(l.get("qty") or "").strip()
        out.append(_clean(l.get("item"), 40) + (f" {q}{l.get('unit') or ''}" if q and q != "1" else ""))
    return " · ".join(out)


def _alt_short(project: dict[str, Any], alt_id: str | None) -> str:
    a = next((a for a in project.get("alternatives") or [] if a.get("id") == alt_id), {})
    return (f"{alt_id}. {a.get('name') or ''}".strip(". ") if comparing(project)
            else f"{SINGLE_LABEL} · {a.get('name') or ''}".strip(" ·"))


def equipment_matrix(project: dict[str, Any]) -> list[list[str]]:
    """주요 항목 비교표(제출본 식): 행 = 로봇 구성·그리퍼·툴·기타 전용 설비, 열 = 컨셉. 견적 줄과 같은 목록."""
    alts = [a["id"] for a in project.get("alternatives") or []] or sorted(
        {l["alt_id"] for l in project.get("quote_lines") or [] if l.get("alt_id")})
    lines = [l for l in project.get("quote_lines") or [] if _group_kind(l) == "alt"]
    head = ["구분"] + [_alt_short(project, a) for a in alts]
    rows = []
    for label, test in (("로봇 구성", lambda t: bool(_ROBOT_RE.search(t)) and not _TOOL_RE.search(t)),
                        ("그리퍼·툴", lambda t: bool(_TOOL_RE.search(t))),
                        ("전용 설비", lambda t: not _ROBOT_RE.search(t) and not _TOOL_RE.search(t))):
        cells = [_join_items([l for l in lines if l.get("alt_id") == a and test(str(l.get("item") or ""))]) or "-"
                 for a in alts]
        if any(c != "-" for c in cells):
            rows.append([label] + cells)
    return [head] + rows if rows else []


def _equipment_page(deck: _Pages, project, page, text: PageText, products: list[tuple[str, bytes]]) -> None:
    """비교표(위) + 제품 사진 후보(왼쪽 아래) + 공통 모듈 수량(오른쪽 아래)."""
    from PIL import Image

    if not equipment_rows(project):
        s = _page(deck, page.get("title") or "", text.headline)
        _columns(deck, s, f"{page['page_no']}쪽", text.cards or [Card("-", [])], Y0, Y1 - Y0)
        return
    s = _page(deck, page.get("title") or "", text.headline)
    y = Y0
    matrix = equipment_matrix(project)
    if matrix:
        k = len(matrix[0]) - 1
        widths = [1.7] + [(CW - 1.7) / k] * k
        y += _grid(s, deck.report, f"{page['page_no']}쪽 비교표", M, y, widths, matrix, 1.9, max_pt=11,
                   merge_first=False) + 0.4
    common = [l for l in project.get("quote_lines") or [] if _group_kind(l) == "common" and l.get("included", True)]
    options = [l for l in project.get("quote_lines") or [] if _group_kind(l) == "option"]
    right_w = 5.4 if products else CW
    rx = M + CW - right_w
    rows = [["공통 모듈", "수량"]] + [[_clean(l.get("item"), 40), _qty(l)] for l in common]
    if not matrix:    # 컨셉이 하나면 컨셉 설비도 같은 표에
        rows += [[_alt_short(project, l.get("alt_id")) + " · " + _clean(l.get("item"), 30), _qty(l)]
                 for l in project.get("quote_lines") or [] if _group_kind(l) == "alt"]
    avail = Y1 - y - (0.4 if options else 0.05)
    if len(rows) > 1:
        _grid(s, deck.report, f"{page['page_no']}쪽 공통 모듈 표", rx, y, [right_w - 1.2, 1.2], rows, avail,
              max_pt=11, merge_first=False)
    if options:
        pp._fitted_text(s, deck.report, f"{page['page_no']}쪽 선택 옵션", rx, Y1 - 0.35, right_w, 0.32,
                        ["선택 옵션: " + _join_items(options)], max_pt=10, min_pt=8.5, color=pp.GREY)
    if products:
        lw = CW - right_w - 0.5
        pp._label(s, M, y, lw, 0.32, "그리퍼·툴 제품 후보", 12, pp.NAVY, bold=True)
        n = len(products)
        cw = (lw - 0.25 * (n - 1)) / n
        ph = min(Y1 - y - 1.05, cw * 1.0)
        for i, (name, blob) in enumerate(products):
            x = M + i * (cw + 0.25)
            try:
                w_px, h_px = Image.open(io.BytesIO(blob)).size
            except Exception:  # noqa: BLE001 — 사진 하나 못 읽어도 쪽은 만든다
                continue
            sc = min(cw / w_px, ph / h_px)
            iw, ih = w_px * sc, h_px * sc
            s.shapes.add_picture(io.BytesIO(blob), Inches(x + (cw - iw) / 2), Inches(y + 0.45 + (ph - ih) / 2),
                                 Inches(iw), Inches(ih))
            pp._fitted_text(s, deck.report, f"{page['page_no']}쪽 제품 이름", x, y + 0.5 + ph, cw, 0.4, [name],
                            max_pt=10.5, min_pt=8, color=pp.NAVY, bold=True, align=PP_ALIGN.CENTER)
        pp._label(s, M, Y1 - 0.3, lw, 0.28, "※ 제품 사진은 외형 참고용이며, 최종 사양은 선정 후 확정합니다.", 8.5,
                  pp.GREY, italic=True)
    note = f"※ {ALT_NOTE} " if comparing(project) else "※ "
    pp._label(s, M, Y1 + 0.02, CW, 0.22, note + "수량은 개념 수량이며 최종 선정에 따라 갱신됩니다.", 8, pp.MUTED,
              italic=True)


def _poc_page(deck: _Pages, page: dict[str, Any], t: PageText) -> None:
    """실증 단계 3개(연결된 상자) + 조건별 검토 방향 표 + 확인 자료 한 줄."""
    s = _page(deck, page.get("title") or "", t.headline)
    phases = [c for c in t.cards if c.heading not in ("", "-") or c.bullets][:3] or [Card("-", [])]
    gap, h = 0.5, 1.35 if t.checks else 2.6
    w = (CW - gap * (len(phases) - 1)) / len(phases)
    size, _ = pp._column_font([c.bullets for c in phases], w - 0.35, h - 0.6, 11.5, 9, bullet=False)
    for i, c in enumerate(phases):
        x = M + i * (w + gap)
        pp._shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, x, Y0, w, h, pp.TINT, radius=0.06)
        pp._fitted_text(s, deck.report, f"{page['page_no']}쪽 단계", x + 0.18, Y0 + 0.1, w - 0.3, 0.42,
                        [f"{i + 1:02d}  {c.heading}"], max_pt=13.5, min_pt=10.5, color=pp.NAVY, bold=True)
        pp._fitted_text(s, deck.report, f"{page['page_no']}쪽 단계 설명", x + 0.18, Y0 + 0.55, w - 0.3, h - 0.62,
                        c.bullets, max_pt=11.5, min_pt=9, color=pp.INK, size=size)
        if i < len(phases) - 1:
            _connector(s, x + w + 0.04, Y0 + h / 2, x + w + gap - 0.04, Y0 + h / 2, pp.NAVY_MID)
    y = Y0 + h + 0.4
    if t.checks:
        avail = Y1 - y - (0.6 if t.confirm else 0.05)
        y += _grid(s, deck.report, f"{page['page_no']}쪽 검토 표", M, y, [3.4, CW - 3.4],
                   [["구분", "검토 방향"]] + t.checks, avail, max_pt=11, merge_first=False) + 0.25
    elif len(t.cards) > 3:
        _columns(deck, s, f"{page['page_no']}쪽 추가", t.cards[3:5], y, Y1 - y - 0.6, max_pt=11)
    if t.confirm:
        pp._label(s, M, Y1 - 0.42, 1.2, 0.34, "확인 자료", 11, pp.NAVY, bold=True, anchor=MSO_ANCHOR.MIDDLE)
        pp._fitted_text(s, deck.report, f"{page['page_no']}쪽 확인 자료", M + 1.25, Y1 - 0.42, CW - 1.25, 0.34,
                        [t.confirm], max_pt=11, min_pt=9, color=pp.GREY, anchor=MSO_ANCHOR.MIDDLE)


def quote_groups(project: dict[str, Any]) -> list[list[str]]:
    """견적 묶음 행(제출본 식) [항목, 구성·공급 범위, 구분, 금액]. 공통은 묶음(group)별, 대안은 컨셉별 한 줄,
    옵션 한 줄, 기본 범위에서 뺀 것 한 줄. 현장 적용비는 표 밖 주석. 묶음 금액은 모든 줄에 근거 단가가 있을 때만."""
    lines = project.get("quote_lines") or []
    cur = next((str(l.get("currency")) for l in lines if l.get("currency")), "원")
    multi = comparing(project)

    def amount(part: list[dict[str, Any]]) -> str:
        vals = [_amount(l) for l in part]
        return _money(None if not part or any(v is None for v in vals) else sum(vals), cur)

    rows: list[list[str]] = []
    groups: dict[str, list[dict[str, Any]]] = {}
    for l in lines:
        if _group_kind(l) == "common" and l.get("included", True):
            groups.setdefault(str(l.get("group") or "공통"), []).append(l)
    for g, part in groups.items():
        rows.append([g if g != "공통" else "공통 설비", _join_items(part), "공통" if g == "공통" else "포함",
                     amount(part)])
    alt_ids = [a["id"] for a in project.get("alternatives") or []] or sorted(
        {l["alt_id"] for l in lines if l.get("alt_id")})
    for aid in alt_ids:
        part = [l for l in lines if _group_kind(l) == "alt" and l.get("alt_id") == aid]
        if part:
            rows.append([_alt_short(project, aid), _join_items(part), "택 1안" if multi else "포함", amount(part)])
    opts = [l for l in lines if _group_kind(l) == "option"]
    if opts:
        rows.append(["선택 옵션", _join_items(opts), "선택 적용", amount(opts)])
    out = [l for l in lines if _group_kind(l) == "common" and not l.get("included", True)]
    if out:
        rows.append(["범위 제외", _join_items(out), "제외", "-"])
    return rows


def _quote_page(deck: _Pages, project, page, text: PageText) -> None:
    """묶음 견적 표 + 기본 구성(택 1) 줄 + 현장 적용비·단가 주석."""
    rows = quote_groups(project)
    if not rows:
        s = _page(deck, page.get("title") or "", text.headline)
        _columns(deck, s, f"{page['page_no']}쪽", text.cards or [Card("-", [])], Y0, Y1 - Y0)
        return
    multi = comparing(project)
    s = _page(deck, page.get("title") or "", QUOTE_HEADLINE if multi else QUOTE_HEADLINE_SINGLE)
    widths = [2.4, CW - 2.4 - 1.35 - 1.6, 1.35, 1.6]
    site = [l for l in project.get("quote_lines") or [] if _group_kind(l) == "site"]
    reserve = 1.25 + (0.35 if site else 0)
    used = _grid(s, deck.report, f"{page['page_no']}쪽 견적 표", M, Y0, widths,
                 [["항목", "구성·공급 범위", "구분", "금액"]] + rows, Y1 - Y0 - reserve, max_pt=11, merge_first=False)
    y = Y0 + used + 0.3
    _, totals = quote_rows(project)
    pp._label(s, M, y, 2.2, 0.34, "기본 구성(택 1)" if multi else "기본 구성", 11.5, pp.NAVY, bold=True,
              anchor=MSO_ANCHOR.MIDDLE)
    # 대안마다 따로(한 글상자에 두 대안을 적으면 합산처럼 읽힌다).
    cw = (CW - 2.3) / max(1, len(totals))
    for k, (label, value) in enumerate(totals):
        pp._fitted_text(s, deck.report, f"{page['page_no']}쪽 기본 구성", M + 2.3 + k * cw, y, cw - 0.15, 0.34,
                        [f"{label.replace('기본 구성 = ', '')}  ·  {value}"], max_pt=11, min_pt=8,
                        color=pp.INK, anchor=MSO_ANCHOR.MIDDLE)
    y += 0.45
    if site:
        pp._label(s, M, y, 2.2, 0.3, "현장 적용비", 11, pp.NAVY, bold=True)
        pp._fitted_text(s, deck.report, f"{page['page_no']}쪽 현장 적용비", M + 2.3, y, CW - 2.3, 0.3,
                        [_join_items(site) + " — 운송·설치·시운전·기존 설비 개조 범위를 확인해 별도 산정합니다."],
                        max_pt=10.5, min_pt=8.5, color=pp.GREY)
        y += 0.4
    pp._fitted_text(s, deck.report, f"{page['page_no']}쪽 견적 기준", M, max(y, Y1 - 0.3), CW, 0.3,
                    [(f"※ 근거 단가가 없는 항목은 '{PRICE_TBD}'로 표시합니다. "
                      + ("대안끼리 합산하지 않습니다. " if multi else "") + "부가세·운송·설치 포함 여부는 협의합니다.")],
                    max_pt=8.5, min_pt=7.5, color=pp.MUTED, italic=True)


def _guard_sizes(prs, report: DeckReport) -> None:
    """크기 0 이하 도형이 하나라도 있으면 PowerPoint 가 파일을 못 연다(실측: 두 줄 공정도 아래 요점 글상자).
    마지막에 한 번 더 막는다 — 연결선은 가로·세로 한쪽이 0 인 게 정상이라 뺀다."""
    for no, slide in enumerate(prs.slides, 1):
        for sh in slide.shapes:
            if sh.shape_type == MSO_SHAPE_TYPE.LINE or sh.width is None:
                continue
            if sh.width <= 0 or sh.height <= 0:
                sh.width, sh.height = max(sh.width, Inches(0.1)), max(sh.height, Inches(0.1))
                report.overflow.append(f"{no}쪽 자리 부족(글상자 크기 보정)")


# ── 조립 ────────────────────────────────────────────────────────────────────

def render(project: dict[str, Any], pages: list[dict[str, Any]], text: DeckText,
           images: dict[int, bytes] | None = None, *, now: datetime | None = None,
           report: DeckReport | None = None, products: list[tuple[str, bytes]] | None = None
           ) -> tuple[bytes, DeckReport]:
    """문구·이미지가 정해진 뒤의 조판(결정적). images: project["images"][].id → 이미지 bytes.
    products: 주요 항목 쪽에 넣을 회사 제품 사진 [(이름, jpeg)]."""
    report = report or DeckReport()
    images = images or {}
    by_img = {i["id"]: i for i in project.get("images") or []}
    alts = [a.get("name") for a in project.get("alternatives") or [] if a.get("name")]
    body = ProposalBody(text.title or "로봇 자동화 제안", text.subtitle or (" · ".join(alts) if alts else ""),
                        text.customer, [])
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(SLIDE_W), Inches(SLIDE_H)
    short = " ".join(x for x in (text.customer, _clean(text.title, 30)) if x)
    deck = _Pages(prs, body, short)
    limit = int((project.get("output") or {}).get("max_pages") or 10)
    if len(pages) > limit:
        report.warnings.append(f"페이지 {len(pages)}쪽 — 상한 {limit}쪽 초과")
    for page in sorted(pages, key=lambda p: p["page_no"]):
        t = text.pages.get(page["page_no"]) or PageText(page.get("title") or "")
        kind = page["type"]
        if kind == "cover":
            _cover(deck, now or datetime.now(), alts)
        elif kind == "overview":
            _overview(deck, page, t)
        elif kind == "flow":
            _flow(deck, page, t)
        elif kind in ("common_concept", "alternative"):
            meta = next((by_img[a] for a in page.get("asset_ids") or [] if a in by_img), None)
            img = images.get(meta["id"]) if meta else None
            if kind == "common_concept" and img is None:
                _common_modules(deck, page, t)
                continue
            tag = f"CONCEPT {page.get('alt_id') or ''}".strip() if kind == "alternative" and comparing(project) else ""
            _concept_page(deck, page, t, img, (meta or {}).get("labels") or [], report, tag)
        elif kind == "equipment":
            _equipment_page(deck, project, page, t, products or [])
        elif kind == "poc":
            _poc_page(deck, page, t)
        elif kind == "quote":
            _quote_page(deck, project, page, t)
        else:
            s = _page(deck, page.get("title") or "", t.headline)
            _columns(deck, s, f"{page['page_no']}쪽", t.cards or [Card("-", [])], Y0, Y1 - Y0)
    report.pages = len(prs.slides)
    report.overflow += deck.report.overflow
    _guard_sizes(prs, report)
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue(), report


async def load_images(project: dict[str, Any], pages: list[dict[str, Any]]) -> dict[int, bytes]:
    """페이지에 쓰는 승인 이미지만 첨부 저장소에서 읽는다."""
    from sqlalchemy import text as sql

    from ..database import SessionLocal
    from ..storage import get_object_stream

    wanted = {a for p in pages for a in p.get("asset_ids") or []}
    metas = [i for i in project.get("images") or [] if i["id"] in wanted and i.get("attachment_id")]
    if not metas:
        return {}
    async with SessionLocal() as db:
        keys = dict((await db.execute(sql("SELECT id, object_key FROM attachments WHERE id = ANY(:ids)"),
                                      {"ids": [m["attachment_id"] for m in metas]})).all())
    out = {}
    for m in metas:
        key = keys.get(m["attachment_id"])
        if not key:
            continue
        resp = get_object_stream(key)
        try:
            out[m["id"]] = b"".join(resp.stream(amt=64 * 1024))
        finally:
            resp.close()
            resp.release_conn()
    return out


PRODUCT_PHOTOS = 3


async def load_products(project: dict[str, Any], pages: list[dict[str, Any]]) -> list[tuple[str, bytes]]:
    """주요 항목 쪽이 있으면, 견적 줄·컨셉에 나오는 회사 제품의 승인 사진(최대 3장). 실패해도 쪽은 만든다."""
    if not any(p["type"] == "equipment" for p in pages):
        return []
    from ..company_knowledge.product_images import photos_for_text

    # 그리퍼·툴 줄에서만 고른다 — 'AMR 양팔형' 같은 부류 말로 컨셉과 다른 이송 로봇 사진이 붙었다(실측).
    lines = " ".join(str(l.get("item") or "") for l in project.get("quote_lines") or []
                     if _group_kind(l) in ("alt", "common", "option") and _TOOL_RE.search(str(l.get("item") or "")))
    # 컨셉 문장은 쓰지 않는다 — 한 문장에 그리퍼와 'AMR 양팔형 로봇'이 함께 나와 이송 로봇 사진이 붙었다(실측).
    if not lines:
        return []
    try:
        photos = await photos_for_text(lines, "", PRODUCT_PHOTOS)
    except Exception as e:  # noqa: BLE001
        _log.warning("[project_deck] 제품 사진 읽기 실패: %r", e)
        return []
    return [(name, blob) for blob, _mime, name in photos]


async def generate_pptx(project: dict[str, Any], *, images: dict[int, bytes] | None = None,
                        now: datetime | None = None, products: list[tuple[str, bytes]] | None = None,
                        ) -> tuple[bytes, DeckReport]:
    """승인된 구성(없으면 기본안)으로 문구를 쓰고 조판한다. images 를 안 주면 저장소에서 읽는다."""
    data, report, _deck = await generate_deck(project, images=images, now=now, products=products)
    return data, report


async def generate_deck(project: dict[str, Any], *, images: dict[int, bytes] | None = None,
                        now: datetime | None = None, products: list[tuple[str, bytes]] | None = None,
                        ) -> tuple[bytes, DeckReport, dict[str, Any]]:
    """generate_pptx + 나중에 말로 고칠 때 쓸 문구·구성(deck_state)."""
    pages = project.get("pages") or build_default_pages(project)
    report = DeckReport()
    if not project.get("pages"):
        report.warnings.append("승인된 페이지 구성이 없어 기본 구성으로 만들었습니다.")
    text = await compose(project, pages, report)
    if images is None:
        images = await load_images(project, pages)
    if products is None:
        products = await load_products(project, pages)
    data, report = render(project, pages, text, images, now=now, report=report, products=products)
    return data, report, deck_state(pages, text)


# ── 완성 후 프롬프트로 수정 ─────────────────────────────────────────────────────

def _page_to_dict(t: PageText) -> dict[str, Any]:
    return {"headline": t.headline, "cards": [{"heading": c.heading, "bullets": list(c.bullets)} for c in t.cards],
            "steps": [{"label": s.label, "note": s.note, "kind": s.kind, "from": s.src} for s in t.steps],
            "kpis": [dict(k) for k in t.kpis], "keywords": list(t.keywords),
            "checks": [{"item": a, "direction": b} for a, b in t.checks], "confirm": t.confirm}


def deck_state(pages: list[dict[str, Any]], text: DeckText) -> dict[str, Any]:
    """출력본에 함께 저장하는 문구·구성(JSON) — 프롬프트 수정의 출발점(수정할 때마다 최신 버전 것을 쓴다)."""
    return {"pages": pages, "title": text.title, "subtitle": text.subtitle, "customer": text.customer,
            "text": {str(n): _page_to_dict(t) for n, t in text.pages.items()}}


def deck_from_state(state: dict[str, Any]) -> tuple[list[dict[str, Any]], DeckText]:
    pages = [dict(p) for p in state.get("pages") or []]
    kinds = {p["page_no"]: p["type"] for p in pages}
    text = DeckText(state.get("title") or "", state.get("subtitle") or "", state.get("customer") or "")
    for n, d in (state.get("text") or {}).items():
        text.pages[int(n)] = _parse_page(d, kinds.get(int(n)) == "flow")
    return pages, text


EDIT_MAX = 1000
_EDIT_SYSTEM = """너는 유엔디로보틱스 제안서 편집자다. [현재 제안서]는 쪽별 문구 JSON 이다. [수정 요청]대로 고칠 쪽만 돌려준다.
규칙:
- 요청과 관계없는 쪽·문장은 돌려주지 않는다(그대로 둔다). 고치는 쪽은 그 쪽의 모든 칸을 완성본으로 돌려준다.
- 수치는 [현재 제안서]·[문항]·[수정 요청]에 있는 것만 쓴다. 새 수치·등급·가격을 만들지 않는다.
- 성능은 단정하지 않는다("대응 가능·보장" 대신 "검토"). 과장 표현 금지.
- 쪽 제목을 바꾸려면 "title". 쪽을 빼려면 {"page_no": n, "remove": true}(표지는 뺄 수 없다).
- 쪽 순서를 바꾸려면 "order": [쪽 번호를 원하는 순서대로 전부]. 표지는 맨 앞.
- 표지 제목·부제·고객사를 바꾸려면 "cover": {"title": "...", "subtitle": "...", "customer": "..."}.
- 견적 표·주요 항목 표의 행·수량·금액과 컨셉 그림은 여기서 바꿀 수 없다. 그런 요청이면 note 에 "구성·견적 단계에서
  고쳐야 합니다"라고 적는다.
- 글 분량 한도: headline 60자, 카드 2~4개, 카드마다 bullets 1~3개(45자 이내).
- 칸 이름(요청에 나오는 말 → 고칠 칸): 제목 아래 한 줄·요지 → headline, 카드·설명 → cards, 핵심 수치 → kpis,
  공정 단계 → steps, 그림 아래 요점 → keywords, 검토 표 → checks, '확인 자료' 줄 → confirm. 요청한 칸을 고친다
  (예: 확인 자료에 넣어 달라 → confirm 을 고친다. checks 에 행을 더하지 않는다).
- 공정 흐름(flow) 쪽 단계 수를 요청받으면 그 수는 kind=main 단계 수다(예외·분기는 세지 않는다). main 은 최대 8개,
  label 은 실제 장치·모듈 이름 10자 이내.
출력은 JSON 하나:
{"cover": {...} 또는 생략, "pages": [{"page_no": 2, "title": "...", "headline": "...", "cards": [...], "steps": [...],
 "kpis": [...], "keywords": [...], "checks": [...], "confirm": "..."}], "order": [...] 또는 생략,
 "note": "무엇을 바꿨는지 한두 문장"}"""


def _edit_user(project: dict[str, Any], pages: list[dict[str, Any]], text: DeckText, request: str) -> str:
    cur = {"cover": {"title": text.title, "subtitle": text.subtitle, "customer": text.customer},
           "pages": [{"page_no": p["page_no"], "type": p["type"], "title": p.get("title") or "",
                      **({k: v for k, v in _page_to_dict(text.pages[p["page_no"]]).items() if v}
                         if p["page_no"] in text.pages else {})}
                     for p in sorted(pages, key=lambda p: p["page_no"]) if p["type"] != "cover"]}
    codes = sorted({c for p in pages for c in p.get("item_codes") or []})
    lines = "\n".join(_item_lines(project, codes))
    return f"[문항]\n{lines}\n\n[현재 제안서]\n{json.dumps(cur, ensure_ascii=False)}\n\n[수정 요청]\n{request}"


@dataclass
class EditResult:
    pages: list[dict[str, Any]]
    text: DeckText
    note: str
    changed: list[str]                    # 바뀐 곳 ["2쪽 문구", "순서", "표지"]
    report: DeckReport


def apply_edit(project: dict[str, Any], pages: list[dict[str, Any]], text: DeckText, data: dict[str, Any],
               request: str) -> EditResult:
    """모델이 돌려준 수정본을 검증해 적용(결정적). 근거 없는 수치는 가리고, 표지는 빼지 않는다."""
    report = DeckReport()
    nums, src = _allowed(project)
    nums |= _numbers_in(request) | _numbers_in("\n".join(_page_text_all(t) for t in text.pages.values()))
    src = f"{src}\n{request}\n" + "\n".join(_page_text_all(t) for t in text.pages.values())
    pages = [dict(p) for p in pages]
    by_no = {p["page_no"]: p for p in pages}
    changed: list[str] = []
    cover = data.get("cover") if isinstance(data.get("cover"), dict) else None
    if cover:
        fields = {"title": 40, "subtitle": 70, "customer": 30}
        for k, lim in fields.items():
            v = _clean(cover.get(k), lim)
            if v and v != getattr(text, k):
                v2, m = mask_unsupported_numbers(v, nums, set())
                report.masked += m
                setattr(text, k, v2)
                if "표지" not in changed:
                    changed.append("표지")
    removed: set[int] = set()
    for d in data.get("pages") or []:
        if not isinstance(d, dict):
            continue
        try:
            no = int(d.get("page_no"))
        except (TypeError, ValueError):
            continue
        p = by_no.get(no)
        if p is None or p["type"] == "cover":
            continue
        if d.get("remove"):
            removed.add(no)
            changed.append(f"{no}쪽 삭제")
            continue
        title = _clean(d.get("title"), 40)
        if title and title != p.get("title"):
            p["title"] = title
        old = text.pages.get(no) or PageText()
        merged = {**_page_to_dict(old), **{k: v for k, v in d.items() if k in (
            "headline", "cards", "steps", "kpis", "keywords", "checks", "confirm")}}
        new = _parse_page(merged, p["type"] == "flow")
        _mask_page(new, nums, src, report)
        if _page_to_dict(new) != _page_to_dict(old) or title:
            text.pages[no] = new
            changed.append(f"{no}쪽")
    order = [o for o in data.get("order") or [] if isinstance(o, int) and o in by_no and o not in removed]
    keep = [p for p in pages if p["page_no"] not in removed]
    if order:
        rank = {n: i for i, n in enumerate(order)}
        cover_first = [p for p in keep if p["type"] == "cover"]
        rest = sorted((p for p in keep if p["type"] != "cover"), key=lambda p: rank.get(p["page_no"], 999))
        if [p["page_no"] for p in rest] != [p["page_no"] for p in keep if p["type"] != "cover"]:
            changed.append("쪽 순서")
        keep = cover_first + rest
    # 쪽 번호를 다시 매기고 문구도 같이 옮긴다.
    new_text = DeckText(text.title, text.subtitle, text.customer)
    for n, p in enumerate(keep, 1):
        if p["page_no"] in text.pages:
            new_text.pages[n] = text.pages[p["page_no"]]
        p["page_no"] = n
    return EditResult(keep, new_text, _clean(data.get("note"), 300), changed, report)


# 요청 글에서 '어느 쪽을 고치라는지' — 쪽 번호("7쪽") 또는 쪽 이름 + 쪽/페이지("실증 쪽", "견적 페이지", "표지").
_PAGE_NAME_RE = re.compile(r"(표지|개요|공정\s*흐름|공정도|공통\s*(?:적용\s*)?컨셉|주요\s*항목|실증|견적)"
                           r"(?:[^\s,.]{0,8})?\s*(쪽|페이지|장\b)|표지")
_PAGE_NAME_TYPE = (("표지", "cover"), ("개요", "overview"), ("공정", "flow"), ("공통", "common_concept"),
                   ("주요", "equipment"), ("실증", "poc"), ("견적", "quote"))


def requested_pages(request: str, pages: list[dict[str, Any]]) -> set[int]:
    nos = {p["page_no"] for p in pages}
    out = {int(n) for n in re.findall(r"(\d{1,2})\s*(?:쪽|페이지)", request) if int(n) in nos}
    for m in _PAGE_NAME_RE.finditer(request):
        word = m.group(0)
        kind = next((t for w, t in _PAGE_NAME_TYPE if w in word), None)
        out |= {p["page_no"] for p in pages if p["type"] == kind}
    return out


def changed_pages(res: "EditResult") -> set[int]:
    return {int(m.group(1)) for c in res.changed for m in [re.match(r"(\d+)쪽", c)] if m}


async def edit_deck(project: dict[str, Any], state: dict[str, Any], request: str, *,
                    images: dict[int, bytes] | None = None, products: list[tuple[str, bytes]] | None = None,
                    now: datetime | None = None) -> tuple[bytes, DeckReport, dict[str, Any], EditResult]:
    """완성본을 프롬프트로 수정: 사내 AI 가 고칠 쪽의 문구만 새로 쓰고 → 검증 → 같은 틀로 다시 조판."""
    request = (request or "").strip()
    if not request:
        raise ValueError("고칠 내용을 적어 주세요.")
    pages, text = deck_from_state(state)
    data = await _ask(_EDIT_SYSTEM, _edit_user(project, pages, text, request[:EDIT_MAX]))
    res = apply_edit(project, pages, text, data, request)
    missing = sorted(requested_pages(request, pages) - changed_pages(res))
    if missing and not any(c in ("쪽 순서",) or c.endswith("삭제") for c in res.changed):
        # 여러 쪽을 한 번에 요청하면 일부만 고치고 끝나기도 한다(QA 4회차: 2쪽만 고치고 실증 쪽 누락) — 빠진 쪽만 한 번 더.
        kinds = {p["page_no"]: p["type"] for p in res.pages}
        extra = ("\n\n[보완] 위 요청 중 다음 쪽을 아직 고치지 않았습니다: "
                 + ", ".join(f"{n}쪽({kinds.get(n, '')})" for n in missing) + ". 그 쪽만 요청대로 고쳐 돌려주세요.")
        data2 = await _ask(_EDIT_SYSTEM, _edit_user(project, res.pages, res.text, request[:EDIT_MAX]) + extra)
        res2 = apply_edit(project, res.pages, res.text, data2, request)
        res2.changed = res.changed + [c for c in res2.changed if c not in res.changed]
        res2.note = " ".join(x for x in (res.note, res2.note) if x)
        res2.report.masked = res.report.masked + res2.report.masked
        res2.report.softened = res.report.softened + res2.report.softened
        res = res2
    if images is None:
        images = await load_images(project, res.pages)
    if products is None:
        products = await load_products(project, res.pages)
    blob, report = render(project, res.pages, res.text, images, now=now, report=res.report, products=products)
    return blob, report, deck_state(res.pages, res.text), res
