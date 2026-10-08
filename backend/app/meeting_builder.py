"""주간 회의록 보고 xlsx 생성기.

설계:
- `app/templates/meeting_template.xlsx` 를 매 요청마다 복제(in-memory)한다. 이 파일은
  실제 운영 중인 누적 회의록 문서로, 한 시트 = 한 주차 회의 기록이다.
- 가장 최근 시트(맨 끝)를 `copy_worksheet` 로 복제하여 셀 스타일/병합/열폭을 그대로
  보존한 새 주차 시트를 만든다. 원본 동작 보존 원칙(req.md §2 원칙 1) — 기존 38개
  시트는 손대지 않고 새 시트만 append 한다.
- 새 시트에 제목(B2), 회의 주요 내용(D5), 지시사항(D13) 세 칸만 채운다.

시트 레이아웃(최근 시트 기준, 2026년 시트들 공통):
- B2:N3  (병합) — 제목  '주간회의보고 (YYYY-MM-DD)'
- B5:C12 (병합) — 라벨  '회의 주요 내용'  (템플릿 그대로 보존)
- D5:N12 (병합) — 회의 주요 내용 본문
- B13:C16(병합) — 라벨  '지시사항'        (템플릿 그대로 보존)
- D13:N16(병합) — 지시사항 본문
"""
from __future__ import annotations

import datetime as _dt
import io
from dataclasses import dataclass
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.styles import Alignment

# ── 템플릿 경로 ────────────────────────────────────────────────────────────────
_TEMPLATE_PATH = Path(__file__).parent / "templates" / "meeting_template.xlsx"

# ── 시트 셀 매핑 ──────────────────────────────────────────────────────────────
_TITLE_CELL = "B2"
_MAIN_CONTENT_CELL = "D5"
_DIRECTIVES_CELL = "D13"

# Excel 시트명 최대 길이.
_SHEET_NAME_MAX = 31


@dataclass
class MeetingReport:
    """주간 회의록 한 건. 한 시트로 출력된다."""

    meeting_date: str  # YYYY-MM-DD
    main_content: str  # 회의 주요 내용 (멀티라인)
    directives: str    # 지시사항 (멀티라인)
    author: str = ""   # 작성자 — 파일명에만 사용, 시트에는 기입하지 않음


def _parse_iso(date_str: str) -> _dt.date:
    """'YYYY-MM-DD' → date. 실패 시 오늘 날짜로 fallback."""
    s = (date_str or "").strip()
    for fmt in ("%Y-%m-%d", "%Y/%m/%d", "%Y.%m.%d"):
        try:
            return _dt.datetime.strptime(s, fmt).date()
        except ValueError:
            continue
    return _dt.date.today()


def _sheet_name_for(d: _dt.date) -> str:
    """시트명 — 기존 2026년 시트 명명 규칙('2026년 04월27일')과 동일."""
    return f"{d.year}년 {d.month:02d}월{d.day:02d}일"


def _unique_sheet_name(wb, base: str) -> str:
    """이미 같은 이름 시트가 있으면 ' (2)', ' (3)' … 접미사로 유일화."""
    if base not in wb.sheetnames:
        return base[:_SHEET_NAME_MAX]
    i = 2
    while True:
        candidate = f"{base} ({i})"[:_SHEET_NAME_MAX]
        if candidate not in wb.sheetnames:
            return candidate
        i += 1


def build_meeting_xlsx(report: MeetingReport) -> bytes:
    """MeetingReport 를 xlsx 바이트로 변환해 반환. I/O 없음(호출자가 저장)."""
    if not _TEMPLATE_PATH.exists():
        raise FileNotFoundError(f"템플릿이 없습니다: {_TEMPLATE_PATH}")

    wb = load_workbook(_TEMPLATE_PATH)
    if not wb.worksheets:
        raise ValueError("템플릿에 시트가 없습니다.")

    # 가장 최근(맨 끝) 시트를 양식 기준으로 삼아 복제 — 스타일/병합/열폭 보존.
    source_ws = wb.worksheets[-1]
    new_ws = wb.copy_worksheet(source_ws)

    d = _parse_iso(report.meeting_date)
    new_ws.title = _unique_sheet_name(wb, _sheet_name_for(d))

    # 제목 + 본문 두 칸 기입. 병합 셀은 좌상단 셀에만 값을 넣는다.
    new_ws[_TITLE_CELL] = f"주간회의보고 ({d.isoformat()})"
    new_ws[_MAIN_CONTENT_CELL] = (report.main_content or "").strip()
    new_ws[_DIRECTIVES_CELL] = (report.directives or "").strip()

    # 멀티라인 본문이 잘리지 않도록 줄바꿈 + 상단 정렬 보장(템플릿 정렬 위에 덧씌움).
    for coord in (_MAIN_CONTENT_CELL, _DIRECTIVES_CELL):
        cell = new_ws[coord]
        cell.alignment = Alignment(
            wrap_text=True, vertical="top", horizontal="left"
        )

    # 새 시트를 활성 시트로 — 파일을 열면 이번 주차가 바로 보이도록.
    wb.active = wb.index(new_ws)

    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def _author_token(author: str, yyyymmdd: str) -> str:
    """파일명에 들어갈 작성자 토큰.

    규칙(2026-05-20 변경):
    - 영문 알파벳 이름(예: Jacob) → 앞 2글자 대문자 + 날짜를 underscore 없이 붙임
      (Jacob, 20260520 → 'JA20260520'). 사내 회의록 파일 명명 규칙.
    - 그 외(한글 등) → 기존 형식 '{이름}_{YYYYMMDD}' 유지.
    - 작성자 없음 → 날짜만.
    """
    name = (author or "").strip()
    if not name:
        return yyyymmdd
    # ASCII 알파벳 1자 이상으로만 구성된 영문 이름
    if name.isascii() and name.replace(" ", "").isalpha():
        prefix = name.replace(" ", "")[:2].upper()
        return f"{prefix}{yyyymmdd}"
    return f"{name}_{yyyymmdd}"


def build_filename(report: MeetingReport) -> str:
    """파일명 규칙: '주간 회의록 보고_{작성자토큰}.xlsx'.

    작성자토큰은 영문이면 앞 2자 대문자 + 날짜 (예: 'JA20260520'),
    그 외면 '{이름}_{YYYYMMDD}'. _author_token 참조.
    """
    d = _parse_iso(report.meeting_date)
    yyyymmdd = d.strftime("%Y%m%d")
    return f"주간 회의록 보고_{_author_token(report.author, yyyymmdd)}.xlsx"


__all__ = ["MeetingReport", "build_meeting_xlsx", "build_filename", "_author_token"]
