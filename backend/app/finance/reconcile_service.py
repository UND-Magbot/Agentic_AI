# -*- coding: utf-8 -*-
"""일일자금수지 비교·검증 서비스 — 챗 첨부(거래내역 + 일일자금수지) → 채워진 파일 + 요약.

흐름:
  1) 첨부 2개(거래내역 xlsx, 일일자금수지 xlsx) 바이트를 MinIO 에서 fetch → 임시 파일.
  2) 시트 구성으로 어느 쪽이 거래내역/일일자금수지인지 자동 분류.
  3) 대상 일자(MM-DD) 자동 탐지(당일잔액 공란 시트) — 또는 호출자가 지정.
  4) run_daily.run_reconcile 로 당일잔액·자금일계표 외과적 기입 + 계획↔실적 검증.
     (자금계획은 일일자금수지 파일에 내장된 '자금계획_원화/외화' 시트에서 읽음 → 별도 첨부 불필요)
  5) 채워진 복사본 + 검증 워크북을 MinIO 저장 + Attachment row 생성.
  6) 다운로드 메타 + 사람이 읽을 검증 요약(가공 점수 없이 사실만) 반환.

원본 첨부는 수정하지 않는다(임시 복사본에만 기입). 모든 수치는 실제 산출값.
"""
from __future__ import annotations

import io
import logging
import re
import shutil
import tempfile
from collections import Counter
from dataclasses import dataclass
from pathlib import Path
from typing import Any

import openpyxl
from sqlalchemy import select

from ..database import SessionLocal
from ..models import Attachment, User, UserRole
from ..storage import get_object_stream, make_object_key, put_object
from . import actual_reader, layout, plan_reader, run_daily, tx_reader

logger = logging.getLogger("finance.reconcile_service")

_XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
_MMDD_RE = re.compile(r"^\d{2}-\d{2}$")
_EPS = 0.5

# 진행 단계 ID — frontend ProgressCard 와 동일 식별자.
PROGRESS_STEPS: tuple[tuple[str, str], ...] = (
    ("fetch", "첨부 파일 준비"),
    ("classify", "거래내역·일일자금수지 인식"),
    ("reconcile", "계획↔실적 대조·검증"),
    ("fill", "당일잔액 기입"),
    ("upload", "파일 저장"),
)


@dataclass
class FundReconcileResult:
    """reconcile_fund_daily 호출 결과(챗 fast-path 가 요약/다운로드에 사용)."""

    ok: bool
    date: str
    filename: str
    download_url: str
    size_bytes: int
    attachment_id: int
    verify_filename: str
    verify_download_url: str
    verify_attachment_id: int
    summary: str           # 사람이 읽을 검증 요약(여러 줄)
    written: int


# ── 첨부 fetch ────────────────────────────────────────────────────────────────
async def _fetch_xlsx_attachments(
    db, attachment_ids: list[int]
) -> tuple[list[tuple[Attachment, bytes]], int | None]:
    """첨부 ID 들에서 xlsx 바이트를 순서 보존하며 가져온다. (목록, 첫 owner)."""
    if not attachment_ids:
        return [], None
    res = await db.execute(select(Attachment).where(Attachment.id.in_(attachment_ids)))
    by_id = {a.id: a for a in res.scalars().all()}

    out: list[tuple[Attachment, bytes]] = []
    owner: int | None = None
    for att_id in attachment_ids:
        att = by_id.get(att_id)
        if att is None:
            logger.warning("[reconcile] 첨부 id=%s 없음 — skip", att_id)
            continue
        name = (att.original_filename or "").lower()
        if not (name.endswith(".xlsx") or "spreadsheet" in (att.mime or "")):
            logger.warning("[reconcile] 첨부 id=%s (%s) xlsx 아님 — skip", att_id, att.mime)
            continue
        if owner is None:
            owner = att.user_id
        try:
            resp = get_object_stream(att.object_key)
            try:
                chunks: list[bytes] = []
                for chunk in resp.stream(amt=256 * 1024):
                    chunks.append(chunk)
                data = b"".join(chunks)
            finally:
                resp.close()
                resp.release_conn()
        except Exception as e:
            raise RuntimeError(f"첨부 다운로드 실패 (id={att_id}): {e}") from e
        out.append((att, data))
    return out, owner


# ── 분류 + 날짜 탐지 ──────────────────────────────────────────────────────────
def _classify(paths: list[Path]) -> tuple[Path, Path]:
    """xlsx 경로들에서 (거래내역, 일일자금수지) 를 시트 구성으로 판별.

    거래내역: '통합 거래내역' 시트 보유. 일일자금수지: 'MM-DD' 시트 또는 '자금계획_원화' 보유.
    """
    tx_path: Path | None = None
    bal_path: Path | None = None
    for p in paths:
        try:
            wb = openpyxl.load_workbook(p, read_only=True)
            names = list(wb.sheetnames)
            wb.close()
        except Exception as e:
            raise RuntimeError(f"엑셀 파일을 열 수 없습니다 ({p.name}): {e}") from e
        has_tx = tx_reader.SHEET_NAME in names
        has_bal = any(_MMDD_RE.match(n) for n in names) or plan_reader.SHEET_NAME in names
        if has_tx and not has_bal:
            tx_path = p
        elif has_bal and not has_tx:
            bal_path = p
        elif has_tx and has_bal:
            # 두 성격을 다 가지면 거래내역 시트 유무로 거래내역 우선 배정(자금실적엔
            # '통합 거래내역' 이 없음). 남은 자리는 일일자금수지로.
            if tx_path is None:
                tx_path = p
            else:
                bal_path = p
    if tx_path is None or bal_path is None:
        raise ValueError(
            "거래내역 엑셀과 일일자금수지 엑셀을 모두 인식하지 못했습니다. "
            "은행 거래내역('통합 거래내역' 시트)과 일일자금수지(MM-DD 시트) 파일을 "
            "함께 첨부했는지 확인해 주세요."
        )
    return tx_path, bal_path


def _target_year(tx_p: Path) -> int | None:
    """거래내역 일자들의 최빈 연도(MM-DD 시트 → YYYY-MM-DD 보정용)."""
    try:
        lines = tx_reader.read_tx_lines(tx_p)
    except Exception:
        return None
    years = Counter(t.dt[:4] for t in lines if t.dt)
    if not years:
        return None
    return int(years.most_common(1)[0][0])


def _sheet_is_blank(ws) -> bool:
    """계좌 List 보통예금 행의 당일잔액(O)이 비어 있으면 '채워야 할' 시트로 본다."""
    blanks = 0
    has_prev = False
    for r in layout.COMMON_ROWS:
        if not ws.cell(r, layout.COL_ACCOUNT).value:
            continue
        if ws.cell(r, layout.COL_PREV).value not in (None, ""):
            has_prev = True
        if ws.cell(r, layout.COL_TODAY).value in (None, ""):
            blanks += 1
    return has_prev and blanks > 0


def _detect_date(bal_p: Path, tx_p: Path, requested: str | None) -> str:
    """대상 일자(YYYY-MM-DD) 결정: 명시값 > 당일잔액 공란 시트 > 단일 거래일.

    MM-DD 시트는 거래내역 최빈 연도로 YYYY-MM-DD 보정. 후보가 여럿이면 거래가 있는
    일자 우선, 그래도 여럿이면 최신.
    """
    wb = openpyxl.load_workbook(bal_p, data_only=True)
    mmdd_sheets = [n for n in wb.sheetnames if _MMDD_RE.match(n)]
    if not mmdd_sheets:
        wb.close()
        raise ValueError("일일자금수지 파일에 MM-DD 형식의 일자 시트가 없습니다.")

    year = _target_year(tx_p)
    tx_dates: set[str] = set()
    try:
        tx_dates = {t.dt[:10] for t in tx_reader.read_tx_lines(tx_p)}
    except Exception:
        pass

    def to_full(mmdd: str) -> str | None:
        if year is None:
            return None
        return f"{year:04d}-{mmdd}"

    # 1) 명시값 — 해당 MM-DD 시트가 있으면 채택.
    if requested:
        sheet = actual_reader.date_to_sheet(requested)
        if sheet in mmdd_sheets:
            wb.close()
            return requested
        wb.close()
        raise ValueError(
            f"지정한 일자 {requested} 의 시트('{sheet}')가 파일에 없습니다. "
            f"보유 일자 시트: {', '.join(mmdd_sheets)}"
        )

    # 2) 당일잔액이 공란인 시트(= 채워야 할 일자).
    blank_dates: list[str] = []
    for n in mmdd_sheets:
        if _sheet_is_blank(wb[n]):
            full = to_full(n)
            if full:
                blank_dates.append(full)
    wb.close()

    if blank_dates:
        with_tx = [d for d in blank_dates if d in tx_dates]
        pool = with_tx or blank_dates
        return sorted(pool)[-1]

    # 3) 공란 시트가 없으면 — 거래가 단일 일자면 그 일자, 아니면 거래 최신일.
    if len(tx_dates) == 1:
        only = next(iter(tx_dates))
        if actual_reader.date_to_sheet(only) in mmdd_sheets:
            return only
    candidates = [d for d in tx_dates if actual_reader.date_to_sheet(d) in mmdd_sheets]
    if candidates:
        return sorted(candidates)[-1]
    raise ValueError(
        "채울 일자를 자동으로 정하지 못했습니다. 메시지에 대상 일자(YYYY-MM-DD)를 "
        "함께 적어 주세요. (예: '5월 28일 일일자금수지 비교·검증')"
    )


# ── 요약 작성 (가공 점수 없이 사실만) ─────────────────────────────────────────
def _fmt(v: Any, fx: bool = False) -> str:
    if v in (None, ""):
        return "-"
    try:
        fv = float(v)
    except (TypeError, ValueError):
        return str(v)
    if fx:
        return ("-$" if fv < 0 else "$") + f"{abs(fv):,.2f}"
    return f"{fv:,.0f}"


def _build_summary(oc: run_daily.ReconcileOutcome, bal_p: Path) -> str:
    """검증 요약(간결판) — 사용자 핵심만: 자금계획 ↔ 자금실적(거래내역 기반 당일잔액)
    비교에서 누락·차이가 있는지, 원화·외화 총액이 각각 일치하는지, 그 결과.
    """
    results = oc.results
    krw = [b for b in results if not b.is_foreign]
    fx = [b for b in results if b.is_foreign]
    krw_delta = sum(b.delta for b in krw)
    krw_today = sum(b.today_balance for b in krw)
    fx_delta = sum(b.delta for b in fx)
    fx_today = sum(b.today_balance for b in fx)

    xch = [b for b in results if b.xcheck_ok is not None]
    xcheck_ok = sum(1 for b in xch if b.xcheck_ok)
    rfail = [b for b in results if b.repro_ok is False]

    pa = oc.pa_recon
    # 외화 계획 순증감(USD) — 일일자금수지 내장 자금계획_외화 시트.
    try:
        cur = plan_reader.read_blocks_both(bal_p, oc.date)
        usd_plan_net = cur.usd_net
        has_fx_plan = bool(cur.usd_income or cur.usd_expense)
    except Exception:
        usd_plan_net, has_fx_plan = 0.0, False

    lines: list[str] = []

    if not oc.db_skipped:
        # 자금일계표 공란 → 자금계획에서 채움. 핵심: 계획 순증감 ↔ 거래내역 당일증감(원화·외화 따로).
        # 같으면 누락·오류 없음, 다르면 계획 누락 또는 거래/당일잔액 오류.
        krw_plan_net = pa.plan_income_total - pa.plan_expense_total
        krw_ok = abs(krw_plan_net - krw_delta) < _EPS
        usd_ok = abs(usd_plan_net - fx_delta) < _EPS
        all_ok = krw_ok and usd_ok
        lines.append(
            f"{'✓' if all_ok else '⚠'} {oc.date} 자금계획 ↔ 자금실적(거래내역 기준) 비교·검증 — "
            f"{'누락·차이 없음, 원화·외화 총액 일치' if all_ok else '총액 불일치(점검 필요)'}"
        )
        lines.append("")
        lines.append(
            f"• 원화 총액: 계획 순증감 {_fmt(krw_plan_net)} {'=' if krw_ok else '≠'} "
            f"실적 당일증감 {_fmt(krw_delta)} → {'일치 ○' if krw_ok else '점검 필요 ✗'}"
            f"  (당일잔액 {_fmt(krw_today)})"
        )
        if has_fx_plan or abs(fx_delta) >= _EPS:
            lines.append(
                f"• 외화 총액(USD): 계획 순증감 {_fmt(usd_plan_net, True)} {'=' if usd_ok else '≠'} "
                f"실적 당일증감 {_fmt(fx_delta, True)} → {'일치 ○' if usd_ok else '점검 필요 ✗'}"
                f"  (당일잔액 {_fmt(fx_today, True)})"
            )
        else:
            lines.append(
                f"• 외화 총액(USD): 계획·거래 없음 → 변동 없음  (당일잔액 {_fmt(fx_today, True)})"
            )
        if not all_ok:
            lines.append(
                "  ※ 계획과 실제(거래내역)가 다릅니다 — 계획 누락 또는 당일잔액·거래 오류 가능. 사람 확정 필요."
            )
    else:
        # 자금일계표에 실적이 이미 있음 — 계획 ↔ 실적 성격별(수입/지출) 총액·누락 대조.
        inc_ok = pa.income_total_ok
        exp_ok = pa.expense_total_ok
        miss_i = pa.counts("income")["missing"]
        miss_e = pa.counts("expense")["missing"]
        f6 = oc.db_actuals.delta_today_fx or 0.0
        usd_ok = abs(f6 - usd_plan_net) < _EPS if has_fx_plan else True
        all_ok = inc_ok and exp_ok and usd_ok
        lines.append(
            f"{'✓' if all_ok else '⚠'} {oc.date} 자금계획 ↔ 자금실적 비교·검증 — "
            f"{'누락 없음, 총액 일치' if all_ok else '차이 있음(점검 필요)'}"
        )
        lines.append("")
        lines.append(
            f"• 원화 수입: 계획 {_fmt(pa.plan_income_total)} vs 실적 {_fmt(pa.actual_income_total)} → "
            + ("일치 ○" if inc_ok else f"차액 {_fmt(pa.actual_income_total - pa.plan_income_total)}")
            + (f"  (누락 {miss_i}건)" if miss_i else "")
        )
        lines.append(
            f"• 원화 지출: 계획 {_fmt(pa.plan_expense_total)} vs 실적 {_fmt(pa.actual_expense_total)} → "
            + ("일치 ○" if exp_ok else f"차액 {_fmt(pa.actual_expense_total - pa.plan_expense_total)}")
            + (f"  (누락 {miss_e}건)" if miss_e else "")
        )
        if has_fx_plan or abs(f6) >= _EPS:
            lines.append(
                f"• 외화 총액(USD): 계획 순증감 {_fmt(usd_plan_net, True)} vs "
                f"당일증감 {_fmt(f6, True)} → {'일치 ○' if usd_ok else '점검 필요 ✗'}"
            )
        else:
            lines.append("• 외화 총액(USD): 계획·거래 없음 → 변동 없음")

    if rfail:
        lines.append(
            f"  ※ {len(rfail)}계좌 전일잔액 불연속 가능(거래후잔액은 정답과 대조됨). 사람 확정 필요."
        )
    lines.append("")
    n_xch = len(xch)
    if n_xch == 0:
        footer = "당일 거래가 있는 계좌가 없어 교차검증 대상이 없습니다. 최종 확정은 검토 후 진행하세요."
    elif xcheck_ok == n_xch:
        footer = (
            f"당일 거래가 있던 {n_xch}개 계좌 모두, 계산한 당일잔액이 은행 거래내역의 "
            "최종 잔액과 일치(교차검증 통과). 최종 확정은 검토 후 진행하세요."
        )
    else:
        footer = (
            f"당일 거래가 있던 {n_xch}개 계좌 중 {xcheck_ok}개만 은행 거래내역의 최종 잔액과 "
            f"일치(나머지 {n_xch - xcheck_ok}개 점검 필요). 최종 확정은 검토 후 진행하세요."
        )
    lines.append(footer)
    return "\n".join(lines)


# ── 결과 owner ────────────────────────────────────────────────────────────────
async def _resolve_owner(db, fallback: int | None) -> int:
    if fallback is not None:
        return fallback
    res = await db.execute(
        select(User.id).where(User.role.in_([UserRole.superadmin, UserRole.domain_admin]))
        .order_by(User.id.asc()).limit(1)
    )
    admin_id = res.scalar_one_or_none()
    if admin_id is not None:
        return admin_id
    res = await db.execute(select(User.id).order_by(User.id.asc()).limit(1))
    any_id = res.scalar_one_or_none()
    if any_id is None:
        raise RuntimeError("저장 가능한 사용자가 없습니다 — users 테이블이 비어 있습니다.")
    return any_id


def _store_key_name(stem: str, suffix: str) -> str:
    """다운로드 파일명용 안전 stem(경로 구분자/제어문자 제거)."""
    safe = re.sub(r"[\\/:*?\"<>|\r\n\t]", "_", stem).strip() or "일일자금수지"
    return f"{safe}{suffix}"


async def _save_attachment(db, owner_id: int, path: Path, filename: str) -> Attachment:
    data = path.read_bytes()
    key = make_object_key(owner_id, filename)
    put_object(key=key, data=io.BytesIO(data), length=len(data), mime=_XLSX_MIME)
    att = Attachment(
        user_id=owner_id, bucket=_xlsx_bucket(), object_key=key,
        original_filename=filename, mime=_XLSX_MIME, size_bytes=len(data),
    )
    db.add(att)
    await db.flush()
    return att


def _xlsx_bucket() -> str:
    from ..config import settings
    return settings.minio_bucket


# ── 핵심 entry ────────────────────────────────────────────────────────────────
async def reconcile_fund_daily(
    *,
    attachment_ids: list[int],
    requested_date: str | None = None,
) -> FundReconcileResult:
    """첨부(거래내역+일일자금수지) → 채워진 일일자금수지 + 검증 워크북 + 요약."""
    async with SessionLocal() as db:
        files, owner = await _fetch_xlsx_attachments(db, attachment_ids)
        if len(files) < 2:
            raise ValueError(
                "엑셀 파일 2개(은행 거래내역, 일일자금수지)를 입력창에 함께 첨부한 뒤 "
                "다시 요청해 주세요."
            )

        tmp = Path(tempfile.mkdtemp(prefix="fund_recon_"))
        try:
            in_dir = tmp / "in"
            out_dir = tmp / "out"
            in_dir.mkdir()
            out_dir.mkdir()
            saved: list[Path] = []
            for i, (att, data) in enumerate(files):
                # 원본 파일명 보존(분류는 시트로 하지만 임시 파일/디버깅에 유용).
                stem = Path(att.original_filename or f"file{i}").stem
                p = in_dir / f"{i}_{stem}.xlsx"
                p.write_bytes(data)
                saved.append(p)

            tx_p, bal_p = _classify(saved)
            date = _detect_date(bal_p, tx_p, requested_date)

            # 자금계획은 일일자금수지 파일에 내장 → plan_p = bal_p.
            # 산출물 명명 — 첨부 파일명이 아니라 사내 자금실적 양식 규칙을 따른다:
            # `자금실적_FY{YY}_{YYMMDD}.xlsx` (예: 2026-05-28 → 자금실적_FY26_260528.xlsx).
            yymmdd = date.replace("-", "")[2:]            # 2026-05-28 → 260528
            label = _store_key_name(f"자금실적_FY{date[2:4]}_{yymmdd}", "")

            oc = run_daily.run_reconcile(
                date, bal_p, bal_p, tx_p, out_dir, label=label,
            )
            if not oc.originals_ok:
                raise RuntimeError("작업 중 입력 파일이 변경되어 중단했습니다.")

            summary = _build_summary(oc, bal_p)

            owner_id = await _resolve_owner(db, owner)
            # 다운로드명은 임시 복사본 경로명이 아니라 양식 규칙(label)을 그대로 사용.
            filled_name = f"{label}.xlsx"
            verify_name = f"{label}_검증.xlsx"
            filled_att = await _save_attachment(db, owner_id, oc.copy_path, filled_name)
            verify_att = await _save_attachment(db, owner_id, oc.verify_path, verify_name)
            await db.commit()
            await db.refresh(filled_att)
            await db.refresh(verify_att)

            return FundReconcileResult(
                ok=oc.ok,
                date=date,
                filename=filled_name,
                download_url=f"/api/attachments/{filled_att.id}/download",
                size_bytes=filled_att.size_bytes,
                attachment_id=filled_att.id,
                verify_filename=verify_name,
                verify_download_url=f"/api/attachments/{verify_att.id}/download",
                verify_attachment_id=verify_att.id,
                summary=summary,
                written=oc.written,
            )
        finally:
            shutil.rmtree(tmp, ignore_errors=True)


__all__ = ["FundReconcileResult", "reconcile_fund_daily", "PROGRESS_STEPS"]
