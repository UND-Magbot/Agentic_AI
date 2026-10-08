# -*- coding: utf-8 -*-
"""자금계획 자동작성 서비스 — 챗 첨부(자금계획 xlsx + 월마감 자료) → 작성된 자금계획.

흐름:
  1) 첨부 2개(자금계획 xlsx, 월마감 자료 xlsx) 바이트를 MinIO 에서 fetch → 임시 파일.
  2) 시트 구성으로 어느 쪽이 자금계획/월마감인지 자동 분류.
  3) plan_generator.generate_from_closing 으로:
     - 마감 기준일 이후(미래) 블록을 복사본에서 절단,
     - 과거 6개월 자동이체/자동출금 반복 지출 학습,
     - 마감 입금예정(→수입)·결제예정(→지출) 확정 라인을 날짜별 블록에 배치,
     - 다음 2개월(예: 6·7월) 영업일 블록을 외과적으로 생성(원화·외화).
  4) 작성된 복사본을 MinIO 저장 + Attachment row 생성.
  5) 다운로드 메타 + 사람이 읽을 작성 요약(가공 점수 없이 사실만) 반환.

원본 첨부는 수정하지 않는다(임시 복사본에만 기입). 모든 수치는 실제 마감/생성값.
"""
from __future__ import annotations

import io
import logging
import re
import shutil
import tempfile
from dataclasses import dataclass
from pathlib import Path

import openpyxl
from sqlalchemy import select

from ..database import SessionLocal
from ..models import Attachment, User, UserRole
from ..storage import get_object_stream, make_object_key, put_object
from . import closing_reader, plan_generator, plan_reader

logger = logging.getLogger("finance.plan_service")

_XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

# 진행 단계 ID — frontend ProgressCard 와 동일 식별자.
PROGRESS_STEPS: tuple[tuple[str, str], ...] = (
    ("fetch", "첨부 파일 준비"),
    ("classify", "자금계획·월마감 자료 인식"),
    ("learn", "과거 자동이체 학습 + 마감 확정 입금/결제 추출"),
    ("generate", "다음 2개월 자금계획 자동작성"),
    ("upload", "파일 저장"),
)


@dataclass
class FundPlanResult:
    """generate_fund_plan 호출 결과(챗 fast-path 가 요약/다운로드에 사용)."""

    ok: bool
    base_date: str
    filename: str
    download_url: str
    size_bytes: int
    attachment_id: int
    summary: str
    blocks: int


# ── 첨부 fetch ────────────────────────────────────────────────────────────────
async def _fetch_xlsx_attachments(db, attachment_ids: list[int]):
    """첨부 ID 들에서 xlsx 바이트를 순서 보존하며 가져온다. (목록, 첫 owner)."""
    if not attachment_ids:
        return [], None
    res = await db.execute(select(Attachment).where(Attachment.id.in_(attachment_ids)))
    by_id = {a.id: a for a in res.scalars().all()}

    out: list[tuple[Attachment, bytes]] = []
    owner: int | None = None
    for att_id in attachment_ids:
        att = by_id.get(att_id)
        if att is None:
            logger.warning("[plan] 첨부 id=%s 없음 — skip", att_id)
            continue
        name = (att.original_filename or "").lower()
        if not (name.endswith(".xlsx") or "spreadsheet" in (att.mime or "")):
            logger.warning("[plan] 첨부 id=%s (%s) xlsx 아님 — skip", att_id, att.mime)
            continue
        if owner is None:
            owner = att.user_id
        try:
            resp = get_object_stream(att.object_key)
            try:
                chunks: list[bytes] = []
                for chunk in resp.stream(amt=256 * 1024):
                    chunks.append(chunk)
                data = b"".join(chunks)
            finally:
                resp.close()
                resp.release_conn()
        except Exception as e:
            raise RuntimeError(f"첨부 다운로드 실패 (id={att_id}): {e}") from e
        out.append((att, data))
    return out, owner


# ── 분류 ──────────────────────────────────────────────────────────────────────
def _classify(paths: list[Path]) -> tuple[Path, Path]:
    """xlsx 경로들에서 (자금계획, 월마감) 을 시트 구성으로 판별.

    자금계획: '자금계획_원화' 시트 보유. 월마감: 입금/결제 예정 섹션 앵커('2-1]','3-1]')
    가 있는 시트 보유(보통 'HQ' 토큰 포함).
    """
    plan_path: Path | None = None
    closing_path: Path | None = None
    for p in paths:
        try:
            wb = openpyxl.load_workbook(p, read_only=True)
            names = list(wb.sheetnames)
            wb.close()
        except Exception as e:
            raise RuntimeError(f"엑셀 파일을 열 수 없습니다 ({p.name}): {e}") from e
        has_plan = plan_reader.SHEET_NAME in names
        has_closing = _looks_like_closing(p)
        if has_plan and not has_closing:
            plan_path = p
        elif has_closing and not has_plan:
            closing_path = p
        elif has_plan and has_closing:
            # 둘 다 닮았으면 자금계획(원화 시트 보유)을 우선, 남은 자리는 마감.
            if plan_path is None:
                plan_path = p
            else:
                closing_path = p
        else:
            # 어느 쪽도 확정 못한 경우 — 빈 자리에 임시 배정(뒤 검증에서 거른다).
            if plan_path is None:
                plan_path = p
            elif closing_path is None:
                closing_path = p
    if plan_path is None or closing_path is None:
        raise ValueError(
            "자금계획 엑셀('자금계획_원화' 시트)과 월마감 자료(입금/결제 예정 표) 파일을 "
            "모두 인식하지 못했습니다. 두 파일을 함께 첨부했는지 확인해 주세요."
        )
    return plan_path, closing_path


_CLOSING_SCAN_ROWS = 200  # 섹션 앵커는 항상 시트 상단 — 대용량 자금계획 시트 전체 스캔 방지.


def _looks_like_closing(path: Path) -> bool:
    """입금/결제 예정 섹션 앵커('2-1]'·'3-1]')가 시트 상단에 보이면 월마감 자료로 판별.

    자금계획 파일에는 100만 행짜리 시트가 있어 전체 스캔은 금물 — A열 상단 일부만 본다.
    """
    try:
        wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    except Exception:
        return False
    try:
        inc_pat = closing_reader._ANCHOR_INC_KRW
        pay_pat = closing_reader._ANCHOR_PAY_KRW
        for name in wb.sheetnames:
            ws = wb[name]
            inc = pay = False
            for i, row in enumerate(ws.iter_rows(min_col=1, max_col=1, values_only=True)):
                if i >= _CLOSING_SCAN_ROWS:
                    break
                a = row[0]
                if a is None:
                    continue
                s = str(a)
                if inc_pat.search(s):
                    inc = True
                elif pay_pat.search(s):
                    pay = True
                if inc and pay:
                    return True
        return False
    finally:
        wb.close()


# ── 요약(가공 점수 없이 사실만) ───────────────────────────────────────────────
def _fmt(v: float, fx: bool = False) -> str:
    unit = "$" if fx else "₩"
    if abs(v - round(v)) < 0.005:
        return f"{unit}{v:,.0f}"
    return f"{unit}{v:,.2f}"


def _build_summary(
    cd: closing_reader.ClosingData, result: plan_generator.PlanFromClosingResult,
) -> str:
    """작성 결과 사람이 읽을 요약 — 마감 반영 건수·대상월·검토 항목을 사실대로."""
    t = result.targets
    months = "·".join(f"{m}월" for _, m in t)
    krw_in = sum(l.amount for l in cd.krw_income)
    krw_out = sum(l.amount for l in cd.krw_expense)
    usd_in = sum(l.amount for l in cd.usd_income)
    usd_out = sum(l.amount for l in cd.usd_expense)
    blocks = sum(w.blocks_written for w in result.sheets)

    lines: list[str] = []
    lines.append(f"**자금계획 자동작성 완료** — {result.base_date} 마감 기준, {months} 계획표를 생성했습니다.")
    lines.append("")
    lines.append("**마감 확정 반영**")
    lines.append(
        f"- 원화: 입금예정 {len(cd.krw_income)}건({_fmt(krw_in)}) → 수입, "
        f"결제예정 {len(cd.krw_expense)}건({_fmt(krw_out)}) → 지출"
    )
    if cd.usd_income or cd.usd_expense:
        lines.append(
            f"- 외화: 입금예정 {len(cd.usd_income)}건({_fmt(usd_in, True)}) → 수입, "
            f"결제예정 {len(cd.usd_expense)}건({_fmt(usd_out, True)}) → 지출"
        )
    lines.append(f"- 미확정(금액 0) {cd.skipped_zero}건은 제외했습니다.")
    lines.append("")
    lines.append("**자동작성 범위**")
    for w in result.sheets:
        kind = "원화" if w.sheet == plan_reader.SHEET_NAME else "외화"
        lines.append(
            f"- {kind} 시트: {w.blocks_written}개 평일 블록(15줄 기준, 공휴일은 빈 블록), "
            f"지출 라인 {w.lines_placed}건 배치"
        )
    lines.append(
        "- 과거 6개월 자동이체/자동출금 반복 지출을 학습해 각 영업일에 드래프트로 배치했습니다 "
        "(고정 금액은 기입, 변동 금액은 공란 → 수기 확정)."
    )
    if cd.undated:
        lines.append("")
        lines.append(
            f"⚠ **검토 필요** — 결제일을 특정할 수 없는 마감 라인 {cd.undated}건은 "
            f"첫 대상월({t[0][1]}월) 첫 영업일에 '[날짜미상-검토]' 표시로 배치했습니다. "
            "실제 결제일을 확인해 이동해 주세요."
        )
    lines.append("")
    lines.append(
        "원본 자금계획은 수정하지 않았으며, 환율 외부연결·서식·테두리를 그대로 보존한 "
        "복사본에 미래 표를 새로 작성했습니다. Excel 에서 열면 잔액 수식이 재계산됩니다."
    )
    return "\n".join(lines)


# ── 저장 ──────────────────────────────────────────────────────────────────────
async def _resolve_owner(db, fallback: int | None) -> int:
    if fallback is not None:
        return fallback
    res = await db.execute(
        select(User.id).where(User.role.in_([UserRole.superadmin, UserRole.domain_admin]))
        .order_by(User.id.asc()).limit(1)
    )
    admin_id = res.scalar_one_or_none()
    if admin_id is not None:
        return admin_id
    res = await db.execute(select(User.id).order_by(User.id.asc()).limit(1))
    any_id = res.scalar_one_or_none()
    if any_id is None:
        raise RuntimeError("저장 가능한 사용자가 없습니다 — users 테이블이 비어 있습니다.")
    return any_id


def _store_key_name(stem: str, suffix: str) -> str:
    safe = re.sub(r"[\\/:*?\"<>|\r\n\t]", "_", stem).strip() or "자금계획"
    return f"{safe}{suffix}"


def _xlsx_bucket() -> str:
    from ..config import settings
    return settings.minio_bucket


async def _save_attachment(db, owner_id: int, path: Path, filename: str) -> Attachment:
    data = path.read_bytes()
    key = make_object_key(owner_id, filename)
    put_object(key=key, data=io.BytesIO(data), length=len(data), mime=_XLSX_MIME)
    att = Attachment(
        user_id=owner_id, bucket=_xlsx_bucket(), object_key=key,
        original_filename=filename, mime=_XLSX_MIME, size_bytes=len(data),
    )
    db.add(att)
    await db.flush()
    return att


# ── 핵심 entry ────────────────────────────────────────────────────────────────
async def generate_fund_plan(*, attachment_ids: list[int]) -> FundPlanResult:
    """첨부(자금계획 + 월마감) → 다음 2개월 자금계획이 작성된 복사본 + 요약."""
    async with SessionLocal() as db:
        files, owner = await _fetch_xlsx_attachments(db, attachment_ids)
        if len(files) < 2:
            raise ValueError(
                "엑셀 파일 2개(자금계획, 월마감 자료)를 입력창에 함께 첨부한 뒤 "
                "다시 요청해 주세요."
            )

        tmp = Path(tempfile.mkdtemp(prefix="fund_plan_"))
        try:
            in_dir = tmp / "in"
            out_dir = tmp / "out"
            in_dir.mkdir()
            out_dir.mkdir()
            saved: list[Path] = []
            for i, (att, data) in enumerate(files):
                stem = Path(att.original_filename or f"file{i}").stem
                p = in_dir / f"{i}_{stem}.xlsx"
                p.write_bytes(data)
                saved.append(p)

            plan_p, closing_p = _classify(saved)
            cd = closing_reader.read_closing(closing_p)

            yymmdd = cd.base_date.strftime("%y%m%d")
            label = _store_key_name(f"자금계획_FY{cd.base_date.strftime('%y')}_자동작성_{yymmdd}", "")
            out_path = out_dir / f"{label}.xlsx"

            result = plan_generator.generate_from_closing(plan_p, closing_p, out_path)
            summary = _build_summary(cd, result)
            blocks = sum(w.blocks_written for w in result.sheets)

            owner_id = await _resolve_owner(db, owner)
            filename = f"{label}.xlsx"
            att = await _save_attachment(db, owner_id, out_path, filename)
            await db.commit()
            await db.refresh(att)

            return FundPlanResult(
                ok=True,
                base_date=str(cd.base_date),
                filename=filename,
                download_url=f"/api/attachments/{att.id}/download",
                size_bytes=att.size_bytes,
                attachment_id=att.id,
                summary=summary,
                blocks=blocks,
            )
        finally:
            shutil.rmtree(tmp, ignore_errors=True)


__all__ = ["FundPlanResult", "generate_fund_plan", "PROGRESS_STEPS"]
