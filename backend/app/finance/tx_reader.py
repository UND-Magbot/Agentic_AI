# -*- coding: utf-8 -*-
"""은행 거래내역 → 정규화 거래 라인(TxLine). 당일잔액의 진짜 소스.

소스: `[클로브]…은행 거래내역…xlsx` 의 `통합 거래내역` 시트.
열: A 거래일시, B 은행, C 계좌번호, E 계좌별칭, F 입금액, G 출금액, H 거래후잔액,
    I 적요, N 통화, O 입금액(KRW), P 출금액(KRW), Q 거래후잔액(KRW).
"""
from __future__ import annotations

import re
from pathlib import Path

from .models import TxLine

SHEET_NAME = "통합 거래내역"

# 1-indexed 열 → iter_rows 는 0-indexed 튜플이므로 -1.
_C_DT, _C_BANK, _C_ACC, _C_ALIAS = 1, 2, 3, 5
_C_DEP_N, _C_WD_N, _C_BAL_N = 6, 7, 8
_C_SUMMARY, _C_CUR = 9, 14
_C_DEP_K, _C_WD_K, _C_BAL_K = 15, 16, 17


def account_key(raw: str) -> str:
    """계좌번호 정규화 키 = 숫자만. 파일별 하이픈 표기 차이 흡수."""
    return re.sub(r"\D", "", str(raw or ""))


def _num(v) -> float:
    if v is None or v == "":
        return 0.0
    if isinstance(v, (int, float)):
        return float(v)
    try:
        return float(re.sub(r"[^\d.\-]", "", str(v)))
    except ValueError:
        return 0.0


def _dt_str(v) -> str:
    """거래일시를 'YYYY-MM-DD HH:MM' 문자열로. 이미 문자열이면 그대로."""
    if hasattr(v, "strftime"):
        return v.strftime("%Y-%m-%d %H:%M")
    return str(v or "")


def read_tx_lines(path: Path) -> list[TxLine]:
    """거래내역 통합 시트의 모든 거래를 TxLine 리스트로 읽는다."""
    import openpyxl

    wb = openpyxl.load_workbook(Path(path), data_only=True, read_only=True)
    if SHEET_NAME not in wb.sheetnames:
        raise ValueError(f"'{SHEET_NAME}' 시트가 없습니다: {wb.sheetnames}")
    ws = wb[SHEET_NAME]
    lines: list[TxLine] = []
    for i, row in enumerate(ws.iter_rows(values_only=True)):
        if i == 0:
            continue  # 헤더
        if len(row) < _C_BAL_K:
            continue
        acc_raw = row[_C_ACC - 1]
        if not acc_raw:
            continue
        bal_n = row[_C_BAL_N - 1]
        bal_k = row[_C_BAL_K - 1]
        lines.append(TxLine(
            dt=_dt_str(row[_C_DT - 1]),
            bank=str(row[_C_BANK - 1] or ""),
            account_raw=str(acc_raw),
            account_key=account_key(acc_raw),
            alias=str(row[_C_ALIAS - 1] or ""),
            currency=str(row[_C_CUR - 1] or "KRW"),
            deposit_native=_num(row[_C_DEP_N - 1]),
            withdraw_native=_num(row[_C_WD_N - 1]),
            balance_native=(None if bal_n is None else _num(bal_n)),
            deposit_krw=_num(row[_C_DEP_K - 1]),
            withdraw_krw=_num(row[_C_WD_K - 1]),
            balance_krw=(None if bal_k is None else _num(bal_k)),
            summary=str(row[_C_SUMMARY - 1] or ""),
        ))
    wb.close()
    return lines


def filter_by_date(lines: list[TxLine], date: str) -> list[TxLine]:
    """대상 일자(YYYY-MM-DD)의 거래만. 거래일시의 날짜 부분으로 필터."""
    return [t for t in lines if t.dt[:10] == date]
