"""1단계 자료 읽기 — 공통 질문 답·고객 요청 원문·첨부 문서에서 P01~P48 값을 뽑는다.

가이드 "AI 가 입력을 읽는 순서"를 따른다:
  - 값마다 근거 문장을 원문에서 그대로 인용하게 하고, 인용이 원문에 없으면 버린다(날조 차단).
    여러 곳을 쉼표로 이어 붙인 근거는 조각마다 원문에 (조사 차이 정도만 두고) 있어야 통과한다.
  - 값에 적힌 숫자가 근거·원문 어디에도 없으면 버린다.
  - 사용자가 대화로 답한 문항(source=dialog)은 다시 읽어도 덮어쓰지 않는다.
  - 가이드 예시 답변·AI 반영 규칙은 다른 고객 사례(카세트·식기·세척 등) 표현이라 프롬프트에 넣지 않는다
    (실측: 넣으면 모델이 헷갈려 원문에 있는 답도 빠뜨리고, 엉뚱한 사례가 섞일 위험이 있다).

모델은 사내 gemma(proposal_llm) 만 쓴다 — 고객 자료를 외부로 보내지 않는다.
"""
from __future__ import annotations

import io
import json
import logging
import re
import zipfile
from dataclasses import dataclass, field
from difflib import SequenceMatcher
from typing import Any

from .. import proposal_llm
from .catalog import FORM_SECTIONS, QUESTIONS, form_question

_log = logging.getLogger("proposal_project.extract")

ASSET_CHARS = 6000        # 첨부 1개에서 읽는 최대 글자
SOURCE_CHARS = 12000      # 한 번 호출에 넣는 원문 합계 상한(24K 컨텍스트 안)
EVIDENCE_MAX = 200
_TEXT_ROLES = ("content", "dimension", "price")
_ALLOWED_STATUS = ("confirmed", "assumed", "conflict")
# 문항 묶음 — 폼 묶음 1개씩(12회 호출). 실측(12b): 3개씩이면 회차마다 8~15개로 흔들리고,
# 1개씩이면 같은 결과가 나오며 시간도 비슷하다(각 호출이 짧다).
_BATCH_SECTIONS = 1


@dataclass
class Source:
    kind: str          # form | request | attachment
    label: str         # 사람에게 보이는 출처 이름
    text: str


@dataclass
class Extracted:
    code: str
    value: str
    unit: str
    status: str
    source: str
    evidence: str


@dataclass
class ExtractResult:
    items: list[Extracted] = field(default_factory=list)
    dropped: list[dict[str, str]] = field(default_factory=list)   # 근거 불일치로 버린 값(로그·테스트용)
    readable: dict[int, bool] = field(default_factory=dict)       # project_assets.id → 읽기 성공


# ---------- 첨부 읽기 ----------

def read_document(data: bytes, filename: str, mime: str) -> str | None:
    """문서 첨부 → 텍스트. 읽을 수 없는 형식이면 None."""
    name = (filename or "").lower()
    try:
        if name.endswith(".pdf") or mime == "application/pdf":
            import pymupdf

            with pymupdf.open(stream=data, filetype="pdf") as doc:
                return "\n".join(page.get_text() for page in doc)
        if name.endswith(".docx"):
            import docx

            d = docx.Document(io.BytesIO(data))
            lines = [p.text for p in d.paragraphs if p.text.strip()]
            for t in d.tables:
                for r in t.rows:
                    lines.append(" | ".join(c.text.strip() for c in r.cells))
            return "\n".join(lines)
        if name.endswith(".pptx"):
            from pptx import Presentation

            prs = Presentation(io.BytesIO(data))
            return "\n".join(
                sh.text_frame.text for s in prs.slides for sh in s.shapes
                if sh.has_text_frame and sh.text_frame.text.strip()
            )
        if name.endswith((".txt", ".md", ".csv")) or (mime or "").startswith("text/"):
            from ..request_files import decode_text

            return decode_text(data)
    except (zipfile.BadZipFile, KeyError, ValueError, RuntimeError) as e:
        _log.warning("첨부 읽기 실패 %s: %s", filename, e)
        return None
    return None


# ---------- 원문 정리·검증 ----------

def normalize(s: str) -> str:
    """인용 대조용 — 공백·따옴표·말줄임 차이를 없앤다."""
    s = re.sub(r"[\"'“”‘’`]", "", s or "")
    s = s.replace("…", "").replace("...", "")
    return re.sub(r"\s+", "", s).lower()


def build_sources(project: dict[str, Any], asset_texts: dict[int, tuple[str, str]]) -> list[Source]:
    """폼 답(묶음별) → 요청 원문 → 첨부 순. 가이드: 사용자가 직접 쓴 답이 가장 강한 근거."""
    out: list[Source] = []
    intake = project.get("intake") or {}
    for s in FORM_SECTIONS:
        if intake.get(s.key):
            out.append(Source("form", f"공통 질문 - {s.title}", intake[s.key]))
    for code, value in intake.items():
        # 필수 질문 답 — 그 문항은 이미 저장됐고, 다른 문항의 단서(예: 대상물 답 속 무게)로 다시 읽는다.
        if code in QUESTIONS and value:
            out.append(Source("form", f"첫 화면 질문 {code} {form_question(code)}", value))
    if project.get("request_text"):
        out.append(Source("request", "고객 요청 원문", project["request_text"]))
    for _aid, (fname, txt) in asset_texts.items():
        if txt.strip():
            out.append(Source("attachment", f"첨부 - {fname}", txt[:ASSET_CHARS]))
    return out


def _source_block(sources: list[Source]) -> str:
    parts, used = [], 0
    for s in sources:
        chunk = s.text[: max(0, SOURCE_CHARS - used)]
        if not chunk:
            break
        parts.append(f"<출처 이름=\"{s.label}\">\n{chunk}\n</출처>")
        used += len(chunk)
    return "\n\n".join(parts)


_NUM_RE = re.compile(r"\d+(?:[.,]\d+)?")


def _numbers(s: str) -> set[str]:
    return {n.replace(",", "") for n in _NUM_RE.findall(s or "")}


PART_MATCH = 0.85
_PART_SPLIT = re.compile(r"[,;/·]|\s-\s|(?<=[.!?。])\s+")   # 쉼표·문장 끝에서 조각을 나눈다


def _part_found(part: str, text: str) -> bool:
    """조각이 원문에 (거의) 그대로 있는가 — 조사 한 글자 차이('1회에'↔'1회') 정도만 허용한다."""
    if part in text:
        return True
    m = SequenceMatcher(None, part, text, autojunk=False).find_longest_match(0, len(part), 0, len(text))
    if m.size < 2:
        return False
    lo = max(0, m.b - m.a - 2)
    window = text[lo: lo + len(part) + 4]
    blocks = SequenceMatcher(None, part, window, autojunk=False).get_matching_blocks()
    return sum(b.size for b in blocks) / len(part) >= PART_MATCH


def _locate(evidence: str, sources: list[Source]) -> Source | None:
    """근거 → 출처. 한 문장을 그대로 인용했으면 그 출처, 모델이 여러 곳을 쉼표로 이어 붙였으면
    조각마다 원문에서 찾아 모두 있을 때만 첫 조각의 출처를 돌려준다."""
    ev = normalize(evidence)
    if not ev:
        return None
    normed = [(s, normalize(s.text)) for s in sources]
    hit = next((s for s, t in normed if ev in t), None)
    if hit is not None:
        return hit
    parts = [normalize(p) for p in _PART_SPLIT.split(evidence)]
    parts = [p for p in parts if len(p) >= 2]
    if not parts:
        return None
    first = None
    for p in parts:
        src = next((s for s, t in normed if _part_found(p, t)), None)
        if src is None:
            return None
        first = first or src
    return first


def verify(raw: dict[str, Any], sources: list[Source], codes: set[str]) -> tuple[Extracted | None, str]:
    """모델이 낸 값 1개 검증 → (통과한 값, 버린 이유)."""
    code = str(raw.get("code", "")).strip().upper()
    value = str(raw.get("value", "")).strip()
    lines = [ln.strip() for ln in str(raw.get("evidence", "")).splitlines() if ln.strip()]
    evidence = (lines[0] if lines else "")[:EVIDENCE_MAX]
    status = str(raw.get("status", "confirmed")).strip()
    if code not in codes:
        return None, "알 수 없는 문항"
    if not value or not evidence:
        return None, "값 또는 근거 없음"
    if status not in _ALLOWED_STATUS:
        status = "assumed"
    hit = _locate(evidence, sources)
    if hit is None:
        return None, "근거 문장이 원문에 없음"
    all_nums = _numbers(" ".join(s.text for s in sources))
    if not _numbers(value) <= all_nums:
        return None, "값의 숫자가 원문에 없음"
    return Extracted(code, value[:2000], str(raw.get("unit", "")).strip()[:40], status,
                     hit.kind, f"[{hit.label}] {evidence}"), ""


# ---------- 모델 호출 ----------

_SYSTEM = """당신은 산업 자동화 기술영업 제안서의 요구사항을 정리하는 담당자입니다.
주어진 <출처> 들에서 각 문항의 답을 찾아 JSON 으로만 답합니다.

규칙:
1. 출처에 적힌 내용만 답합니다. 출처에 없는 내용을 추측하거나 일반 상식으로 채우지 않습니다.
2. 답을 찾은 문항만 items 에 넣습니다. 답이 없으면 그 문항은 빼십시오(빈 값 금지).
   문항을 하나씩 모두 확인하십시오. 한 문장이 여러 문항의 근거가 될 수 있습니다.
   "검토 중", "확인이 필요", "예정" 처럼 확정되지 않은 내용도 빼지 말고 status "assumed" 로 넣습니다.
3. evidence 에는 답의 근거가 된 출처 문장을 한 글자도 바꾸지 말고 그대로 복사합니다.
   출처의 한 줄 안에서만, 200자 이내로 복사하고 줄바꿈 문자를 넣지 않습니다.
4. status: 출처에 확정된 사실로 적혀 있으면 "confirmed", "검토 중·예정·가능하면·추정" 처럼
   확정되지 않은 표현이면 "assumed", 출처끼리 서로 다른 값을 말하면 "conflict"(value 에 양쪽을 함께 적음).
5. 숫자는 출처의 숫자와 단위를 그대로 씁니다. 계산하거나 환산하지 않습니다.
6. value 는 제안서에 옮겨 적을 수 있게 짧고 명확한 한국어 문장이나 구로 씁니다.

출력 형식: {"items": [{"code": "P08", "value": "...", "unit": "...", "status": "confirmed", "evidence": "..."}]}"""


# 문항별 "원문에서 찾을 것" — 질문이 추상적이면 모델이 원문과 잇지 못한다(실측: 'Class 100 클린룸' 을
# 환경 문항과, '협동로봇 검토 중' 을 로봇 대안 문항과 잇지 못함). 특정 고객 사례 표현은 쓰지 않는다.
_HINTS: dict[str, str] = {
    "P01": "고객사 이름, 요청 부서, 프로젝트 이름",
    "P02": "제안서로 결정할 일: 실증, 양산 투자, 예산 확보, 업체 선정 등",
    "P03": "자동화가 시작하는 지점과 끝나는 지점, 공급 범위의 앞뒤 경계",
    "P04": "공간·비용·확장성·속도 중 우선순위",
    "P05": "지금 사람이 하는 작업 순서, 인원, 소요 시간, 병목",
    "P06": "다루는 제품·대상물의 이름과 종류 수",
    "P07": "대상물의 크기, 용량, 무게, 재질, 상태(젖음·뜨거움 등)",
    "P08": "처리 수량, 택트타임, 시간당·일당 목표, 운영 시간·교대",
    "P09": "대상물이 들어오는 방식(컨베이어, 트레이, 사람 투입 등)이나 앞 설비",
    "P10": "잡기 전 정렬, 자세 바꾸기, 기울어짐·방향 확인",
    "P11": "로봇이 다루지 않는 물품과 그 처리 경로",
    "P12": "처리가 끝난 대상물을 어디에 어떻게 두거나 내보내는지(트레이, 박스, 컨베이어, 적재)",
    "P13": "검토 중이거나 지정된 로봇 형태·제품(협동로봇, 산업용 다관절, 스카라, 갠트리, 브랜드·가반하중), 비교를 원하는 형태 수",
    "P14": "첫 번째(또는 유일한) 로봇 후보의 형태·브랜드·모델·가반하중·구성",
    "P15": "두 번째 로봇 후보의 형태·브랜드·모델·구성",
    "P16": "세 번째 로봇 후보의 형태·브랜드·모델·구성",
    "P17": "공정 앞단 카메라·비전이 판단할 것(종류 구분, 불량, 위치)",
    "P18": "잡기 직전 카메라가 확인할 것(위치, 자세, 기울어짐)",
    "P19": "앞 단계 인식 결과를 뒤 단계 대상물과 연결하는 방법(추적, 순번, 바코드)",
    "P20": "카메라·비전 제품이나 모델",
    "P21": "그리퍼 종류·제품(흡착, 전동 그리퍼, 전용 지그 등)",
    "P22": "팔·툴마다 그리퍼를 같게 또는 다르게 쓰는지",
    "P23": "툴체인저·툴스탠드 필요 여부",
    "P24": "그리퍼 선정에 필요한 자료(치수, 무게, 접촉 가능 면)",
    "P25": "종류별 보관 위치·목적지 구분",
    "P26": "적재 높이·단수·용량 관리",
    "P27": "가득 찬 용기·빈 용기 교체 담당과 방법",
    "P28": "턴테이블·버퍼 같은 보조 장치",
    "P29": "로봇이 놓친 대상물 처리",
    "P30": "뒤 공정이 멈출 때 앞 공정 처리, 버퍼",
    "P31": "정지 후 확인·재개 절차와 담당",
    "P32": "미인식·잡기 실패·가득 참 같은 예외 처리",
    "P33": "설치 공간 크기(폭·길이·높이), 도면 유무",
    "P34": "기존 공간·설비를 바꿀 수 있는 정도",
    "P35": "설치 환경 조건: 클린룸 등급, 습기, 먼지, 온도, 방폭, 방수 등급",
    "P36": "기구 재질, 배수, 세척·청소 조건",
    "P37": "작업자가 맡을 일, 사람 동선, 접근 위치, 안전",
    "P38": "연동할 설비와 주고받을 신호(PLC, MES 등)",
    "P39": "첨부 파일마다 반영할 부분",
    "P40": "자료와 대화가 다를 때 우선 기준",
    "P41": "최소 실증(PoC)을 시작할 범위·모듈·시기",
    "P42": "다음 단계로 넘어갈 판단 기준(실증 결과, 성공 조건)",
    "P43": "견적에 포함할 범위와 제외 항목(기존 설비 재사용 포함), 금액 표기 방식",
    "P44": "ROI·인력 절감 효과 계산 필요 여부",
    "P45": "제안서 페이지 수·필수 페이지",
    "P46": "컨셉도 시점·상세 수준",
    "P47": "확인·승인 순서",
    "P48": "최종 파일 형식과 수정 관리",
}


def _question_block(codes: list[str]) -> str:
    return "\n".join(f"- {c}: {QUESTIONS[c].question}\n  (원문에서 찾을 것: {_HINTS[c]})" for c in codes)


def _batches() -> list[list[str]]:
    out = []
    for i in range(0, len(FORM_SECTIONS), _BATCH_SECTIONS):
        codes = [c for s in FORM_SECTIONS[i:i + _BATCH_SECTIONS] for c in s.codes]
        out.append(codes)
    return out


def _parse(raw: str) -> list[dict[str, Any]]:
    """모델 응답 → 항목 목록. 생성 한도에 걸려 JSON 이 중간에 잘려도 완성된 항목은 건진다(실측:
    근거를 복사하다 줄바꿈을 반복하며 한도를 다 쓰는 경우가 있다)."""
    try:
        data = json.loads(raw)
        items = data.get("items") if isinstance(data, dict) else data
        return [x for x in (items or []) if isinstance(x, dict)]
    except (json.JSONDecodeError, AttributeError):
        pass
    dec = json.JSONDecoder()
    out, i = [], 0
    raw = raw or ""
    while (i := raw.find("{", i)) != -1:
        try:
            obj, end = dec.raw_decode(raw, i)
        except json.JSONDecodeError:
            i += 1
            continue
        if isinstance(obj, dict) and "code" in obj:
            out.append(obj)
        elif isinstance(obj, dict) and isinstance(obj.get("items"), list):
            return [x for x in obj["items"] if isinstance(x, dict)]
        i = end if isinstance(obj, dict) and "code" in obj else i + 1
    return out


async def extract(project: dict[str, Any], asset_texts: dict[int, tuple[str, str]], *,
                  chat=proposal_llm.chat) -> ExtractResult:
    """원문 → 검증된 문항 값 목록. chat 은 테스트에서 바꿔 끼운다."""
    sources = build_sources(project, asset_texts)
    result = ExtractResult()
    if not sources:
        return result
    block = _source_block(sources)
    seen: set[str] = set()
    for codes in _batches():
        user = f"{block}\n\n다음 문항의 답을 출처에서 찾으십시오:\n{_question_block(codes)}"
        raw = await chat([{"role": "system", "content": _SYSTEM}, {"role": "user", "content": user}],
                         fmt="json", num_predict=3000, temperature=0.1)
        for r in _parse(raw):
            item, why = verify(r, sources, set(codes))
            if item is None:
                result.dropped.append({"code": str(r.get("code", "")), "value": str(r.get("value", ""))[:200],
                                       "evidence": str(r.get("evidence", ""))[:200], "reason": why})
            elif item.code not in seen:
                seen.add(item.code)
                result.items.append(item)
    return result
