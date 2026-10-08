"""회사 경험 카드 — 과거 사례·제안서를 '학습된 경험' 으로 정리한다(사용자·의뢰인 요구 2026-09-30).

의뢰인 요구: 저장했다가 비슷한 걸 찾아 붙이는 백과사전이 아니라, 회사가 겪은 경험을 알고 있다가
"예전에 이런 공정을 이렇게 처리했고 효과가 좋았습니다. 이렇게 해 보시겠습니까?" 가 먼저 나오게 할 것.

1단계(이 모듈): 원자료를 한 건씩 읽어 경험 카드로 정리한다.
  카드 = 무슨 공정이었나 / 무엇이 문제였나 / 어떻게 풀었나 / 다른 공정에도 쓸 수 있는 핵심 아이디어 /
        효과(원문에 있는 수치만) / 주의점 / 이 경험이 맞는 조건.
  - evidence: '실적' = 실제 수행 결과(컨설팅 사례집), '제안' = 제안서의 기대 효과(검증 전).
    "효과가 좋았다" 는 실적 카드에만 쓸 수 있다 — 제안 카드는 "이렇게 제안한 적이 있다" 까지만.
  - 원자료(회사 제안서 원문 포함)는 사내 gemma 로만 읽는다(외부 전송 없음). 효과 수치는 원문에 있는 것만 남긴다.
2단계(다음): 카드를 공정 분야별 노하우 요약으로 묶고, 공정 컨셉 제안 때 통째로 읽혀 '경험 기반 제안' 을 먼저 낸다.

실행(카드 만들기·다시 만들기):
    docker compose run --rm --no-deps -v ./docs:/docs backend python -m app.proposal_project.experience build
"""
from __future__ import annotations

import asyncio
import json
import logging
import re
import sys
from dataclasses import dataclass
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

from sqlalchemy import text

from .. import proposal_llm
from ..database import SessionLocal
from .concept import _clean, _json, _nums

_log = logging.getLogger("proposal_project.experience")

DOCS = Path("/docs") if Path("/docs").is_dir() else Path(__file__).resolve().parents[3] / "docs"
# 회사 제안서(docs 폴더) — 검토용 초안·확장본도 우리 회사가 실제로 만든 제안이다.
COMPANY_PROPOSALS = (
    "(주)유엔디_GS칼텍스_주유소_4족로봇_자율운영_제안서_ma20260624.pptx",
    "Hanwha_Ocean_Engine_Room_Robot_PoC_Proposal_UND.pdf",
    "KEC_프로젝트_제안서_확장본.pptx",
    "넥시스_드릴링_체결기_자동화_제안서_유엔디_최종본 (1).pptx",
)
GUIDE = "UND_영업AI_제안서작성_질문답변_가이드.docx"
SOURCE_CHARS = 9000


@dataclass
class Source:
    key: str              # 고유 키 — 다시 만들 때 같은 카드를 갱신
    type: str             # casebook | company_proposal | guide
    title: str
    evidence: str         # 실적 | 제안
    text: str


# ---------- 원자료 읽기 ----------

def _pptx_text(path: Path) -> str:
    from pptx import Presentation

    out = []
    for i, s in enumerate(Presentation(str(path)).slides, 1):
        parts = []
        for sh in s.shapes:
            if sh.has_text_frame and sh.text_frame.text.strip():
                parts.append(sh.text_frame.text.strip())
            if getattr(sh, "has_table", False) and sh.has_table:
                parts.append(" | ".join(c.text.strip() for r in sh.table.rows for c in r.cells))
        out.append(f"[{i}쪽] " + " / ".join(parts))
    return "\n".join(out)


def _pdf_text(path: Path) -> str:
    import pymupdf

    with pymupdf.open(str(path)) as d:
        return "\n".join(f"[{i}쪽] {p.get_text()}" for i, p in enumerate(d, 1))


def company_sources() -> list[Source]:
    out = []
    for name in COMPANY_PROPOSALS:
        path = DOCS / "examples" / "proposals" / name
        if not path.exists():
            _log.warning("회사 제안서 없음: %s", path)
            continue
        body = _pdf_text(path) if path.suffix.lower() == ".pdf" else _pptx_text(path)
        out.append(Source(f"proposal:{name}", "company_proposal", path.stem, "제안", body))
    return out


def guide_source() -> list[Source]:
    """가이드의 삼성웰스토리 사례 — P 문항 예시 답변이 곧 채택된 컨셉(검증 전이라 '제안')."""
    notes = json.loads((Path(__file__).parent / "guide_items.json").read_text(encoding="utf-8"))
    body = "\n".join(f"{n['code']} {n['question']}: {n['example']}" for n in notes if n["code"].startswith("P"))
    return [Source("guide:samsung_welstory", "guide", "삼성웰스토리 식기세척 후단 자동화 실증 제안(가이드 사례)",
                   "제안", body)]


async def casebook_sources() -> list[Source]:
    from ..ingest_casebook import SOURCE_PATHS

    async with SessionLocal() as db:
        rows = (await db.execute(text(
            "SELECT source_path, content, metadata FROM documents WHERE source_path = ANY(:p) "
            "AND metadata ? 'case_no' ORDER BY source_path, (metadata->>'case_no')::int, id"),
            {"p": list(SOURCE_PATHS)})).mappings().all()
    cases: dict[tuple[str, int], dict[str, Any]] = {}
    for r in rows:
        m = r["metadata"]
        c = cases.setdefault((r["source_path"], int(m["case_no"])), {"meta": m, "parts": []})
        c["parts"].append(r["content"])
    out = []
    for (_path, no), c in cases.items():
        m = c["meta"]
        kpis = "; ".join(f"{k.get('label', '')} {k.get('value', '')}{k.get('unit', '')} {k.get('direction', '')}".strip()
                         for k in m.get("kpis") or [])
        body = f"업종: {m.get('industry', '')}\n" + "\n".join(c["parts"]) + (f"\nKPI: {kpis}" if kpis else "")
        out.append(Source(f"casebook:{m['year']}-{no:02d}", "casebook",
                          f"{m.get('doc', '컨설팅 사례집')} #{no:02d} {m.get('title', '')} ({m.get('company', '')})",
                          "실적", body))
    return out


# ---------- 카드 만들기 ----------

_CARD_SYSTEM = """너는 로봇 자동화 회사의 수석 엔지니어다. 아래 [자료]는 우리 회사(또는 컨설팅 사례)가 겪은 자동화 프로젝트 하나다.
다음 프로젝트에 경험으로 꺼내 쓸 수 있도록 핵심만 정리해 JSON 으로만 답한다.

규칙:
1. 자료에 적힌 내용만 쓴다. 자료에 없는 사실·수치·효과를 지어내지 않는다.
2. key_ideas 는 다른 고객·다른 공정에도 옮겨 쓸 수 있는 설계 아이디어(1~4개, 각 한 문장).
   예: "기존 설비는 그대로 두고 I/O 인터페이스만 추가해 비용·기간을 줄인다"
3. effect_numbers 는 자료에 수치로 적힌 효과만(개선 전·후, 절감률, 시간 등). 없으면 빈 목록.
4. cautions 는 이 프로젝트에서 드러난 위험·제약·조건(0~3개).
5. applies_when 은 이 경험을 다시 쓸 만한 조건(대상물·공정·환경)을 한 문장으로.
6. process 는 공정을 짧은 명사구로(예: "웨이퍼 카세트 이송", "식기 분류·적재", "설비 순찰 점검").

{"title": "한 줄 요약", "industry": "업종", "process": "...", "problem": "기존 문제 1~2문장",
 "solution": "적용(또는 제안)한 해법 1~3문장 — 로봇 형태·주변 설비·방법", "key_ideas": ["..."],
 "effect": "효과 한두 문장", "effect_numbers": [{"metric": "...", "value": "..."}],
 "cautions": ["..."], "applies_when": "..."}"""


def validate_card(raw: dict[str, Any], source: Source) -> tuple[dict[str, Any], list[str]]:
    """카드 검증 — 원문에 없는 효과 수치는 버린다. 반환 (카드, 버린 사유)."""
    src_nums = _nums(source.text)
    dropped = []
    effects = []
    for e in raw.get("effect_numbers") or []:
        if not isinstance(e, dict):
            continue
        metric, value = _clean(e.get("metric"), 60), _clean(e.get("value"), 60)
        if not metric or not value:
            continue
        if _nums(value) - src_nums:
            dropped.append(f"원문에 없는 수치: {metric} {value}")
            continue
        effects.append({"metric": metric, "value": value})
    effect = _clean(raw.get("effect"), 300)
    if _nums(effect) - src_nums:
        dropped.append(f"효과 문장에 원문에 없는 수치 → 숫자 빼고 보관: {effect[:60]}")
        effect = re.sub(r"\d+(?:[.,]\d+)?\s*%?", "", effect).strip()
    card = {
        "title": _clean(raw.get("title"), 80) or source.title[:80],
        "industry": _clean(raw.get("industry"), 60),
        "process": _clean(raw.get("process"), 40),
        "problem": _clean(raw.get("problem"), 400),
        "solution": _clean(raw.get("solution"), 500),
        "key_ideas": [_clean(x, 200) for x in (raw.get("key_ideas") or []) if _clean(x, 200)][:4],
        "effect": effect,
        "effect_numbers": effects[:6],
        "cautions": [_clean(x, 200) for x in (raw.get("cautions") or []) if _clean(x, 200)][:3],
        "applies_when": _clean(raw.get("applies_when"), 200),
    }
    return card, dropped


async def make_card(source: Source, *, chat=None) -> tuple[dict[str, Any], list[str]]:
    chat = chat or proposal_llm.chat
    kind = "실제 수행 결과(컨설팅 사례)" if source.evidence == "실적" else "우리 회사가 고객에게 낸 제안서(검증 전 기대 효과)"
    user = f"[자료 성격] {kind}\n[자료 제목] {source.title}\n[자료]\n{source.text[:SOURCE_CHARS]}"
    raw = await chat([{"role": "system", "content": _CARD_SYSTEM}, {"role": "user", "content": user}],
                     fmt="json", num_predict=2500, temperature=0.2)
    return validate_card(_json(raw), source)


async def save_card(source: Source, card: dict[str, Any]) -> None:
    async with SessionLocal() as db:
        await db.execute(text(
            "INSERT INTO experience_cards (source_key, source_type, source_title, evidence, process, card) "
            "VALUES (:k, :t, :ti, :e, :p, CAST(:c AS jsonb)) "
            "ON CONFLICT (source_key) DO UPDATE SET source_title = EXCLUDED.source_title, evidence = EXCLUDED.evidence, "
            "process = EXCLUDED.process, card = EXCLUDED.card, updated_at = NOW()"),
            {"k": source.key, "t": source.type, "ti": source.title, "e": source.evidence,
             "p": card.get("process", ""), "c": json.dumps(card, ensure_ascii=False)})
        await db.commit()


async def list_cards() -> list[dict[str, Any]]:
    async with SessionLocal() as db:
        rows = (await db.execute(text(
            "SELECT id, source_key, source_type, source_title, evidence, process, card FROM experience_cards "
            "WHERE active ORDER BY source_type, id"))).mappings().all()
    return [dict(r) for r in rows]


def card_markdown(c: dict[str, Any]) -> str:
    k = c["card"]
    lines = [f"### [{c['evidence']}] {k['title']}",
             f"- 출처: {c['source_title']}",
             f"- 업종·공정: {k['industry']} / {k['process']}",
             f"- 문제: {k['problem']}",
             f"- 해법: {k['solution']}"]
    lines += [f"- 핵심 아이디어: {x}" for x in k["key_ideas"]]
    lines.append(f"- 효과: {k['effect'] or '(원문에 효과 기술 없음)'}")
    if k["effect_numbers"]:
        lines.append("- 효과 수치: " + "; ".join(f"{e['metric']} {e['value']}" for e in k["effect_numbers"]))
    lines += [f"- 주의점: {x}" for x in k["cautions"]]
    lines.append(f"- 다시 쓸 조건: {k['applies_when']}")
    return "\n".join(lines)


# ---------- 회상: 공정 컨셉 때 회사 경험을 먼저 떠올린다 ----------
# 5c 의 회상 엔진(company_knowledge.recall)이 준비되면 그걸 쓰고, 그 전에는 아래 임시 경로를 쓴다.
# 임시 경로도 "비슷한 것만 골라 붙이기" 가 아니라, 전체 카드 한 줄 목록을 통째로 읽힌 뒤 모델이 떠올리게 한다.

HITS_MAX = 3
SKIP_HIDE = 2          # 빼기가 적용보다 이만큼 많아진 지식은 더 떠올리지 않는다


async def card_scores() -> dict[int, int]:
    """카드별 (적용 수 − 빼기 수)."""
    async with SessionLocal() as db:
        rows = (await db.execute(text(
            "SELECT card_id, SUM(CASE WHEN action = 'apply' THEN 1 WHEN action = 'skip' THEN -1 ELSE 0 END) "
            "FROM experience_feedback WHERE card_table = 'experience_cards' GROUP BY card_id"))).all()
    return {int(r[0]): int(r[1]) for r in rows}


def memory_index(cards: list[dict[str, Any]], scores: dict[int, int]) -> str:
    """전체 카드 한 줄 목록 — 적용받은 지식이 앞에. 모델이 '회사가 아는 것' 전체를 한 번에 본다."""
    live = [c for c in cards if scores.get(c["id"], 0) > -SKIP_HIDE]
    live.sort(key=lambda c: (-scores.get(c["id"], 0), c["evidence"] != "실적", c["id"]))
    lines = []
    for c in live:
        k = c["card"]
        nums = "; ".join(f"{e['metric']} {e['value']}" for e in k.get("effect_numbers", [])[:3])
        mark = f" (사용자 적용 {scores[c['id']]}회)" if scores.get(c["id"], 0) > 0 else ""
        lines.append(f"E{c['id']} [{c['evidence']}]{mark} {k.get('process', '')} — {k.get('title', '')} / 핵심: "
                     + "; ".join(k.get("key_ideas", [])[:2]) + (f" / 효과: {nums}" if nums else ""))
    return "\n".join(lines)


_RECALL_SYSTEM = """너는 로봇 자동화 회사의 수석 엔지니어다. 지금 [이번 프로젝트]의 공정 컨셉을 만드는 중이다.
[회사 경험 목록]은 우리 회사가 알고 있는 경험 전체다. 목록 전체를 훑어보고, 이번 프로젝트에 도움이 될 경험을
최대 3개 떠올려 JSON 으로만 답한다. 맞는 경험이 없으면 빈 목록이 정답이다(억지로 고르지 않는다).

규칙:
1. id 는 목록의 E번호 그대로.
2. suggestion 은 사용자에게 건넬 한두 문장: "예전에 ○○ 공정에서 △△ 하였습니다(효과). 이번에도 □□ 해 보시겠습니까?"
   [실적] 경험만 "효과가 있었다/좋았다" 고 쓸 수 있다. [제안] 경험은 "이렇게 제안한 적이 있다" 까지만 쓴다.
3. 수치는 목록에 적힌 수치만 쓴다. 지어내지 않는다.
4. why 는 이번 프로젝트의 어떤 조건 때문에 떠올렸는지 한 문장. fit 은 "맞음" 또는 "주의",
   caution 은 이번 조건과 다른 점·주의점 한 문장(없으면 빈 값).
{"hits": [{"id": "E12", "suggestion": "...", "why": "...", "fit": "맞음", "caution": ""}]}"""


def validate_hits(raw: dict[str, Any], cards: dict[int, dict[str, Any]], project_text: str) -> list[dict[str, Any]]:
    """모델이 떠올린 경험 검증 — 없는 카드·근거 없는 수치·제안 카드의 효과 단정은 걸러 낸다."""
    out = []
    seen = set()
    for h in raw.get("hits") or []:
        if not isinstance(h, dict):
            continue
        m = re.fullmatch(r"E?(\d+)", str(h.get("id", "")).strip())
        cid = int(m.group(1)) if m else -1
        c = cards.get(cid)
        if c is None or cid in seen:
            continue
        k = c["card"]
        suggestion = _clean(h.get("suggestion"), 400)
        allowed = _nums(json.dumps(k, ensure_ascii=False)) | _nums(project_text)
        if not suggestion or _nums(suggestion) - allowed:
            continue
        if c["evidence"] != "실적" and re.search(r"효과가\s*(좋|있)", suggestion):
            suggestion = re.sub(r"효과가\s*(좋았|있었)[^.。]*", "이렇게 제안한 적이 있", suggestion)
        seen.add(cid)
        out.append({
            "card_id": cid, "card_table": "experience_cards", "evidence": c["evidence"],
            "title": k.get("title", ""), "source": c["source_title"], "process": k.get("process", ""),
            "effect_numbers": k.get("effect_numbers", []), "suggestion": suggestion,
            "why": _clean(h.get("why"), 200), "fit": "주의" if str(h.get("fit")).strip() == "주의" else "맞음",
            "caution": _clean(h.get("caution"), 200), "status": "suggested",
        })
        if len(out) >= HITS_MAX:
            break
    return out


async def recall(situation: str, *, brain: str = "gpt", user_id: int | None = None, chat=None) -> list[dict[str, Any]]:
    """이번 프로젝트 상황 → 떠올린 회사 경험(최대 3). situation 에는 고객사명을 넣지 않는다(호출부 책임)."""
    try:
        from ..company_knowledge import recall as ck          # 5c 회상 엔진(상황 판단·연상·번짐·되짚기·근거 확인)
    except ImportError:
        ck = None
    if ck is not None and hasattr(ck, "recall"):
        res = await ck.recall(situation, brain=brain, user_id=user_id, chat=chat)
        return await _from_engine(res)
    from .brain import ask_json

    cards = {c["id"]: c for c in await list_cards()}
    if not cards:
        return []
    index = memory_index(list(cards.values()), await card_scores())
    data = f"[이번 프로젝트]\n{situation}\n\n[회사 경험 목록]\n{index}"
    res = await ask_json(_RECALL_SYSTEM, data, brain=brain, user_id=user_id, chat=chat, num_predict=2000)
    return validate_hits(_json(res.raw), cards, situation)


DETAIL_KEYS = ("process", "problem", "solution", "how_it_works", "applies_when", "industry", "summary")


def hit_detail(card: dict[str, Any], *, slides: list[int] | None = None, source_doc: str = "") -> dict[str, Any]:
    """떠올린 경험이 '정확히 무슨 공정이었나' — 반영할지 판단할 때 보여 준다(사용자 요청 2026-09-30:
    제목만으로는 무슨 공정인지 몰라 반영 여부가 헷갈림)."""
    d = {k: _clean(card.get(k), 300) for k in DETAIL_KEYS if _clean(card.get(k), 300)}
    if slides:
        d["pages"] = [int(x) for x in slides if str(x).isdigit()][:4]
    if source_doc:
        d["source_doc"] = source_doc
    return d


async def knowledge_details(card_ids: list[int]) -> dict[int, dict[str, Any]]:
    """회사 소개서 카드(knowledge_cards) 상세."""
    if not card_ids:
        return {}
    async with SessionLocal() as db:
        rows = (await db.execute(text("SELECT id, card, slides, source_doc FROM knowledge_cards WHERE id = ANY(:ids)"),
                                 {"ids": list(card_ids)})).mappings().all()
    return {r["id"]: hit_detail(r["card"] or {}, slides=list(r["slides"] or []), source_doc=r["source_doc"] or "")
            for r in rows}


async def attach_details(hits: list[dict[str, Any]]) -> list[dict[str, Any]]:
    """상세가 없는 경험(예전 프로젝트)에 상세를 붙인다 — 화면 표시용, 저장하지 않는다."""
    need_k = [h["card_id"] for h in hits if not h.get("detail") and h.get("card_table") == "knowledge_cards"]
    need_e = [h["card_id"] for h in hits if not h.get("detail") and h.get("card_table") != "knowledge_cards"]
    if not need_k and not need_e:
        return hits
    kd = await knowledge_details(need_k)
    mine = {c["id"]: c for c in await list_cards()} if need_e else {}
    for h in hits:
        if h.get("detail"):
            continue
        if h.get("card_table") == "knowledge_cards":
            h["detail"] = kd.get(h["card_id"], {})
        elif h["card_id"] in mine:
            h["detail"] = hit_detail(mine[h["card_id"]].get("card") or {})
        if not h.get("process") and h.get("detail", {}).get("process"):
            h["process"] = h["detail"]["process"]
    return hits


async def _from_engine(res: dict[str, Any]) -> list[dict[str, Any]]:
    """5c 회상 결과 → 작업 화면의 '경험 기반 제안' 형식. '안 맞음' 판정은 제안하지 않는다."""
    mine = {c["id"]: c for c in await list_cards()}
    kd = await knowledge_details([h["card_id"] for h in res.get("hits") or []
                                  if h.get("card_table") == "knowledge_cards" and h.get("card_id")])
    out = []
    for h in res.get("hits") or []:
        if h.get("fit") == "안 맞음" or not h.get("why"):
            continue
        own = mine.get(h.get("card_id")) if h.get("card_table") == "experience_cards" else None
        k = (own or {}).get("card", {})
        out.append({
            "card_id": h["card_id"], "card_table": h.get("card_table", "experience_cards"),
            "evidence": h.get("evidence", ""), "title": h.get("title", ""),
            "source": (own or {}).get("source_title", "회사 소개서"),
            "process": k.get("process", "") or kd.get(h.get("card_id"), {}).get("process", ""),
            "detail": hit_detail(k) if own else kd.get(h.get("card_id"), {}),
            "effect_numbers": k.get("effect_numbers", []), "facts": h.get("facts", []),
            "suggestion": h["why"], "why": "", "fit": h.get("fit", "주의"), "caution": h.get("caution", ""),
            "status": "suggested",
        })
        if len(out) >= HITS_MAX:
            break
    return out


async def record_feedback(card_id: int, action: str, *, project_id: int | None, user_id: int | None,
                          card_table: str = "experience_cards") -> None:
    if action not in ("apply", "skip"):
        raise ValueError(action)
    async with SessionLocal() as db:
        await db.execute(text("INSERT INTO experience_feedback (card_table, card_id, project_id, user_id, action) "
                              "VALUES (:t, :c, :p, :u, :a)"),
                         {"t": card_table, "c": card_id, "p": project_id, "u": user_id, "a": action})
        await db.commit()


def project_card(project: dict[str, Any], alt: dict[str, Any], applied: list[dict[str, Any]]) -> dict[str, Any]:
    """사용자가 '회사 지식으로 저장' 한 컨셉 → 경험 카드(근거 수준 '제안' — 실증 전)."""
    items = project["items"]

    def v(code: str) -> str:
        it = items.get(code) or {}
        return it.get("value", "") if it.get("status") in ("confirmed", "adopted", "assumed") else ""

    ideas = [h["suggestion"] for h in applied][:2] + alt["structure"][:4 - min(2, len(applied))]
    return {
        "title": f"{alt['name']} — {v('P06') or '공정 컨셉'}"[:80],
        "industry": "", "process": (v("P06") or "")[:40],
        "problem": (v("P05") or v("P02"))[:400],
        "solution": (" / ".join(x for x in (alt.get("robot"), alt.get("summary")) if x) + " — "
                     + " ".join(alt["structure"][:3]))[:500],
        "key_ideas": [_clean(x, 200) for x in ideas][:4],
        "effect": "", "effect_numbers": [],
        "cautions": [_clean(x, 200) for x in alt.get("exclude", [])][:3],
        "applies_when": _clean(" / ".join(x for x in (v("P06"), v("P07"), v("P35")) if x), 200),
    }


def is_knowledge_approver(user: Any) -> bool:
    """회사 지식 승인권자 — 영업 도메인 관리자(사용자 결정 2026-09-30: '일단 sales_admin 만', superadmin 제외).
    계정 이름이 아니라 역할·도메인으로 본다 — 영업 관리자가 늘어도 그대로 동작한다."""
    role, domain = getattr(user, "role", None), getattr(user, "domain", None)
    return getattr(role, "value", role) == "domain_admin" and getattr(domain, "value", domain) == "sales"


async def save_project_card(project: dict[str, Any], alt: dict[str, Any], applied: list[dict[str, Any]], *,
                            user_id: int | None, approver: bool) -> tuple[int, str]:
    """완성된 컨셉 → 경험 카드. 승인권자가 저장하면 바로 회상에 쓰이고, 그 외는 '승인 대기'(active=FALSE).
    다시 저장(내용 수정)하면 승인권자가 아닌 한 다시 승인 대기로 돌린다 — 승인된 내용이 몰래 바뀌지 않게.
    반환: (카드 id, 상태 approved|pending)."""
    card = project_card(project, alt, applied)
    status = "approved" if approver else "pending"
    async with SessionLocal() as db:
        card_id = int((await db.execute(text(
            "INSERT INTO experience_cards (source_key, source_type, source_title, evidence, process, card, active, "
            "review_status, submitted_by, reviewed_by, reviewed_at, review_note) "
            "VALUES (:k, 'project', :ti, '제안', :p, CAST(:c AS jsonb), :a, :s, :u, :rb, :ra, '') "
            "ON CONFLICT (source_key) DO UPDATE SET source_title = EXCLUDED.source_title, process = EXCLUDED.process, "
            "card = EXCLUDED.card, active = EXCLUDED.active, review_status = EXCLUDED.review_status, "
            "submitted_by = EXCLUDED.submitted_by, reviewed_by = EXCLUDED.reviewed_by, "
            "reviewed_at = EXCLUDED.reviewed_at, review_note = '', updated_at = NOW() RETURNING id"),
            {"k": f"project:{project['id']}:{alt['id']}", "ti": f"제안서 작업 #{project['id']} {alt['name']}",
             "p": card.get("process", ""), "c": json.dumps(card, ensure_ascii=False), "a": approver, "s": status,
             "u": user_id, "rb": user_id if approver else None,
             "ra": datetime.now(timezone.utc) if approver else None})).scalar_one())
        await db.commit()
    if approver:
        await refresh_cues(card_id, card)
    return card_id, status


async def card_reviews(card_ids: list[int]) -> dict[int, str]:
    """카드 id → 승인 상태(approved|pending|rejected). 작업 화면이 저장한 컨셉의 현재 상태를 보여 줄 때."""
    if not card_ids:
        return {}
    async with SessionLocal() as db:
        rows = (await db.execute(text("SELECT id, review_status FROM experience_cards WHERE id = ANY(:i)"),
                                 {"i": card_ids})).all()
    return {int(r[0]): r[1] for r in rows}


async def pending_reviews() -> list[dict[str, Any]]:
    """승인 대기 지식 — 카드 내용·제출자·원본 프로젝트."""
    async with SessionLocal() as db:
        rows = (await db.execute(text(
            "SELECT e.id, e.source_key, e.source_title, e.evidence, e.card, e.updated_at, "
            "u.username AS submitted_by FROM experience_cards e LEFT JOIN users u ON u.id = e.submitted_by "
            "WHERE e.review_status = 'pending' ORDER BY e.updated_at"))).mappings().all()
    out = []
    for r in rows:
        m = re.match(r"project:(\d+):", r["source_key"])
        out.append({"id": r["id"], "title": r["source_title"], "evidence": r["evidence"], "card": r["card"],
                    "submitted_by": r["submitted_by"] or "", "submitted_at": r["updated_at"].isoformat(),
                    "project_id": int(m.group(1)) if m else None})
    return out


class ReviewError(Exception):
    pass


async def review_card(card_id: int, action: str, *, reviewer_id: int, note: str = "") -> str:
    """승인(approve) → active + 회상 단서 생성, 반려(reject) → 비활성 + 사유. 승인 대기 카드만."""
    if action not in ("approve", "reject"):
        raise ReviewError("알 수 없는 판정입니다.")
    note = _clean(note, 300)
    if action == "reject" and not note:
        raise ReviewError("반려 사유를 적어 주세요.")
    status = "approved" if action == "approve" else "rejected"
    async with SessionLocal() as db:
        row = (await db.execute(text(
            "UPDATE experience_cards SET review_status = :s, active = :a, reviewed_by = :r, reviewed_at = NOW(), "
            "review_note = :n WHERE id = :i AND review_status = 'pending' RETURNING card"),
            {"s": status, "a": action == "approve", "r": reviewer_id, "n": note, "i": card_id})).first()
        if row is None:
            raise ReviewError("승인 대기 중인 지식이 아닙니다(이미 판정했거나 없음).")
        await db.commit()
    if action == "approve":
        await refresh_cues(card_id, row[0])
    return status


async def refresh_cues(card_id: int, card: dict[str, Any]) -> int:
    """카드 한 장의 회상 단서만 다시 만든다(5c 회상 엔진의 knowledge_cues). 새로 저장한 지식이 바로 연상되게.
    실패해도 카드는 남는다 — 전체 목록(되짚기)으로는 떠오르고, 다음 전체 재생성 때 단서가 채워진다."""
    try:
        from ..company_knowledge.cards import refresh_experience_cues
    except ImportError:
        return 0
    try:
        return await refresh_experience_cues(card_id, card)
    except Exception as e:  # noqa: BLE001 — 단서 실패가 지식 저장을 막지 않게
        _log.warning("회상 단서 생성 실패(E%s): %r", card_id, e)
        return 0


async def build(only: str = "") -> None:
    sources = company_sources() + guide_source() + await casebook_sources()
    if only:
        sources = [s for s in sources if only in s.key]
    print(f"원자료 {len(sources)}건 → 경험 카드", flush=True)
    for n, s in enumerate(sources, 1):
        try:
            card, dropped = await make_card(s)
        except Exception as e:  # noqa: BLE001 — 한 건 실패가 전체를 막지 않게
            print(f"  {n:02d} 실패 {s.key}: {e}", flush=True)
            continue
        await save_card(s, card)
        print(f"  {n:02d} [{s.evidence}] {card['process'] or '?'} — {card['title'][:50]}"
              + (f" (버림 {len(dropped)})" if dropped else ""), flush=True)
    cards = await list_cards()
    out = DOCS / "experience_cards"
    out.mkdir(parents=True, exist_ok=True)
    md = ["# 회사 경험 카드", "",
          f"총 {len(cards)}장 — 실적 {sum(1 for c in cards if c['evidence'] == '실적')}장(컨설팅 사례집 수행 결과), "
          f"제안 {sum(1 for c in cards if c['evidence'] == '제안')}장(회사 제안서·가이드 사례, 검증 전).", ""]
    for title, t in (("회사 제안서", "company_proposal"), ("가이드 사례", "guide"), ("컨설팅 사례집", "casebook")):
        group = [c for c in cards if c["source_type"] == t]
        if group:
            md += [f"## {title} ({len(group)}장)", ""] + [card_markdown(c) + "\n" for c in group]
    (out / "cards.md").write_text("\n".join(md), encoding="utf-8")
    print(f"저장: {out / 'cards.md'}", flush=True)


if __name__ == "__main__":
    if len(sys.argv) > 1 and sys.argv[1] == "build":
        asyncio.run(build(sys.argv[2] if len(sys.argv) > 2 else ""))
    else:
        print(__doc__)
