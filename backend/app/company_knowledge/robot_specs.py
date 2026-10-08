"""로봇 스펙 DB — 타사 6축 협동로봇 비교표(docs/sources/6축 로봇 비교표_Ma260929.pptx, 기술영업팀) → robot_specs 표.

회사 지식 카드(knowledge_cards)와 따로 둔다: 그쪽은 '회사가 가진 제품·겪은 경험'을 회상하는 곳이라 타사 로봇
스펙을 섞으면 회사 제품처럼 제안될 수 있다. 여기는 정형 수치(가반하중·도달거리·반복정밀도·본체 중량·IP 등급)라
조건으로 후보를 고르는 데 쓴다(candidates).

- 표의 "확인 필요" 칸은 값을 비우고 unverified 에 칸 이름을 남긴다(추측으로 채우지 않는다).
- "해당 비교급 없음" 행은 모델이 아니므로 넣지 않는다.
- 같은 (제조사, 모델)은 다시 적재하면 갱신한다. 비교표가 바뀌면 ingest 를 다시 돌리면 된다.

적재:
    docker compose run --rm --no-deps -v ./docs:/docs backend python -m app.company_knowledge.robot_specs ingest
    (정리 문서: docs/company_knowledge/robot_specs.md 를 함께 다시 쓴다)
"""
from __future__ import annotations

import asyncio
import json
import logging
import re
import sys
from pathlib import Path
from typing import Any

from sqlalchemy import text

_log = logging.getLogger("company_knowledge.robot_specs")

DOCS = Path("/docs") if Path("/docs").is_dir() else Path(__file__).resolve().parents[3] / "docs"
DECK_NAME = "6축 로봇 비교표_Ma260929.pptx"
HEADER = ["제조사", "모델", "타입", "가반하중", "도달거리", "반복정밀도", "본체 중량", "IP 등급"]
FIELDS = {"가반하중": "payload_kg", "도달거리": "reach_mm", "반복정밀도": "repeatability_mm", "본체 중량": "weight_kg"}
UNVERIFIED = "확인 필요"
_NUM_RE = re.compile(r"(\d[\d,]*(?:\.\d+)?)")
_CLASS_RE = re.compile(r"(\d+\s*kg\s*이상급|\d+\s*~\s*\d+\s*kg급|\d+\s*kg급)", re.I)
_IP_RE = re.compile(r"IP\s*(\d{2})", re.I)


def _num(cell: str) -> float | None:
    m = _NUM_RE.search(cell or "")
    return float(m.group(1).replace(",", "")) if m else None


def parse_rows(table: list[list[str]], *, payload_class: str = "", slide: int | None = None) -> list[dict[str, Any]]:
    """비교표 한 장(첫 행 머리) → 모델 행. 머리가 HEADER 와 다르면 빈 목록(요약표 등)."""
    if not table or [c.strip() for c in table[0]] != HEADER:
        return []
    out = []
    for cells in table[1:]:
        cells = [c.strip() for c in cells] + [""] * (len(HEADER) - len(cells))
        maker, model, kind = cells[0], cells[1], cells[2]
        if not maker or not model or "없음" in model or _num(cells[3]) is None:
            continue
        row: dict[str, Any] = {"maker": maker, "model": model, "robot_type": kind or "협동로봇", "axes": 6,
                               "payload_class": payload_class, "source_slide": slide, "unverified": []}
        for label, key in FIELDS.items():
            cell = cells[HEADER.index(label)]
            row[key] = None if UNVERIFIED in cell else _num(cell)
            if UNVERIFIED in cell:
                row["unverified"].append(label)
        ip = cells[HEADER.index("IP 등급")]
        m = _IP_RE.search(ip)
        row["ip_rating"] = f"IP{m.group(1)}" if m and UNVERIFIED not in ip else None
        if UNVERIFIED in ip:
            row["unverified"].append("IP 등급")
        out.append(row)
    return out


def parse_deck(path: Path | None = None) -> list[dict[str, Any]]:
    """비교표 PPT → 모델 행 목록(가반하중급 = 그 쪽 제목)."""
    from pptx import Presentation

    path = path or DOCS / "sources" / DECK_NAME
    prs = Presentation(str(path))
    rows: list[dict[str, Any]] = []
    for no, slide in enumerate(prs.slides, 1):
        title = " ".join(sh.text_frame.text for sh in slide.shapes if sh.has_text_frame)
        m = _CLASS_RE.search(title)
        cls = re.sub(r"\s+", "", m.group(1)).replace("KG", "kg") if m else ""     # 쪽 머리 "01  3KG급"
        for sh in slide.shapes:
            if sh.has_table:
                table = [[c.text for c in r.cells] for r in sh.table.rows]
                rows += parse_rows(table, payload_class=cls, slide=no)
    return rows


def _fmt(v: Any, unit: str = "", pm: str = "") -> str:
    """1300 → '1,300mm', 927.7 → '927.7mm', 0.02 → '±0.02mm', None → '확인 필요'."""
    if v is None:
        return UNVERIFIED
    v = float(v)
    return f"{pm}{int(v):,}{unit}" if v == int(v) else f"{pm}{v:g}{unit}"


def spec_line(r: dict[str, Any]) -> str:
    """사람·모델이 읽는 한 줄: 'DOBOT CR10A — 가반 10kg · 도달 1,300mm · 반복 ±0.03mm · 중량 40kg · IP54'."""
    return (f"{r['maker']} {r['model']} — 가반 {_fmt(r.get('payload_kg'), 'kg')} · 도달 {_fmt(r.get('reach_mm'), 'mm')}"
            f" · 반복 {_fmt(r.get('repeatability_mm'), 'mm', '±')} · 중량 {_fmt(r.get('weight_kg'), 'kg')}"
            f" · {r.get('ip_rating') or 'IP ' + UNVERIFIED}")


def ip_level(rating: str | None) -> int | None:
    """IP54 → 54 (두 자리 비교용). 첫째 자리(분진)·둘째 자리(물)를 따로 보지 않는 단순 비교."""
    m = _IP_RE.search(rating or "")
    return int(m.group(1)) if m else None


REACH_SLACK = 0.10      # 도달거리 추정이 조금 크게 잡혀 맞는 급이 통째로 빠지지 않게(표시하고 후보에 둠)


def rank(rows: list[dict[str, Any]], payload_kg: float, *, reach_mm: float | None = None,
         ip_min: int | None = None, makers: list[str] | None = None, exclude_makers: list[str] | None = None,
         mobile: bool = False, limit: int = 8, need_full: bool = True) -> list[dict[str, Any]]:
    """조건에 맞는 모델을 '과하지 않은 순'으로. 점수 = 가반하중 여유 + 도달거리 여유(모자라면 벌점) + 본체 무게.
    - 도달거리가 REACH_SLACK 안에서 모자란 모델은 reach_short 표시로 후보에 둔다(AI 추정 도달거리가 커서
      5kg급이 통째로 빠지고 무거운 6kg급만 남았다 — 실측 2026-09-30).
    - mobile(AMR 위 탑재)이면 본체 무게를 더 무겁게 본다 — 팔 무게가 곧 AMR 적재 부담.
    - 결과에 조건을 다 만족하는 모델이 하나는 들어가게 한다.
    확인 필요로 값이 빈 조건은 만족으로 치지 않는다. 가반하중·IP 는 모자라면 뺀다."""
    excluded = {m.strip().lower() for m in exclude_makers or [] if m.strip()}
    scored = []
    for r in rows:
        if r.get("payload_kg") is None or float(r["payload_kg"]) < payload_kg:
            continue
        short = False
        if reach_mm is not None:
            if r.get("reach_mm") is None or float(r["reach_mm"]) < reach_mm * (1 - REACH_SLACK):
                continue
            short = float(r["reach_mm"]) < reach_mm
        if ip_min is not None and (ip_level(r.get("ip_rating")) or 0) < ip_min:
            continue
        if makers and r["maker"] not in makers:
            continue
        if any(x in r["maker"].lower() or x in r["model"].lower() for x in excluded):
            continue
        reach_ratio = float(r["reach_mm"]) / reach_mm if reach_mm else 1.0
        reach_term = 1 + (1 - reach_ratio) * 3 if short else reach_ratio
        weight = float(r["weight_kg"]) if r.get("weight_kg") is not None else 60.0
        score = float(r["payload_kg"]) / payload_kg + reach_term + weight / (20.0 if mobile else 40.0)
        scored.append((score, r["maker"], r["model"], {**r, "reach_short": short}))
    scored.sort(key=lambda t: t[:3])
    out = [t[3] for t in scored[:limit]]
    if need_full and out and all(r["reach_short"] for r in out):
        full = next((t[3] for t in scored if not t[3]["reach_short"]), None)
        if full:
            out = out[:limit - 1] + [full]
    return out


# 회사 협동로봇 우선순위 — 가이드 K03(2026-09-28): 가격 중심은 DOBOT 우선, 국내 후보로 한화·레인보우,
# 레인보우는 국내 신뢰성을 중시할 때의 후보. 한화의 세부 선정 기준은 미정 — 임의로 순위를 올리지 않는다.
# 프로젝트에서 제조사를 지정하면(예: DOBOT 양팔) 이 일반 순위보다 지정 구성을 우선한다(policy_pick 의 makers).
POLICY: dict[str, list[tuple[str, str]]] = {
    "price": [("DOBOT", "가격 중심 우선 후보(K03)"), ("레인보우로보틱스", "국내 후보 — 국내 신뢰성 중시 시"),
              ("한화로보틱스", "국내 후보(세부 선정 기준 미정)")],
    "domestic": [("레인보우로보틱스", "국내 신뢰성 중시 우선 후보(K03)"), ("한화로보틱스", "국내 후보(세부 선정 기준 미정)"),
                 ("DOBOT", "가격 중심 후보")],
}
FILL_NOTE = "보충 후보 — 우선 제조사에 맞는 모델이 없음"


def policy_pick(rows: list[dict[str, Any]], payload_kg: float, *, priority: str = "price",
                makers: list[str] | None = None, limit: int = 3, **kw: Any) -> list[dict[str, Any]]:
    """회사 우선순위(K03)로 후보 고르기: 우선 제조사마다 가장 알맞은 모델 1개씩, 모자라면 다른 제조사로 보충.
    프로젝트가 제조사를 지정했으면(makers) 그 제조사 모델만 알맞은 순으로."""
    if makers:
        return [{**r, "policy_note": "프로젝트 지정 제조사"} for r in
                rank(rows, payload_kg, makers=makers, limit=limit, **kw)]
    out: list[dict[str, Any]] = []
    for maker, note in POLICY.get(priority, POLICY["price"]):
        best = rank(rows, payload_kg, makers=[maker], limit=1, need_full=False, **kw)   # 제조사별 가장 알맞은 것
        if best:
            out.append({**best[0], "policy_note": note})
    if out and all(r["reach_short"] for r in out):
        # 모두 도달거리가 조금 모자라면, 우선 제조사 중 조건을 다 만족하는 가장 알맞은 모델 하나를 끝에 둔다.
        full = next((r for maker, _n in POLICY.get(priority, POLICY["price"])
                     for r in rank(rows, payload_kg, makers=[maker], limit=1, **kw) if not r["reach_short"]), None)
        if full and full["model"] not in {r["model"] for r in out}:
            note = dict(POLICY.get(priority, POLICY["price"])).get(full["maker"], "")
            out = out[:limit - 1] + [{**full, "policy_note": f"{note} · 도달거리 조건 충족"}]
    if len(out) < limit:
        seen = {r["model"] for r in out}
        out += [{**r, "policy_note": FILL_NOTE} for r in rank(rows, payload_kg, limit=limit + len(out), **kw)
                if r["model"] not in seen][:limit - len(out)]
    return out[:limit]


# ── DB ─────────────────────────────────────────────────────────────────────

_COLS = ("maker", "model", "robot_type", "axes", "payload_kg", "reach_mm", "repeatability_mm", "weight_kg",
         "ip_rating", "payload_class", "source_slide")


async def upsert(rows: list[dict[str, Any]], *, source_doc: str = DECK_NAME) -> dict[str, int]:
    """(제조사, 모델) 기준 넣기/갱신. 반환 {inserted, updated}."""
    from ..database import SessionLocal

    ins = upd = 0
    async with SessionLocal() as db:
        for r in rows:
            params = {k: r.get(k) for k in _COLS} | {"unverified": json.dumps(r.get("unverified") or [],
                                                                               ensure_ascii=False),
                                                     "src": source_doc}
            new = (await db.execute(text(
                "INSERT INTO robot_specs (maker, model, robot_type, axes, payload_kg, reach_mm, repeatability_mm, "
                "weight_kg, ip_rating, payload_class, source_slide, unverified, source_doc) VALUES (:maker, :model, "
                ":robot_type, :axes, :payload_kg, :reach_mm, :repeatability_mm, :weight_kg, :ip_rating, :payload_class, "
                ":source_slide, CAST(:unverified AS jsonb), :src) "
                "ON CONFLICT (maker, model) DO UPDATE SET robot_type=EXCLUDED.robot_type, axes=EXCLUDED.axes, "
                "payload_kg=EXCLUDED.payload_kg, reach_mm=EXCLUDED.reach_mm, repeatability_mm=EXCLUDED.repeatability_mm, "
                "weight_kg=EXCLUDED.weight_kg, ip_rating=EXCLUDED.ip_rating, payload_class=EXCLUDED.payload_class, "
                "source_slide=EXCLUDED.source_slide, unverified=EXCLUDED.unverified, source_doc=EXCLUDED.source_doc, "
                "active=TRUE, updated_at=NOW() RETURNING (xmax = 0)"), params)).scalar_one()
            ins, upd = ins + bool(new), upd + (not new)
        await db.commit()
    return {"inserted": ins, "updated": upd}


async def list_specs(*, active_only: bool = True) -> list[dict[str, Any]]:
    from ..database import SessionLocal

    async with SessionLocal() as db:
        rows = (await db.execute(text(
            "SELECT * FROM robot_specs" + (" WHERE active" if active_only else "") +
            " ORDER BY payload_kg, maker, model"))).mappings().all()
    out = []
    for r in rows:
        d = dict(r)
        for k in ("payload_kg", "reach_mm", "repeatability_mm", "weight_kg"):
            d[k] = float(d[k]) if d[k] is not None else None
        out.append(d)
    return out


async def candidates(payload_kg: float, *, reach_mm: float | None = None, ip_min: int | None = None,
                     makers: list[str] | None = None, exclude_makers: list[str] | None = None, mobile: bool = False,
                     limit: int = 8, priority: str | None = None) -> list[dict[str, Any]]:
    """조건(가반하중·도달거리·IP)에 맞는 협동로봇 후보 — 공정 컨셉·장비 선정 때 쓴다.
    priority 를 주면 회사 우선순위(K03)로 제조사별 대표 후보를 고른다(policy_pick)."""
    kw = {"reach_mm": reach_mm, "ip_min": ip_min, "exclude_makers": exclude_makers, "mobile": mobile}
    rows = await list_specs()
    if priority:
        return policy_pick(rows, payload_kg, priority=priority, makers=makers, limit=limit, **kw)
    return rank(rows, payload_kg, makers=makers, limit=limit, **kw)


def markdown(rows: list[dict[str, Any]]) -> str:
    """사람용 정리(가반하중급별 표 + 확인 필요 목록)."""
    lines = ["# 6축 협동로봇 스펙 정리 (robot_specs)", "",
             f"원본: docs/sources/{DECK_NAME} (기술영업팀) · 모델 {len(rows)}개 · "
             f"제조사 {len({r['maker'] for r in rows})}곳. DB 표 robot_specs 와 같은 내용입니다.", ""]
    classes: dict[str, list[dict[str, Any]]] = {}
    for r in rows:
        classes.setdefault(r.get("payload_class") or "기타", []).append(r)
    for cls, part in classes.items():
        lines += [f"## {cls} ({len(part)}개)", "", "| 제조사 | 모델 | 가반하중 | 도달거리 | 반복정밀도 | 본체 중량 | IP |",
                  "|---|---|---|---|---|---|---|"]
        for r in part:
            lines.append(f"| {r['maker']} | {r['model']} | {_fmt(r['payload_kg'], ' kg')} | {_fmt(r['reach_mm'], ' mm')} | "
                         f"{_fmt(r['repeatability_mm'], ' mm', '±')} | {_fmt(r['weight_kg'], ' kg')} | "
                         f"{r.get('ip_rating') or '확인 필요'} |")
        lines.append("")
    todo = [f"- {r['maker']} {r['model']}: {', '.join(r['unverified'])}" for r in rows if r.get("unverified")]
    if todo:
        lines += ["## 확인 필요 (원본 표에 값 없음 — 비워 둠)", ""] + todo + [""]
    return "\n".join(lines)


async def ingest(path: Path | None = None) -> dict[str, Any]:
    rows = parse_deck(path)
    res = await upsert(rows)
    doc = DOCS / "company_knowledge" / "robot_specs.md"
    try:
        doc.parent.mkdir(parents=True, exist_ok=True)
        doc.write_text(markdown(await list_specs()), encoding="utf-8")
    except OSError as e:          # docs 를 마운트하지 않고 돌린 경우 — DB 적재는 끝났다
        _log.warning("정리 문서를 쓰지 못함: %r", e)
    return {"parsed": len(rows), **res, "unverified": sum(1 for r in rows if r["unverified"])}


if __name__ == "__main__":
    logging.basicConfig(level=logging.INFO)
    cmd = sys.argv[1] if len(sys.argv) > 1 else ""
    if cmd == "ingest":
        print(asyncio.run(ingest(Path(sys.argv[2]) if len(sys.argv) > 2 else None)), flush=True)
    elif cmd == "list":
        for r in asyncio.run(list_specs()):
            print(spec_line(r))
    else:
        print(__doc__)
