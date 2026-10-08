"""맥봇 제품 공식 사양 → 회사 제품 DB(지식 카드 kind=product) 다시 세팅(사용자 요청 2026-10-01).

원본: docs/sources/magbot_spec_imgver_Ma260930.pptx — 소개서에서 자동 추출한 카드보다 정확한 자료.
  - 1~14쪽 요약표: ATC 6 · MG 10 · D-Air 5 · SMG 3 · 기타 그리퍼 3 = 27모델 + ATC 액세서리 8종.
    모델마다 사진 · 9개 사양(크기·가반하중·무게·IP·온습도·전기·IoT·실린더 센서) · 적용 조건(값이 성립하는 전제).
  - 15~24쪽: M-LTC 공압 툴체인저 17종 사양서(모델마다 사진 1장).
  - 25~39쪽: M-LTC 구성 모듈 표 — 넣지 않는다(사용자: 모듈 사진까지는 필요 없음).
추가 원본: docs/sources/MGP,MGM.pptx(1쪽) + 회사 소개서 13·14쪽 — MGP·MGM 시리즈 모델 카드(panel_models).
사진: ATC 사진은 Master TC(로봇 측)와 Tool Plate(툴 측)가 함께 찍힌 한 세트 — 나누지 않고 한 장으로 저장한다.

카드는 모델 단위(card_key 'magbot:<모델>'), 다시 적재하면 갱신. 이 자료가 다루는 예전 소개서 카드(SUPERSEDED_KEYS)는
비활성화한다 — 추천·제안서·회상 어디에도 안 쓰이고, 소개서를 다시 적재해도 살아나지 않는다(cards.save 가 확인).
'확인 필요'·'자료 없음' 은 원본 그대로 두고, 확인 필요 항목은 limits 에 모은다(추측으로 채우지 않는다).

    docker compose run --rm --no-deps -v ./docs:/docs backend python -m app.company_knowledge.magbot_catalog ingest
"""
from __future__ import annotations

import asyncio
import json
import logging
import re
import sys
from pathlib import Path
from typing import Any

from sqlalchemy import text

_log = logging.getLogger("company_knowledge.magbot_catalog")

DOCS = Path("/docs") if Path("/docs").is_dir() else Path(__file__).resolve().parents[3] / "docs"
DECK_NAME = "magbot_spec_imgver_Ma260930.pptx"
SOURCE_DOC = "맥봇 제품 스펙 요약표 (2026-09-30)"
EVIDENCE = "공식 사양"

# 이 모듈의 카드가 대체하는 소개서 카드(같은 제품을 다룬 예전 추출본). 대체 카드가 없는 MTC·DTC·TCV4-Vision 은 남긴다.
SUPERSEDED_KEYS = (
    "intro:product:mdpgsseries", "intro:product:mdpgcseries", "intro:product:magbotdcpnuematicgrippermdpg시리즈",
    # 맥봇 플랫폼 네 축 소개 카드 — 제품이 아니라 플랫폼 구성. '맥봇 플랫폼 구성' 회사 역량 카드 1장으로 합침
    "intro:product:magbotatc", "intro:product:magboteoat", "intro:product:magbotfix", "intro:product:magbotcore",
    # 스위칭 마그네틱 소개 카드 2장(소개서 8쪽) — 제품이 아니라 ATC 의 툴 교체 기술. 회사 역량 카드 1장으로 합침(사용자 2026-10-01)
    "intro:product:스위칭마그네틱기술", "intro:product:magbotatc스위칭마그네틱",
    # mSMG-1S25D 의 다른 이름(Shift Lock Gripper) 카드 — 별칭으로 합침(사용자 2026-10-01)
    "intro:product:시프트락그리퍼smg1s25d",
    # mDPG 시리즈 개념 카드(철자만 다른 중복 포함) — 제품군·시리즈로만 남기고 카드는 두지 않는다(사용자 2026-10-01)
    "intro:product:magbotdcpneumaticgrippermdpg시리즈", "intro:product:magbotdcpneumaticgrippermdpg",
    "intro:product:tcv1supersimple", "intro:product:tcv2superpower", "intro:product:tcw1superfast",
    "intro:product:맥봇자동툴체인져mtc시리즈",
    "intro:product:2fingergripper", "intro:product:3fingergripper", "intro:product:2jawparallelgripper",
    "intro:product:시프트락그리퍼smg1s25d", "intro:product:shapememorygripper",
    "intro:product:마그네틱그리퍼", "intro:product:마그네틱그리퍼mgseries",
    # MGP·MGM 시리즈 소개 카드 — 모델 카드(panel_models)로 대체, 시리즈는 제품군으로만(사용자 2026-10-01)
    "intro:product:magboteoatmgmseries", "intro:product:magboteoatmgpseries", "intro:product:맥봇eoatmgmseries",
    # H시리즈 묶음 카드 — TCHK100·150·220 모델 카드(h_models)로 대체(사용자 2026-10-01)
    "intro:product:맥봇자동툴체인져h시리즈", "intro:product:magbotatch시리즈",
    # 리프팅 이송 로봇(소개서 24쪽) — 리프팅 AMR 과 같은 제품. 사양·사진은 리프팅 AMR 카드로 합침(사용자 2026-10-01)
    "intro:product:리프팅이송로봇",
    # 소개서 20쪽 End Tool 카드 2장 — 형상기억 그리퍼 제품군의 Magnet Gripper 1·2 카드(endtool_models)로
    # (Air Gripper 는 카드 없이 없앰, 2026-10-01)
    # 대체, 'End Tool' 제품군은 없앤다(사용자 2026-10-01)
    "intro:product:endtoolairgripper", "intro:product:endtoolmagnetgripper", "magbot:airgripper",
)

# 공식 자료에 없는 소개서 카드의 제품군을 맥봇 제품군 체계에 맞춘다(사용자 지정). 소개서를 다시 적재해도 유지 —
# cards.save 와 ingest 가 apply_family_overrides 를 부른다.
ATC_AUTO = "Magbot ATC · 자동 툴체인저"
FAMILY_OVERRIDES = {
    "intro:product:cylindricalgripper": "Magbot EOAT · 전동 핑거 그리퍼",     # 2026-10-01: 24V 원통형, Ø30~60
    # 2026-10-01: 맥봇 ATC 한 묶음 아래 자동·수동·듀얼·공압·액세서리로
    "intro:product:tcv4vision": ATC_AUTO,
    "intro:product:mtcmanualtc": "Magbot ATC · 수동 툴체인저",
    "intro:product:dtcdualtc": "Magbot ATC · 듀얼 툴체인저",
    # 2026-10-01: '4족 로봇' 제품군 폐기 — X30 기반 커스터마이징도 4족 보행 로봇군으로
    "intro:product:outdoorselfdrivingrobotmpayload": "4족 보행 로봇",
}

# 소개서 카드 이름 바로잡기(사용자 지정) — 제품군과 같이 소개서를 다시 적재해도 유지된다
NAME_OVERRIDES = {
    "intro:product:outdoorselfdrivingrobotcustomizing": "Outdoor Self-Driving Robot(M20 Customizing)",
    "intro:product:outdoorselfdrivingrobotmpayload": "Outdoor Self-Driving Robot(X30 Customizing)",
}


async def apply_family_overrides(db) -> int:
    n = 0
    for key, fam in FAMILY_OVERRIDES.items():
        n += (await db.execute(text(
            "UPDATE knowledge_cards SET card = jsonb_set(card, '{family}', to_jsonb(CAST(:f AS text))), updated_at = NOW() "
            "WHERE card_key = :k AND card->>'family' IS DISTINCT FROM :f"), {"k": key, "f": fam})).rowcount
    for key, name in NAME_OVERRIDES.items():
        n += (await db.execute(text(
            "UPDATE knowledge_cards SET name = :n, card = jsonb_set(card, '{name}', to_jsonb(CAST(:n AS text))), "
            "updated_at = NOW() WHERE card_key = :k AND card->>'name' IS DISTINCT FROM :n"), {"k": key, "n": name})).rowcount
    return n


# 같은 제품의 다른 이름 — 별도 카드로 빼지 않고 모델 카드에 붙인다(사용자 2026-10-01).
# 추천·공정 계획에서 AI 가 다른 이름으로 써도 이 카드로 연결된다(product_recommend.match_product).
ALIASES = {
    "mSMG-1S25D": ["Shift Lock Gripper", "시프트락 그리퍼", "SMG1S25D"],
}


# 예전 소개서 카드(SUPERSEDED_KEYS)에만 있던 내용 — 카드를 지우기 전에 모델 카드로 옮김(사용자 2026-10-01).
# 공식 사양과 값이 다른 것(TCV1 전기 사양 300ms, 3 Finger 가반하중 5kg)은 공식 사양을 따르고 옮기지 않는다.
_ATC_AUTO_MODELS = ("TCC1", "TCV1", "TCV2", "TCV3", "TCV4")
_MG_MODELS = ("MG5", "MG10", "MG16", "MG25", "MG30")
INTRO_EXTRAS: dict[str, dict[str, Any]] = {
    **{m: {"specs": [("반복 정밀도", "±0.05mm (회사 소개서 9쪽)")],
           "strengths": ["마그넷 동작 센서 및 근접 센서 내장", "Safety Lock 기능 포함"]} for m in _ATC_AUTO_MODELS},
    **{m: {"specs": [("반복 정밀도", "±0.05mm (회사 소개서 12쪽)")]} for m in _MG_MODELS},
}
for _m, _nick in (("TCV1", "Super-Simple"), ("TCV2", "Super-Power"), ("TCW1", "Super-Fast")):
    INTRO_EXTRAS.setdefault(_m, {})["aliases"] = [f"{_m} ({_nick})", _nick]
# 예전 카드가 이어 주던 회사 수행 사례 — 회상에서 제품 ↔ 사례로 번지는 연결(recall.spread)
FAMILY_RELATED = {
    "Magbot ATC · 자동 툴체인저": ["intro:reference:amrvisionatc시스템표면sanding", "intro:reference:디스플레이패널멀티검사공정자동화",
                               "intro:reference:magbotatc적용소방용amr호스자동탈부착시스템"],
    "Magbot EOAT · MG 마그네틱 그리퍼": ["intro:reference:lm가이드toolchanger", "intro:reference:자동차차체프레스물이송",
                                   "intro:reference:도어힌지이송공정", "intro:reference:customizing등속조인트이송그리퍼",
                                   "intro:reference:dualarmvisionhandlingsystem"],
    "Magbot EOAT · 전동 핑거 그리퍼": ["intro:reference:dualarmvisionhandlingsystem"],
}


def apply_intro_extras(card: dict[str, Any]) -> dict[str, Any]:
    ex = INTRO_EXTRAS.get(card["name"], {})
    card["specs"] = card["specs"] + [{"item": k, "value": v, "model": card["name"], "source": "소개서"}
                                     for k, v in ex.get("specs", [])]
    card["strengths"] = list(card.get("strengths") or []) + ex.get("strengths", [])
    if ex.get("aliases"):
        card["aliases"] = list(card.get("aliases") or []) + ex["aliases"]
    if card.get("family") in FAMILY_RELATED:
        card["related"] = FAMILY_RELATED[card["family"]]
    return card


# 맥봇 플랫폼 구성(사용자 2026-10-01: 플랫폼 이름은 제품 카드가 아니다) — 제품 DB 가 아니라 회사 역량 카드 1장.
# 내용은 소개서 6~7쪽의 네 축 설명 그대로. 제안서 회상에는 떠오르고, 제품 추천·그리퍼 후보·제품 사진에는 안 나온다.
PLATFORM_KEY = "magbot:platform"
PLATFORM_CARD = {
    "name": "맥봇 플랫폼 구성 (ATC · EOAT · FIX · CORE)", "kind": "capability", "family": "Magbot Platform",
    "summary": "맥봇 플랫폼은 자동 툴체인저(ATC)·엔드이펙터(EOAT)·스마트 클램프(FIX)·통합 제어(CORE) 네 축으로 "
               "로봇이 툴을 자동으로 바꾸고, 대상물을 잡고, 고정하고, 통합 제어하는 DC 구동 자동화 플랫폼이다.",
    "how_it_works": "Magbot ATC: 범용 모듈 구조의 초절전 DC 구동·AIoT 연동 자동 툴체인저(스위칭 마그네틱, 로봇 전원 24V 2A). "
                    "Magbot EOAT: DC 구동 기반 비정형 소재·제품 대응 모듈형 엔드이펙터(센서·비전·공압·전기 통합). "
                    "Magbot FIX: AI·IoT 기반 DC 스마트 클램프, ATC/EOAT 와 연동되는 자동 고정. "
                    "Magbot CORE: mHB 구동/제어 + On-Device AI SoC 기반 통합 제어(전원·구동·제어·통신 통합).",
    "key_ideas": ["공압 없이 DC 로 툴 교체·파지·고정·제어를 하나의 플랫폼으로", "모델은 제품 DB 의 ATC·EOAT 제품군 참고"],
    "cues": ["로봇 한 대로 여러 툴을 자동으로 바꿔 가며 여러 공정을 하고 싶다", "공압 설비 없이 전기(DC)로 그리퍼와 툴을 쓰고 싶다",
             "지그·클램프 고정까지 로봇과 연동해 자동화하고 싶다", "툴 교체·파지·고정·제어를 한 회사 플랫폼으로 통합하고 싶다"],
    "related": [],
}

# 스위칭 마그네틱(사용자 2026-10-01: ATC 에서 쓰는 전자석 이용 툴 교체 기술 — 제품이 아니다). 내용은 소개서 8쪽 카드 2장을 합친 것.
SWITCHING_KEY = "magbot:switching-magnetic"
SWITCHING_CARD = {
    "name": "스위칭 마그네틱 툴 교체 기술 (Magbot ATC)", "kind": "capability", "family": "Magbot ATC",
    "summary": "맥봇 ATC 에서 쓰는 전자석(스위칭 마그네틱) 기반 툴 교체 기술 — Master T/C 와 Tool Plate 를 마그네틱으로 "
               "결합해 초간편·초고속·초강력으로 툴을 바꾸며, 공압식/기계식 대비 신뢰성과 확장성이 높다.",
    "how_it_works": "Master T/C 와 Tool Plate 간의 마그네틱 결합. 배터리 내장 시스템(DC24V@2A, 0.1초 구동)으로 공압 유선 없이 완전 무선 결합.",
    "strengths": ["공압 유선이 필요 없는 완전 무선 결합", "공압 오링이나 볼 마모 문제 없는 내구성",
                  "고 신뢰성 및 내구성 (방수/방진 IP68, -40도~+200도 견딤)",
                  "강력한 가격 경쟁력 (부품 90% 이상 철(SS400), 알루미늄 일반 가공)",
                  "확장성 극대화 (낮은 가격의 초간단 구조의 툴사이드 툴체인저)"],
    "key_ideas": ["모델은 제품 DB 의 Magbot ATC · 자동 툴체인저 참고"],
    "cues": ["공압식 툴체인저의 볼 마모나 오링 마모로 인한 유지보수 문제가 빈번한 현장",
             "공압 설비 장치가 불가능하거나 최소화해야 하는 환경", "공압 컴프레셔나 구동 시스템을 설치할 공간이 부족한 현장",
             "공압 유선으로 인한 배선 복잡성을 해결해야 하는 경우", "고온이나 저온 등 극한 환경에서 내구성이 필요한 경우",
             "방수 및 방진 기능이 필수적인 작업 환경", "고가의 초정밀 가공 부품 대신 경제적인 솔루션을 찾는 경우",
             "로봇사이드 T/C보다 툴사이드 T/C를 통해 툴 확장성을 확보하려는 경우"],
    "products_used": ["magbot ATC"],
    "related": ["intro:reference:amrvisionatc시스템표면sanding", "intro:reference:다품종이송플랫폼",
                "intro:reference:magbotatc적용소방용amr호스자동탈부착시스템"],
}

# 제품 DB 가 아니라 회사 역량으로 두는 맥봇 카드 — (키, 카드, 소개서 쪽)
COMPANY_CARDS = ((PLATFORM_KEY, PLATFORM_CARD, [6, 7]), (SWITCHING_KEY, SWITCHING_CARD, [8]))


# 요약표 쪽 제목(앞부분) → (제품군, 동작 방식). 동작 방식은 각 제품군 소개 쪽 문장에서.
FAMILIES: list[tuple[str, str, str]] = [
    ("ATC 액세서리", "Magbot ATC · 액세서리",
     "Master TC(로봇 측)와 Tool Plate(툴 측) 사이에서 전원·통신(Pogo Pin)과 공압(PneuMatic)을 툴 교체 때 자동 연결·분리하는 커플러"),
    ("ATC", ATC_AUTO,
     "DC 24V 전기식 스위칭 마그네틱 결합 — 공압 없이 그리퍼·툴을 자동 체결·분리(Master TC + Tool Plate 한 세트)"),
    ("MG", "Magbot EOAT · MG 마그네틱 그리퍼",
     "스위칭 마그네틱 — 자력으로 철계(강자성체) 소재를 흡착·이송, 진공·공압 불필요"),
    ("D-Air", "Magbot EOAT · mDPG 시리즈 (DC Pneumatic Gripper · D-Air)",
     "DC 전동 진공 — 별도 공압 컴프레서 없이 자체적으로 진공을 만들어 흡착 패드로 파지"),
    ("SMG", "Magbot EOAT · SMG 형상기억 그리퍼",
     "형상기억 — 대상 형상에 맞춰 그리퍼(흡착컵·스트로크)가 변형되며 비정형 대상을 파지"),
    ("기타 그리퍼", "Magbot EOAT · 전동 핑거 그리퍼", "전동 기계식 핑거 파지"),
]
LTC_FAMILY = "Magbot ATC · M-LTC 공압 툴체인저"
LTC_HOW = "공압(4~7bar) 볼 록킹 결합 — 공압이 낮아지거나 끊겨도 자동 잠금으로 분리되지 않음(Master TC + Tool Plate 한 세트)"
_NOT_GIVEN = ("자료 없음", "미정", "없음")


def _clean(v: Any, limit: int = 400) -> str:
    return re.sub(r"\s+", " ", str(v or "")).strip()[:limit]


def family_of(title: str) -> tuple[str, str] | None:
    t = title.replace(" ", "")
    for prefix, fam, how in FAMILIES:
        if t.startswith(prefix.replace(" ", "")):
            return fam, how
    return None


def _texts(slide) -> list[tuple[float, float, str]]:
    out = []
    for sh in slide.shapes:
        if sh.shape_type == 6:          # 그룹 속 글(머리띠)은 쓰지 않는다
            continue
        if sh.has_text_frame and sh.text_frame.text.strip():
            out.append((sh.left / 914400, sh.top / 914400, sh.text_frame.text.strip()))
    return out


def _pictures(slide) -> list:
    """제품 사진만 — 쪽 머리띠(가로로 긴 그림)·쪽 전체 그림은 뺀다."""
    pics = []
    for sh in slide.shapes:
        if sh.shape_type != 13:
            continue
        w, h = sh.width / 914400, sh.height / 914400
        if w > 9 or (h < 0.6 and w > 3):
            continue
        pics.append(sh)
    return pics


def _row_values(table) -> list[list[str]]:
    return [[_clean(c.text, 600) for c in r.cells] for r in table.rows]


def summary_models(prs) -> list[dict[str, Any]]:
    """1~14쪽 요약표 → 모델 행(사진 포함). 소개 쪽 문장은 같은 제품군 카드의 설명으로."""
    intros: dict[str, str] = {}
    out: list[dict[str, Any]] = []
    for no, slide in enumerate(prs.slides, 1):
        texts = _texts(slide)
        title = next((t for x, y, t in texts if abs(y - 0.5) < 0.15 and x < 1), "")
        fam = family_of(title)
        tables = [sh for sh in slide.shapes if getattr(sh, "has_table", False) and sh.has_table]
        if fam and not tables:
            body = next((t for x, y, t in texts if 1.3 < y < 3 and x < 1 and len(t) > 40), "")
            if body:
                intros[fam[0]] = _clean(body, 600)
            continue
        if not fam or not tables:
            continue
        rows = _row_values(tables[0].table)
        head = rows[0]
        name_col = 1
        pics = sorted(_pictures(slide), key=lambda p: p.top)
        body_rows = [r for r in rows[1:] if r[name_col]]
        for i, r in enumerate(body_rows):
            spec = {head[j]: r[j] for j in range(2, len(head)) if j < len(r)}
            pic = pics[i] if len(pics) == len(body_rows) else None
            out.append({"model": r[name_col], "family": fam[0], "how": fam[1], "head": head, "spec": spec,
                        "slide": no, "image": pic.image.blob if pic is not None else None,
                        "accessory": head[1] == "품목"})
    for m in out:
        m["intro"] = intros.get(m["family"], "")
    return out


def _ltc_spec(rows: list[list[str]]) -> tuple[str, dict[str, str]]:
    """M-LTC 사양표(세 칸: 항목 · 세부 · 값) → (모델, {항목: 값})."""
    model, spec, last = "", {}, ""
    for r in rows:
        r = (r + ["", "", ""])[:3]
        key = r[0] or last
        last = key
        sub = r[1]
        if key == "Model" and "Master" in sub:
            model = re.sub(r"M$", "", r[2])
            if model.startswith("LTC-"):        # 원본 오타(0120D 쪽 'M-' 누락)
                model = "M-" + model
            continue
        if key == "Model":
            continue
        label = f"{key} — {sub}" if sub else key
        spec[label] = (spec[label] + " / " + r[2]) if label in spec and r[2] else r[2]
    return model, spec


def ltc_models(prs) -> list[dict[str, Any]]:
    """15~24쪽 M-LTC 사양 쪽 → 모델(사진은 같은 쪽 왼/오 반쪽의 제품 그림). 모듈 표 쪽(25~)은 건너뛴다."""
    out = []
    for no, slide in enumerate(prs.slides, 1):
        tables = [sh for sh in slide.shapes if getattr(sh, "has_table", False) and sh.has_table]
        specs = []
        for t in tables:
            rows = _row_values(t.table)
            if rows and rows[0][0] == "Model":
                model, spec = _ltc_spec(rows)
                if model:
                    specs.append((t.left / 914400, model, spec))
        if not specs:
            continue
        pics = _pictures(slide)
        texts = _texts(slide)
        for left, model, spec in specs:
            side = [p for p in pics if (p.left / 914400 < 6.6) == (left < 6.6)]
            feat = next((t for x, y, t in texts if t.startswith("Connector") and (x < 6.6) == (left < 6.6)), "")
            out.append({"model": model, "family": LTC_FAMILY, "how": LTC_HOW, "spec": spec, "slide": no,
                        "features": _clean(feat, 500),
                        "image": max(side, key=lambda p: p.width * p.height).image.blob if side else None})
    return out


# ── 카드 ──────────────────────────────────────────────────────────────────────

def _value_known(v: str) -> bool:
    return bool(v) and not any(v.startswith(x) for x in _NOT_GIVEN)


def card_key(model: str) -> str:
    return "magbot:" + re.sub(r"[^0-9a-z가-힣]+", "", model.lower())


def series_of(name: str) -> str:
    """mDPG 시리즈 안의 세부 시리즈(사용자 확인 2026-10-01): mDPG1S1·4S3·8S7 = S-series, mDPG-C5·C25 = C-series.
    'mDPG 시리즈'는 따로 카드로 두지 않고 다섯 모델 카드의 제품군·시리즈로만 남긴다."""
    if re.match(r"mDPG-C", name):
        return "C-series"
    if re.match(r"mDPG\d", name):
        return "S-series"
    return ""


def summary_card(m: dict[str, Any]) -> dict[str, Any]:
    spec = m["spec"]
    if m.get("accessory"):
        name = f"{m['model']} ({spec.get('코드', '')})".replace(" ()", "")
        specs = [{"item": k, "value": v, "model": name, "source": "원문"} for k, v in spec.items() if v]
        return {"name": name, "kind": "product", "family": m["family"],
                "summary": _clean(f"{m['family']} — {spec.get('기능', '')} ({spec.get('장착 위치', '')})", 200),
                "how_it_works": m["how"], "applies_when": _clean(f"적용 제품: {spec.get('적용 제품', '')}", 200),
                "specs": specs, "strengths": [], "limits": [], "not_when": [], "intro": m.get("intro", "")}
    m_ = re.match(r"(.+?)\s*\((.+)\)$", m["model"])
    name, variant = (m_.group(1), m_.group(2)) if m_ else (m["model"], "")
    payload = spec.get("가반하중", "")
    specs = [{"item": k, "value": v, "model": name, "source": "원문"} for k, v in spec.items()
             if k != "적용 조건" and _value_known(v)]
    limits = [f"{k}: {v}" for k, v in spec.items() if "확인 필요" in v]
    unknown = [k for k, v in spec.items() if not _value_known(v) and k != "적용 조건"]
    if unknown:
        limits.append("자료 없음: " + ", ".join(unknown))
    series = series_of(name)
    summary = f"{m['family']}" + (f" {series}" if series else "") + f" {name}" + (f"({variant})" if variant else "") + (
        f" — 가반하중 {payload}" if _value_known(payload) else "")
    size = spec.get("제품 사이즈", "")
    if _value_known(size):
        summary += f", 크기 {size}"
    aliases = ALIASES.get(name, [])
    if aliases:
        summary += " · 다른 이름: " + ", ".join(aliases)
    return {"name": name, "kind": "product", "family": m["family"], **({"series": series} if series else {}),
            **({"aliases": aliases} if aliases else {}), "summary": _clean(summary, 300),
            "how_it_works": m["how"], "applies_when": _clean(spec.get("적용 조건", ""), 300),
            "specs": specs, "strengths": [], "limits": limits[:4], "not_when": [], "intro": m.get("intro", "")}


def ltc_card(m: dict[str, Any]) -> dict[str, Any]:
    spec = m["spec"]
    payload = next((v for k, v in spec.items() if k.startswith("Payload")), "")
    size = next((v for k, v in spec.items() if "When coupled" in k and "Size" in k), "")
    specs = [{"item": k, "value": v, "model": m["model"], "source": "원문"} for k, v in spec.items() if v]
    if payload:
        specs.insert(0, {"item": "가반하중", "value": payload, "model": m["model"], "source": "원문"})
    return {"name": m["model"], "kind": "product", "family": LTC_FAMILY,
            "summary": _clean(f"{LTC_FAMILY} {m['model']} — 가반하중 {payload}, 결합 크기 {size}", 240),
            "how_it_works": LTC_HOW, "applies_when": "공압(4~7bar) 공급이 되는 로봇 셀의 자동 툴 교체",
            "specs": specs, "strengths": [_clean(m.get("features"), 200)] if m.get("features") else [],
            "limits": [], "not_when": [], "intro": ""}


def parse(path: Path | None = None) -> list[dict[str, Any]]:
    """원본 → [{card, image, slide}] (모델 단위)."""
    from pptx import Presentation

    prs = Presentation(str(path or DOCS / "sources" / DECK_NAME))
    out = [{"card": apply_intro_extras(summary_card(m)), "image": m["image"], "slide": m["slide"]}
           for m in summary_models(prs)]
    out += [{"card": ltc_card(m), "image": m["image"], "slide": m["slide"]} for m in ltc_models(prs)]
    return out


# ── MGP·MGM 시리즈 (docs/sources/MGP,MGM.pptx 1쪽 + 회사 소개서 13·14쪽) ─────────────────────────
# 원본이 그림 위주(표는 소개서 14쪽 MGP 두 모델뿐)라 값은 원문을 옮겨 적는다. 원본의 'MG3O6O'·'MGM1Oq'(숫자 0 자리에
# 영문 O)는 MG3060·MGM10Q 로. MGM 은 형태(D·T·Q·H)별 모델이고 소개서 13쪽에 형태마다 변형 4종이 있다.
PANEL_DECK = "MGP,MGM.pptx"
MGP_FAMILY = "Magbot EOAT · MGP 시리즈"
MGM_FAMILY = "Magbot EOAT · MGM 시리즈"
_PANEL_HOW = "스위칭 마그네틱 기술 — 고객 제품과 공정에 최적화 설계된 고객향 제품(가반하중 선택적 제작)"
_MGP_WHEN = "다양한 형태의 소형 부품(예: 자동차 베어링)을 다량 이동할 때. 자석 흡착이라 철계(강자성체) 대상"
_MGM_WHEN = "넓은 철판 이동 공정(예: 자동차 섀시)에서 제품에 맞는 모델을 선정해 쓸 때. 자석 흡착이라 철계(강자성체) 대상"
_INTRO_SRC = "회사 소개서 3Q"
PANEL_MODELS: list[dict[str, Any]] = [
    {"name": "TCT1", "family": MGP_FAMILY, "when": _MGP_WHEN, "pic": 14,
     "specs": [("가반하중", "5kg")], "limits": ["자료 없음: 크기, 중량, 전기사양 (MGP,MGM.pptx 에 사진·가반하중만)"]},
    {"name": "MG3060", "family": MGP_FAMILY, "when": _MGP_WHEN, "pic": 15, "aliases": ["MG3O6O"], "intro_slide": 14,
     "specs": [("가반하중", "10kg / 20kg / 30kg / 50kg (가반하중 선택적 제작)"),
               ("마스터 TC 크기", "600 X 300 X 110mm (가반하중 선택적 제작에 따른 사이즈 제작)"),
               ("제품 중량", "38kg (제작 사양에 따른 중량 변화 예상)"), ("결합방식", "스위칭 마그네틱 기술"),
               ("전기사양", "24V 3A X 6EA")],
     "limits": ["허용온도·습도·IP 등급: 고객 사양에 따른 등급 측정", "자료 없음: 위치 재현 정도, 툴 플레이트 크기"]},
    {"name": "MG7179", "family": MGP_FAMILY, "when": _MGP_WHEN, "pic": 16, "intro_slide": 14,
     "specs": [("가반하중", "20kg / 30kg / 80kg (가반하중 선택적 제작)"),
               ("마스터 TC 크기", "790 X 710 X 110mm (가반하중 선택적 제작에 따른 사이즈 제작)"),
               ("제품 중량", "120kg (제작 사양에 따른 중량 변화 예상)"), ("결합방식", "스위칭 마그네틱 기술"),
               ("전기사양", "24V 3A X 20EA")],
     "limits": ["허용온도·습도·IP 등급: 고객 사양에 따른 등급 측정",
                "확인 필요: 소개서 14쪽 '툴 플레이트 크기' 칸에 무게 값 1,057g [2.33lb] 이 적혀 있음",
                "자료 없음: 위치 재현 정도"]},
    {"name": "MGM3D", "family": MGM_FAMILY, "when": _MGM_WHEN, "pic": 12, "intro_slide": 13,
     "variants": ["MGM3D", "MGM10D", "MGM16D", "MGM30D"]},
    {"name": "MGM5T", "family": MGM_FAMILY, "when": _MGM_WHEN, "pic": 13, "intro_slide": 13,
     # 소개서 13쪽 변형 목록엔 5T 가 없지만 대표 이름으로 이미 적어 둔 것 — 변형에 포함(사용자 2026-10-01)
     "variants": ["MGM5T", "MGM3T", "MGM10T", "MGM16T", "MGM30T"]},
    {"name": "MGM10Q", "family": MGM_FAMILY, "when": _MGM_WHEN, "pic": 17, "intro_slide": 13, "aliases": ["MGM1OQ"],
     "variants": ["MGM3Q", "MGM10Q", "MGM16Q", "MGM30Q"]},
    {"name": "MGM3H", "family": MGM_FAMILY, "when": _MGM_WHEN, "intro_slide": 13, "crop": (13, (1165, 385, 1490, 515)),
     "variants": ["MGM3H", "MGM10H", "MGM16H", "MGM30H"], "limits": ["MGP,MGM.pptx 에는 없음 — 소개서 13쪽에만"]},
]
_MGP_RELATED = ["intro:reference:베어링이송마그네틱그리퍼패널타입", "intro:reference:dualarmvisionhandlingsystem"]
_MGM_RELATED = ["intro:reference:자동차차체프레스물이송", "intro:reference:dualarmvisionhandlingsystem"]


def panel_card(m: dict[str, Any]) -> dict[str, Any]:
    mgm = m["family"] == MGM_FAMILY
    specs = [{"item": k, "value": v, "model": m["name"], "source": "원문"} for k, v in m.get("specs", [])]
    limits = list(m.get("limits", []))
    if mgm:
        limits.append("자료 없음: 가반하중·크기·중량 등 사양 표 (소개서 13쪽은 사진·변형 목록만)")
    aliases = [v for v in m.get("variants", []) if v != m["name"]] + m.get("aliases", [])
    payload = next((v for k, v in m.get("specs", []) if k == "가반하중"), "")
    summary = f"{m['family']} {m['name']}" + (f" — 가반하중 {payload}" if payload else "") + (
        f" · 변형: {' / '.join(m['variants'])}" if m.get("variants") else "")
    return {"name": m["name"], "kind": "product", "family": m["family"], **({"aliases": aliases} if aliases else {}),
            "summary": _clean(summary, 300), "how_it_works": _PANEL_HOW, "applies_when": m["when"], "specs": specs,
            "strengths": ["SUPER POWERFUL", "SUPER FAST", "POWER SAVING"] if mgm else ["가반하중에 따른 맞춤형 제작 가능"],
            "limits": limits[:4], "not_when": [], "related": _MGM_RELATED if mgm else _MGP_RELATED}


def panel_models(path: Path | None = None) -> list[dict[str, Any]]:
    """MGP·MGM 모델 → [{card, image, slide, source_doc, source_ref}]. 사진은 MGP,MGM.pptx 의 개별 그림,
    거기 없는 MGM3H 는 소개서 13쪽 배경 그림에서 잘라 쓴다."""
    import io
    import zipfile

    from PIL import Image
    from pptx import Presentation

    from .intro_deck import DECK_NAME as INTRO_DECK

    deck = path or DOCS / "sources" / PANEL_DECK
    pics = {sh.shape_id: sh.image.blob for sh in Presentation(str(deck)).slides[0].shapes if sh.shape_type == 13}
    out = []
    for m in PANEL_MODELS:
        image = pics.get(m.get("pic"))
        if m.get("crop"):
            no, box = m["crop"]
            with zipfile.ZipFile(DOCS / "sources" / INTRO_DECK) as z:
                rels = z.read(f"ppt/slides/_rels/slide{no}.xml.rels").decode()
                media = re.search(r'Target="\.\./media/([^"]+\.jpe?g)"', rels).group(1)
                buf = io.BytesIO()
                Image.open(io.BytesIO(z.read(f"ppt/media/{media}"))).convert("RGB").crop(box).save(buf, "JPEG", quality=92)
                image = buf.getvalue()
        in_deck = "pic" in m
        # 표 사양이 있는 건 소개서 14쪽 MGP 두 모델뿐 — '공식 사양'(요약표 자료)과 구분한다
        out.append({"card": panel_card(m), "image": image, "slide": m.get("intro_slide") or 1, "evidence": "제품 사양",
                    "source_doc": f"{PANEL_DECK} · {_INTRO_SRC}" if in_deck else _INTRO_SRC,
                    "source_ref": f"{PANEL_DECK} 1쪽" if in_deck else f"소개서 {m['crop'][0]}쪽 배경 그림"})
    return out


# ── H시리즈 TCHK (회사 소개서 10쪽) ─────────────────────────────────────────────
# 표 그대로. 'Positioning Repeatability' 부터 아래 공통 행은 세 모델에 걸친 병합 칸이라 모두에 넣는다.
# '†' 는 원본 표시(미확정 사양 — 변경될 수 있음)를 그대로 둔다. 사진은 모델별이 아니라 H시리즈 그림 2장을
# 세 모델 모두에 연결한다(사용자 2026-10-01).
H_SLIDE = 10
# 산업용 로봇용(사용자 2026-10-01: 산업용 툴체인저로 추천). 협동로봇용 DC 전기식 TCC·TCV·TCW(5~30kgf)와
# 체급(100~220kgf)·구동(공압 Ball-Locking 융복합)이 달라 제품군을 나눈다 — 같은 제품군이면 그리퍼 후보·추천에서 서로 형제로 묶인다
H_FAMILY = "Magbot ATC · 산업용 툴체인저 (H시리즈)"
H_HOW = "magbot 스위칭 마그네틱과 DC Ball Lock/Unlock 융복합 — Ball-Locking 방식, 작동 공기압 4~7bar"
_H_COMMON = [("반복 정밀도 (Positioning Repeatability)", "±0.05mm"), ("결합 방식", "Ball-Locking 방식"),
             ("작동 공기압", "4~7bar †"), ("최대 공기압", "8bar †"), ("사용 온도 / 습도", "0~60℃ / 0~95% †"),
             ("Safety Lock", "O")]
H_MODELS = [
    ("TCHK100", [("가반하중 (Payload 정격 하중)", "100kgf †"), ("Housing Diameter", "Ø200mm †"),
                 ("Master TC (Robot-Side) 크기", "200×200×103.5mm †"), ("Tool Plate (Tool-Side) 크기", "200×200×33.2mm †"),
                 ("Tool Changer (결합 시) 크기", "200×200×111.3mm †"), ("Master TC 중량", "13.6 kg †"),
                 ("Tool Plate 중량", "6.4kg †"), ("Tool Changer (결합 시) 중량", "20 kg †")]),
    ("TCHK150", [("가반하중 (Payload 정격 하중)", "150kgf"), ("Housing Diameter", "Ø220mm †"),
                 ("Master TC (Robot-Side) 크기", "220×220×107.5mm †"), ("Tool Plate (Tool-Side) 크기", "220×220×37.2mm †"),
                 ("Tool Changer (결합 시) 크기", "220×220×115.3mm †"), ("Master TC 중량", "14kg †"),
                 ("Tool Plate 중량", "6.7kg †"), ("Tool Changer (결합 시) 중량", "20.7 Kg †")]),
    ("TCHK220", [("가반하중 (Payload 정격 하중)", "220kgf"), ("Housing Diameter", "Ø260mm †"),
                 ("Master TC (Robot-Side) 크기", "260×260×109mm †"), ("Tool Plate (Tool-Side) 크기", "260×260×38.7mm †"),
                 ("Tool Changer (결합 시) 크기", "260×260×116.8mm"), ("Master TC 중량", "15.1kg †"),
                 ("Tool Plate 중량", "7kg †"), ("Tool Changer (결합 시) 중량", "22.1kg †")]),
]
_H_RELATED = ["intro:reference:amrvisionatc시스템표면sanding", "intro:reference:디스플레이패널멀티검사공정자동화"]


def h_models() -> list[dict[str, Any]]:
    """TCHK100·150·220 → [{card, images, slide, ...}]. 사진은 소개서 10쪽의 그림 전부(H시리즈 공통)."""
    from pptx import Presentation

    from .intro_deck import DECK_NAME as INTRO_DECK

    slide = Presentation(str(DOCS / "sources" / INTRO_DECK)).slides[H_SLIDE - 1]
    from .product_images import MIN_SIDE_PX

    # 머리 장식 그림(268×85)은 빼고 제품 그림만
    pics = sorted((sh for sh in slide.shapes if sh.shape_type == 13 and min(sh.image.size) >= MIN_SIDE_PX),
                  key=lambda sh: sh.top)
    images = [sh.image.blob for sh in pics]
    out = []
    for name, own in H_MODELS:
        specs = [{"item": k, "value": v, "model": name, "source": "원문"} for k, v in own + _H_COMMON]
        payload = own[0][1]
        card = {"name": name, "kind": "product", "family": H_FAMILY, "series": "H시리즈",
                "summary": _clean(f"{H_FAMILY} {name} — 가반하중 {payload}, 산업용 로봇용. 스위칭 마그네틱과 "
                                  "DC Ball Lock/Unlock 융복합 Hybrid·High-Payload 자동 툴체인저", 300),
                "how_it_works": H_HOW,
                "applies_when": "산업용 로봇 툴체인저로 추천 — 고중량(100~220kgf) 툴 자동 교체. 공압(4~7bar) 필요",
                "specs": specs, "strengths": ["산업용 로봇용 High-Payload", "Hybrid(스위칭 마그네틱 + DC Ball Lock/Unlock)",
                                              "Ball-Locking 방식", "Safety Lock"],
                "limits": ["† 표시 값은 미확정 사양 — 변경될 수 있음 (소개서 10쪽)"],
                "not_when": ["협동로봇·소형 툴(5~30kgf) 교체 — 자동 툴체인저 TCC1·TCV1~4·TCW1 쪽"], "related": _H_RELATED}
        out.append({"card": card, "images": images, "image": None, "slide": H_SLIDE, "evidence": "제품 사양",
                    "source_doc": _INTRO_SRC, "source_ref": f"소개서 {H_SLIDE}쪽 H시리즈 그림(세 모델 공통)"})
    return out


# ── 형상기억 그리퍼 End Tool (회사 소개서 20쪽 아래 박스 3개) ──────────────────────────────
# 박스마다 그림 3장(조립 모습 + 세부 2장)을 슬라이드 배치 그대로 한 장으로 합쳐 모델 사진으로 쓴다(사용자 2026-10-01).
ENDTOOL_SLIDE = 20
SMG_FAMILY = "Magbot EOAT · SMG 형상기억 그리퍼"
# 첫 박스(Air Gripper)는 카드로 두지 않는다(사용자 2026-10-01: 카드 자체를 없앰).
ENDTOOL_MODELS = [
    ("Magnet Gripper 1", (16, 17, 18), "형상기억 그리퍼 End Tool — 자력을 이용한 파지 및 이동이 가능한 마그네틱 그리퍼 툴",
     "형상기억 그리퍼 구조에 마그네틱을 단 엔드 툴 — 자력으로 고정·이동"),
    ("Magnet Gripper 2", (42, 43, 44), "형상기억 그리퍼 End Tool — 자력을 이용한 파지 및 이동이 가능한 마그네틱 그리퍼 툴",
     "형상기억 그리퍼 구조에 마그네틱을 단 엔드 툴 — 자력으로 고정·이동"),
]


def _compose(shapes: list[Any], scale: float = 3.0) -> bytes:
    """슬라이드 위 그림 여러 장 → 배치 그대로 흰 바탕 한 장(JPEG)."""
    import io

    from PIL import Image

    emu = 12700 / scale
    x0, y0 = min(sh.left for sh in shapes), min(sh.top for sh in shapes)
    x1, y1 = max(sh.left + sh.width for sh in shapes), max(sh.top + sh.height for sh in shapes)
    canvas = Image.new("RGBA", (int((x1 - x0) / emu), int((y1 - y0) / emu)), "white")
    for sh in shapes:
        im = Image.open(io.BytesIO(sh.image.blob)).convert("RGBA")
        im = im.resize((max(1, int(sh.width / emu)), max(1, int(sh.height / emu))), Image.LANCZOS)
        canvas.alpha_composite(im, (int((sh.left - x0) / emu), int((sh.top - y0) / emu)))
    buf = io.BytesIO()
    canvas.convert("RGB").save(buf, "JPEG", quality=92)
    return buf.getvalue()


def endtool_models() -> list[dict[str, Any]]:
    from pptx import Presentation

    from .intro_deck import DECK_NAME as INTRO_DECK

    slide = Presentation(str(DOCS / "sources" / INTRO_DECK)).slides[ENDTOOL_SLIDE - 1]
    pics = {sh.shape_id: sh for sh in slide.shapes if sh.shape_type == 13}
    out = []
    for name, ids, summary, how in ENDTOOL_MODELS:
        card = {"name": name, "kind": "product", "family": SMG_FAMILY, "summary": summary, "how_it_works": how,
                "applies_when": "", "specs": [], "strengths": [], "not_when": [],
                "limits": ["자료 없음: 가반하중·크기 등 사양 (소개서 20쪽은 그림만)"]}
        out.append({"card": card, "image": _compose([pics[i] for i in ids]), "slide": ENDTOOL_SLIDE,
                     "evidence": "제품 사양", "source_doc": _INTRO_SRC,
                     "source_ref": f"소개서 {ENDTOOL_SLIDE}쪽 '{name}' 박스 그림 3장(한 장으로)"})
    return out


# ── DB ─────────────────────────────────────────────────────────────────────

async def ingest(path: Path | None = None) -> dict[str, Any]:
    """카드 넣기/갱신 + 예전 카드 비활성화 + 사진(승인) + 회상 단서 다시 만들기."""
    from ..database import SessionLocal
    from . import product_images as pi

    items = parse(path) + panel_models() + h_models() + endtool_models()
    added = photos = 0
    async with SessionLocal() as db:
        off = (await db.execute(text("UPDATE knowledge_cards SET active = FALSE, updated_at = NOW() "
                                     "WHERE card_key = ANY(:k) AND active RETURNING id"),
                                {"k": list(SUPERSEDED_KEYS)})).scalars().all()
        await apply_family_overrides(db)
        for key, card, slides in COMPANY_CARDS:
            await db.execute(text(
                "INSERT INTO knowledge_cards (card_key, kind, name, evidence, source_doc, slides, card, active) "
                "VALUES (:k, 'capability', :n, '회사 소개', '회사 소개서 3Q', :s, CAST(:c AS jsonb), TRUE) "
                "ON CONFLICT (card_key) DO UPDATE SET kind = 'capability', name = EXCLUDED.name, card = EXCLUDED.card, "
                "slides = EXCLUDED.slides, active = TRUE, updated_at = NOW()"),
                {"k": key, "n": card["name"], "s": slides, "c": json.dumps(card, ensure_ascii=False)})
        for it in items:
            c = it["card"]
            cid = (await db.execute(text(
                "INSERT INTO knowledge_cards (card_key, kind, name, evidence, source_doc, slides, card, active) "
                "VALUES (:k, 'product', :n, :e, :d, :s, CAST(:c AS jsonb), TRUE) "
                "ON CONFLICT (card_key) DO UPDATE SET name = EXCLUDED.name, evidence = EXCLUDED.evidence, "
                "source_doc = EXCLUDED.source_doc, slides = EXCLUDED.slides, card = EXCLUDED.card, active = TRUE, "
                "updated_at = NOW() RETURNING id"),
                {"k": card_key(c["name"]), "n": c["name"], "e": it.get("evidence", EVIDENCE), "d": it.get("source_doc", SOURCE_DOC),
                 "s": [it["slide"]],
                 "c": json.dumps(c, ensure_ascii=False)})).scalar_one()
            added += 1
            for blob in it.get("images") or ([it["image"]] if it["image"] else []):
                jp = pi.to_jpeg(blob)
                if jp and await pi._store(db, cid, *jp, source="spec", source_ref=it.get("source_ref") or f"{DECK_NAME} {it['slide']}쪽",
                                         status="approved", user_id=None):
                    photos += 1
        await db.commit()
    from .cards import rebuild_cues
    cues = await rebuild_cues()
    return {"cards": added, "photos_new": photos, "superseded_off": len(off), "cues": cues}


def markdown(items: list[dict[str, Any]]) -> str:
    lines = ["# 맥봇 제품 공식 사양 (회사 제품 DB)", "",
             f"원본: docs/sources/{DECK_NAME} · 모델 {len(items)}개 · 카드 키 'magbot:<모델>'. 소개서 추출 카드 중 같은 제품 "
             f"{len(SUPERSEDED_KEYS)}장은 비활성화.", ""]
    fam = ""
    for it in items:
        c = it["card"]
        if c["family"] != fam:
            fam = c["family"]
            lines += ["", f"## {fam}", "", "| 모델 | 가반하중 | 적용 조건 | 확인 필요 | 사진 |", "|---|---|---|---|---|"]
        pay = next((s["value"] for s in c["specs"] if s["item"] in ("가반하중",)), "-")
        lines.append(f"| {c['name']} | {pay} | {c['applies_when'][:80]} | {'; '.join(c['limits'])[:80]} | "
                     f"{'○' if it['image'] else '-'} |")
    return "\n".join(lines) + "\n"


if __name__ == "__main__":
    logging.basicConfig(level=logging.INFO)
    cmd = sys.argv[1] if len(sys.argv) > 1 else ""
    if cmd == "ingest":
        res = asyncio.run(ingest())
        doc = DOCS / "company_knowledge" / "magbot_catalog.md"
        doc.write_text(markdown(parse()), encoding="utf-8")
        print(res, flush=True)
    elif cmd == "parse":
        for it in parse():
            c = it["card"]
            print(c["name"], "|", c["family"], "|", len(c["specs"]), "specs |", "img" if it["image"] else "-",
                  "|", c["applies_when"][:60])
    else:
        print(__doc__)
