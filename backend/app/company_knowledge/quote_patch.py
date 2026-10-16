"""올린 견적서 파일 위에서 이어서 고치기(사용자 2026-10-15) — 엑셀에서 한 서식(글꼴·색·채우기·셀 병합·행 높이·그림·도형)을 지킨다.

작성 상태를 올린 파일로 맞춘 뒤(quote_readback) 사용자가 [견적서 고치기]·AI 대화로 더 고치면, 미리보기·발행 파일을 회사 양식으로
새로 그리지 않고 올린 파일을 열어 바뀐 칸만 고친다.
- 짝짓기: 올린 파일의 표 줄(quote_readback.parse) ↔ 그 파일에 맞춘 작성 상태(base)가 그리는 줄 — 종류가 한 줄씩 같아야 한다.
- 바뀐 점: base 가 그리는 줄 ↔ 지금 상태가 그리는 줄(줄 id·글로 짝지음). 같은 줄은 바뀐 칸 값만 쓰고(서식 그대로),
  없어진 줄은 행을 지우고, 새 줄은 행을 넣어 회사 양식으로 만든 같은 줄의 서식·병합·높이를 옮긴다.
  새 줄의 겉모습(글꼴·색·채우기·정렬·테두리)은 올린 파일의 같은 종류 줄을 따른다(사용자 2026-10-16) — 칸마다 가장 흔한 모양,
  같으면 위쪽 줄. 한 줄만 강조한 서식은 퍼지지 않는다. 같은 종류 줄이 없으면 회사 양식 그대로.
- 금액 수식(수량×단가·소계·합계)은 행이 움직인 뒤 다시 쓴다. 머리·조건·비고·주석 칸은 바뀐 것만.
- openpyxl 은 열었다 저장할 때 도형·텍스트 상자를 버린다 — 원본 그림판(drawing) XML 에서 도형만 꺼내 행 이동만큼 옮겨 다시 넣는다.
- 맞지 않으면 PatchError — 부르는 쪽이 회사 양식으로 새로 그리고 이유를 알린다.
"""
from __future__ import annotations

import copy
import difflib
from collections import Counter
import io
import re
import zipfile
from datetime import date
from typing import Any

from openpyxl import load_workbook

from . import quote_readback as rb
from . import quote_xlsx as qx

_SHAPE_RE = re.compile(r"<((?:\w+:)?)(twoCellAnchor|oneCellAnchor|absoluteAnchor)\b.*?</\1\2>", re.S)
_NS = {"xdr": "http://schemas.openxmlformats.org/drawingml/2006/spreadsheetDrawing",
       "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
       "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships"}


class PatchError(ValueError):
    """올린 파일 위에서 고칠 수 없는 이유(사람이 읽는 문장)."""


# ── 줄 비교용 표시 ─────────────────────────────────────────────────────────────

def _kind(r: dict[str, Any]) -> str:
    k = r.get("row") or "item"
    return "item" if k in ("item", "option") else k


def _ident(r: dict[str, Any]) -> str:
    """같은 줄인지 — 품목은 상태 줄 id, 서술형 줄은 note id, 나머지는 종류+글."""
    if r.get("id"):
        return f"i:{r['id']}"
    if r.get("note_id"):
        return f"n:{r['note_id']}"
    k = r.get("row")
    if k == "spec":
        return f"s:{rb._norm(r.get('label') or '')}"
    if k in ("section", "text"):
        return f"{k}:{rb._norm(r.get('text') or '')}"
    return k or "item"


class _Sheet:
    """행 넣기·지우기 — openpyxl 이 옮기지 않는 병합·행 높이·그림 위치와 기억해 둔 행 번호(marks)를 함께 옮긴다."""

    def __init__(self, ws) -> None:
        self.ws = ws
        self.marks: dict[str, int] = {}
        self.log: list[tuple[str, int, int]] = []           # (insert|delete, 행, 개수) — 도형 옮길 때

    def _shift_dims(self, at: int, n: int) -> None:
        dims = dict(self.ws.row_dimensions)
        self.ws.row_dimensions.clear()
        for k, d in dims.items():
            if n < 0 and at <= k < at - n:
                continue
            nk = k + n if k >= at else k
            d.index = nk
            self.ws.row_dimensions[nk] = d

    def _shift_images(self, at: int, n: int) -> None:
        for img in self.ws._images:
            a = img.anchor
            for m in (getattr(a, "_from", None), getattr(a, "to", None)):
                if m is not None and m.row + 1 >= at:
                    m.row = max(at - 1, m.row + n) if n < 0 and m.row + 1 < at - n else m.row + n

    def _shift(self, at: int, n: int) -> None:
        for rng in list(self.ws.merged_cells.ranges):
            if n < 0 and at <= rng.min_row and rng.max_row < at - n:
                self.ws.merged_cells.remove(rng)
            elif rng.min_row >= at:
                rng.shift(row_shift=n)
        self._shift_dims(at, n)
        self._shift_images(at, n)
        for k, v in self.marks.items():
            if v >= at:
                self.marks[k] = v + n
        self.log.append(("insert" if n > 0 else "delete", at, abs(n)))

    def insert(self, at: int, n: int = 1) -> None:
        self.ws.insert_rows(at, n)
        self._shift(at, n)

    def delete(self, at: int, n: int = 1) -> None:
        for rng in list(self.ws.merged_cells.ranges):           # 지울 행에 걸친 병합은 먼저 푼다
            if at <= rng.min_row < at + n or at <= rng.max_row < at + n:
                if not (at <= rng.min_row and rng.max_row < at + n):
                    self.ws.merged_cells.remove(rng)
        self.ws.delete_rows(at, n)
        self._shift(at, -n)

    def new_row(self, orig: int) -> int | None:
        """원본 행 번호 → 지금 행 번호(지워졌으면 None)."""
        r = orig
        for op, at, n in self.log:
            if op == "insert" and r >= at:
                r += n
            elif op == "delete":
                if at <= r < at + n:
                    return None
                if r >= at + n:
                    r -= n
        return r


def _copy_row(src_ws, sr: int, ws, tr: int) -> None:
    """회사 양식으로 만든 줄(sr) → 대상 행(tr): 값·서식·병합·높이."""
    for c in range(1, 8):
        s, t = src_ws.cell(sr, c), ws.cell(tr, c)
        if type(t).__name__ == "MergedCell":
            continue
        t.value = s.value if type(s).__name__ != "MergedCell" else None
        if s.has_style:
            t._style = copy.copy(s._style)
    for rng in src_ws.merged_cells.ranges:
        if rng.min_row == sr == rng.max_row:
            ws.merge_cells(start_row=tr, start_column=rng.min_col, end_row=tr, end_column=rng.max_col)
    h = src_ws.row_dimensions[sr].height
    if h:
        ws.row_dimensions[tr].height = h


def _looks(ws, kinds: list[str], rows: list[int]) -> dict[str, list[Any]]:
    """줄 종류 → 칸(A~G)마다 그 종류 줄에서 가장 흔한 서식(같으면 위쪽 줄). 행을 넣고 지우기 전에 부른다."""
    out: dict[str, list[Any]] = {}
    for k in dict.fromkeys(kinds):
        rs = [r for kk, r in zip(kinds, rows) if kk == k]
        look = []
        for c in range(1, 8):
            cells = [ws.cell(r, c) for r in rs]
            styles = {tuple(x._style): x._style for x in cells if type(x).__name__ != "MergedCell" and x.has_style}
            seen = Counter(tuple(x._style) for x in cells if type(x).__name__ != "MergedCell" and x.has_style)
            look.append(copy.copy(styles[seen.most_common(1)[0][0]]) if seen else None)
        out[k] = look
    return out


def _dress(ws, tr: int, look: list[Any], ref_ws, sr: int) -> None:
    """새 줄(tr)에 올린 파일의 같은 종류 줄 모양을 입힌다 — 숫자 표시 형식만 회사 양식 줄(sr) 것."""
    for c, st in enumerate(look, start=1):
        t = ws.cell(tr, c)
        if st is None or type(t).__name__ == "MergedCell":
            continue
        t._style = copy.copy(st)
        t.number_format = ref_ws.cell(sr, c).number_format


def _set(ws, r: int, c: int, value: Any, fmt: str | None = None) -> None:
    cell = ws.cell(r, c)
    if type(cell).__name__ == "MergedCell":
        return
    cell.value = value
    if fmt:
        cell.number_format = fmt


def _qty_fmt(r: dict[str, Any], en: bool) -> str:
    return f'#,##0_ "{r.get("unit") or "ea"}"' if en else (f'0" {r["unit"]}"' if r.get("unit") else qx.EA)


# ── 고치기 ───────────────────────────────────────────────────────────────────

def patch(base: bytes, base_lines: list[dict[str, Any]], base_header: dict[str, Any],
          new_lines: list[dict[str, Any]], new_header: dict[str, Any]) -> tuple[bytes, list[str]]:
    """올린 파일(base) + 그 파일에 맞춘 줄·머리(base_*) + 지금 줄·머리(new_*) → (고친 파일, 알림)."""
    parsed = rb.parse(base)
    if parsed is None:
        raise PatchError("올린 파일이 회사 견적서 양식이 아닙니다.")
    if parsed["lang"] != new_header["lang"]:
        raise PatchError("올린 파일과 지금 견적서의 언어(국내·해외)가 다릅니다.")
    if [_kind(r) for r in parsed["rows"]] != [_kind(r) for r in base_lines]:
        raise PatchError("올린 파일의 표 줄이 작성 상태와 맞지 않습니다(엑셀에서 표 구조를 바꾼 것 같습니다).")
    en = new_header["lang"] == "en"
    money = qx.USD_CENTS if en and any(round(float(ln.get(k) or 0) * 100) % 100 for ln in new_lines for k in ("unit_price", "amount")) \
        else qx.USD_WHOLE if en else qx.WON
    notes: list[str] = []

    wb = load_workbook(io.BytesIO(base), rich_text=True)    # 글 조각별 서식(흰 글자로 숨긴 글 등)을 지킨다
    ws = wb.worksheets[0]
    auto = _auto_heights(base, ws)
    sh = _Sheet(ws)
    pos = parsed["pos"]
    for k in ("total", "customer", "comments", "by", "date", "no", "delivery", "validity", "note", "order", "bank", "account_no",
              "account_holder"):
        if pos.get(k):
            sh.marks[k] = pos[k]
    for i, r in enumerate(pos.get("foot") or []):
        sh.marks[f"foot{i}"] = r
    for i, r in enumerate(parsed["rownums"]):
        sh.marks[f"row{i}"] = r

    ref_x = qx.build_xlsx(new_header, new_lines)
    ref = load_workbook(io.BytesIO(ref_x))                  # 새 줄 서식·값의 출처(회사 양식)
    ref_ws = ref.worksheets[0]
    ref_p = rb.parse(ref_x)
    ref_rows = ref_p["rownums"]

    def grow(r: int, ref_row: int | None, new: tuple, old: tuple) -> None:
        """줄을 더 바꾼 칸(Alt+Enter, 사용자 2026-10-16) — 행이 낮으면 회사 양식의 같은 줄 높이만큼(더 높은 행은 그대로)."""
        if ref_row is not None and qx.nlines(*new) > qx.nlines(*old):
            _grow(ws, r, ref_ws.row_dimensions[ref_row].height, auto, sh)
    looks = _looks(ws, [_kind(r) for r in base_lines], parsed["rownums"])

    # 표 — base 줄 ↔ 지금 줄
    plan: list[tuple[int | None, int | None]] = []
    sm = difflib.SequenceMatcher(a=[_ident(r) for r in base_lines], b=[_ident(r) for r in new_lines], autojunk=False)
    for op, i1, i2, j1, j2 in sm.get_opcodes():
        if op == "equal":
            plan += [(i1 + k, j1 + k) for k in range(i2 - i1)]
        else:
            plan += [(i, None) for i in range(i1, i2)] + [(None, j) for j in range(j1, j2)]
    for bi in sorted((bi for bi, ni in plan if ni is None), reverse=True):       # 없어진 줄부터 지운다(아래부터)
        sh.delete(sh.marks.pop(f"row{bi}"))
    at = pos["head"] + 1
    final: list[int] = []                                    # 지금 줄 순서대로 행 번호
    for bi, ni in [p for p in plan if p[1] is not None]:
        if bi is not None:
            r = sh.marks[f"row{bi}"]
            _update(ws, r, base_lines[bi], new_lines[ni], en, money, notes)
            grow(r, ref_rows[ni], _texts(new_lines[ni]), _texts(base_lines[bi]))
        else:
            r = max([at, *[x + 1 for x in final]])
            sh.insert(r)
            _copy_row(ref_ws, ref_rows[ni], ws, r)
            if _kind(new_lines[ni]) in looks:
                _dress(ws, r, looks[_kind(new_lines[ni])], ref_ws, ref_rows[ni])
            if new_lines[ni].get("photo"):
                notes.append(f"새 줄 '{str(new_lines[ni].get('desc') or '').splitlines()[0][:20]}'의 사진은 올린 파일에 넣지 못했습니다.")
        final.append(r)
    _formulas(ws, new_lines, final, sh.marks["total"], money)

    # 머리·조건·비고·주석
    b, n = base_header, new_header
    if n["quote_no"] != b["quote_no"] and sh.marks.get("no"):
        _set(ws, sh.marks["no"], 5, n["quote_no"])
    if n["date"] != b["date"] and sh.marks.get("date"):
        _set(ws, sh.marks["date"], 2, date.fromisoformat(n["date"]))
    if sh.marks.get("by") and qx.who_text(n) != qx.who_text(b):
        _set(ws, sh.marks["by"] + 1, 1, qx.who_text(n))
    if sh.marks.get("customer") and qx.customer_text(n) != qx.customer_text(b):
        _set(ws, sh.marks["customer"], 1, qx.customer_text(n))
    if sh.marks.get("comments") and n["comments"] != b["comments"]:
        _set(ws, sh.marks["comments"], 5, n["comments"])
    for k in ("bank", "account_no", "account_holder"):
        if sh.marks.get(k) and n["issuer"].get(k) != b["issuer"].get(k):
            _set(ws, sh.marks[k], 5, n["issuer"].get(k))
    lab = {**qx.pp.QUOTE_LABELS[n["lang"]], **n["labels"]}
    rp = ref_p["pos"]
    if n.get("total_note") != b.get("total_note"):
        _set(ws, sh.marks["total"], 1, "\n".join(x for x in (lab["total"], n.get("total_note") or "") if x))
        grow(sh.marks["total"], rp.get("total"), (lab["total"], n.get("total_note")), (lab["total"], b.get("total_note")))
    if sh.marks.get("delivery"):
        if n["delivery"] != b["delivery"]:
            _set(ws, sh.marks["delivery"], 2, n["delivery"])
        if n["place"] != b["place"]:
            _set(ws, sh.marks["delivery"], 7, n["place"])
        grow(sh.marks["delivery"], rp.get("delivery"), (n["delivery"], n["place"]), (b["delivery"], b["place"]))
    if sh.marks.get("validity") and n["payment"] != b["payment"]:
        _set(ws, sh.marks["validity"], 7, n["payment"])
        grow(sh.marks["validity"], rp.get("validity"), (n["payment"],), (b["payment"],))
    if sh.marks.get("note") and n["notes"] != b["notes"]:
        _set(ws, sh.marks["note"], 2, "\n".join(n["notes"]))
        grow(sh.marks["note"], rp.get("note"), ("\n".join(n["notes"]),), ("\n".join(b["notes"]),))
    _foot(ws, sh, qx.foot_text(b), qx.foot_text(n), len(pos.get("foot") or []))

    last = ws.max_row
    ws.print_area = f"A1:G{last}"
    wb.calculation.fullCalcOnLoad = True
    buf = io.BytesIO()
    wb.save(buf)
    out = _put_shapes(base, _restore_auto_heights(buf.getvalue(), auto, sh), sh, notes)
    return out, notes


def _texts(r: dict[str, Any]) -> tuple:
    """줄 하나에서 칸에 들어가는 글(줄 수를 셀 때)."""
    return tuple(r.get(k) for k in ("desc", "commodity", "remark", "text", "label", "value"))


def _grow(ws, r: int, need: float | None, auto: dict[int, str], sh: _Sheet) -> None:
    """행 높이를 need 까지 늘린다. 자동 높이 행은 다시 쓸 ht 만 키운다(자동 높이 표시는 그대로)."""
    if not need:
        return
    orig = next((o for o in auto if sh.new_row(o) == r), None)
    if orig is not None:
        if float(auto[orig]) < need:
            auto[orig] = f"{need:g}"
    elif (ws.row_dimensions[r].height or 0) < need:
        ws.row_dimensions[r].height = need


def _update(ws, r: int, b: dict[str, Any], n: dict[str, Any], en: bool, money: str, notes: list[str]) -> None:
    """같은 줄 — 바뀐 칸 값만(서식은 올린 파일 그대로)."""
    k = n.get("row")
    if k in ("section", "text"):
        if n.get("text") != b.get("text"):
            _set(ws, r, 1, n.get("text") or "")
        return
    if k == "spec":
        if n.get("label") != b.get("label"):
            _set(ws, r, 1, n.get("label") or "")
        if n.get("value") != b.get("value"):
            _set(ws, r, 3, n.get("value") or "")
        return
    if k == "nego":
        if n.get("desc") != b.get("desc"):
            _set(ws, r, 1, n["desc"])
        if not rb._same_num(n.get("amount"), b.get("amount")):
            _set(ws, r, 6, n["amount"], money)
        return
    if k in ("gap", "subtotal"):
        return
    if (n.get("commodity") or "") != (b.get("commodity") or ""):
        _set(ws, r, 1, n.get("commodity") or "")
    if n["desc"] != b["desc"]:
        _set(ws, r, 2, n["desc"])
    if n["qty"] != b["qty"] or (n.get("unit") or "") != (b.get("unit") or ""):
        _set(ws, r, 4, n["qty"], _qty_fmt(n, en) if (n.get("unit") or "") != (b.get("unit") or "") else None)
    if not rb._same_num(n["unit_price"], b["unit_price"]):
        _set(ws, r, 5, n["unit_price"], money)
    if (n.get("remark") or "") != (b.get("remark") or ""):
        _set(ws, r, 7, n.get("remark") or "")
    if (n.get("photo") or "") != (b.get("photo") or ""):
        notes.append(f"'{str(n['desc']).splitlines()[0][:20]}' 사진 변경은 올린 파일에 반영하지 못했습니다.")
    if (k == "option") != (b.get("row") == "option"):
        _set(ws, r, 6, None if k == "option" else f"=D{r}*E{r}", money)


def _formulas(ws, lines: list[dict[str, Any]], rows: list[int], total_row: int, money: str) -> None:
    """행이 움직인 뒤 금액 수식 — 품목 공급가액(수량×단가)·소계·합계. 사람이 수식 대신 값을 적은 칸은 그대로."""
    billed, block = [], []
    for ln, r in zip(lines, rows):
        k = ln.get("row")
        if k in (None, "item"):
            f = ws.cell(r, 6)
            if f.value is None or (isinstance(f.value, str) and f.value.startswith("=")):
                _set(ws, r, 6, f"=D{r}*E{r}")
            billed.append(f"F{r}")
            block.append(f"F{r}")
        elif k == "nego":
            billed.append(f"F{r}")
        elif k == "subtotal":
            _set(ws, r, 6, f"=SUM({','.join(block)})" if block else 0)
            block = []
        elif k in ("section", "gap"):
            block = [] if k == "section" else block
    _set(ws, total_row, 7, f"=SUM({','.join(billed)})" if billed else 0)


def _foot(ws, sh: _Sheet, old: list[str], new: list[str], n_rows: int) -> None:
    """표 아래 주석 칸 — 바뀐 글만, 줄 수가 달라지면 마지막 주석 행을 본떠 넣거나 지운다."""
    if old == new or not n_rows:
        return
    rows = [sh.marks[f"foot{i}"] for i in range(n_rows)]
    for i in range(min(len(rows), len(new))):
        if i >= len(old) or old[i] != new[i]:
            _set(ws, rows[i], 1, new[i])
    for i in range(len(rows) - 1, len(new) - 1, -1):          # 줄어든 만큼 아래부터 지운다
        sh.delete(rows[i])
    last = rows[-1] if len(new) <= len(rows) else rows[-1]
    for i in range(len(rows), len(new)):
        r = last + (i - len(rows) + 1)
        sh.insert(r)
        _copy_style_row(ws, last, r)
        _set(ws, r, 1, new[i])


def _copy_style_row(ws, sr: int, tr: int) -> None:
    for c in range(1, 8):
        s, t = ws.cell(sr, c), ws.cell(tr, c)
        if s.has_style and type(t).__name__ != "MergedCell":
            t._style = copy.copy(s._style)
    for rng in list(ws.merged_cells.ranges):
        if rng.min_row == sr == rng.max_row:
            ws.merge_cells(start_row=tr, start_column=rng.min_col, end_row=tr, end_column=rng.max_col)
    if ws.row_dimensions[sr].height:
        ws.row_dimensions[tr].height = ws.row_dimensions[sr].height


_SHEET = "xl/worksheets/sheet1.xml"


def _auto_heights(base: bytes, ws) -> dict[int, str]:
    """엑셀이 자동 높이로 둔 행(ht 는 있지만 customHeight 없음) → {원본 행: ht}. openpyxl 은 높이가 있으면 고정 높이로
    저장하므로 여기서 높이를 비우고, 저장한 뒤 _restore_auto_heights 가 ht 만 다시 쓴다
    (실제 고객 파일의 결제조건 두 줄 행이 고정 27pt 가 돼 아래 행·서명 그림 위치가 어긋나던 것, 사용자 2026-10-15)."""
    with zipfile.ZipFile(io.BytesIO(base)) as z:
        if _SHEET not in z.namelist():
            return {}
        xml = z.read(_SHEET).decode("utf-8", "replace")
    out: dict[int, str] = {}
    for m in re.finditer(r"<(?:\w+:)?row\b([^>]*)>", xml):
        attrs = m.group(1)
        r, ht = re.search(r'\br="(\d+)"', attrs), re.search(r'\bht="([^"]+)"', attrs)
        if r and ht and not re.search(r'customHeight="(?:1|true)"', attrs):
            out[int(r.group(1))] = ht.group(1)
            ws.row_dimensions[int(r.group(1))].height = None
    return out


def _restore_auto_heights(data: bytes, heights: dict[int, str], sh: _Sheet) -> bytes:
    """저장한 파일의 자동 높이 행에 원래 ht 를 customHeight 없이 다시 쓴다(행 이동 반영)."""
    want = {r: ht for o, ht in heights.items() if (r := sh.new_row(o)) is not None}
    if not want:
        return data
    zi = zipfile.ZipFile(io.BytesIO(data))
    if _SHEET not in zi.namelist():
        return data

    def fix(m: re.Match) -> str:
        r = re.search(r'\br="(\d+)"', m.group(0))
        ht = want.get(int(r.group(1))) if r else None
        if ht is None or " ht=" in m.group(0):
            return m.group(0)
        return m.group(0)[:-1].rstrip("/") + f' ht="{ht}"' + ("/>" if m.group(0).endswith("/>") else ">")

    xml = re.sub(r"<row\b[^>]*>", fix, zi.read(_SHEET).decode("utf-8"))
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as z:
        for item in zi.infolist():
            z.writestr(item, xml.encode("utf-8") if item.filename == _SHEET else zi.read(item.filename))
    return buf.getvalue()


# ── 도형 되살리기 ─────────────────────────────────────────────────────────────

def _drawing_of_sheet(z: zipfile.ZipFile) -> str | None:
    rels = "xl/worksheets/_rels/sheet1.xml.rels"
    if rels not in z.namelist():
        return None
    m = re.search(r'Target="([^"]*drawings/drawing\d+\.xml)"', z.read(rels).decode("utf-8", "replace"))
    if not m:
        return None
    t = m.group(1)                                           # 엑셀 '../drawings/…', openpyxl '/xl/drawings/…'
    return t.lstrip("/") if t.startswith("/") else "xl/" + t.replace("../", "")


def _put_shapes(base: bytes, out: bytes, sh: _Sheet, notes: list[str]) -> bytes:
    """원본 그림판의 도형·텍스트 상자(그림·차트 아님)를 행 이동만큼 옮겨 새 파일 그림판에 다시 넣는다.
    엑셀은 'xdr:' 접두어, openpyxl 은 접두어 없이 쓴다 — 도형마다 이름공간을 직접 달아 어느 쪽에 넣어도 같은 뜻이 되게."""
    zb = zipfile.ZipFile(io.BytesIO(base))
    src = _drawing_of_sheet(zb)
    if not src:
        return out
    shapes: list[tuple[str, str]] = []
    for m in _SHAPE_RE.finditer(zb.read(src).decode("utf-8")):
        a, pre = m.group(0), m.group(1)
        if re.search(rf"<{pre}(?:pic|graphicFrame)\b", a):
            continue                                         # 그림은 openpyxl 이 이미 옮겼다, 차트는 다루지 않는다
        end = a.index(">")
        decl = "".join(f' xmlns:{k}="{v}"' for k, v in _NS.items()
                       if f"xmlns:{k}=" not in a[:end] and (k != "xdr" or pre == "xdr:"))
        if not pre and "xmlns=" not in a[:end]:
            decl += f' xmlns="{_NS["xdr"]}"'
        shapes.append((a[:end] + decl + a[end:], pre))
    if not shapes:
        return out
    zo = zipfile.ZipFile(io.BytesIO(out))
    dst = _drawing_of_sheet(zo)
    if not dst:
        notes.append(f"도형·텍스트 상자 {len(shapes)}개를 다시 넣지 못했습니다(그림판 없음).")
        return out

    def move(a: str, pre: str) -> str | None:
        row_re = re.compile(rf"(<{pre}row>)(\d+)(</{pre}row>)")
        rows = [int(x.group(2)) + 1 for x in row_re.finditer(a)]
        if rows and sh.new_row(rows[0]) is None:
            return None                                      # 지운 행에 있던 도형
        return row_re.sub(lambda m: f"{m.group(1)}{(sh.new_row(int(m.group(2)) + 1) or int(m.group(2)) + 1) - 1}{m.group(3)}", a)

    moved = [m for m in (move(a, pre) for a, pre in shapes) if m]
    if len(moved) < len(shapes):
        notes.append(f"지운 줄에 있던 도형 {len(shapes) - len(moved)}개는 뺐습니다.")
    xml = zo.read(dst).decode("utf-8")
    close = re.search(r"</(?:\w+:)?wsDr>\s*$", xml)
    if not close:
        notes.append("도형·텍스트 상자를 다시 넣지 못했습니다.")
        return out
    xml = xml[: close.start()] + "".join(moved) + xml[close.start():]
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as z:
        for item in zo.infolist():
            z.writestr(item, xml.encode("utf-8") if item.filename == dst else zo.read(item.filename))
    return buf.getvalue()
