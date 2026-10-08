"""수주 뒤 진행(sales_deals/flow.py · statement_xlsx.py) 점검 — 결제 조건(회차 %)의 단계·출하 조건·지남, 분기 수주·매출, 리스트 엑셀, 거래명세서.

    docker compose run --rm --no-deps -e PYTHONPATH=/app -v ./backend/app:/app/app -v ./scripts:/scripts -v ./docs:/docs backend \\
        python /scripts/test_sales_flow.py
"""
from __future__ import annotations

import io
import sys

import openpyxl

from app.sales_deals import flow
from app.sales_deals import statement_xlsx as sx

PASS: list[str] = []
FAIL: list[tuple[str, str]] = []


def check(ok: bool, label: str, detail: str = "") -> None:
    (PASS.append(label) if ok else FAIL.append((label, detail)))


SPLIT = [{"label": "선금", "pct": 30, "expected_date": "2026-09-10"},
         {"label": "중도금", "pct": 30, "expected_date": None},
         {"label": "잔금", "pct": 40, "expected_date": "2026-10-31"}]


def deal(terms: list | None, confirmed: tuple = (), shipped: bool = False, statement: bool = True) -> dict:
    return {"kind": "quote", "won": True, "won_date": "2026-09-01", "customer": "(주)다라전자",
            "pay_case": "terms" if terms else None, "pay_terms": flow.clean_terms(terms) if terms else [],
            "quote": {"amount": 3_700_000, "vat_included": False, "currency": "KRW", "date": "2026-08-20"},
            "statement": {"date": "2026-09-02", "buyer": {"name": "㈜다라전자"},
                          "totals": {"supply": 3_700_000, "vat": 370_000, "total": 4_070_000}} if statement else None,
            "payments": [{"amount": 1, "term": i, "date": "2026-09-09"} for i in confirmed],
            "shipments": [{"date": "2026-09-20"}] if shipped else [],
            "shipped_note": False, "paid_full": False, "invoices": []}


def keys(d: dict) -> list[tuple[str, bool, bool]]:
    return [(s["key"], s["done"], s["ready"]) for s in flow.steps(d)]


# 1) 금액은 공급가액(사용자 2026-10-08 — 부가세는 거래명세서에서만)
check(flow.expected_total(deal(SPLIT)) == 3_700_000, "받을 금액 = 공급가액(명세서 공급가액)")
check(flow.supply_total({**deal(SPLIT, statement=False), "quote": {"amount": 4_070_000, "vat_included": True}}) == 3_700_000,
      "부가세 포함 견적이면 ÷1.1")
ts = flow.terms(deal(SPLIT))
check([t["amount"] for t in ts] == [1_110_000, 1_110_000, 1_480_000] and sum(t["amount"] for t in ts) == 3_700_000,
      "회차 금액 = 공급가액 × %(끝전은 마지막 회차)", str([t["amount"] for t in ts]))

# 2) 결제 조건 검사
for bad, why in ((SPLIT[:2], "합 100"), ([{"label": "", "pct": 100}], "구분"),
                 ([{"label": "선금", "pct": 100, "expected_date": "10월"}], "날짜"), ([], "빈 조건")):
    try:
        flow.clean_terms(bad)
        check(False, f"결제 조건 거절: {why}")
    except ValueError:
        check(True, f"결제 조건 거절: {why}")
check(flow.clean_terms([{"label": "선금", "pct": 60, "when": "  발주   시 "}, {"label": "잔금", "pct": 40}])
      == [{"label": "선금", "pct": 60, "when": "발주 시", "expected_date": None}, {"label": "잔금", "pct": 40, "when": "", "expected_date": None}],
      "시점은 자유 글(공백 정리), 안 적으면 빈칸 — 출하 조건과는 무관")

# 3) 단계 — 출하·입금 확인은 순서 없음(사용자 2026-10-08: 출하는 입금 확인 없이도)
d = deal(SPLIT)
check(keys(d) == [("statement", True, False), ("case", True, False), ("ship", False, True),
                  ("pay:0", False, True), ("pay:1", False, True), ("pay:2", False, True)],
      "단계: 명세서 → 조건 → 출하·회차 입금(모두 바로 가능)", str(keys(d)))
check(flow.next_due(d) == 1_110_000, "다음 확인할 금액 = 첫 미확인 회차(선금)")
d = deal(SPLIT, confirmed=(2,))
check(("pay:2", True, False) in keys(d) and flow.next_due(d) == 1_110_000, "잔금 먼저 확인해도 됨(순서 없음)")
check(("ship", False, True) in keys(deal(None)), "결제 조건 전이라도 명세서만 있으면 출하 가능")
check(flow.steps({**deal(SPLIT), "won": None}) == [], "수주 확인 전에는 진행 단계 없음")
check(flow.steps({**deal(SPLIT), "dropped": True}) == [], "드랍한 건은 진행 단계 없음")
check(flow.steps({**deal(None, statement=False), "imported": True}) == [] and flow.steps({**deal(SPLIT), "imported": True}),
      "엑셀에서 옮긴 건은 명세서·결제 조건을 시작한 경우만 새 절차")
check(keys(deal(None, statement=False))[:2] == [("statement", False, True), ("case", False, False)], "명세서 발급이 먼저")

# 4) 예정 입금일 지남(팝업·리스트 강조) · 단계
from app.sales_deals.status import derive  # noqa: E402
from datetime import date  # noqa: E402
od = flow.overdue(deal(SPLIT), date(2026, 9, 15))
check([(t["label"], t["days"]) for t in od] == [("선금", 5)], "선금 예정일(9/10) 지남 5일", str(od))
check(flow.overdue(deal(SPLIT, confirmed=(0,)), date(2026, 9, 15)) == [], "확인한 회차는 지남 아님")
v = derive(deal(SPLIT), date(2026, 9, 15))
check(any(a["code"] == "OVERDUE" for a in v["alerts"]), "지난 회차는 알림(OVERDUE)")
check(derive({**deal(SPLIT), "dropped": True}, date(2026, 9, 15))["stage"] == "드랍"
      and not any(a["code"] == "OVERDUE" for a in derive({**deal(SPLIT), "dropped": True}, date(2026, 9, 15))["alerts"]),
      "드랍 단계, 드랍한 건은 지남 알림 없음")
v = derive(deal(SPLIT, shipped=True), date(2026, 9, 5))
check(v["stage"] == "출하 후 입금 대기" and not any(a["code"] == "SHIP_BEFORE_PAY" for a in v["alerts"]),
      "입금 확인 없이 출하해도 경고 없음(출하 후 입금 대기)", str(v["alerts"]))
check(derive({**deal(SPLIT, confirmed=(0, 1, 2)), "paid_full": True}, date(2026, 10, 7))["stage"] == "수주"
      and derive({**deal(SPLIT, confirmed=(0, 1, 2), shipped=True), "paid_full": True}, date(2026, 10, 7))["stage"] == "완료 (출하·완납)",
      "결제 조건 건: 완납이어도 출하 전에는 '완료'가 아님")

# 5) 분기별 수주금액·올해 매출(출하일 기준) — 공급가액, 드랍 제외(예전 실주 기록 포함)
from app.sales_deals.service import kpi, NEW_NO  # noqa: E402
rows = [{**deal(SPLIT, shipped=True), "won_date": "2026-02-10"},
        {**deal(SPLIT), "won_date": "2026-08-01"},
        {**deal(SPLIT, shipped=True), "won_date": "2026-05-01", "dropped": True},
        {**deal(SPLIT), "won_date": "2025-12-01", "shipments": [{"date": "2026-01-05"}]},
        {**deal(SPLIT), "won": False, "won_date": "2026-03-01"}]
k = kpi(rows, date(2026, 10, 8))
check(k["won_by_quarter"] == [3_700_000, 0, 3_700_000, 0] and k["won_total"] == 7_400_000 and k["sales"] == 7_400_000
      and k["quarter"] == 4 and k["won_estimated"] == 0, "분기별 수주·올해 매출(출하일 기준, 작년 수주 올해 출하 포함, 드랍 제외(예전 실주 기록 포함))", str(k))
legacy = {**deal(None, statement=False), "won_date": None, "quote": None, "imported": True,
          "invoices": [{"date": "2026-05-03", "supply": 2_000_000}, {"date": "2026-07-01", "supply": 1_000_000}]}
k = kpi([legacy], date(2026, 10, 8))
check(k["won_by_quarter"] == [0, 3_000_000, 0, 0] and k["won_estimated"] == 1,
      "수주일 없는 옮긴 건: 첫 계산서 발행일 분기 · 공급가액 = 계산서 공급가액 합, 대신 셈 1건", str(k))
check(NEW_NO.match("AL20261008") and NEW_NO.match("AL20261008B") and not NEW_NO.match("S26-0182"), "새 건번호 형식(이니셜+날짜)")

# 6) 리스트 엑셀
import openpyxl as _ox  # noqa: E402
from app.sales_deals.export import build_xlsx  # noqa: E402
rr = [{**deal(SPLIT), "deal_no": "AL20261008", "display_no": "AL20261008-03", "title": "ATC", "owner": "Alex",
       "base_date": "2026-08-20", "stage": "수주", "alerts": [{"text": "선금 지남"}],
       "flow": {"terms": flow.terms(deal(SPLIT)), "pay_case_label": "선금 30%", "paid": 0, "overdue": [1]},
       "customer": "한빛플라스틱",
       # 히스토리(최신 먼저) — 엑셀에 전부 들어가야 함(사용자 2026-10-08)
       "history": [
           {"ver_no": "AL20261010-03", "action": "payment", "detail": {"note": "선금 30% 입금 확인", "date": "2026-10-10", "amount": 1_110_000},
            "created_at": "2026-10-10T01:00:00+00:00", "user_name": "Alex", "stage": "수주", "quote_rev": 2, "amount": 3_700_000, "superseded": False},
           {"ver_no": "AL20261008-02", "action": "quote_revision", "detail": {"revision": 2, "amount": 3_700_000},
            "created_at": "2026-10-08T05:00:00+00:00", "user_name": "Alex", "stage": "견적", "quote_rev": 2, "amount": 3_700_000, "superseded": False},
           {"ver_no": "AL20261008-01", "action": "quote", "detail": {"revision": 1, "amount": 3_500_000},
            "created_at": "2026-10-08T01:00:00+00:00", "user_name": "Alex", "stage": "견적", "quote_rev": 1, "amount": 3_500_000, "superseded": True}]}]
wb = _ox.load_workbook(io.BytesIO(build_xlsx({"rows": rr, "kpi": k, "summary": {"수주": 1}, "today": "2026-10-08"})))
ws = wb["영업 건"]
check(ws["A2"].value == "한빛플라스틱" and ws["C2"].value == "AL20261008-03" and ws["F2"].value == 3_700_000 and ws["K2"].value == "2026-09-10"
      and ws["A2"].fill.fgColor.rgb.endswith("FDE2E1") and "요약" in wb.sheetnames, "엑셀: 프로젝트 먼저·표시 번호·공급가액·다음 예정일·지남 강조")
check([ws[f"C{i}"].value for i in (3, 4, 5)] == ["AL20261010-03", "AL20261008-02", "AL20261008-01"]
      and all(ws.row_dimensions[i].outlineLevel == 1 for i in (3, 4, 5)) and "입금 확인" in ws["B3"].value
      and ws["E3"].value == "2026-10-10 10:00" and ws["G5"].value == "견적 V1 (버전 업데이트됨)",
      "엑셀 시트1: 프로젝트 아래 히스토리 3줄(그룹으로 접힘, 한국 시각, 옛 판 표시)", str([ws[f"G{i}"].value for i in (3, 4, 5)]))
hs = wb["히스토리 전체"]
check(hs.max_row == 4 and hs["D2"].value == "AL20261010-03" and hs["A2"].value == "한빛플라스틱" and hs["G3"].value == "견적 V2"
      and hs["J4"].value == "버전 업데이트됨", "엑셀 시트2: 히스토리 전체 한 표", str([c.value for c in hs[4]]))

# 7) 거래명세서
check(sx.split_name("TCV1\n- MASTER T.C (Payload 10Kg)") == ("TCV1 MASTER T.C", "Payload 10Kg")
      and sx.split_name("PPM\n-Pogo Pin Male") == ("PPM", "Pogo Pin Male"), "견적 품목 → 품명·규격(양식 예시와 같게)")
items = [{"name": "TCV1 MASTER T.C", "spec": "Payload 10Kg", "qty": 1, "unit_price": 2_300_000},
         {"name": "PPM", "spec": "Pogo Pin Male", "qty": 2, "unit_price": 250_000}]
check(sx.totals(items) == {"supply": 2_800_000, "vat": 280_000, "total": 3_080_000}, "합계·부가세 계산")
try:
    sx.check({"buyer": {"name": "x"}, "items": items})
    check(False, "공급받는자 등록번호·주소 없으면 거절")
except sx.StatementError as e:
    check("등록번호" in str(e) and "주소" in str(e), "공급받는자 등록번호·주소 없으면 거절", str(e))
x = sx.build_xlsx({"no": "S26-0182", "date": "2026-09-02", "items": items * 6,
                   "buyer": {"reg_no": "314-86-54321", "name": "㈜다라전자", "ceo": "홍길동", "address": "경기도 포천시"}})
ws = openpyxl.load_workbook(io.BytesIO(x)).active
vals = [c.value for row in ws.iter_rows() for c in row if c.value is not None]
check(vals.count("거래명세서 (납 품 표)") == 2 and "공급자보관용" in vals and "공급받는자보관용" in vals
      and vals.count("314-86-54321") == 2 and ws.page_setup.fitToHeight == 1, "두 장(공급자·공급받는자 보관용)·12줄도 한 장에")

# 세트로 묶어서 발급(사용자 2026-10-08) — 한 줄 'TCV1 툴체인저(ATC) 1 SET', 공급가액은 그대로
set_items, parts = sx.bundle(items, "TCV1 툴체인저(ATC)", 1)
check(len(set_items) == 1 and set_items[0]["unit"] == "SET" and set_items[0]["unit_price"] == 2_800_000 and len(parts) == 2
      and sx.totals(set_items) == sx.totals(items), "세트로 묶기: 한 줄 + SET, 합계 그대로", str(set_items))
check(sx.bundle(items, "세트", 2)[0][0]["unit_price"] == 1_400_000, "세트 2개면 단가 = 합계 ÷ 2")
check(sx.unit_of({"unit": 'SE"T'}) == "SET" and sx.unit_of({}) == "", "단위는 글자·숫자만(서식 깨짐 방지)")
x = sx.build_xlsx({"no": "AL20261008", "date": "2026-10-08", "items": set_items, "parts": parts,
                   "buyer": {"reg_no": "314-86-54321", "name": "㈜다라전자", "ceo": "홍길동", "address": "경기도 포천시"}})
ws = openpyxl.load_workbook(io.BytesIO(x)).active
qty_cells = [c for row in ws.iter_rows() for c in row if c.value == 1 and c.column_letter == "F"]
names = [c.value for row in ws.iter_rows() for c in row if isinstance(c.value, str) and "PPM" in c.value]
check(len(qty_cells) == 2 and all('" SET"' in c.number_format for c in qty_cells) and not names,
      "세트 명세서: 수량 칸 '1 SET'(숫자 그대로), 구성 품목은 양식에 안 나옴", str([c.number_format for c in qty_cells]))

print(f"PASS {len(PASS)} / FAIL {len(FAIL)}")
for label, detail in FAIL:
    print(f"  ✗ {label}: {detail}")
sys.exit(1 if FAIL else 0)
