"""올린 견적서 파일 → 작성 상태 되읽기(company_knowledge/quote_readback.py) 점검 — 사용자 2026-10-15(고객 지적:
AI 01 견적서를 엑셀에서 고쳐 같은 번호로 교체한 뒤 [견적서 고치기]·미리보기가 예전 01 로 나옴).

    MSYS_NO_PATHCONV=1 docker compose run --rm --no-deps -e PYTHONPATH=/app -v ./backend/app:/app/app -v ./scripts:/scripts \\
        -v ./docs:/docs backend python /scripts/test_quote_readback.py

방식: 지금 상태 S 로 만든 견적서를 '사람이 고친 결과' T 로 바꿔 파일을 만들고(build(T)), S 에 되읽어(merge) 얻은 S' 로 다시
만든 견적서가 T 의 견적서와 칸 값이 같은지 본다. 엑셀에서 값만 고친 파일(openpyxl 로 칸 수정)도 따로 본다.
"""
from __future__ import annotations

import copy
import io
import sys
from datetime import date

import openpyxl

from app.company_knowledge import product_prices as pp
from app.company_knowledge import quote_readback as rb
from app.company_knowledge import quote_session as qs
from app.company_knowledge import quote_xlsx as qx

PASS: list[str] = []
FAIL: list[tuple[str, str]] = []
DAY = date(2026, 10, 15)


def check(ok: bool, label: str, detail: str = "") -> None:
    (PASS.append(label) if ok else FAIL.append((label, detail)))


def build(state: dict) -> bytes:
    f = state["form"]
    lang = f.get("lang") or "ko"
    lines = qs.render_lines(state, lang, f.get("fx_rate"))
    pogo = any(ln.get("key") in ("PPM", "PPF") for ln in lines if not ln.get("row"))
    return qx.build_xlsx(qx.quote_header(f, lang=lang, quote_no="QT-LU20261015-01", day=DAY, pogo=pogo), lines)


def edit(state: dict, payload: dict) -> dict:
    new, _ = qs.apply_ops(copy.deepcopy(state), qs.edit_ops(state, payload), None)
    return new


def same_file(a: bytes, b: bytes) -> tuple[bool, str]:
    pa, pb = rb.parse(a), rb.parse(b)
    if pa["rows"] != pb["rows"]:
        diff = [(x, y) for x, y in zip(pa["rows"], pb["rows"]) if x != y][:2] or (len(pa["rows"]), len(pb["rows"]))
        return False, f"줄 다름 {diff}"
    if pa["form"] != pb["form"]:
        return False, f"머리 다름 {[(k, pa['form'].get(k), pb['form'].get(k)) for k in pa['form'] if pa['form'].get(k) != pb['form'].get(k)]}"
    return True, ""


# ── 기본 상태: 툴체인저 AI 견적(본체·액세서리) ──
book = pp.parse_book()
items = [
    {"name": "TCV1 MASTER T.C", "code": None, "qty": 1, "price_model": "TCV1", "price_item": "MASTER T.C", "optional": False},
    {"name": "TCV1 T.P", "code": None, "qty": 2, "price_model": "TCV1", "price_item": "T.P", "optional": False},
    {"name": "Pogo Pin Male (PPM)", "code": "PPM", "qty": 2, "price_model": "TCV1", "price_item": "PPM/PPF", "optional": False},
    {"name": "Pogo Pin Female (PPF)", "code": "PPF", "qty": 2, "price_model": "TCV1", "price_item": "PPM/PPF", "optional": False},
    {"name": "Cable 0.3m", "code": None, "qty": 4, "price_model": "TCV1", "price_item": "Cable 0.3m", "optional": False},
    {"name": "PneuMatic Male (PMM)", "code": "PMM", "qty": 1, "price_model": "TCV1", "price_item": "PMM/PMF", "optional": False},
]
pp.attach_prices(items, book)
S = qs.draft_state({"applicable": True, "items": items}, model="TCV1", payload_kg=10, form={
    "customer": "가나테크", "to": "김철수 대표님", "contact_name": "홍길동", "contact_title": "이사", "initials": "LU",
    "contact_mobile": "010-0000-0000", "contact_email": "name@example.com"})
S["delivery"] = {"case": 1, "freight": None}
S["asked"] = ["extra"]

# 1. 고치지 않은 파일은 바뀐 점 없음
p0 = rb.parse(build(S))
check(p0 is not None and p0["lang"] == "ko" and p0["quote_no"] == "QT-LU20261015-01", "회사 양식 인식·발행번호 읽기", str(p0 and p0["quote_no"]))
S1, ch, wn = rb.merge(S, p0)
check(ch == [] and wn == [] and S1["lines"] == S["lines"], "고치지 않은 파일 → 바뀐 점 없음", f"{ch} {wn}")
check([x["name"] for x in rb.items_of(p0)][:2] == ["Magbot ATC", "Magbot ATC"] and len(rb.items_of(p0)) == 6,
      "확인 창 품목 = 표의 품목 줄 그대로(AI 없이)", str(rb.items_of(p0)[:2]))

# 2. 사람이 고친 견적서 T — 수량·단가·비고·규격·품목·추가·삭제·머리·조건·주석·NEGO
T = copy.deepcopy(S)
ids = {ln["key"]: ln["id"] for ln in T["lines"]}
T = edit(T, {"form": {"customer": "가나테크(주)", "to": "박영희 과장님", "comments": "www.magbot.kr\n- 툴체인저 공급 건(수정)",
                      "payment": "발주시 30%, 납품전 70%", "note": "설치 별도 협의", "total_note": "* 부가세 별도",
                      "foot_notes": ["납기는 발주 후 협의", "운송 보험 별도"]},
             "lines": [{**{k: v for k, v in ln.items() if k in ("id", "name", "commodity", "qty", "unit_price", "remark")},
                        **({"qty": 3} if ln["id"] == ids["T.P"] else {}),
                        **({"unit_price": 450000, "remark": "포고핀 8핀(특가)"} if ln["id"] == ids["PPM"] else {}),
                        **({"commodity": "Magbot ATC Option"} if ln["id"] == ids["Cable 0.3m"] else {})}
                       for ln in T["lines"] if ln["id"] != ids["PMM"]]
                      + [{"name": "출장 교육(1일)", "commodity": "Service", "qty": 1, "unit_price": 300000, "remark": "현장"}],
             "nego": {"mode": "percent", "value": 10, "label": ""}})
fileT = build(T)
wb = openpyxl.load_workbook(io.BytesIO(fileT))                 # 본체 규격 글을 직접 고친 것(코드가 만들던 글)
ws = wb.active
r = next(r for r in range(1, ws.max_row + 1) if str(ws.cell(r, 2).value or "").startswith("TCV1\n- MASTER"))
ws.cell(r, 2).value = "TCV1\n- MASTER T.C (Payload 10Kg)\n- Control Box 포함"
buf = io.BytesIO()
wb.save(buf)
fileT = buf.getvalue()
S2, ch, wn = rb.merge(S, rb.parse(fileT))
ok, why = same_file(build(S2), fileT)
check(ok, "고친 견적서를 되읽은 상태로 다시 만들면 고친 파일과 칸 값이 같음", why)
master = next(ln for ln in S2["lines"] if ln["id"] == ids["MASTER T.C"])
check(master["key"] == "custom" and "Control Box 포함" in master["name"], "코드가 만들던 규격을 고치면 적힌 글 그대로(이후 유지)", str(master))
check(next(ln for ln in S2["lines"] if ln["id"] == ids["T.P"])["qty"] == 3 and not any(ln["id"] == ids["PMM"] for ln in S2["lines"]),
      "같은 줄은 id 유지(수량만 변경), 지운 줄은 빠짐")
check(any(ln.get("name") == "출장 교육(1일)" and ln["unit_price"] == 300000 for ln in S2["lines"]), "새로 넣은 줄은 새 품목으로")
check(S2["nego"] == {"mode": "percent", "value": 10.0, "label": ""}, "NEGO 10% 를 %로 되읽음", str(S2["nego"]))
check(S2["form"]["payment"] == "발주시 30%, 납품전 70%" and S2["form"]["note"] == "설치 별도 협의"
      and S2["form"]["foot_notes"] == ["납기는 발주 후 협의", "운송 보험 별도"] and S2["form"]["total_note"] == "* 부가세 별도",
      "조건·비고(포고핀 자동 문구 제외)·주석·Total 아래 문구", str({k: S2["form"][k] for k in ("payment", "note", "foot_notes", "total_note")}))
check(any("빠진 품목" in c for c in ch) and any("새 품목" in c for c in ch) and any("NEGO" in c for c in ch),
      "바뀐 점 목록(빠진·새 품목·NEGO)", str(ch))
check(qs.problems(S2) == [], "되읽은 상태는 바로 다시 발행 가능", str(qs.problems(S2)))

# 3. 묶음 — 사용자 묶음(소제목·소계)·< OPTION >·사양 묶음을 엑셀에서 만든 견적서
B = edit(S, {"blocks": [{"id": "B0", "type": "items"}, {"id": "t1", "type": "items", "title": "설치 및 시운전", "subtotal": True},
                        {"id": "t2", "type": "option", "title": ""},
                        {"id": "t3", "type": "spec", "title": "사양", "rows": [{"label": "전원", "value": "DC 24V"}]}],
             "lines": [*[{k: v for k, v in ln.items() if k in ("id", "name", "commodity", "qty", "unit_price", "remark")} for ln in S["lines"]],
                       {"name": "시운전", "qty": 1, "unit_price": 500000, "block": "t1"},
                       {"name": "IB 브라켓", "qty": 1, "unit_price": 100000, "block": "t2"}]})
fileB = build(B)
S3, ch, wn = rb.merge(S, rb.parse(fileB))
ok, why = same_file(build(S3), fileB)
check(ok, "묶음(소제목·Sub total·OPTION·사양) 견적서도 다시 만들면 같음", why)
check([b["type"] for b in qs.blocks(S3)] == ["items", "items", "option", "spec"], "묶음 종류 되읽음", str(qs.blocks(S3)))
check(qs.subtotal(S3) == qs.subtotal(B), "옵션 제외 합계 같음", f"{qs.subtotal(S3)} vs {qs.subtotal(B)}")

# 4. 엑셀에서 값만 고친 파일(행 그대로, 칸만 수정)
wb = openpyxl.load_workbook(io.BytesIO(build(S)))
ws = wb.active
head = next(r for r in range(1, ws.max_row + 1) if ws.cell(r, 1).value == "Commodity")
ws.cell(head + 2, 4).value = 5                                # T.P 수량
ws.cell(head + 3, 5).value = 480000                           # PPM 단가
ws.cell(head + 1, 7).value = "컨트롤박스·4m 케이블 포함(무상 1년 보증)"
ws.cell(next(r for r in range(1, head) if ws.cell(r, 4).value == "Comments"), 5).value = "www.magbot.kr\n- 수정 공급 건"
buf = io.BytesIO()
wb.save(buf)
S4, ch, wn = rb.merge(S, rb.parse(buf.getvalue()))
tp = next(ln for ln in S4["lines"] if ln["id"] == ids["T.P"])
ppm = next(ln for ln in S4["lines"] if ln["id"] == ids["PPM"])
check(tp["qty"] == 5 and ppm["unit_price"] == 480000 and "무상 1년" in S4["lines"][0]["remark"]
      and S4["form"]["comments"] == "www.magbot.kr\n- 수정 공급 건", "엑셀에서 칸 값만 고친 파일", str(ch))
check(len(ch) == 4, "바뀐 점 4개(수량·단가·비고·Comments)", str(ch))

# 5. 물류비·설치 인건비 — 금액·수량만 납품 방식 값으로, 이름 변경은 경고
W = copy.deepcopy(S)
W["delivery"] = {"case": 4, "freight": 150000}
W["work"] = {"needed": True, "people": 2, "days": 3, "rate": 800000}
qs._sync_work(W)
W2 = copy.deepcopy(W)
W2["delivery"]["freight"] = 200000
W2["work"]["rate"] = 700000
qs._sync_work(W2)
S5, ch, wn = rb.merge(W, rb.parse(build(W2)))
check(S5["delivery"]["freight"] == 200000 and S5["work"]["rate"] == 700000 and wn == [], "물류비·인건비 단가 변경 → 납품 방식 값", f"{ch} {wn}")

# 6. 해외(영문) — 달러 단가는 usd 로
E = edit(S, {"form": {"lang": "en", "fx_rate": 1380, "contact_name_en": "Gildong Hong", "contact_honorific": "Mr."}})
E2 = copy.deepcopy(E)
E2["lines"][0]["usd"] = 1500.0
S6, ch, wn = rb.merge(E, rb.parse(build(E2)))
check(rb.parse(build(E2))["lang"] == "en" and S6["lines"][0].get("usd") == 1500.0, "해외 견적 달러 단가 → usd", f"{ch} {wn}")
S7, ch, wn = rb.merge(S, rb.parse(build(E2)))
check(ch == [] and wn and S7 == S, "언어가 다른 파일은 작성 상태를 바꾸지 않고 알림", str(wn))

# 7. 회사 양식이 아닌 파일 → None, 확인한 품목만 맞추기
other = openpyxl.Workbook()
other.active.append(["품명", "수량", "단가"])
buf = io.BytesIO()
other.save(buf)
check(rb.parse(buf.getvalue()) is None and rb.parse(b"not a file") is None, "회사 양식이 아닌 파일·깨진 파일 → None")
conf = [{"name": x["name"], "spec": x["spec"], "qty": x["qty"], "unit_price": x["unit_price"]} for x in rb.items_of(p0)]
conf[1]["qty"] = 4
S8, ch, wn = rb.merge_items(S, conf)
check(next(ln for ln in S8["lines"] if ln["id"] == ids["T.P"])["qty"] == 4 and len(ch) == 1 and wn,
      "다른 양식·PDF: 확인한 품목의 수량·단가만, 규격 글은 그대로", f"{ch} {wn}")

# 글자로 적은 수량·단가(사용자 2026-10-16, 고객 파일 변형본 13번)
check([rb._num(x) for x in ("3 ea", "2SET", "5개", "₩400,000", "-500,000", "1,000.5")] == [3, 2, 5, 400000, -500000, 1000.5]
      and rb._num("TCV1") is None and rb._num("") is None, "글자로 적은 수량·단가를 숫자로", "")

print(f"PASS {len(PASS)} / FAIL {len(FAIL)}")
for label, detail in FAIL:
    print(f"  ✗ {label}: {detail}")
sys.exit(1 if FAIL else 0)
