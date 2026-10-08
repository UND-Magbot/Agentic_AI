"""제안서 작업 3~6단계(공정 제안·컨셉 이미지·구성·제작·바로 만들기) 검증.

- 기본: 제안 검증(근거 없는 수치·설계 문항만)·그림 지시·비전 검수 파싱·견적/구성/라벨 검증 — DB 불필요.
- --db : 실제 DB·저장소로 전 단계 진행. 모델·Codex 는 가짜로 바꿔 끼운다(PPT 조판은 실제 5c 코드).
    docker compose run --rm --no-deps -v ./scripts:/scripts backend python /scripts/test_proposal_workflow.py --db
"""
import asyncio
import io
import json
import sys
from pathlib import Path

sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8", errors="replace")
ROOT = Path(__file__).resolve().parents[1]
for cand in (ROOT / "backend", Path("/app")):
    if (cand / "app").is_dir():
        sys.path.insert(0, str(cand))
        break

from app.proposal_project import concept, workflow  # noqa: E402
from app.proposal_project.catalog import CODES  # noqa: E402
from app.proposal_project.service import ProjectError  # noqa: E402

PASS: list[str] = []
FAIL: list[tuple[str, str]] = []


def check(ok: bool, label: str, detail: str = "") -> None:
    (PASS.append(label) if ok else FAIL.append((label, detail)))


def expect_error(label: str, fn) -> None:
    try:
        fn()
        check(False, label)
    except ProjectError:
        check(True, label)


def items(**vals):
    it = {c: {"value": "", "status": "empty", "source": "", "evidence": "", "unit": ""} for c in CODES}
    for c, (v, st) in vals.items():
        it[c].update(value=v, status=st)
    return it


PROJECT = {"items": items(P06=("500ml 유리 보틀 2종", "confirmed"), P08=("시간당 17병", "confirmed"),
                          P14=("협동로봇(가반하중 14kg) 검토 중", "assumed"), P21=("흡착 그리퍼", "confirmed"))}

PROPOSAL = {
    "alternatives": [
        {"name": "협동로봇 셀", "robot": "협동로봇 1대(14kg)", "summary": "시간당 17병 처리", "reason": "고객 검토안",
         "structure": ["협동로봇 1대가 컨베이어 옆에 선다", "비전 카메라 1대가 컨베이어 위에 있다",
                       "주입 노즐 250mm 앞에 둔다"],       # 250 은 근거 없음 → 이 문장만 버림
         "exclude": ["AMR"]},
        {"name": "속도 셀", "robot": "고속 로봇 3000mm/s", "summary": "", "reason": "",
         "structure": ["속도 3000mm/s"]},                   # 구성 문장이 전부 버려져 대안째 빠짐
        {"name": "스카라 셀", "robot": "스카라 2대", "summary": "", "reason": "공간 우선",
         "structure": ["스카라 2대가 좌우에 배치된다"]},
    ],
    "fills": [{"code": "P21", "value": "다른 그리퍼", "reason": "x"},        # 이미 확인됨 → 무시
              {"code": "P06", "value": "3종", "reason": "x"},                # 고객 사실 문항 → 무시
              {"code": "P29", "value": "놓친 보틀은 컨베이어 끝 수거함으로", "reason": "단순"},
              {"code": "P30", "value": "버퍼 40병", "reason": "x"}],         # 40 근거 없음
}

# ── 1) 제안 검증 ────────────────────────────────────────────────────────────
alts, fills, dropped = concept.validate_proposal(PROPOSAL, PROJECT)
check([a["id"] for a in alts] == ["A", "B"] and alts[1]["name"] == "스카라 셀", "구성 없는 대안은 빠지고 id 를 A부터 다시",
      str([(a["id"], a["name"]) for a in alts]))
check([a["item_codes"] for a in alts] == [["P14"], ["P15"]], "대안 문항 코드 A→P14, B→P15")
check(len(alts[0]["structure"]) == 2 and not any("250" in s for s in alts[0]["structure"]),
      "근거 없는 수치 문장만 버림(14·17·1~9 는 허용)", str(alts[0]["structure"]))
check(alts[0]["robot"] == "협동로봇 1대(14kg)" and alts[0]["summary"] == "시간당 17병 처리", "근거 있는 수치는 유지")
check([f["code"] for f in fills] == ["P29"], "설계 문항·빈 문항·근거 있는 수치만 채움", str(fills))
check(any("250" in d for d in dropped) and any("40" in d for d in dropped), "버린 사유 기록", str(dropped))

prompt = concept.image_prompt(PROJECT, alts[0], refs=[{"role": "appearance", "note": "로봇 외형"},
                                                      {"role": "dimension", "note": ""}])
check("협동로봇 1대가 컨베이어 옆에 선다" in prompt and "[넣지 않을 것]" in prompt and "AMR" in prompt,
      "그림 지시: 구성·제외 문장 그대로")
check("글자·숫자·로고" in prompt and "1번: 로봇·장치의 모양" in prompt and "(로봇 외형)" in prompt
      and "2번: 배치·비율" in prompt and len(prompt) <= 2000, "글자 없는 그림·참고 이미지 순서별 용도·길이 한도", prompt[-300:])
check("대상물: 500ml 유리 보틀 2종" in prompt, "공정 요약에 확인된 문항", prompt[:400])

checks, labels = concept.parse_vision(json.dumps({
    "checks": [{"key": "count", "ok": True, "note": "1대"}, {"key": "text", "ok": False, "note": "글자 보임"}],
    "labels": [{"text": "협동로봇", "x": 0.4, "y": 0.5}, {"text": "밖", "x": 1.4, "y": 0.2},
               {"text": "", "x": 0.1, "y": 0.1}]}))
check([c["key"] for c in checks] == ["count", "layout", "flow", "text", "hidden"], "검수 5항목 고정 순서")
check([c["ok"] for c in checks] == [True, None, None, False, None], "빠진 항목은 미판정(None)", str(checks))
check(labels == [{"text": "협동로봇", "x": 0.4, "y": 0.5}], "범위 밖·빈 라벨 버림", str(labels))

q = concept.validate_quote({"lines": [
    {"group": "공통", "item": "비전", "qty": "1", "unit": "식"},
    {"group": "대안", "alt_id": "A", "item": "협동로봇", "qty": "2대", "unit": "대"},
    {"group": "대안", "alt_id": "Z", "item": "없는 대안"},
    {"group": "기타", "item": "모르는 구분"},
    {"group": "현장 적용", "item": "컨베이어", "included": False}]}, {"A", "B"})
check([x["item"] for x in q] == ["비전", "협동로봇", "컨베이어"], "모르는 구분·대안 행 버림", str(q))
check(q[1]["qty"] == "1" and all(x["unit_price"] is None for x in q) and q[2]["included"] is False,
      "수량 숫자만·단가 없음·제외 표시", str(q))

reuse_proj = {"items": items(P43=("로봇·그리퍼·비전까지. 컨베이어는 기존 설비 재사용.", "confirmed"),
                            P11=("이물 보틀은 제외 후 수작업", "confirmed"))}
r = concept.apply_reuse([dict(x) for x in [
    {"group": "공통", "item": "입구 컨베이어", "included": True, "basis": ""},
    {"group": "공통", "item": "비전 검사 시스템", "included": True, "basis": ""},
    {"group": "대안", "item": "보틀 그리퍼", "included": True, "basis": ""}]], reuse_proj)
check([x["included"] for x in r] == [False, True, True] and r[0]["basis"] == "고객 기존 설비 재사용",
      "고객이 재사용한다는 설비는 견적 미포함(다른 문장의 '제외' 는 무시)", str(r))

# ── 2) 구성·견적·라벨 입력 검증 ───────────────────────────────────────────────
pages = [{"type": "cover", "title": "표지"}, {"type": "alternative", "alt_id": "A", "asset_ids": [5, 99, "x"]},
         {"type": "equipment"}, {"type": "quote"}]
got = workflow.clean_pages(pages, 10, {"A"}, {5})
check(got[1]["asset_ids"] == [5] and got[1]["alt_id"] == "A" and [p["page_no"] for p in got] == [1, 2, 3, 4],
      "페이지 번호 다시 매김·승인 이미지만", str(got))
for drop, why in (("equipment", "주요 항목"), ("quote", "견적")):
    try:
        workflow.clean_pages([x for x in pages if x["type"] != drop], 10, {"A"}, {5})
        check(False, f"K08 필수 쪽({why}) 삭제 거부")
    except ProjectError as e:
        check(why in str(e), f"K08 필수 쪽({why}) 삭제 거부", str(e))
try:
    workflow.clean_pages([x for x in pages if x["type"] != "alternative"], 10, {"A"}, {5})
    check(False, "K08 공정 컨셉 쪽 삭제 거부")
except ProjectError as e:
    check("공정 컨셉" in str(e), "K08 공정 컨셉 쪽 삭제 거부", str(e))
try:
    workflow.clean_pages(pages, 10, {"A", "B"}, {5})
    check(False, "비교안마다 컨셉 쪽 필요")
except ProjectError as e:
    check("컨셉 B" in str(e), "비교안마다 컨셉 쪽 필요", str(e))
over = workflow.clean_pages([{"type": "cover"}] * 9 + [{"type": "equipment"}, {"type": "quote"}], 10, set(), set())
check(len(over) == 11, "10쪽(권장 상한)을 넘어도 저장은 거부하지 않음(경고만)")
expect_error("안전 상한(30쪽) 초과만 거부", lambda: workflow.clean_pages([{"type": "cover"}] * 31, 10, set(), set()))
expect_error("모르는 페이지 유형 거부", lambda: workflow.clean_pages([{"type": "hack"}], 10, set(), set()))
expect_error("없는 대안 페이지 거부", lambda: workflow.clean_pages([{"type": "alternative", "alt_id": "C"}], 10, {"A"}, set()))
expect_error("근거 없는 단가 거부", lambda: workflow.clean_quote([{"group": "공통", "item": "비전", "unit_price": 1000}], set()))
expect_error("음수 단가 거부", lambda: workflow.clean_quote(
    [{"group": "공통", "item": "비전", "unit_price": -1, "basis": "견적서"}], set()))
ok_q = workflow.clean_quote([{"group": "공통", "item": "비전", "unit_price": "1200000", "basis": "A사 견적서 9/1"}], set())
check(ok_q[0]["unit_price"] == 1200000.0 and ok_q[0]["basis"] == "A사 견적서 9/1", "근거 있는 단가 저장")
expect_error("대안 행은 대안 필수", lambda: workflow.clean_quote([{"group": "대안", "item": "로봇"}], {"A"}))
expect_error("대안은 1~3개", lambda: workflow.clean_alternatives([]))
expect_error("대안 구성 문장 필수", lambda: workflow.clean_alternatives([{"name": "A", "structure": [" "]}]))
check(workflow.clean_labels([{"text": "로봇", "x": "0.5", "y": 0.2}, {"text": "밖", "x": 2, "y": 0}])
      == [{"text": "로봇", "x": 0.5, "y": 0.2}], "라벨: 범위 밖 버림")
expect_error("라벨 위치 숫자 아니면 거부", lambda: workflow.clean_labels([{"text": "a", "x": "?", "y": 0}]))


# ── 2a) 수정 요청 검증 ───────────────────────────────────────────────────────
alt0 = alts[0]
new, note, dr = concept.validate_revision(
    {"structure": ["협동로봇 2대가 좌우에 선다", "폭 3500mm 공간에 둔다"], "robot": "협동로봇 2대", "note": "2대로"},
    alt0, PROJECT, "로봇을 2대로 늘려 주세요")
check(new["structure"] == ["협동로봇 2대가 좌우에 선다"] and new["robot"] == "협동로봇 2대" and note == "2대로"
      and any("3500" in d for d in dr), "수정 반영: 요청 속 수치는 허용, 근거 없는 수치는 버림", str(new["structure"]))
kept, _n, dr = concept.validate_revision({}, alt0, PROJECT, "아무거나")
check(kept["structure"] == alt0["structure"] and kept["name"] == alt0["name"] and dr,
      "수정 결과가 비면 기존 구성 유지")
rp = concept.image_prompt(PROJECT, alt0, refs=[{"role": "previous", "note": ""}], revision="로봇을 오른쪽으로")
check("[이번 수정" in rp and "로봇을 오른쪽으로" in rp and "1번: 이전 버전 그림" in rp, "그림 지시: 수정 요청·이전 버전 안내")
check("공정 컨셉은 기본 1개" in concept._PROPOSE_SYSTEM and "대안은 1~3개" not in concept._PROPOSE_SYSTEM,
      "제안 규칙: 기본은 공정 컨셉 1개(비교 요청이 있을 때만 여러 개)")


# ── 2b) 두뇌 선택(gemma|gpt) ─────────────────────────────────────────────────
from app import codex_client as _cc  # noqa: E402
from app.proposal_project import brain as _brain  # noqa: E402

_asked: list[tuple[str, str]] = []


async def _fake_ask(question, context="", *, user_id=None):
    _asked.append((question, context))
    return "```json\n" + json.dumps(PROPOSAL, ensure_ascii=False) + "\n```"


async def _fake_gemma(messages, **_kw):
    return json.dumps({"alternatives": [{"name": "젬마 셀", "robot": "협동로봇 1대", "structure": ["협동로봇 1대"]}]},
                      ensure_ascii=False)

_orig = (_cc.ask, _cc.is_configured)
_cc.ask, _cc.is_configured = _fake_ask, (lambda: True)
r = asyncio.run(concept.propose_with_brain(PROJECT, brain="gpt", user_id=1, chat=_fake_gemma))
check(r["used"] == "gpt" and [a["name"] for a in r["alternatives"]] == ["협동로봇 셀", "스카라 셀"]
      and any("250" in d for d in r["dropped"]), "GPT 두뇌: 마크다운 감싼 JSON 도 읽고 같은 검증 적용", str(r)[:300])
check(len(_asked) == 1 and "JSON 객체 하나만" in _asked[0][0] and "P06" in _asked[0][1]
      and len(_asked[0][0]) <= 4000, "GPT 로는 규칙=question, 문항 값=context 로 보냄(길이 한도)")


async def _boom(*_a, **_k):
    raise _cc.CodexError("브리지 꺼짐")
_cc.ask = _boom
r = asyncio.run(concept.propose_with_brain(PROJECT, brain="gpt", user_id=1, chat=_fake_gemma))
check(r["used"] == "gemma" and "브리지 꺼짐" in r["fallback_reason"] and r["alternatives"][0]["name"] == "젬마 셀",
      "GPT 실패 시 gemma 로 대신하고 사유 남김", str(r)[:200])
_asked.clear()
r = asyncio.run(concept.propose_with_brain(PROJECT, brain="gemma", user_id=1, chat=_fake_gemma))
check(r["used"] == "gemma" and not _asked, "gemma 두뇌는 외부로 보내지 않음")

# 그리퍼(P21)를 비우면 회사 그리퍼·툴 제품 목록을 근거로 준다 — 정했으면 고객 지정을 따른다(목록 안 줌).
_tool = ["- mDPG-S series [Magbot EOAT] 표준형 DC 공압 그리퍼 / 맞는 조건: 흡착 파지"]
_open = {**PROJECT, "items": {**PROJECT["items"], "P21": {**PROJECT["items"]["P21"], "value": "", "status": "empty"}}}
check(concept.gripper_open(_open) and "[회사 그리퍼·툴 제품]" in concept._propose_user(_open, tooling=_tool)
      and "mDPG-S series" in concept._propose_user(_open, tooling=_tool),
      "그리퍼 미지정: 회사 그리퍼·툴 제품 목록을 제안 자료에 넣음")
check(not concept.gripper_open(PROJECT) and "[회사 그리퍼·툴 제품]" not in concept._propose_user(PROJECT, tooling=_tool),
      "그리퍼 지정(P21 흡착 그리퍼): 회사 목록을 넣지 않고 고객 지정을 따름")
_asked.clear()
_cc.ask = _fake_ask
r = asyncio.run(concept.propose_with_brain(_open, brain="gpt", user_id=1, chat=_fake_gemma, tooling=_tool))
check("mDPG-S series" in _asked[0][1] and "툴체인저·툴스탠드는 회사 제품이라는 이유만으로" in _asked[0][0],
      "GPT 로 회사 제품 목록(context)과 우선 추천·툴체인저 옵션 규칙(question)을 보냄", _asked[0][0][-300:])
line = concept.tooling_line("mDPG-C series", {"family": "magbot EOAT", "summary": "커스텀 공압 그리퍼",
                                              "strengths": ["컴프레서 불필요", "맞춤 사양", "셋째"],
                                              "applies_when": "에어 설비가 없는 현장", "limits": [], "not_when": "젖은 표면"})
check(line.startswith("- mDPG-C series [magbot EOAT] 커스텀 공압 그리퍼") and "셋째" not in line
      and "맞는 조건: 에어 설비가 없는 현장" in line and "한계: 젖은 표면" in line, "제품 카드 한 줄 요약", line)
_cc.ask, _cc.is_configured = _orig


# ── 3) DB 통합 (--db) ───────────────────────────────────────────────────────
async def db_flow() -> None:
    import base64

    import pymupdf  # noqa: F401 — 컨테이너 의존성 확인
    from pptx import Presentation
    from sqlalchemy import text

    from app import codex_client, proposal_llm
    from app.database import SessionLocal
    from app.migrations import run_migrations
    from app.proposal_project import dialog, jobs, service
    from app.storage import get_object_stream, make_object_key, put_object

    png = base64.b64decode("iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJAAAADUlEQVR42mNk+M9QDwADhgGAWjR9awAAAABJRU5ErkJggg==")
    sent: list[int] = []

    async def fake_codex(prompt, refs, *, user_id=None):
        sent.append(len(refs))
        return png, "image/png", "그림", 0

    async def fake_vision(image, alt, **_kw):
        return concept.parse_vision(json.dumps({"checks": [{"key": "count", "ok": True, "note": "맞음"}],
                                                "labels": [{"text": "협동로봇", "x": 0.3, "y": 0.4}]}))

    async def fake_propose(project, **_kw):
        alts, fills, dropped = concept.validate_proposal(PROPOSAL, project)
        return {"alternatives": alts, "fills": fills, "dropped": dropped, "used": "gemma", "fallback_reason": ""}

    async def fake_quote(project, **_kw):
        return concept.validate_quote({"lines": [
            {"group": "공통", "item": "비전 카메라", "qty": "1", "unit": "식"},
            {"group": "대안", "alt_id": "A", "item": "협동로봇", "qty": "1", "unit": "대"},
            {"group": "대안", "alt_id": "B", "item": "스카라", "qty": "2", "unit": "대"},
            {"group": "통합·실증", "item": "시스템 통합·시운전", "qty": "1", "unit": "식"}]},
            {a["id"] for a in project["alternatives"]})

    async def fake_llm(messages, **_kw):      # 5c compose → 문구 실패 경로(문항 값 그대로)로 조판만 확인
        return "모델 없음"

    async def fake_extract_chat(messages, **_kw):
        return '{"items": []}'

    codex_client.generate_image_with_refs = fake_codex
    codex_client.is_configured = lambda: True

    async def no_external_ask(*_a, **_k):
        # 시험 중에는 GPT(브리지)로 아무것도 보내지 않는다 — '만드는' 단계는 gemma 대체 경로로 검증한다.
        raise codex_client.CodexError("시험: 외부 호출 없음")
    codex_client.ask = no_external_ask
    workflow.BRIDGE_TAKES_REFS = True    # 참고 이미지 규칙 검증용(운영 기본값도 True — 2026-09-30 브리지 확장)
    concept.vision_check = fake_vision
    concept.propose_with_brain = fake_propose
    concept.draft_quote_lines = fake_quote
    proposal_llm.chat = fake_llm

    async def wait_job(pid):
        for _ in range(200):
            if not jobs.is_running(pid):
                await workflow.wait_vision()      # 이미지 검수 초안은 작업이 끝난 뒤에도 뒤에서 채워진다
                return
            await asyncio.sleep(0.1)

    async with SessionLocal() as db:
        await run_migrations(db)
        uid = (await db.execute(text("SELECT id FROM users WHERE username='alex'"))).scalar_one()
        key = make_object_key(uid, "robot_ref.png")
        put_object(key=key, data=io.BytesIO(png), length=len(png), mime="image/png")
        ref = int((await db.execute(text(
            "INSERT INTO attachments (user_id, bucket, object_key, original_filename, mime, size_bytes) "
            "VALUES (:u, 'und-cortex-attachments', :k, 'robot_ref.png', 'image/png', :s) RETURNING id"),
            {"u": uid, "k": key, "s": len(png)})).scalar_one())
        key2 = make_object_key(uid, "private_ref.png")
        put_object(key=key2, data=io.BytesIO(png), length=len(png), mime="image/png")
        ref2 = int((await db.execute(text(
            "INSERT INTO attachments (user_id, bucket, object_key, original_filename, mime, size_bytes) "
            "VALUES (:u, 'und-cortex-attachments', :k, 'private_ref.png', 'image/png', :s) RETURNING id"),
            {"u": uid, "k": key2, "s": len(png)})).scalar_one())
        before = (await db.execute(text("SELECT COALESCE(MAX(id), 0) FROM external_calls"))).scalar_one()
        await db.commit()
        pid = await service.create_project(
            db, user_id=uid, title="워크플로 시험", request_text="시간당 17병, 협동로봇(가반하중 14kg) 검토 중",
            intake={"target": "500ml 유리 보틀 2종"},
            assets=[service.AssetInput(ref, "appearance", external_ok=True),
                    service.AssetInput(ref2, "appearance", external_ok=False),
                    service.AssetInput(ref2, "dimension", external_ok=True),
                    service.AssetInput(ref2, "price", external_ok=False)])
    made_atts = [ref, ref2]
    try:
        await dialog.run_reading(pid, chat=fake_extract_chat)
        async with SessionLocal() as db:
            await service.set_item(db, pid, "P14", value="협동로봇(가반하중 14kg) 검토 중", status="assumed",
                                   source="request", evidence="요청")
            await service.set_item(db, pid, "P06", value="500ml 유리 보틀 2종", status="confirmed", source="form")
            await service.set_item(db, pid, "P08", value="시간당 17병", status="confirmed", source="request")
            await db.commit()
            p = await service.get_project(db, pid)
            try:
                await workflow.save_alternatives(db, p, [{"name": "x", "structure": ["y"]}])
                check(False, "질문 단계에선 대안 저장 거부")
            except ProjectError:
                check(True, "질문 단계에선 대안 저장 거부")

        # 3단계: 컨셉 단계로 넘어가 제안 → 이미지 → 승인 → 확정
        async with SessionLocal() as db:
            await db.execute(text("UPDATE proposal_projects SET stage='concept' WHERE id=:p"), {"p": pid})
            await db.commit()
        check(workflow.MAKER == "gpt", "만드는 단계는 선택 없이 GPT(실패 시 gemma)")
        msg = await workflow.run_propose(pid, lambda s: asyncio.sleep(0))
        async with SessionLocal() as db:
            p = await service.get_project(db, pid)
        check([a["id"] for a in p["alternatives"]] == ["A", "B"], "제안 대안 저장", msg)
        check(p["items"]["P14"]["source"] == "request" and p["items"]["P15"]["source"] == "ai"
              and p["items"]["P15"]["status"] == "assumed" and p["items"]["P13"]["value"].startswith("비교 컨셉 2종"),
              "빈 대안 문항만 '가정' 으로 채우고 고객 요청 값(P14)은 유지",
              str({c: p["items"][c]["value"] for c in ("P13", "P14", "P15")}))
        check(p["items"]["P29"]["status"] == "assumed" and p["items"]["P29"]["evidence"].startswith("AI 제안"),
              "설계 문항 AI 제안 표시")
        # 제안 직후는 계획 확인 단계 — 이미지 전에 로봇·그리퍼를 정하고 확정한다(2026-09-30).
        from app.proposal_project import plan as plan_mod
        check(plan_mod.pending(p) and all(a.get("plan") and not a["plan"]["confirmed"] for a in p["alternatives"]),
              "제안 직후는 계획 확인 단계(이미지 전)")
        async with SessionLocal() as db:
            try:
                await workflow.finish_concept(db, p, uid)
                check(False, "계획 확정 전엔 컨셉 확정 거부")
            except ProjectError as e:
                await db.rollback()
                check("계획" in str(e), "계획 확정 전엔 컨셉 확정 거부", str(e))
            await workflow.plan_select(db, p, "A", "DOBOT CR10A", "시험용 흡착 패드")
            p = await service.get_project(db, pid)
            pa = p["alternatives"][0]["plan"]
            check(pa["robot_pick"] == "DOBOT CR10A" and pa["gripper_pick"] == "시험용 흡착 패드", "계획: 로봇·그리퍼 선택 저장")
            await db.execute(text("UPDATE proposal_projects SET alternatives=CAST(:a AS jsonb) WHERE id=:p"),
                             {"a": json.dumps([plan_mod.confirm(a) for a in p["alternatives"]], ensure_ascii=False),
                              "p": pid})
            await db.commit()
            p = await service.get_project(db, pid)
        check(not plan_mod.pending(p) and p["alternatives"][0]["structure"][-1].startswith("선택 모델: 로봇 팔 DOBOT CR10A"),
              "계획 확정 → 고른 모델이 구성 끝 한 줄로(그림·견적에 반영)", p["alternatives"][0]["structure"][-1])

        vision_gate = asyncio.Event()

        async def slow_vision(image, alt, **_kw):
            await vision_gate.wait()
            return await fake_vision(image, alt)
        concept.vision_check = slow_vision
        iid = await workflow.run_image(pid, "A", lambda s: asyncio.sleep(0))
        async with SessionLocal() as db:
            p = await service.get_project(db, pid)
        img = next(i for i in p["images"] if i["id"] == iid)
        check(img["status"] == "draft" and img["attachment_id"] and img["ref_sent"] == 2 and img["ref_used"] == 0,
              "이미지 초안 저장·참고 이미지 보냄·반영 0", str(img))
        check(workflow.vision_pending(pid) and all(c["note"] == workflow.VISION_PENDING_NOTE for c in img["checks"])
              and img["labels"] == [],
              "이미지는 검수 초안을 기다리지 않고 먼저 저장(검수는 뒤에서 작성 중)", str(img["checks"])[:120])
        vision_gate.set()
        await workflow.wait_vision()
        concept.vision_check = fake_vision
        async with SessionLocal() as db:
            p = await service.get_project(db, pid)
        img = next(i for i in p["images"] if i["id"] == iid)
        check(not workflow.vision_pending(pid) and img["checks"][0]["note"] == "맞음" and img["labels"],
              "뒤에서 검수 초안·라벨 후보가 채워짐", str(img["checks"])[:120])
        check(sent == [2], "외부 전송 허용 이미지만 보냄(외형+치수 2장, 미허용·가격 제외)", str(sent))
        async with SessionLocal() as db:
            rows = (await db.execute(text("SELECT status, payload_masked FROM external_calls WHERE id > :b "
                                          "AND status = 'blocked'"), {"b": before})).all()
            await db.execute(text("DELETE FROM external_calls WHERE id > :b AND status = 'blocked' "
                                  "AND payload_masked ? 'excluded_reference_images'"), {"b": before})
            await db.commit()
        reasons = sorted((r[1].get("reason"), [x["role"] for x in r[1].get("excluded_reference_images", [])])
                         for r in rows)
        check(reasons == [("가격 자료 전송 금지", ["price"]), ("외부 전송 미허용", ["appearance"])],
              "뺀 이미지는 사유별(미허용·가격)로 blocked 기록", str(reasons))
        for rp in (workflow.ref_block_reason({"role": "price", "external_ok": True}),):
            check(rp == workflow.BLOCK_PRICE, "가격 근거는 허용 체크해도 차단")
        check(len(img["checks"]) == 5 and img["labels"][0]["text"] == "협동로봇", "검수 5항목·라벨 후보 저장")
        check("참고 이미지 2장을 보냈지만 반영되지 않았습니다" in p["messages"][-1]["content"],
              "참고 이미지 미반영 안내", p["messages"][-1]["content"])

        async def boom(*_a, **_k):
            raise codex_client.CodexError("브리지 꺼짐")
        codex_client.generate_image_with_refs = boom
        try:
            await workflow.run_image(pid, "B", lambda s: asyncio.sleep(0))
            check(False, "이미지 실패 전파")
        except codex_client.CodexError:
            check(True, "이미지 실패 전파")
        codex_client.generate_image_with_refs = fake_codex
        async with SessionLocal() as db:
            p = await service.get_project(db, pid)
            failed = [i for i in p["images"] if i["alt_id"] == "B"]
            check(failed and failed[0]["status"] == "failed" and "브리지 꺼짐" in failed[0]["error"], "실패 이미지 기록")
            try:
                await workflow.approve_image(db, p, failed[0]["id"], uid)
                check(False, "실패 이미지 승인 거부")
            except ProjectError:
                await db.rollback()
                check(True, "실패 이미지 승인 거부")
            await workflow.save_labels(db, p, iid, [{"text": "협동로봇", "x": 0.35, "y": 0.45},
                                                    {"text": "비전", "x": 0.6, "y": 0.2}])
            await workflow.approve_image(db, p, iid, uid)
            p = await service.get_project(db, pid)
            img = next(i for i in p["images"] if i["id"] == iid)
            check(img["approved"] and len(img["labels"]) == 2, "라벨 수정·이미지 승인")

        # 수정 요청(여러 번) — 말로 한 요청 → 구성 문장 수정 → 이전 버전을 참고로 새 버전
        async def fake_revise_llm(messages, **_kw):
            top = "위에서 내려다본다" in messages[-1]["content"] or "내려다보게" in messages[-1]["content"]
            return json.dumps({"name": "협동로봇 셀", "robot": "협동로봇 1대(14kg)",
                               "structure": ["협동로봇 1대가 컨베이어 오른쪽에 선다",
                                             "비전 카메라 1대가 컨베이어 위에서 내려다본다" if top else "비전 카메라 1대가 컨베이어 위에 있다",
                                             "주입 노즐은 999mm 앞에 둔다"],
                               "exclude": ["AMR"], "note": "로봇을 컨베이어 오른쪽으로 옮겼습니다."}, ensure_ascii=False)
        proposal_llm.chat = fake_revise_llm
        sent.clear()
        await workflow.run_revise(pid, "A", "로봇을 컨베이어 오른쪽으로 옮겨 주세요", lambda s: asyncio.sleep(0))
        await workflow.run_revise(pid, "A", "카메라는 위에서 내려다보게", lambda s: asyncio.sleep(0))
        proposal_llm.chat = fake_llm
        async with SessionLocal() as db:
            p = await service.get_project(db, pid)
        alt_a = p["alternatives"][0]
        check(alt_a["structure"][0] == "협동로봇 1대가 컨베이어 오른쪽에 선다"
              and not any("999" in x for x in alt_a["structure"]),
              "수정 요청이 구성 문장에 반영(근거 없는 수치 문장은 버림)", str(alt_a["structure"]))
        va = sorted((i for i in p["images"] if i["alt_id"] == "A"), key=lambda i: i["version"])
        check([i["version"] for i in va] == [1, 2, 3] and va[1]["revision"] == "로봇을 컨베이어 오른쪽으로 옮겨 주세요"
              and va[2]["revision"] == "카메라는 위에서 내려다보게", "수정할 때마다 새 버전(v2, v3)과 요청 기록",
              str([(i["version"], i["revision"]) for i in va]))
        check("[이번 수정" in va[2]["prompt"] and "이전 버전 그림" in va[2]["prompt"]
              and "협동로봇 1대가 컨베이어 오른쪽에 선다" in va[2]["prompt"],
              "그림 지시: 이번 수정 + 이전 버전 참고 + 고친 구성", va[2]["prompt"][-400:])
        check(sent == [3, 3], "이전 버전 1장 + 참고 이미지 2장(최대 3장)", str(sent))
        workflow.BRIDGE_TAKES_REFS = False
        sent.clear()
        off = await workflow.run_image(pid, "A", lambda s_: asyncio.sleep(0), revision="시험")
        async with SessionLocal() as db:
            off_img = next(i for i in (await service.get_project(db, pid))["images"] if i["id"] == off)
        check(sent == [0] and off_img["ref_sent"] == 0 and "[참고 이미지" not in off_img["prompt"],
              "브리지 미지원으로 끄면 참고 이미지를 보내지 않음 — 거절 후 재전송 낭비 없음", str(sent))
        workflow.BRIDGE_TAKES_REFS = True
        check(va[0]["approved"] and not va[2]["approved"], "새 버전은 초안 — 이전 승인본은 그대로")
        async with SessionLocal() as db:
            n_before = len([i for i in (await service.get_project(db, pid))["images"] if i["alt_id"] == "A"])
        await workflow.run_revise(pid, "A", "카메라는 위에서 내려다보게", lambda s_: asyncio.sleep(0))
        async with SessionLocal() as db:
            n_after = len([i for i in (await service.get_project(db, pid))["images"] if i["alt_id"] == "A"])
        check(n_after == n_before, "구성이 바뀌지 않는 수정 요청은 이미지를 다시 그리지 않음", f"{n_before}→{n_after}")
        ch = alt_a.get("changes") or {}
        check(ch.get("request") == "카메라는 위에서 내려다보게"
              and ch.get("added") == ["비전 카메라 1대가 컨베이어 위에서 내려다본다"]
              and ch.get("removed") == ["비전 카메라 1대가 컨베이어 위에 있다"],
              "수정마다 구성에서 바뀐 점(추가·삭제 문장) 기록", str(ch))
        er = workflow.experience_request({"title": "양팔 비전 분류", "suggestion": "취출 직전 카메라로 분류",
                                          "facts": ["양팔 로봇이 비전으로 식기를 분류"],
                                          "detail": {"process": "식기 분류 자동화", "solution": "양팔 로봇 + 비전"}})
        check("식기 분류 자동화" in er and "양팔 로봇 + 비전" in er and "최소 1줄" in er,
              "경험 반영 요청: 그 경험의 공정·방식 + 구성에 장치로 넣으라는 지시", er)
        check(any("수정 요청: 로봇을" in m["content"] for m in p["messages"])
              and any("로봇을 컨베이어 오른쪽으로 옮겼습니다" in m["content"] for m in p["messages"]),
              "요청과 반영 내용이 진행 기록에 남음")
        async with SessionLocal() as db:
            await workflow.approve_image(db, p, va[2]["id"], uid)
            p = await service.get_project(db, pid)
            va = sorted((i for i in p["images"] if i["alt_id"] == "A"), key=lambda i: i["version"])
            check({i["id"] for i in va if i["approved"]} == {va[2]["id"]}, "다른 버전을 승인하면 승인본 교체(하나만)")
            await workflow.approve_image(db, p, va[0]["id"], uid)
            p = await service.get_project(db, pid)
            va = sorted((i for i in p["images"] if i["alt_id"] == "A"), key=lambda i: i["version"])
            check({i["id"] for i in va if i["approved"]} == {va[0]["id"]}, "예전 버전으로 되돌려 승인도 가능")
            warn = await workflow.finish_concept(db, p, uid)
            p = await service.get_project(db, pid)
            check(p["stage"] == "structure" and warn == ["컨셉 B 에 승인된 이미지가 없습니다."],
                  "컨셉 확정 → 구성 단계, 이미지 없는 컨셉 경고", str(warn))
            check(p["items"]["P14"]["status"] == "adopted" and p["items"]["P15"]["status"] == "adopted",
                  "확정 대안은 '채택 컨셉'")

        # 5단계: 구성 초안·저장·승인
        await workflow.run_structure(pid, lambda s: asyncio.sleep(0))
        async with SessionLocal() as db:
            p = await service.get_project(db, pid)
            types = [pg["type"] for pg in p["pages"]]
            check(types.count("alternative") == 2 and len(p["pages"]) <= 10 and len(p["quote_lines"]) == 4,
                  "기본 구성(대안 2쪽)·견적 초안", str(types))
            alt_a = next(pg for pg in p["pages"] if pg.get("alt_id") == "A")
            check(alt_a["asset_ids"] == [iid], "대안 A 쪽에 승인 이미지")
            ql = p["quote_lines"]
            ql[0].update(unit_price=1200000, basis="시험 견적서")
            pages = [pg for pg in p["pages"] if pg["type"] != "poc"]
            await workflow.save_structure(db, p, pages, ql)
            p = await service.get_project(db, pid)
            check("poc" not in [pg["type"] for pg in p["pages"]] and p["quote_lines"][0]["unit_price"] == 1200000.0,
                  "구성·견적 수정 저장")
            await workflow.approve_structure(db, p, uid)
            p = await service.get_project(db, pid)
            check(p["stage"] == "producing" and any(a["target"] == "structure" for a in p["approvals"]), "구성 승인 기록")

        # 6단계: 제작(실제 5c 조판)
        msg = await workflow.run_produce(pid, lambda s: asyncio.sleep(0))
        async with SessionLocal() as db:
            p = await service.get_project(db, pid)
            out = p["outputs"][-1]
            key = (await db.execute(text("SELECT object_key FROM attachments WHERE id=:a"),
                                    {"a": out["attachment_id"]})).scalar_one()
        made_atts += [out["attachment_id"]] + [i["attachment_id"] for i in p["images"] if i["attachment_id"]]
        resp = get_object_stream(key)
        data = b"".join(resp.stream(amt=65536))
        resp.close()
        resp.release_conn()
        prs = Presentation(io.BytesIO(data))
        check(p["stage"] == "done" and out["version"] == 1 and len(prs.slides) == out["report"]["pages"] == len(p["pages"]),
              "제작 완료·쪽수 일치", f"{msg} / {out['report']}")
        parts = []
        for sl in prs.slides:
            for sh in sl.shapes:
                if sh.has_text_frame:
                    parts.append(sh.text_frame.text)
                if getattr(sh, "has_table", False) and sh.has_table:
                    parts += [c.text for r in sh.table.rows for c in r.cells]
        all_text = "  ".join(parts)
        check("비전" in all_text, "PPT 에 이미지 라벨(편집 가능 글상자) 반영", all_text[:300])
        check("1,200,000" in all_text.replace(" ", ""), "PPT 견적에 근거 있는 단가 반영",
              [t for t in all_text.split("  ") if "비전" in t or "원" in t][:10])

        # 완성본 프롬프트로 수정 → 새 버전(v2), 바꿀 게 없으면 버전을 만들지 않음
        async with SessionLocal() as db:
            deck_row = (await db.execute(text("SELECT deck FROM project_outputs WHERE id=:i"), {"i": out["id"]})).scalar_one()
        check(bool(deck_row and deck_row.get("pages") and deck_row.get("text") is not None),
              "제작 때 쪽 문구·구성을 함께 저장(프롬프트로 수정 출발점)")

        async def fake_edit_llm(messages, **_kw):
            if "편집자" in messages[0]["content"]:
                if "요지" in messages[-1]["content"]:
                    return json.dumps({"pages": [{"page_no": 2, "headline": "실증으로 확인할 핵심 3가지"}],
                                       "note": "2쪽 요지를 바꿨습니다."}, ensure_ascii=False)
                return json.dumps({"pages": [], "note": "견적 표는 구성 단계에서 고쳐야 합니다."}, ensure_ascii=False)
            return await fake_llm(messages, **_kw)
        proposal_llm.chat = fake_edit_llm
        emsg = await workflow.run_deck_edit(pid, "2쪽 요지를 바꿔 줘", lambda s_: asyncio.sleep(0))
        nmsg = await workflow.run_deck_edit(pid, "견적 금액을 올려 줘", lambda s_: asyncio.sleep(0))
        proposal_llm.chat = fake_llm
        async with SessionLocal() as db:
            p = await service.get_project(db, pid)
        o2 = p["outputs"][-1]
        made_atts.append(o2["attachment_id"])
        check(len(p["outputs"]) == 2 and o2["version"] == 2 and o2["report"]["edit"]["changed"] == ["2쪽"]
              and o2["report"]["edit"]["from"] == 1, "프롬프트로 수정 → v2(어디를 고쳤는지 기록)", f"{emsg} / {o2['report'].get('edit')}")
        check("새 버전을 만들지 않았습니다" in nmsg and "구성 단계" in nmsg, "바꿀 곳이 없으면 버전 안 늘림 + 안내", nmsg)

        # 되돌려 다시 만들기 → v3
        async with SessionLocal() as db:
            p = await service.get_project(db, pid)
            await workflow.back_to(db, p, "structure")
            p = await service.get_project(db, pid)
            check(p["stage"] == "structure", "완료 → 구성 단계로 되돌리기")
            try:
                await workflow.back_to(db, p, "questioning")
                check(False, "질문 단계로는 되돌리기 거부")
            except ProjectError:
                check(True, "질문 단계로는 되돌리기 거부")

        # 검수 초안을 쓰는 사이 사용자가 라벨을 고쳤으면 그 라벨은 덮어쓰지 않는다(검수 결과만 채움).
        vision_gate = asyncio.Event()
        concept.vision_check = slow_vision
        iid2 = await workflow.run_image(pid, "A", lambda s: asyncio.sleep(0))
        async with SessionLocal() as db:
            await db.execute(text("UPDATE project_images SET labels = CAST(:l AS jsonb) WHERE id = :i"),
                             {"l": json.dumps([{"text": "사용자 라벨", "x": 0.5, "y": 0.5}], ensure_ascii=False), "i": iid2})
            await db.commit()
        vision_gate.set()
        await workflow.wait_vision()
        concept.vision_check = fake_vision
        async with SessionLocal() as db:
            p = await service.get_project(db, pid)
        img2 = next(i for i in p["images"] if i["id"] == iid2)
        check([lb["text"] for lb in img2["labels"]] == ["사용자 라벨"] and img2["checks"][0]["note"] == "맞음",
              "검수 중 고친 라벨은 보존, 검수 결과만 채움", str(img2["labels"]))

        # 질문을 마친 뒤 자동 진행 — 공정 컨셉 1개 제안 → 이미지, 비교안은 요청할 때만
        async def fake_one(project, **_kw):
            one = {"alternatives": [PROPOSAL["alternatives"][0]], "fills": []}
            alts, fills, dropped = concept.validate_proposal(one, project)
            return {"alternatives": alts, "fills": fills, "dropped": dropped, "used": "gemma", "fallback_reason": ""}

        async def fake_extra(project, **_kw):
            alts, _f, dropped = concept.validate_proposal({"alternatives": [PROPOSAL["alternatives"][2]]}, project)
            return {"alternatives": alts, "fills": [], "dropped": dropped, "used": "gemma", "fallback_reason": ""}

        concept.propose_with_brain, concept.propose_extra = fake_one, fake_extra
        async with SessionLocal() as db:
            cid = await service.create_project(db, user_id=uid, title="자동 진행 시험",
                                               request_text="시간당 17병 보틀 주입", intake={}, assets=[])
            await db.execute(text("UPDATE proposal_projects SET stage='concept' WHERE id=:p"), {"p": cid})
            await db.commit()
        try:
            msg = await workflow.run_concept_plan(cid, lambda s: asyncio.sleep(0))
            async with SessionLocal() as db:
                c = await service.get_project(db, cid)
            check(len(c["alternatives"]) == 1 and not c["images"] and plan_mod.pending(c)
                  and "공정 컨셉 '협동로봇 셀'" in " ".join(m["content"] for m in c["messages"]),
                  "질문 완료 → 계획까지만(이미지는 아직, '대안' 표현 없음)", msg)
            check(workflow.concept_label(c, "A") == "공정 컨셉", "컨셉이 하나면 '공정 컨셉' 으로 부름")
            msg2 = await workflow.run_plan_confirm(cid, [{"alt_id": "A", "robot_pick": "DOBOT CR5A",
                                                          "gripper_pick": ""}], lambda s: asyncio.sleep(0))
            async with SessionLocal() as db:
                c = await service.get_project(db, cid)
            check(len(c["images"]) == 1 and c["images"][0]["status"] == "draft" and not plan_mod.pending(c)
                  and "로봇 팔 DOBOT CR5A" in c["items"]["P14"]["value"] and "선택 모델" in c["images"][0]["prompt"],
                  "계획 확정 → 고른 로봇으로 이미지 1장, 컨셉 문항에도 기록", f"{msg2} / {c['items']['P14']['value']}")
            try:
                await workflow.run_plan_confirm(cid, [], lambda s: asyncio.sleep(0))
                check(False, "확정할 계획이 없으면 다시 그리지 않음")
            except ProjectError:
                check(True, "확정할 계획이 없으면 다시 그리지 않음")
            await workflow.run_extra(cid, lambda s: asyncio.sleep(0))
            async with SessionLocal() as db:
                c = await service.get_project(db, cid)
            check([a["id"] for a in c["alternatives"]] == ["A", "B"] and c["alternatives"][1]["item_codes"] == ["P15"]
                  and not [i for i in c["images"] if i["alt_id"] == "B"] and plan_mod.pending(c)
                  and workflow.concept_label(c, "B") == "컨셉 B",
                  "비교안 추가: 컨셉 B 는 계획부터(이미지 전), 둘 이상이면 '컨셉 A/B'", str(c["alternatives"])[:200])
            await workflow.run_plan_confirm(cid, [], lambda s: asyncio.sleep(0))
            async with SessionLocal() as db:
                c = await service.get_project(db, cid)
            check(len([i for i in c["images"] if i["alt_id"] == "B"]) == 1 and len(c["images"]) == 2,
                  "비교안 계획 확정 → 비교안만 새로 그림")
            made_atts += [i["attachment_id"] for i in c["images"] if i["attachment_id"]]
        finally:
            async with SessionLocal() as db:
                await db.execute(text("DELETE FROM proposal_projects WHERE id=:p"), {"p": cid})
                await db.commit()

        # 공정 컨셉 응답이 한 번 비면(형식 깨짐·수치 탈락) 한 번 더 묻고, 두 번 다 비면 원인을 알려 준다
        calls: list[int] = []

        async def flaky(project, **_kw):
            calls.append(1)
            if len(calls) == 1:
                return {"alternatives": [], "fills": [], "dropped": [], "used": "gpt", "fallback_reason": "",
                        "raw": "형식 깨진 응답"}
            return await fake_one(project)

        async def always_empty(project, **_kw):
            calls.append(1)
            return {"alternatives": [], "fills": [], "dropped": ["대안 A: 구성 문장이 없어 뺌"], "used": "gpt",
                    "fallback_reason": "", "raw": "{}"}

        async with SessionLocal() as db:
            rid = await service.create_project(db, user_id=uid, title="컨셉 재시도 시험",
                                               request_text="시간당 17병 보틀 주입", intake={}, assets=[])
            await db.execute(text("UPDATE proposal_projects SET stage='concept' WHERE id=:p"), {"p": rid})
            await db.commit()
        try:
            concept.propose_with_brain = flaky
            msg = await workflow.run_propose(rid, lambda s: asyncio.sleep(0))
            check(len(calls) == 2 and "한 번 더 요청" in msg, "첫 응답이 비면 한 번 재시도해 성공", msg)
            calls.clear()
            concept.propose_with_brain = always_empty
            try:
                await workflow.run_propose(rid, lambda s: asyncio.sleep(0))
                check(False, "두 번 다 비면 실패해야 함")
            except service.ProjectError as e:
                check(len(calls) == 2 and "두 번 시도" in str(e) and "수치" in str(e),
                      "두 번 다 비면 실제 원인(수치 탈락)을 알림", str(e))
        finally:
            async with SessionLocal() as db:
                await db.execute(text("DELETE FROM proposal_projects WHERE id=:p"), {"p": rid})
                await db.commit()
        concept.propose_with_brain = fake_propose

        # 바로 만들기(시험용) — 새 프로젝트, 질문 단계에서 시작, 필수 미답인 채로 제작까지
        async with SessionLocal() as db:
            qid = await service.create_project(db, user_id=uid, title="바로 만들기 시험",
                                               request_text="시간당 17병 보틀 주입", intake={}, assets=[])
        try:
            await dialog.run_reading(qid, chat=fake_extract_chat)
            ok = jobs.start(qid, "quick", lambda step: workflow.run_quick(qid, step, with_images=True))
            check(ok and not jobs.start(qid, "quick", lambda step: asyncio.sleep(0)), "작업 중복 실행 막음")
            await wait_job(qid)
            async with SessionLocal() as db:
                q = await service.get_project(db, qid)
            check(q["job"].get("status") == "done" and q["stage"] == "done" and len(q["outputs"]) == 1,
                  "바로 만들기: 질문 건너뛰고 제작까지", json.dumps(q["job"], ensure_ascii=False))
            check(all(i["approved"] for i in q["images"]) and len(q["images"]) == 2, "시험 진행: 이미지 자동 승인")
            check("필수 문항 미답" in q["job"].get("message", "")
                  and any("필수 문항 미확정" in w for w in q["outputs"][0]["report"]["warnings"]),
                  "필수 미답은 결과 경고에 남김", q["job"].get("message", ""))
            made_atts += [q["outputs"][0]["attachment_id"]] + [i["attachment_id"] for i in q["images"]]
        finally:
            async with SessionLocal() as db:
                await db.execute(text("DELETE FROM proposal_projects WHERE id=:p"), {"p": qid})
                await db.commit()
    finally:
        async with SessionLocal() as db:
            await db.execute(text("DELETE FROM proposal_projects WHERE id=:p"), {"p": pid})
            await db.execute(text("DELETE FROM attachments WHERE id = ANY(:ids)"), {"ids": [a for a in made_atts if a]})
            # 시험이 남긴 차단 기록(참고 이미지 제외 감사)은 운영 기록에 남기지 않는다.
            await db.execute(text("DELETE FROM external_calls WHERE id > :b"), {"b": before})
            await db.commit()


if "--db" in sys.argv:
    asyncio.run(db_flow())

print(f"PASS {len(PASS)} / FAIL {len(FAIL)}")
for label, detail in FAIL:
    print("  FAIL", label, detail)
sys.exit(1 if FAIL else 0)
