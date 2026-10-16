"""툴체인저 견적서 엑셀(company_knowledge/quote_xlsx.py) 점검 — 머리·양식(한글/영문)·파일 내용(수식·그림·인쇄).
품목 줄은 견적서 작성(quote_session) 초안에서 만든다 — 대화·수정은 scripts/test_quote_session.py.

    docker compose run --rm --no-deps -e PYTHONPATH=/app -v ./backend/app:/app/app -v ./scripts:/scripts -v ./docs:/docs backend \\
        python /scripts/test_quote_xlsx.py
"""
from __future__ import annotations

import io
import sys
from datetime import date

import openpyxl

from app.company_knowledge import product_prices as pp
from app.company_knowledge import quote_session as qs
from app.company_knowledge import quote_xlsx as qx

PASS: list[str] = []
FAIL: list[tuple[str, str]] = []


def check(ok: bool, label: str, detail: str = "") -> None:
    (PASS.append(label) if ok else FAIL.append((label, detail)))


rows = pp.parse_book()
model = "TCV1"
items = [
    {"name": f"{model} MASTER T.C", "code": None, "qty": 1, "price_model": model, "price_item": "MASTER T.C", "optional": False},
    {"name": "IB (Interface Bracket)", "code": None, "qty": 1, "price_model": model, "price_item": "IB", "optional": True, "included": True},
    {"name": f"{model} T.P", "code": None, "qty": 2, "price_model": model, "price_item": "T.P", "optional": False},
    {"name": "Pogo Pin Male (PPM)", "code": "PPM", "qty": 2, "price_model": model, "price_item": "PPM/PPF", "optional": False},
    {"name": "Pogo Pin Female (PPF)", "code": "PPF", "qty": 2, "price_model": model, "price_item": "PPM/PPF", "optional": False},
    {"name": "Cable 0.3m", "code": None, "qty": 4, "price_model": model, "price_item": "Cable 0.3m", "optional": False},
    {"name": "PneuMatic Male (PMM)", "code": "PMM", "qty": 1, "price_model": model, "price_item": "PMM/PMF", "optional": False},
    {"name": "PneuMatic Female (PMF)", "code": "PMF", "qty": 1, "price_model": model, "price_item": "PMM/PMF", "optional": False},
]
pp.attach_prices(items, rows)
st = qs.draft_state({"applicable": True, "items": items}, model=model, payload_kg=10)
lines = qs.render_lines(st, "ko", None)
check([ln["key"] for ln in lines] == ["MASTER T.C", "T.P", "PPM", "PPF", "Cable 0.3m", "PMM", "PMF", "IB"],
      "품목 순서: 본체 → 액세서리 → 체크한 옵션(IB) 맨 뒤", str([ln["key"] for ln in lines]))
check(lines[0]["desc"] == "TCV1\n- MASTER T.C (Payload 10Kg)" and lines[1]["desc"] == "TCV1\n- Tool side T.C (Payload 10Kg)"
      and lines[2]["desc"] == "PPM\n-Pogo Pin Male" and lines[0]["remark"] == "CONTROLBOX·4m 케이블 기본 포함"
      and lines[2]["remark"] == "포고핀 수, 8pin, 각 핀당 1A 전원 혹은 DIO로 사용",
      "품목 문구 = 최근 견적(가나테크) 표기 + 마스터 비고 'CONTROLBOX·4m 케이블 기본 포함'")
check(qx.total(lines) == 4_800_000, "합계 4.8M(마스터 2.3M + T.P 0.8M + PPM·PPF 1.0M + 케이블 0.2M + PMM·PMF 0.4M + IB 0.1M)", str(qx.total(lines)))
en_lines = qs.render_lines(st, "en", 1380)
check(en_lines[0]["remark"] == "Control box & 4m cable included" and en_lines[0]["unit_price"] == round(2_300_000 / 1380, 2)
      and en_lines[0]["unit_price_krw"] == 2_300_000, "영문(USD): 비고 영어, 원화 ÷ 환율(센트 반올림), 원화 단가도 남김")

form = {"customer": "가나테크", "to": "김 철 수 대표님", "contact_name": "홍길동", "contact_title": "이사",
        "contact_mobile": "010-0000-0000", "contact_email": "name@example.com"}
day = date(2026, 10, 6)
h = qx.quote_header(form, lang="ko", quote_no="S26-0182", day=day, pogo=True)
check(h["delivery"] == "통상 3 ~ 4주" and h["payment"] == "발주시 50%, 납품전 50%" and "고객사 수행 범위" in h["notes"][0],
      "빈 조건은 최근 견적 기본값, 포고핀이 있으면 배선 범위 비고 자동(C17)")
try:
    qx.quote_header({**form, "contact_email": ""}, lang="ko", quote_no="x", day=day, pogo=False)
    check(False, "필수 값 검사")
except qx.QuoteError as e:
    check("담당자 이메일" in str(e), "필수 값(고객사·담당자 이름·휴대폰·이메일)이 비면 거절")
en_form = pp.quote_form("en")
check(en_form["missing"] == [] and en_form["issuer"] is pp.QUOTE_ISSUER_EN,
      "영문 회사 정보 = 회사 프로필(quote_issuer_en)")

wb = openpyxl.load_workbook(io.BytesIO(qx.build_xlsx(h, lines)))
ws = wb.active
vals = {c.coordinate: c.value for row in ws.iter_rows() for c in row if c.value is not None}
total_cell = next(k for k, v in vals.items() if isinstance(v, str) and v.startswith("=SUM(F17:F"))
check(vals["E2"] == "S26-0182" and vals["A4"].startswith(pp.QUOTE_ISSUER["company"])
      and "담당자 : 홍 길 동 이사" in vals["A4"] and pp.QUOTE_ISSUER["company_en"] in vals.values()
      and vals["E6"] == pp.QUOTE_ISSUER["account_no"] and "가나테크" in vals["A9"],
      "머리: 견적번호, 위쪽 회사명(담당자 '홍 길 동 이사'), 아래쪽 영문 회사명, 계좌, 고객")
check(vals["F17"] == "=D17*E17" and vals["E17"] == 2_300_000 and vals["D17"] == 1 and total_cell == f"G{17 + len(lines)}"
      and vals[total_cell] == f"=SUM(F17:F{16 + len(lines)})" and wb.calculation.fullCalcOnLoad,
      "금액 = 수량×단가 수식, 합계 = SUM 수식(열 때 다시 계산)", str((vals.get("F17"), total_cell)))
check(vals["A16"] == "Commodity" and vals["B16"] == "Specifications" and "B16:C16" in {str(m) for m in ws.merged_cells.ranges}
      and vals["A17"] == "Magbot ATC" and vals["B17"] == "TCV1\n- MASTER T.C (Payload 10Kg)"
      and vals["A19"] == "Magbot ATC Accessory" and "B17:C17" in {str(m) for m in ws.merged_cells.ranges},
      "품목(Commodity A열)·규격(Specifications B:C열) 칸을 나눔 — 최근 TCV1 견적서 양식", str((vals.get("A16"), vals.get("A17"))))
mine = qs.draft_state({"applicable": True, "items": items}, model=model, payload_kg=10)
mine, _ = qs.apply_ops(mine, qs.edit_ops(mine, {"lines": [
    {"id": "L1", "name": mine["lines"][0]["name"], "commodity": "Magbot Tool Changer", "qty": 1, "unit_price": 2_300_000, "remark": ""},
    *[{k: ln.get(k) for k in ("id", "name", "qty", "unit_price", "remark")} for ln in mine["lines"][1:]],
    {"id": None, "name": "MG10", "commodity": "Magbot Gripper", "qty": 1, "unit_price": 1_000_000, "remark": ""}]}), None)
ml = qs.render_lines(mine, "ko", None)
check(ml[0]["commodity"] == "Magbot Tool Changer" and ml[1]["commodity"] == "Magbot ATC"
      and next(x for x in ml if x["desc"] == "MG10")["commodity"] == "Magbot Gripper",
      "품목 칸은 화면에서 고칠 수 있음(안 고친 줄은 기본값, 새 항목은 적은 값)", str([x["commodity"] for x in ml]))
# Comments 줄바꿈(사용자 2026-10-13) — 화면에서 띄운 줄이 저장·견적서에 그대로, 가운데 정렬(국내·해외 모두)
cm = qs.draft_state({"applicable": True, "items": items}, model=model, payload_kg=10)
cm, _ = qs.apply_ops(cm, qs.edit_ops(cm, {"form": {"comments": "www.magbot.kr\n\n  - 그리퍼   공급 건 \n맥봇\n4\n5\n6\n7"}}), None)
check(cm["form"]["comments"] == "www.magbot.kr\n- 그리퍼 공급 건\n맥봇\n4\n5\n6", "Comments 줄바꿈 저장(줄마다 공백 정리·빈 줄 제거, 6줄까지)", repr(cm["form"]["comments"]))
for lg in ("ko", "en"):
    hc = qx.quote_header({**form, "fx_rate": 1380, "comments": cm["form"]["comments"]}, lang=lg, quote_no="x", day=day, pogo=False)
    wc = openpyxl.load_workbook(io.BytesIO(qx.build_xlsx(hc, lines))).active
    check(wc["E9"].value == cm["form"]["comments"] and wc["E9"].alignment.horizontal == "center" and wc["E9"].alignment.wrap_text,
          f"견적서 Comments 칸에 줄바꿈 그대로·가운데 정렬({lg})", repr(wc["E9"].value))
# 품목 사진(사용자 2026-10-12) — Remark 칸에 작게, 행 높이 늘림
from PIL import Image as _Img  # noqa: E402
_buf = io.BytesIO()
_Img.new("RGB", (400, 300), "red").save(_buf, "JPEG")
pl = [dict(lines[0], photo="k1"), *lines[1:]]
wsp = openpyxl.load_workbook(io.BytesIO(qx.build_xlsx(h, pl, {"k1": _buf.getvalue()}))).active
check(len(wsp._images) == 5 and wsp.row_dimensions[17].height >= 48 and any(getattr(i.anchor, '_from', None) is not None and i.anchor._from.col == 6 and i.anchor._from.row == 16 for i in wsp._images),
      "품목 사진 → 견적서 Remark(G) 칸에 그림 하나 더, 그 줄 높이 늘림", str((len(wsp._images), wsp.row_dimensions[17].height)))
wsn = openpyxl.load_workbook(io.BytesIO(qx.build_xlsx(h, pl, {}))).active
check(len(wsn._images) == 4, "사진을 못 읽으면(없으면) 그림 없이 견적서는 그대로")
check(len(ws._images) == 4 and ws.print_area and int(ws.page_setup.paperSize) == int(ws.PAPERSIZE_A4),
      "제목(양식의 WordArt 그림)·로고·서명·직인 그림 4개, A4 인쇄 영역", str(len(ws._images)))
check(not any(c.font and c.font.name == "Book Antiqua" for row in ws.iter_rows() for c in row),
      "제목을 글꼴(Book Antiqua)로 쓰지 않음 — 서버 PDF 에서 다른 글꼴로 바뀌어 보였다")

he = qx.quote_header({**form, "fx_rate": 1380, "contact_name_en": "Gildong Hong", "contact_honorific": "Mr."}, lang="en", quote_no="S26-0183", day=day, pogo=True)
we = openpyxl.load_workbook(io.BytesIO(qx.build_xlsx(he, en_lines))).active
ev = {c.value for row in we.iter_rows() for c in row if isinstance(c.value, str)}
check({"Delivery Schedule", "Payment Terms", "Total Amount", "Purchase Order Confirmation", pp.QUOTE_ISSUER_EN["company_en"],
       pp.QUOTE_ISSUER_EN["account_no"], pp.QUOTE_ISSUER_EN["account_holder"]} <= ev and not any("납기" in v or "부가세" in v for v in ev) and we.title == "Quotation"
      and we["E17"].number_format == qx.USD_CENTS and he["currency"] == "USD" and he["fx_rate"] == 1380
      and "Mr. Gildong Hong (" in we["A4"].value and "이사" not in we["A4"].value,
      "영문 양식: 같은 구성, 문구 영어, 달러 서식, 영문 계좌·예금주, 담당자 'Mr. 영문 이름'(해외 기준 견적서, 2026-10-14)")
# 해외 기준 견적서 배치(사용자 2026-10-14) — 줄마다 앞 한 칸·국가명 앞 줄바꿈, 고객사 13칸 들여쓰기, Comments 14pt, 수량 서식, A열 15
a4 = we["A4"].value.split("\n")
check(all(x.startswith(" ") for x in a4) and any(x.strip().startswith("Republic of Korea") for x in a4)
      and we["A9"].value.startswith(" " * 13 + form["customer"]) and we["E9"].font.sz == 14
      and we["D17"].number_format == '#,##0_ "ea"' and we.column_dimensions["A"].width == 15.0
      and ws["A9"].value.startswith(" " * 8 + form["customer"]) and ws["E9"].font.sz == 12 and ws.column_dimensions["A"].width == 13.58,
      "영문 배치 = 해외 기준 견적서(주소 줄바꿈·들여쓰기·14pt·'#,##0_ \"ea\"'·A열 15), 국내 양식은 그대로", str(a4))
whole = [{**en_lines[0], "unit_price": 1000.0, "amount": 1000.0, "qty": 2, "unit": "sets"}]
ww = openpyxl.load_workbook(io.BytesIO(qx.build_xlsx(he, whole))).active
check(ww["E17"].number_format == qx.USD_WHOLE and ww["D17"].number_format == '#,##0_ "sets"',
      "센트가 없으면 '$ 1,000'(해외 기준 견적서 회계 서식), 수량 '2 sets'")

# 묶음·NEGO 줄(사용자 2026-10-13) — 소제목·소계·옵션(공급가액 빈칸·합계 제외)·사양·문구·NEGO, Total 아래는 적은 문구만
bl_lines = [lines[0],
            {"row": "section", "text": "주요 사양"}, {"row": "spec", "label": "무게", "value": "55kg"},
            {"row": "gap"}, {"row": "section", "text": "설치 및 시운전"},
            {**lines[1], "row": None}, {"row": "subtotal", "text": "Sub total", "amount": lines[1]["amount"]},
            {"row": "gap"}, {"row": "section", "text": "< OPTION >"}, {**lines[2], "row": "option"},
            {"row": "text", "text": "렌탈 조건은 별도 계약서"},
            {"row": "nego", "key": "nego", "commodity": "", "desc": "NEGO (특별 할인)", "qty": 1, "unit_price": -100_000, "amount": -100_000, "remark": ""}]
hb = qx.quote_header({**form, "total_note": "월 995,000원", "foot_notes": ["본 견적은 1개월 유효", "※ 중도 해지 시 위약금\n일시불 납부"]}, lang="ko",
                     quote_no="QT-X", day=day, pogo=False)
wb_ = openpyxl.load_workbook(io.BytesIO(qx.build_xlsx(hb, bl_lines))).active
cv = {c.coordinate: c.value for row in wb_.iter_rows() for c in row if c.value is not None}
tot = next(k for k, v in cv.items() if isinstance(v, str) and v.startswith("Total"))
tot_row = int(tot[1:])
check(cv["A18"] == "주요 사양" and cv["A19"] == "무게" and cv["C19"] == "55kg" and "A19:B19" in {str(m) for m in wb_.merged_cells.ranges} and "G19" not in cv and cv["A21"] == "설치 및 시운전"
      and cv["F23"] == "=SUM(F22)" and cv["A23"] == "Sub total" and cv["A25"] == "< OPTION >" and "F26" not in cv and cv["E26"] == lines[2]["unit_price"]
      and cv["A27"] == "렌탈 조건은 별도 계약서" and cv["A28"] == "NEGO (특별 할인)" and cv["F28"] == -100_000,
      "엑셀 줄: 소제목·사양·구분·소계(그 묶음 SUM)·< OPTION >(공급가액 빈칸)·문구·NEGO", str({k: cv.get(k) for k in ("A18", "F23", "F26", "A28")}))
check(cv[f"G{tot_row}"] == "=SUM(F17,F22,F28)" and cv[tot] == "Total (부가세 별도)\n월 995,000원"
      and wb_["A23"].fill.fgColor.rgb == "FFFFC000" and wb_["A28"].font.color.rgb == "FFFF0000" and wb_["F28"].font.color.rgb == "FFFF0000",
      "Total = 품목·NEGO 칸만 SUM(옵션·소계 제외), Total 아래는 직접 적은 문구만(옵션 미포함 자동 문구 없음), 소계는 주황", str((cv.get(f"G{tot_row}"), cv.get(tot))))
foot_ref = next(k for k, v in cv.items() if isinstance(v, str) and v.startswith("* 견적 외"))
fr = int(foot_ref[1:])
check(cv.get(f"A{fr + 1}") == "* 본 견적은 1개월 유효" and cv.get(f"A{fr + 2}") == "※ 중도 해지 시 위약금\n일시불 납부"
      and f"A{fr + 1}:G{fr + 1}" in {str(m) for m in wb_.merged_cells.ranges} and wb_.row_dimensions[fr + 2].height >= 34,
      "표 아래 주석: 고정 문구 다음에 칸마다 한 줄씩('*' 없으면 붙임, 칸 안 줄바꿈이면 높이 늘림)", str((cv.get(f"A{fr + 1}"), cv.get(f"A{fr + 2}"))))
check(qx.total(bl_lines) == lines[0]["amount"] + lines[1]["amount"] - 100_000, "견적 합계(영업 건 금액)도 품목·NEGO 만")

# 견적서 칸 안 줄바꿈(Alt+Enter, 사용자 2026-10-16) — 직접 수정으로 저장해도 줄이 남고(빈 줄·앞뒤 공백 정리), 엑셀 칸·행 높이에 반영
nl_st = qs.draft_state({"applicable": True, "items": items}, model=model, payload_kg=10)
l0 = nl_st["lines"][0]
nl_st, _ = qs.apply_ops(nl_st, qs.edit_ops(nl_st, {
    "form": {"place": "공장 인도\n(화물 비용 별도)", "payment": "  발주시 50% \n\n 납품전 50%\n잔금 없음", "note": "설치 일정\n발주 후 협의",
             "customer": "가나\n테크", "to": "김철수\n대표님"},
    "lines": [{**x, **({"commodity": "Magbot\nATC", "remark": "기본 포함\n케이블 4m"} if x["id"] == l0["id"] else {})}
              for x in nl_st["lines"] if x["kind"] not in ("work", "freight")]}), None)
nf = nl_st["form"]
nl0 = next(x for x in nl_st["lines"] if x["id"] == l0["id"])
check(nf["place"] == "공장 인도\n(화물 비용 별도)" and nf["payment"] == "발주시 50%\n납품전 50%\n잔금 없음" and nf["note"] == "설치 일정\n발주 후 협의"
      and nl0["commodity"] == "Magbot\nATC" and nl0["remark"] == "기본 포함\n케이블 4m",
      "칸 안 줄바꿈 저장: 납기·장소·결제·비고·품목·품목 비고는 줄 그대로(빈 줄·공백 정리)", str((nf.get("payment"), nl0.get("commodity"))))
check(nf["customer"] == "가나 테크" and nf["to"] == "김철수 대표님", "고객사·받는 분은 영업 건·파일명에도 쓰여 한 줄로", str((nf["customer"], nf["to"])))
check(qs.clean_lines("a\nb\nc\nd\ne", 200) == "a\nb\nc\nd" and qs.clean_lines("x" * 300, 60) == "x" * 60, "칸 안 줄은 4줄·글자 수 제한")
nl_lines = qs.render_lines(nl_st, "ko", None)
nl_h = qx.quote_header({**form, "place": nf["place"], "payment": nf["payment"], "note": nf["note"]}, lang="ko", quote_no="QT-NL", day=day, pogo=False)
nws = openpyxl.load_workbook(io.BytesIO(qx.build_xlsx(nl_h, nl_lines))).active
nr = next(c.row for c in nws["A"] if c.value == "Magbot\nATC")
dv = next(c.row for c in nws["A"] if c.value == {**pp.QUOTE_LABELS["ko"], **nl_h["labels"]}["delivery"])
nt = next(c.row for c in nws["A"] if c.value == {**pp.QUOTE_LABELS["ko"], **nl_h["labels"]}["note"])
check(nws[f"G{nr}"].value == "기본 포함\n케이블 4m" and nws[f"G{nr}"].alignment.wrap_text and nws.row_dimensions[nr].height >= 31
      and nws[f"G{dv}"].value == "공장 인도\n(화물 비용 별도)" and nws[f"G{dv + 1}"].value == "발주시 50%\n납품전 50%\n잔금 없음"
      and nws.row_dimensions[dv + 1].height > 34.4 and nws.row_dimensions[dv].height == 34.4
      and nws[f"B{nt}"].value == "설치 일정\n발주 후 협의" and nws.row_dimensions[nt].height > 34.4,
      "엑셀: 칸 안 줄바꿈 그대로(자동 줄바꿈), 3줄 결제조건·2줄 비고 행은 높게, 2줄까지는 원래 높이",
      str((nws.row_dimensions[dv].height, nws.row_dimensions[dv + 1].height, nws.row_dimensions[nt].height)))

print(f"PASS {len(PASS)} / FAIL {len(FAIL)}")
for label, detail in FAIL:
    print(f"  ✗ {label}: {detail}")
sys.exit(1 if FAIL else 0)
