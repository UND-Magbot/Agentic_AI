"""제안서 본문 원클릭(S5) 검증 — 서버·LLM·MinIO·DB 불필요(가짜 응답 주입).

제안서에서 가장 중요한 것 = 요청 사항이 전부 제대로 들어갔는가 + 엉뚱한 사례가 방향을 끌고 가지 않는가
+ 표절이 없는가. 이 세 가지를 중심으로 검증한다.
"""
import asyncio
import io
import json
import sys
from pathlib import Path

sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8", errors="replace")
ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "backend"))

from pptx import Presentation  # noqa: E402

from app import main, proposal_llm  # noqa: E402
from app.casebook import CaseHit  # noqa: E402
from app.proposal import (  # noqa: E402
    body as pb, case_guard, concept_image, pptx_premium, pptx_render, quality, render, requirements, service as ps,
)

PASS: list[str] = []
FAIL: list[tuple[str, str]] = []


def check(ok: bool, label: str, detail: str = "") -> None:
    (PASS.append(label) if ok else FAIL.append((label, detail)))


SYS_WITH = "- 도구가 attachment_ids 인자를 받으면 **위 1개 ID 전부를 빠짐없이** 위 순서 그대로 전달: [42]."
user = lambda t: [{"role": "user", "content": t}]  # noqa: E731
PROMPT = "제안서 본문 작성 기능 수행"

# ── 1) 의도 감지 ────────────────────────────────────────────────────────────
check(main._detect_proposal_body_intent(user(PROMPT), SYS_WITH), "표준 프롬프트+첨부 → 감지")
check(not main._detect_proposal_body_intent(user(PROMPT), ""), "첨부 없으면 미감지")
check(not main._detect_proposal_body_intent(user("제안서가 뭐야? 설명해줘"), SYS_WITH), "정보 질문 미감지")
check(not main._detect_concept_map_intent(user(PROMPT), SYS_WITH), "개념도 fast-path 가 가로채지 않음")
check(main._detect_concept_map_intent(user("제안서용 공정 개념도 만들어줘"), SYS_WITH),
      "'제안서용 개념도' 는 개념도가 먼저")
check(not main._detect_expense_reconcile_intent(user(PROMPT), SYS_WITH), "영수증 대조가 가로채지 않음")
check(not main._detect_fund_plan_intent(user(PROMPT), SYS_WITH), "자금계획이 가로채지 않음")
check(not main._detect_meeting_intent(user(PROMPT)), "회의록이 가로채지 않음")

# ── 2) 수치 근거 가림 ───────────────────────────────────────────────────────
REQ = ("AMR 로 700kg 파레트를 6분 걸리던 F→D 구간 이송. 로봇 2대 검토. 1,200mm 통로. "
       "문턱 높이 확인 필요. ERP 연동 희망.")
req_nums = pb._numbers_in(REQ)
case_nums = pb._numbers_in("- 일일생산량 (KPI): 16.6% 증가 / 투자회수 (연간): 3.7년")
t, m = pb.mask_unsupported_numbers("700kg 파레트를 6분 구간에서 자동 이송", req_nums, case_nums)
check(m == [] and "700kg" in t, "요청에 있는 수치는 유지", str(m))
t, m = pb.mask_unsupported_numbers("사이클타임 30% 단축, 인원 2명 절감", req_nums, case_nums)
check(m == ["30%", "2명"] and t.count(pb.NUMBER_MASK) == 2, "근거 없는 수치 가림(요청 2대 ≠ 본문 2명)", f"{m} {t}")
t, m = pb.mask_unsupported_numbers("유사 사례(사례 2024-02)에서는 일일생산량 16.6% 증가", req_nums, case_nums)
check(m == [], "사례 번호를 밝힌 사례 수치는 유지", str(m))
t, m = pb.mask_unsupported_numbers("일일생산량 16.6% 증가 기대", req_nums, case_nums)
check(m == ["16.6%"], "출처 없이 옮긴 사례 수치는 가림", str(m))
t, m = pb.mask_unsupported_numbers("통로 1200mm 확인, 1단계 맵핑, 로봇2대", req_nums, case_nums)
check(m == [], "천단위 쉼표 정규화·'1단계' 제외·붙은 수치 인식", str(m))

t, m = pb.mask_unsupported_specs("고온/유분 환경에 적합한 IP69K 등급 로봇, SUS316L 그리퍼", "어묵 포장, HACCP")
check(m == ["IP69K", "SUS316L"] and t.count(pb.SPEC_MASK) == 2, "요청에 없는 등급·규격 표기 가림(실측 IP69K)", f"{m} {t}")
t, m = pb.mask_unsupported_specs("한화로봇 14kg Class 100 적합성 확인", "한화로봇 14kg Class 100 되는지 확인")
check(m == [], "요청에 있는 등급(Class 100)은 유지", str(m))
t, m = pb.mask_unsupported_specs("PLC/IPC 기반 통합 제어, 상위 PLC 연동", "어묵 포장, HACCP")
check(m == [], "PLC 는 등급 표기가 아님(실측 오탐: 'PL'+'C' 를 안전 등급으로 봄)", str(m))
t, m = pb.mask_unsupported_specs("안전 등급 PL d, Cat 3 구성", "어묵 포장")
check(m == ["PL d", "Cat 3"], "안전 등급 표기(PL d, Cat 3)는 요청에 없으면 가림", str(m))

# ── 가짜 LLM — 호출 종류를 프롬프트로 구분 ─────────────────────────────────
CASE_TEXT = "물류이송용 AMR을 이용하여 생산된 제품을 파레트에 적재하여 반제품 창고에 이송한다."


def hit(year, no, title, company, text=CASE_TEXT, kpis=None):
    return CaseHit(year=year, case_no=no, source_path=f"casebook/kitech_robot_consulting_{year}",
                   doc_title=f"{year} 로봇 엔지니어링 컨설팅 사례집(KITECH)", title=title, company=company,
                   industry="자동차 부품", printed_pages="20-23", score=0.69, matched_sections=["컨설팅 결과"],
                   kpis=kpis or [], sections={"개요": text, "컨설팅 결과": text})


H_AMR = hit(2024, 2, "AMR을 이용한 물류 자동화", "㈜ 대광소결 금속",
            kpis=[{"label": "일일생산량 (KPI)", "value": "16.6", "unit": "%", "direction": "증가"}])
H_DRILL = hit(2024, 19, "임플란트 시술용 툴-드릴 생산공정", "오스템임플란트㈜", text="드릴 연마 공정")
H_BOLT = hit(2024, 11, "복합 보호계전기 조립 볼팅 공정", "비케이전자㈜", text="볼팅 공정")


def sec(sid, bullet="요청 원문 기준 검토 필요"):
    """섹션 — 코드 구성 검사(quality.MIN_CARDS)를 통과하도록 최소 카드 수를 채운다. 첫 카드에 bullet."""
    n = quality.MIN_CARDS.get(sid, 1)
    cards = [{"heading": "카드", "bullets": [bullet]}] + [
        {"heading": f"카드{i}", "bullets": [f"{sid} 보조 항목 {i}"]} for i in range(2, n + 1)]
    return {"id": sid, "title": sid, "headline": f"{sid} 요지", "cards": cards,
            "banner_label": "핵심", "banner": "배너"}


EXTRACT = {"items": [
    {"kind": "fact", "text": "700kg 파레트 이송", "evidence": "700kg 파레트를 6분 걸리던 F→D 구간 이송"},
    {"kind": "question", "text": "로봇 대수 검토", "evidence": "로봇 2대 검토"},
    {"kind": "question", "text": "문턱 높이 확인", "evidence": "문턱 높이 확인 필요"},
    {"kind": "constraint", "text": "ERP 연동", "evidence": "ERP 연동 희망"},
    {"kind": "goal", "text": "무인 24시간 운영", "evidence": "24시간 무인으로 돌리고 싶다"},   # 원문에 없음 → 버림
]}
JUDGE = {"cases": [
    {"ref": "2024-02", "verdict": "same_process", "aspect": "문턱·경사 조건에서 AMR 가반하중 선정",
     "request_quote": "700kg 파레트를 6분 걸리던", "case_quote": "파레트에 적재하여 반제품 창고에 이송", "reason": ""},
    {"ref": "2024-19", "verdict": "unrelated", "aspect": "", "reason": "드릴 가공"},
    {"ref": "2024-11", "verdict": "shared_problem", "aspect": "", "reason": "관점 없음"},  # 관점 없음 → 제외
]}
DRAFT = {
    "title": "F→D 구간 AMR 물류 자동화 제안", "subtitle": "안전 경로부터 검증", "customer": "대광",
    "sections": [sec("key_message", "사이클타임 30% 단축"),
                 sec("situation", "700kg 파레트를 6분 걸리던 F→D 구간 수동 이송, 1,200mm 통로"),
                 sec("risk"), sec("solution", "AMR을 이용하여 생산된 제품을 파레트에 적재하여 이송"), sec("system"),
                 {"id": "effects", "headline": "", "cards": []}],   # effects 비어 있음, next_steps 없음
    "case_insights": [{"case_ref": "2024-02", "insight": "문턱·경사 확인 후 기종 선정"},
                      {"case_ref": "2023-02", "insight": "검색에 없는 사례"}],
    "open_questions": ["문턱 높이 확인", "로봇 2대 운용 동선"],
}
REPAIR = {"sections": [sec("effects"), sec("next_steps")]}
REVIEW_1 = {"items": [
    {"id": "a1", "covered": True, "where": "고객 현황과 과제", "response": "700kg 파레트 F→D 이송"},
    {"id": "a2", "covered": True, "where": "시스템 구성", "response": "대수 산정 방법"},   # 본문에 근거 토큰 없음 → 뒤집힘
    {"id": "a3", "covered": False},
    {"id": "a4", "covered": False},
], "issues": [{"section": "risk", "problem": "일반론", "text": "요청 원문 기준 검토 필요", "fix": "공정 조건 명시"}]}
REVISE = {"sections": [
    sec("system", "로봇 2대 검토: 동선·충전 주기로 대수 산정, ERP 연동 인터페이스 확인"),
    sec("risk", "문턱 높이 확인 전 주행 시험으로 전복 위험 검증"),
]}
REVIEW_2 = {"items": [
    {"id": "a1", "covered": True, "where": "고객 현황과 과제", "response": "700kg 파레트 F→D 이송"},
    {"id": "a2", "covered": True, "where": "시스템 구성(안)", "response": "동선·충전 주기로 대수 산정"},
    {"id": "a3", "covered": True, "where": "기술 리스크", "response": "주행 시험으로 검증"},
    {"id": "a4", "covered": True, "where": "시스템 구성(안)", "response": "ERP 연동 인터페이스 확인"},
], "issues": []}
REWRITE = {"rewrites": [{"i": 0, "text": "완성 파레트를 창고까지 무인 반송하는 AMR 구성"}]}

calls: list[str] = []
script: dict[str, list] = {}


def kind_of(messages):
    head = messages[0]["content"]
    if "요구 항목을 원문에서" in head:
        return "extract"
    if "참고 자료로 적절한지" in head:
        return "judge"
    if "최종 검토를 한다" in head:
        return "review"
    if "아래 제안서 섹션을 고쳐라" in head:
        return "revise"
    if "참고 자료(다른 기업 사례" in head:
        return "rewrite"
    if "다음 섹션이 빠졌거나" in messages[-1]["content"]:
        return "repair"
    return "draft"


async def fake_chat(messages, **kw):
    k = kind_of(messages)
    calls.append(k)
    seq = script.get(k)
    if not seq:
        raise AssertionError(f"예상 밖 호출: {k}")
    resp = seq.pop(0) if len(seq) > 1 else seq[0]
    if isinstance(resp, Exception):
        raise resp
    return resp if isinstance(resp, str) else json.dumps(resp, ensure_ascii=False)


proposal_llm.chat = fake_chat


def reset(**kw):
    calls.clear()
    script.clear()
    base = {"extract": [EXTRACT], "judge": [JUDGE], "draft": [DRAFT], "repair": [REPAIR],
            "review": [REVIEW_1, REVIEW_2], "revise": [REVISE], "rewrite": [REWRITE]}
    base.update(kw)
    script.update({k: list(v) for k, v in base.items()})


# ── 3) 요구 항목 추출 — 원문 근거 없는 항목 버림 + 원문 표기 항목 ─────────────
reset()
reqs, dropped = asyncio.run(requirements.extract(REQ))
ai = [r for r in reqs if r.source == "ai"]
rule = [r for r in reqs if r.source == "rule"]
check([r.text for r in ai] == ["700kg 파레트 이송", "로봇 대수 검토", "문턱 높이 확인", "ERP 연동"],
      "AI 추출 항목", str([r.text for r in ai]))
check(dropped == ["무인 24시간 운영"], "원문에 없는 인용 → 항목 버림(지어낸 요구 차단)", str(dropped))
check({"700kg", "ERP"} <= {r.text for r in rule}, "원문 표기 항목(수치·영문)도 함께", str([r.text for r in rule]))

# ── 4) 사례 관련성 판정 ─────────────────────────────────────────────────────
reset()
kept, drop, ok = asyncio.run(case_guard.judge(REQ, [H_DRILL, H_AMR, H_BOLT]))
check(ok and [v.hit.ref for v in kept] == ["2024-02"], "같은 공정만 채택", str([v.hit.ref for v in kept]))
check({v.hit.ref for v in drop} == {"2024-19", "2024-11"}, "무관·관점 없는 사례 제외", str([v.hit.ref for v in drop]))
Q = {"request_quote": "700kg 파레트를 6분 걸리던", "case_quote": "드릴 연마 공정"}
reset(judge=[{"cases": [
    {"ref": "2024-02", "verdict": "shared_problem", "aspect": "파레트 단위 중량물 AMR 가반하중 선정",
     "request_quote": "700kg 파레트", "case_quote": "파레트에 적재하여"},
    {"ref": "2024-19", "verdict": "same_process", "aspect": "드릴 연마 방식", **Q}]}])
kept, drop, ok = asyncio.run(case_guard.judge(REQ, [H_AMR, H_DRILL, H_BOLT]))
check([v.hit.ref for v in kept] == ["2024-19", "2024-02"], "같은 공정 판정이 앞", str([v.hit.ref for v in kept]))
check([v.verdict for v in drop] == ["unjudged"], "판정 누락 사례도 제외", str([v.verdict for v in drop]))
# 근거 인용 검증 — 2026-09-28 실측 사고(가반하중 14kg 를 제품 무게로, 피로도를 공통 문제로 보고 연마 사례 채택)
reset(judge=[{"cases": [
    {"ref": "2024-02", "verdict": "same_process", "aspect": "AMR 선정", "request_quote": "",
     "case_quote": "파레트에 적재하여"},                                   # 요청 인용 없음
    {"ref": "2024-19", "verdict": "shared_problem", "aspect": "드릴 공정 비전 검사",
     "request_quote": "비전으로 드릴 검사", "case_quote": "드릴 연마 공정"},  # 요청에 없는 인용(지어냄)
    {"ref": "2024-11", "verdict": "shared_problem", "aspect": "중량물 반복 작업 피로 해결",
     "request_quote": "700kg 파레트", "case_quote": "복합 보호계전기 조립 볼팅 공정"},  # 일반 문제만 공유
]}])
kept, drop, ok = asyncio.run(case_guard.judge(REQ, [H_AMR, H_DRILL, H_BOLT]))
reasons = {v.hit.ref: v.reason for v in drop}
check(kept == [], "근거 없는 채택 판정은 전부 제외", str([v.hit.ref for v in kept]))
check(reasons["2024-02"].startswith("요청 원문 인용 없음"), "요청 인용 없음 → 제외", reasons["2024-02"])
check(reasons["2024-19"].startswith("요청 원문 인용 없음·불일치"), "지어낸 요청 인용 → 제외", reasons["2024-19"])
check(reasons["2024-11"].startswith("일반 문제만 공유"), "피로·반복작업 같은 일반 문제만 → 제외", reasons["2024-11"])
reset(judge=[{"cases": [{"ref": "2024-02", "verdict": "same_process", "aspect": "AMR 선정",
                         "request_quote": "700kg 파레트", "case_quote": "무인 지게차로 이송"}]}])
kept, drop, ok = asyncio.run(case_guard.judge(REQ, [H_AMR]))
check(kept == [] and drop[0].reason.startswith("사례 원문 인용 없음"), "사례에 없는 인용 → 제외",
      drop[0].reason if drop else "")
W = ("현재 가접 후 4면 본용접을 작업자 1인이 수동으로 하고 있고, 용접사 숙련도에 따라 품질 편차가 큼.\n"
     "예열은 토치로 약 180℃, 제품 무게 350kg")
check(case_guard._quoted("가접 후 4면 본용접을 ... 용접사 숙련도에 따라 품질 편차가 크며", W),
      "인용: '…'로 이은 조각 + 어미만 바뀐 경우 인정(실측 형태)")
check(not case_guard._quoted("비전으로 드릴 검사 요청", W), "인용: 원문에 없는 조각은 불인정")
check(not case_guard._quoted("용접", W), "인용: 너무 짧은 단어는 불인정")
check(case_guard._quoted("2단계 예열 공정: 가접 후 4면, 예열은 토치로 약 180℃ / 1인 작업", W),
      "인용: 쉼표·콜론·슬래시로 재배열한 인용도 원문 조각이 있으면 인정(실측 형태)")
check(case_guard._quoted("약액 주입", "약액 주입 210초"), "인용: 짧지만 정확한 핵심 구절 인정")
reset(judge=["이건 JSON 이 아님"])
kept, drop, ok = asyncio.run(case_guard.judge(REQ, [H_AMR]))
check(not ok and kept == [], "판정 실패 → 사례 미사용(엉뚱한 사례보다 사례 없음)")


# ── 5) 전 과정(write_body) — 가짜 검색·확정 예시 ─────────────────────────────
async def fake_search(query):
    fake_search.q = query
    return [H_DRILL, H_AMR, H_BOLT]


async def no_examples(request):
    return ""


ps._search, ps._confirmed_examples = fake_search, no_examples
reset()
body, report, reqs, verdicts, _ = asyncio.run(ps.write_body(REQ))
ids = [s.id for s in body.sections]
check(fake_search.q == REQ[:600], "사례 질의 = 요청 원문(의역 항목 아님)")
check(calls[:4] == ["extract", "judge", "draft", "repair"], "호출 순서", str(calls))
check([v.hit.ref for v in verdicts] == ["2024-02"], "본문에는 판정 통과 사례만", str([v.hit.ref for v in verdicts]))
check(report.dropped_cases == ["2024-19 임플란트 시술용 툴-드릴 생산공정", "2024-11 복합 보호계전기 조립 볼팅 공정"],
      "제외 사례 보고", str(report.dropped_cases))
check(ids == ["key_message", "situation", "risk", "solution", "system", "effects",
              "requirements", "references", "open_items", "next_steps"], "최종 섹션 순서", str(ids))
check(report.repaired == ["effects", "next_steps"], "누락 섹션 보완", str(report.repaired))
check(report.review_rounds == 1 and calls.count("review") == 2 and calls.count("revise") == 1,
      "검토 → 보완 1회 → 재검토 통과", str(calls))
check(report.review_issues == ["일반론: 요청 원문 기준 검토 필요"], "고친 지적 사항 기록", str(report.review_issues))
by_text = {r.text: r for r in reqs}
check(all(r.covered for r in reqs if r.source == "ai"), "보완 후 AI 항목 전부 반영",
      str([(r.text, r.covered) for r in reqs if r.source == "ai"]))
check(report.uncovered == [], "미반영 없음", str(report.uncovered))
check(by_text["로봇 대수 검토"].where == "시스템 구성(안)", "반영 위치 기록")
check(report.masked_numbers == ["key_message: 30%"], "근거 없는 수치 가림", str(report.masked_numbers))
check(report.rewritten == 1 and report.copied == [], "사례 문장 복사 → 재작성", f"{report.rewritten} {report.copied}")
check(not any("적재하여" in b for s in body.sections for c in s.cards for b in c.bullets), "표절 문장이 남지 않음")
req_sec = next(s for s in body.sections if s.id == "requirements")
check(req_sec.table[0] == ["요구사항", "원문", "제안 반영", "상태"] and len(req_sec.table) == 5
      and not any("1.200" in c for row in req_sec.table for c in row),
      "대응표: 머리글 + AI 항목 4행(반영된 원문 표기 항목은 생략)", str(req_sec.table))
check(all(row[3] == "반영" for row in req_sec.table[1:]), "대응표 상태 = 반영", str(req_sec.table))
check(req_sec.headline.startswith("주요 요구사항 4건 모두 반영 · 원문 수치·장비 표기"),
      "대응표 머리말 = 표 행(주요 요구사항)과 원문 표기 대조를 나눠 표기(숫자 불일치 방지)", req_sec.headline)
oi_all = next(s for s in body.sections if s.id == "open_items").cards[0].bullets
check(oi_all == ["문턱 높이 확인", "로봇 2대 운용 동선"], "본문에서 답한 요청은 확인 항목에 되풀이하지 않음", str(oi_all))
ref = next(s for s in body.sections if s.id == "references")
check(len(ref.cards) == 1 and ref.cards[0].heading.startswith("사례 2024-02"), "참고 사례 = 판정 통과 1건")
check(any(b.startswith("참고 관점(같은 공정): 문턱") for b in ref.cards[0].bullets), "참고 관점 표시")
check(any("16.6%" in b for b in ref.cards[0].bullets), "사례 KPI 는 코드가 채움")
check(set(body.case_insights) == {"2024-02"}, "검색에 없는 사례 시사점 버림")

# 프롬프트 — 사례 판정·관점, 요구 항목 원문 인용, 전문가 역할
msgs = pb.build_messages(REQ, reqs, verdicts)
check("참고 관점: 문턱·경사 조건에서 AMR 가반하중 선정" in msgs[1]["content"], "초안 프롬프트에 사례 참고 관점")
check("임플란트" not in msgs[1]["content"], "제외 사례는 프롬프트에 없음")
check("〔원문: 로봇 2대 검토〕" in msgs[1]["content"], "요구 항목에 원문 인용")
check("수석 기술영업" in msgs[0]["content"] and "표절 금지" in msgs[0]["content"]
      and "제안 방향은 오직 [고객 요청 원문]에서" in msgs[0]["content"], "전문가 역할·표절 금지·방향 원칙")
check("맥봇(Magbot)" in msgs[0]["content"], "자사 제품 정식 표기·언급 조건")

# ── 6) 반영 판정 교차 확인 · 보완 한도 · 미반영 보고 ──────────────────────────
reset(review=[REVIEW_1], revise=[{"sections": []}])  # 보완이 아무것도 못 고침
body2, report2, reqs2, _, _ = asyncio.run(ps.write_body(REQ))
r2 = {r.text: r for r in reqs2}
check(not r2["로봇 대수 검토"].covered,
      "모델이 '반영'이라 해도 원문 수치(2대)가 본문에 없으면 미반영(확인 질문에만 있는 건 답이 아님)",
      str(r2["로봇 대수 검토"]))
check(report2.uncovered == ["ERP 연동", "원문 표기 '로봇 2대 검토'", "원문 표기 'ERP 연동 희망'"],
      "미반영 보고(질문 항목은 확인 필요로 이관, 원문 표기는 원문 구절로)", str(report2.uncovered))
t2 = next(s for s in body2.sections if s.id == "requirements").table
status = {row[0]: row[3] for row in t2[1:]}
check(status.get("(제약) ERP 연동") == "미반영 — 담당자 보완 필요", "대응표에 미반영 표시", str(status))
check(status.get("(확인 요청) 로봇 대수 검토") == "확인 필요", "미답 질문은 확인 필요", str(status))
oi = next(s for s in body2.sections if s.id == "open_items").cards[0].bullets
check(oi[:2] == ["로봇 대수 검토", "문턱 높이 확인"], "미답 질문이 확인 항목 맨 앞", str(oi))
cls = requirements.Requirement("a9", "question", "Class 100 가능 여부", "한화로봇 14kg Class 100 되는지 확인")
check(not quality._tokens_present(cls, pb.rulecheck.normalize("로봇 14kg Class 사양 확인")),
      "교차 확인: 인증 등급(Class 100)이 본문에 없으면 '반영' 판정 뒤집음(실측 누락)")
check(quality._tokens_present(cls, pb.rulecheck.normalize("한화 14kg 로봇의 Class 100 클린룸 대응 확인")),
      "교차 확인: Class 100 이 있으면 통과")
reset(review=[REVIEW_1], revise=[REVISE])
body3, report3, *_ = asyncio.run(ps.write_body(REQ))
check(calls.count("review") == quality.MAX_ROUNDS + 1 and report3.review_rounds == quality.MAX_ROUNDS,
      "보완은 최대 MAX_ROUNDS 회", str(calls))

# ── 7) 표절 — 재작성이 또 베끼면 문장 삭제, 요청 원문 인용은 표절 아님 ─────────
b = pb.ProposalBody("t", "", "", [pb.Section("system", "구성", "요지", [pb.Card("카드", [
    "AMR을 이용하여 생산된 제품을 파레트에 적재하여 이송", "700kg 파레트를 6분 걸리던 F→D 구간 이송"])])])
rep = pb.BodyReport()
reset(rewrite=[{"rewrites": [{"i": 0, "text": "AMR을 이용하여 생산된 제품을 파레트에 적재하여 옮김"}]}])
asyncio.run(quality.remove_plagiarism(b, [CASE_TEXT, "700kg 파레트를 6분 걸리던 F→D 구간 이송 사례"], REQ, rep))
check(b.sections[0].cards[0].bullets == ["700kg 파레트를 6분 걸리던 F→D 구간 이송"],
      "재작성도 겹치면 삭제, 요청 원문과 같은 문장은 유지", str(b.sections[0].cards[0].bullets))
check(len(rep.copied) == 1 and rep.rewritten == 0, "삭제 보고", str(rep.copied))

# ── 8) 출력 ─────────────────────────────────────────────────────────────────
md = render.to_markdown(body)
check("초안(Draft)" in md and "| 요구사항 | 원문 | 제안 반영 | 상태 |" in md, "마크다운 대응표")
data, rrep = pptx_render.to_pptx(body)
prs = Presentation(io.BytesIO(data))
slides = list(prs.slides)
all_text = ["".join(r.text for p in sh.text_frame.paragraphs for r in p.runs)
            for s in slides for sh in s.shapes if sh.has_text_frame]
check(round(prs.slide_width / 914400, 2) == 13.33 and round(prs.slide_height / 914400, 1) == 7.5, "16:9 슬라이드")
check(len(slides) == 1 + len(body.sections), "표지 + 섹션당 1장", f"{len(slides)} vs {1 + len(body.sections)}")
check(any("초안(Draft)" in t for t in all_text[:5]), "표지 초안 표기")
check(rrep.overflow == [], "넘친 글상자 없음", str(rrep.overflow))
tables = [sh.table for s in slides for sh in s.shapes if sh.has_table]
check(len(tables) == 1 and tables[0].cell(0, 0).text == "요구사항", "대응표 = 표")
check(any(t.startswith("출처:") for t in all_text), "사례 출처 표시")
check(any("고객사 대광" in t for t in all_text), "바닥글 고객사")
runs = [r for s in slides for sh in s.shapes if sh.has_text_frame for p in sh.text_frame.paragraphs for r in p.runs]
check(runs and all(r._r.get_or_add_rPr().get("lang") == "ko-KR" for r in runs),
      "모든 글자 ko-KR — 한글 단어 단위 줄바꿈(PowerPoint 실측)")
# 표 페이지 나눔 — 행이 많으면 (1/n) 슬라이드로
many = pb.Section("requirements", "고객 요구사항 대응표", "요지", [], table=[["요구사항", "원문", "제안 반영", "상태"]] + [
    [f"항목 {i}", "원문 " * 12, "반영 위치 설명 " * 6, "반영"] for i in range(30)])
data2, _ = pptx_render.to_pptx(pb.ProposalBody("t", "", "", [many]))
titles = ["".join(r.text for p in sh.text_frame.paragraphs for r in p.runs)
          for s in Presentation(io.BytesIO(data2)).slides for sh in s.shapes if sh.has_text_frame]
check(sum("고객 요구사항 대응표 (" in t for t in titles) >= 2, "긴 대응표는 여러 장으로 나눔")
# 글자 크기 맞춤
size_short, over_s = pptx_render.fit_font(["짧은 문장"], 3.0, 1.0, 16, 9)
size_long, over_l = pptx_render.fit_font(["아주 긴 문장입니다 " * 20], 3.0, 1.0, 16, 9)
check(size_short == 16 and not over_s and size_long < 16, "긴 글은 글자 크기를 줄임", f"{size_short} {size_long}")
check(pptx_render.fit_font(["넘치는 글 " * 400], 2.0, 0.5, 16, 9)[1], "최소 크기로도 넘치면 경고")

# ── 8-0) 디자인 버전 pptx ────────────────────────────────────────────────────
pdata, prep = pptx_premium.to_pptx_premium(body)
pp = Presentation(io.BytesIO(pdata))
ptext = ["".join(r.text for p in sh.text_frame.paragraphs for r in p.runs)
         for s in pp.slides for sh in s.shapes if sh.has_text_frame]
groups = pptx_premium._chapters(body)
check(len(pp.slides) == 2 + len(groups) + len(body.sections) + 1,
      "디자인판: 표지·목차 + 장 구분 + 섹션 + 마무리", f"{len(pp.slides)}")
check(prep.overflow == [], "디자인판 넘친 글상자 없음", str(prep.overflow))
check("목차" in ptext and "감사합니다" in ptext, "목차·마무리 슬라이드")
check(any(t.startswith("Ⅲ.  근거와 검증") for t in ptext), "장 위치 표시(breadcrumb)")
check("PREPARED FOR" in ptext, "고객사가 있으면 표지에 PREPARED FOR")
nb = pb.ProposalBody("t", "", "", body.sections)
check("PREPARED FOR" not in ["".join(r.text for p in sh.text_frame.paragraphs for r in p.runs)
                             for s in Presentation(io.BytesIO(pptx_premium.to_pptx_premium(nb)[0])).slides
                             for sh in s.shapes if sh.has_text_frame], "고객사 없으면 PREPARED FOR 생략")
tiles_total = next(t for t in ptext if t.endswith("건") and t[:-1].isdigit())
check(int(tiles_total[:-1]) == len(req_sec.table) - 1, "숫자 타일 = 대응표 행 수", tiles_total)
bad = [sh for s in pp.slides for sh in s.shapes if sh.width <= 0 or sh.height <= 0]
check(not bad, "크기 0 이하 도형 없음(PowerPoint 열기 실패 원인)")
# 좁은 슬라이드(하단 강조 띠 + 긴 관점)에서도 사례 도형 높이가 음수가 되지 않음
long_ref = next(s for s in body.sections if s.id == "references")
long_ref.cards[0].bullets = ["참고 관점(같은 공정): " + "아주 긴 참고 관점 설명 " * 15] + long_ref.cards[0].bullets[1:]
pdata2, _ = pptx_premium.to_pptx_premium(body)
bad2 = [sh for s in Presentation(io.BytesIO(pdata2)).slides for sh in s.shapes if sh.height <= 0]
check(not bad2, "긴 사례 관점에서도 음수 크기 도형 없음(실측 열기 실패 회귀)")

# 컨셉 이미지: 있으면 구성(system) 뒤에 한 장, 없으면 생략
from PIL import Image  # noqa: E402
_buf = io.BytesIO()
Image.new("RGB", (160, 90), (40, 80, 120)).save(_buf, "PNG")
PNG = _buf.getvalue()
cdata, crep = pptx_premium.to_pptx_premium(body, concept_image=PNG)
cp = Presentation(io.BytesIO(cdata))
ctext = [["".join(r.text for p in sh.text_frame.paragraphs for r in p.runs) for sh in s.shapes if sh.has_text_frame]
         for s in cp.slides]
cidx = next((i for i, t in enumerate(ctext) if pptx_premium.CONCEPT_TITLE in t), -1)
check(len(cp.slides) == len(pp.slides) + 1 and cidx > 0, "이미지가 있으면 컨셉 슬라이드 1장 추가", f"{len(cp.slides)} {cidx}")
sys_title = next(s.title for s in body.sections if s.id == "system")
check(cidx > 0 and sys_title in ctext[cidx - 1], "컨셉 슬라이드는 구성 섹션 바로 뒤", f"{sys_title} | {ctext[cidx - 1][:6]}")
check(cidx > 0 and any(t == pptx_premium.CONCEPT_NOTE for t in ctext[cidx]), "실제 설비와 다를 수 있다는 주석")
check(cidx > 0 and any(sh.shape_type == 13 for sh in cp.slides[cidx].shapes), "그림 도형 포함")
check(crep.overflow == [], "컨셉 슬라이드 넘침 없음", str(crep.overflow))

prompt = concept_image.build_prompt(body)
check(len(prompt) <= concept_image.PROMPT_MAX and "글자·숫자·로고" in prompt, "그림 지시: 한도 안 + 글자 없는 그림 요청")
check("[제안 구성]" in prompt, "그림 지시에 제안 구성 포함", prompt[-300:])


sent: list[str] = []


async def fake_gen(prompt, *, user_id=None):
    sent.append(prompt)
    return PNG, "image/png", ""


_orig = (concept_image.codex_client.is_configured, concept_image.codex_client.generate_image)
concept_image.codex_client.is_configured = lambda: False
img, note = asyncio.run(concept_image.generate(body))
check(img is None and "설정되지 않아" in note, "브리지 미설정이면 이미지 없이 진행")
concept_image.codex_client.is_configured = lambda: True
concept_image.codex_client.generate_image = fake_gen
img, note = asyncio.run(concept_image.generate(body))
check(img == PNG and sent == [prompt], "그림 지시를 이미지 생성에 넘김(가림은 codex_client._post 관문)", str(sent)[:200])


async def gen_fail(prompt, *, user_id=None):
    raise concept_image.codex_client.CodexError("외부 전송 전 가림 단계가 실패해 전송하지 않았습니다")

concept_image.codex_client.generate_image = gen_fail
img, note = asyncio.run(concept_image.generate(body))
check(img is None and "생략" in note, "생성·가림 실패해도 예외 없이 생략", note)
concept_image.codex_client.is_configured, concept_image.codex_client.generate_image = _orig

check(pptx_premium._parse_kpis("생산성 12.5% 증가 · 공정불량률 N/A · 작업인원 3명 감소")
      == [("생산성", "12.5%", "증가"), ("작업인원", "3명", "감소")], "사례 KPI 해석(압축형, N/A 제외)")
check(pptx_premium._parse_kpis("생산성 (생산C/T): 12.5% 증가, 투자회수 (연간): 4년")
      == [("생산성", "12.5%", "증가"), ("투자회수", "4년", "")], "사례 KPI 해석(구형 표기)")

# ── 8-1) 코드 구성·사례 용어 검사 ─────────────────────────────────────────────
thin = pb.ProposalBody("t", "", "", [pb.Section("risk", "리스크", "요지", [pb.Card("리스크", ["a", "b", "c"])])])
check([d["section"] for d in quality.structure_issues(thin)] == ["risk"], "리스크를 카드 하나에 몰아 쓰면 지적")
full = pb.ProposalBody("t", "", "", [pb.Section("system", "구성", "요지", [
    pb.Card("로봇", ["a", "b"]), pb.Card("그리퍼", ["c", "d"]), pb.Card("비전", ["e"]), pb.Card("주변 설비", ["f", "g"])])])
check(quality.shrinks(full, pb.Section("system", "구성", "요지", [pb.Card("주변 설비", ["f", "g", "h"])])),
      "보완본이 카드 4→1장으로 줄면 거부(실측 내용 손실)")
check(not quality.shrinks(full, pb.Section("system", "구성", "요지", [
    pb.Card("로봇", ["a", "b"]), pb.Card("그리퍼", ["c", "d2"]), pb.Card("비전", ["e"]), pb.Card("주변", ["f", "g"])])),
      "같은 분량으로 고친 보완본은 받음")
cer = case_guard.CaseVerdict(hit(2024, 15, "세라믹 제품 검사공정", "㈜신한세라믹"), "shared_problem", "비전 검사", "")
leak = pb.ProposalBody("t", "", "", [pb.Section("system", "구성", "요지", [
    pb.Card("비전", ["세라믹 카트리지 정렬용 비전 구성", "보틀 목 기울어짐 검출 비전"])])])
li = quality.case_term_issues(leak, [cer], "보틀 목 기울어짐 비전 검출, 약액 주입")
check(len(li) == 1 and "세라믹" in li[0]["fix"], "참고 사례 고유 용어(세라믹) 유출 지적", str(li))
check(quality.case_term_issues(leak, [cer], "세라믹 카트리지 비전 정렬") == [], "요청에도 있는 말은 유출 아님")
bbc = case_guard.CaseVerdict(hit(2024, 10, "로봇 활용 칫솔모 적재 공정", "BBC㈜"), "shared_problem", "파지", "")
bbc.hit.industry = "모노필라멘트, 기능성에어필터, 덴탈 케어용 소재(테이퍼모 등) 제조"
gen = pb.ProposalBody("t", "", "", [pb.Section("system", "구성", "요지", [
    pb.Card("로봇", ["고온/유분 환경 대응형 로봇(비부식성 소재 확인 필요)"])])])
check(quality.case_term_issues(gen, [bbc], "어묵 포장 라인, 고온 유분") == [],
      "업종의 일반 명사('소재')는 유출로 보지 않음(실측 오탐)")

# ── 9) 서비스 + 챗 스트림 ───────────────────────────────────────────────────
stages: list[str] = []
saved: list[tuple[int, str, int]] = []
drafts: list[dict] = []


async def fake_fetch(ids, *, what):
    return REQ, 7


async def fake_save(owner, filename, data, mime):
    saved.append((owner, filename, len(data)))
    return ps.ResultFile(filename, "/api/attachments/9/download", 9)


EXAMPLE = "[확정된 과거 제안서 본문 예시] 구성·표기만 참고"


async def fake_examples(request):
    return EXAMPLE


async def fake_record(owner, request, body, files):
    drafts.append({"owner": owner, "atts": [f.attachment_id for f in files]})
    return 31


ps.fetch_request_text, ps.save_result = fake_fetch, fake_save
ps._confirmed_examples, ps._create_record = fake_examples, fake_record


async def fake_concept(body, *, user_id=None):
    concept_users.append(user_id)
    return PNG, "컨셉 이미지 생성됨(테스트)"

concept_users: list = []
ps.concept_image.generate = fake_concept
reset()
res = asyncio.run(ps.build_proposal_body(attachment_ids=[42], on_stage=stages.append))
check(stages == ["fetch", "extract", "cases", "write", "review", "save", "image", "files"], "진행 단계 순서", str(stages))
check("컨셉 이미지 생성됨(테스트)" in ps.result_message(res), "요약에 이미지 결과 한 줄")
check(concept_users == [7], "이미지 전송 기록에 요청자 전달", str(concept_users))
check(len(saved) == 2 and all(o == 7 for o, _, _ in saved), "요청자 소유 pptx 2개 저장", str(saved))
check(saved[0][1].startswith("제안서_초안_대광_") and saved[0][1].endswith("_디자인.pptx")
      and saved[1][1].endswith("_심플.pptx"), "디자인판 먼저, 심플판 다음", str(saved))
check(drafts == [{"owner": 7, "atts": [9, 9]}], "확정 대기 기록(두 파일 첨부)", str(drafts))
msg = ps.result_message(res)
check("요청 항목" in msg and "본문 반영" in msg, "요약에 반영률")
check("관련이 약해 제외한 사례: 2024-19" in msg, "요약에 제외 사례", msg[:600])
check("[이 결과를 확정](/api/proposal-records/31/confirm)" in msg, "확정 카드 링크")
check("확정 제안서를 구성·표기 참고 예시로 반영" in msg, "확정 예시 사용 안내")
check(calls.count("draft") == 1, "초안 1회")


async def collect():
    return [c async for c in main._stream_proposal_body(attachment_ids=[42])]


reset()
chunks = asyncio.run(collect())
check(chunks[0].startswith(main.PROGRESS_SENTINEL) and "관련성 판정" in chunks[0], "스트림 첫 마커")
check("제안서_초안_" in chunks[-1] and "_디자인.pptx" in chunks[-1], "스트림 마지막 = 결과(두 파일 링크)")


async def boom(ids, *, what):
    raise ValueError("제안서 요청 내용이 비어 있습니다.")

ps.fetch_request_text = boom
chunks = asyncio.run(collect())
check(chunks[-1] == "⚠ 제안서 본문 작성 실패: 제안서 요청 내용이 비어 있습니다.", "입력 오류 문구", chunks[-1])

print(f"PASS {len(PASS)} / FAIL {len(FAIL)}")
for label, detail in FAIL:
    print(f"  ✗ {label}: {detail}")
sys.exit(1 if FAIL else 0)
