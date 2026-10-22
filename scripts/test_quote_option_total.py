# -*- coding: utf-8 -*-
"""< OPTION > 묶음 줄마다 Total 반영/미반영(사용자 2026-10-19) 점검.

    MSYS_NO_PATHCONV=1 docker compose run --rm -T --no-deps -e PYTHONPATH=/app -e PYTHONIOENCODING=utf-8 \\
        -v ./backend/app:/app/app -v ./scripts:/scripts -v ./docs:/docs backend python /scripts/test_quote_option_total.py

반영 = 옵션 Sub total·Total(·NEGO % 기준·영업 건 금액)에 들어감, 미반영 = 공급가액만 보이고 Sub total·Total 제외.
화면 수정 → 상태 → 견적서 엑셀 수식 → 고친 파일 되읽기 → 올린 파일 위에서 이어 고치기 → 영업 건 품목·시리얼 대상까지 본다.
"""
from __future__ import annotations

import copy
import io
import re
import sys
from datetime import date

import openpyxl

from app.company_knowledge import product_prices as pp
from app.company_knowledge import quote_patch as qp
from app.company_knowledge import quote_readback as rb
from app.company_knowledge import quote_session as qs
from app.company_knowledge import quote_xlsx as qx
from app.sales_deals import service as svc
from app.sales_deals import units

PASS: list[str] = []
FAIL: list[tuple[str, str]] = []
DAY = date(2026, 10, 19)


def check(ok: bool, label: str, detail: str = "") -> None:
    (PASS.append(label) if ok else FAIL.append((label, detail)))


def parts(state: dict, no: str = "QT-TS20261019-01") -> tuple[list, dict]:
    f = state["form"]
    lines = qs.render_lines(state, "ko", None)
    return lines, qx.quote_header(f, lang="ko", quote_no=no, day=DAY, pogo=False)


def build(state: dict) -> bytes:
    return qx.build_xlsx(*reversed(parts(state)))


def edit(state: dict, payload: dict) -> dict:
    new, _ = qs.apply_ops(copy.deepcopy(state), qs.edit_ops(state, payload), None)
    return new


def screen_lines(state: dict, flags: dict[str, bool]) -> list[dict]:
    """화면이 저장 때 보내는 줄 — in_total 은 flags 에 있으면 그 값."""
    return [{"id": ln["id"], "name": ln["name"], "commodity": qx.commodity_of(ln), "qty": ln["qty"], "unit_price": ln["unit_price"],
             "remark": ln.get("remark") or "", "block": ln.get("block") or qs.MAIN_BLOCK, "in_total": flags.get(ln["id"], bool(ln.get("in_total")))}
            for ln in state["lines"] if ln["kind"] not in ("work", "freight")]


def refs(formula: str) -> set[int]:
    return rb._sum_rows(formula) or set()


def row_of(ws, text: str) -> int:
    return next(r for r in range(1, ws.max_row + 1) if text in str(ws.cell(r, 2).value or ""))


# ── 상태: 기본 품목 2줄 + 옵션 2줄(충전 스테이션 1,200,000 · 케이블 2 × 100,000) ──
S = qs.draft_state({"applicable": True, "items": []}, model="", payload_kg=None, form={
    "customer": "[테스트] 옵션고객", "to": "김철수 대표님", "contact_name": "홍길동", "contact_title": "이사", "initials": "TS",
    "contact_mobile": "010-0000-0000", "contact_email": "name@example.com"})
S.update(manual=True, asked=["extra"], delivery={"case": 1, "freight": None}, next_id=1)
S["lines"] = [{"id": "L1", "kind": "extra", "key": "extra", "name": "TCV1 MASTER T.C", "commodity": "Magbot ATC", "qty": 1,
               "unit_price": 2_300_000, "remark": "", "price_status": "set"},
              {"id": "L2", "kind": "extra", "key": "extra", "name": "TCV1 Tool Plate", "commodity": "Magbot ATC", "qty": 2,
               "unit_price": 400_000, "remark": "", "price_status": "set"}]
S["next_id"] = 3
S, _ = qs.apply_ops(S, [{"op": "add", "name": "CS-01 충전 스테이션", "commodity": "Option", "qty": 1, "unit_price": 1_200_000, "option": True},
                        {"op": "add", "name": "연장 케이블 5m", "commodity": "Option", "qty": 2, "unit_price": 100_000, "option": True}], None)
opt_ids = [ln["id"] for ln in S["lines"] if ln.get("block") and ln["block"] != qs.MAIN_BLOCK]
BASE = 2_300_000 + 2 * 400_000
check(len(opt_ids) == 2 and qs.subtotal(S) == BASE and not any(ln.get("in_total") for ln in S["lines"]),
      "옵션으로 넣으면 기본은 미반영 — Total = 기본 품목만", str((opt_ids, qs.subtotal(S))))

# 1. 화면에서 한 줄만 [반영]
S1 = edit(S, {"lines": screen_lines(S, {opt_ids[0]: True})})
check([bool(ln.get("in_total")) for ln in S1["lines"] if ln["id"] in opt_ids] == [True, False] and qs.subtotal(S1) == BASE + 1_200_000,
      "[반영] 한 줄 → Total 에 그 옵션 금액 포함", str(qs.subtotal(S1)))
lines1 = qs.render_lines(S1, "ko", None)
opt_rows = [r for r in lines1 if r.get("row") == "option"]
sub_rows = [r for r in lines1 if r.get("row") == "subtotal"]
check([(r["in_total"], r["amount"]) for r in opt_rows] == [(True, 1_200_000), (False, 200_000)],
      "옵션 줄 = 반영·미반영 모두 공급가액 있음", str([(r.get("in_total"), r.get("amount")) for r in opt_rows]))
check(sub_rows and sub_rows[-1]["amount"] == 1_200_000 and lines1.index(sub_rows[-1]) == lines1.index(opt_rows[-1]) + 1,
      "옵션 묶음 끝에 Sub total = 반영한 줄 합만", str(sub_rows))
check(qx.total(lines1) == BASE + 1_200_000, "발행 기록 합계(qx.total) = 반영 옵션 포함")

# 2. 엑셀 — 두 옵션 줄 모두 공급가액 수식, Total·옵션 Sub total 에는 반영 줄만
ws = openpyxl.load_workbook(io.BytesIO(build(S1))).active
r_in, r_out = row_of(ws, "CS-01"), row_of(ws, "연장 케이블")
total_r = next(r for r in range(17, ws.max_row + 1) if str(ws.cell(r, 1).value or "").startswith("Total"))
sub_r = next(r for r in range(r_out, total_r) if str(ws.cell(r, 1).value or "") == "Sub total")
check(ws.cell(r_in, 6).value == f"=D{r_in}*E{r_in}" and ws.cell(r_out, 6).value == f"=D{r_out}*E{r_out}",
      "엑셀 — 반영·미반영 옵션 모두 공급가액(수량×단가) 표시", str((ws.cell(r_in, 6).value, ws.cell(r_out, 6).value)))
check(r_in in refs(ws.cell(total_r, 7).value) and r_out not in refs(ws.cell(total_r, 7).value),
      "엑셀 Total = 반영 옵션 포함, 미반영 옵션 제외", str(ws.cell(total_r, 7).value))
check(refs(ws.cell(sub_r, 6).value) == {r_in}, "엑셀 옵션 Sub total = 반영 줄만", str(ws.cell(sub_r, 6).value))

# 3. 하나도 반영 안 하면 옵션 Sub total 없음, 금액은 보임
lines0 = qs.render_lines(S, "ko", None)
ws0 = openpyxl.load_workbook(io.BytesIO(build(S))).active
r0 = row_of(ws0, "CS-01")
check(not any(r.get("row") == "subtotal" for r in lines0) and ws0.cell(r0, 6).value == f"=D{r0}*E{r0}",
      "전부 미반영 → 옵션 Sub total 없음, 공급가액은 보임")

# 4. NEGO % 기준 = 반영 옵션 포함 합계
SN = copy.deepcopy(S1)
SN["nego"] = {"mode": "percent", "value": 10, "label": ""}
check(qs.subtotal(SN) == round((BASE + 1_200_000) * 0.9), "NEGO % 기준 = 반영 옵션 포함 합계", str(qs.subtotal(SN)))

# 5. 예전 발행본(in_total 없는 옵션 줄)은 예전처럼 공급가액 칸 비움·합계 제외
legacy = [{**r, "in_total": None} if r.get("row") == "option" else r for r in lines1 if r.get("row") != "subtotal"]
legacy = [{k: v for k, v in r.items() if not (k == "in_total" and v is None)} for r in legacy]
wl = openpyxl.load_workbook(io.BytesIO(qx.build_xlsx(parts(S1)[1], legacy))).active
rl = row_of(wl, "CS-01")
check(wl.cell(rl, 6).value is None and qx.total(legacy) == BASE, "예전 옵션 줄 = 공급가액 칸 비움·합계 제외 그대로")

# 6. 되읽기 — 그대로면 바뀐 점 없음, 엑셀에서 반영을 바꾼 파일은 그 값으로
M, ch, wn = rb.merge(S1, rb.parse(build(S1)))
check(ch == [] and wn == [] and [bool(x.get("in_total")) for x in M["lines"]] == [bool(x.get("in_total")) for x in S1["lines"]],
      "고치지 않은 파일 되읽기 → 바뀐 점 없음(반영 상태 유지)", f"{ch} {wn}")
T = edit(S1, {"lines": screen_lines(S1, {opt_ids[0]: False, opt_ids[1]: True})})
M2, ch2, _ = rb.merge(S1, rb.parse(build(T)))
check([bool(x.get("in_total")) for x in M2["lines"] if x["id"] in opt_ids] == [False, True] and qs.subtotal(M2) == BASE + 200_000
      and any("Total 미반영" in c for c in ch2) and any("Total 반영" in c for c in ch2),
      "반영을 바꾼 견적서를 되읽으면 그 반영 상태로", str(ch2))

# 7. 올린 파일 위에서 이어 고치기 — 반영을 바꾸면 수식·Total 이 따라감
bl, bh = parts(S1)
nl, nh = parts(T, no="QT-TS20261019-02")
bh["date"] = ""
out, _ = qp.patch(build(S1), bl, bh, nl, nh)
wp = openpyxl.load_workbook(io.BytesIO(out)).active
pi, po = row_of(wp, "CS-01"), row_of(wp, "연장 케이블")
tp = next(r for r in range(17, wp.max_row + 1) if str(wp.cell(r, 1).value or "").startswith("Total"))
check(po in refs(wp.cell(tp, 7).value) and pi not in refs(wp.cell(tp, 7).value) and wp.cell(pi, 6).value == f"=D{pi}*E{pi}",
      "이어 고치기 — 바꾼 반영 상태대로 Total 수식, 미반영도 공급가액 유지", str((wp.cell(tp, 7).value, wp.cell(pi, 6).value)))
secs, _ = rb._sections(rb.parse(out)["rows"])
opt_sec = next(x for x in secs if x["type"] == "option")
check([r.get("in_total") for r in opt_sec["rows"]] == [False, True], "이어 고친 파일도 반영 상태로 되읽힘",
      str([(r["row"], r.get("in_total")) for r in opt_sec["rows"]]))

# 8. 영업 건 품목·시리얼 대상 — 반영 옵션은 들어감, 미반영은 빠짐
items = svc.snapshot_items(lines1)
check(sum(it["unit_price"] * it["qty"] for it in items) == BASE + 1_200_000 and not any("연장 케이블" in it["name"] for it in items)
      and next(it for it in items if "CS-01" in it["name"])["statement"] is False,
      "영업 건 품목 = 반영 옵션 포함(거래명세서 줄은 아님), 미반영 제외", str(items))
names = [t["name"] for t in units.target_lines(lines1)]
check(any("CS-01" in n for n in names) and not any("연장 케이블" in n for n in names), "시리얼 대상 = 반영 옵션 포함", str(names))

# 9. 대화로 반영
S9, d9 = qs.apply_ops(copy.deepcopy(S), [{"op": "update", "line": opt_ids[1], "in_total": True}], "케이블 옵션도 합계에 넣어 줘")
check(qs.subtotal(S9) == BASE + 200_000 and "Total 반영" in d9[-1], "대화 — 옵션 줄 Total 반영", str(d9))

for label, detail in FAIL:
    print(f"FAIL {label} :: {detail}")
print(f"PASS {len(PASS)} / FAIL {len(FAIL)}")
sys.exit(1 if FAIL else 0)
