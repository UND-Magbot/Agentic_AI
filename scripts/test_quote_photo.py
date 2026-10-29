"""올린 견적서의 Remark 칸 사진 잇기 점검 — 사용자 2026-10-29.

    MSYS_NO_PATHCONV=1 docker compose run --rm --no-deps -e PYTHONPATH=/app -v ./backend/app:/app/app -v ./scripts:/scripts \\
        backend python /scripts/test_quote_photo.py

① 읽기: 회사 양식 파일의 Remark(G) 칸 사진 → 줄마다 지문(quote_readback.parse)
② 맞추기: 엑셀에서 사진을 그대로 둠·뺌·새로 넣음·바꿈 → 작성 상태 품목 사진(merge + product_recommend._take_photos)
③ 이어 고치기: 앱에서 바꾼 사진·새 줄 사진은 올린 파일 Remark 칸에 들어가고, 지운 줄의 사진은 빠짐(quote_patch)
저장소는 흉내(메모리) — 실제 MinIO 에 쓰지 않는다.
"""
from __future__ import annotations

import copy
import io
import sys
from datetime import date

import openpyxl
from openpyxl.drawing.image import Image as XLImage
from PIL import Image

from app import storage
from app.company_knowledge import product_recommend as pr
from app.company_knowledge import quote_patch as qp
from app.company_knowledge import quote_readback as rb
from app.company_knowledge import quote_session as qs
from app.company_knowledge import quote_xlsx as qx

PASS: list[str] = []
FAIL: list[tuple[str, str]] = []
DAY = date(2026, 10, 29)


def check(ok: bool, label: str, detail: str = "") -> None:
    (PASS.append(label) if ok else FAIL.append((label, detail)))


def pic(color: str, fmt: str = "JPEG") -> bytes:
    b = io.BytesIO()
    Image.new("RGB", (120, 90), color).save(b, fmt)
    return b.getvalue()


# 저장소 흉내 — 키 → 바이트
STORE: dict[str, bytes] = {}


def fake_put(*, key: str, data, length: int, mime: str) -> None:
    STORE[key] = data.read()


storage.put_object = fake_put
pr.line_photos = lambda lines: {k: STORE[k] for ln in lines if (k := str(ln.get("photo") or "")) in STORE}


def parts(state: dict, no: str = "QT-CH20261029-01") -> tuple[list, dict]:
    f = state["form"]
    lines = qs.render_lines(state, "ko", None)
    return lines, qx.quote_header(f, lang="ko", quote_no=no, day=DAY, pogo=False)


def g_pics(data: bytes) -> dict[str, str]:
    """파일의 품목 줄(규격 첫 줄) → Remark 칸 사진 지문."""
    p = rb.parse(data)
    return {r["desc"].split("\n")[0]: r["pic"] for r in p["rows"] if r.get("pic")}


def row_of(ws, text: str) -> int:
    return next(r for r in range(1, ws.max_row + 1) if str(ws.cell(r, 2).value or "").startswith(text))


def drop(ws, r: int) -> None:
    ws._images = [i for i in ws._images if not (i.anchor._from.col == 6 and i.anchor._from.row + 1 == r)]


# ── ① 앱이 만든 견적서(L1·L2 사진 있음) ──
STORE["product-quotes/1/photos/a.jpg"] = A = pic("red")
STORE["product-quotes/1/photos/b.jpg"] = B = pic("blue")
S = qs.draft_state({}, model="", payload_kg=None, manual=True, form={
    "customer": "가나테크", "to": "김철수 대표님", "contact_name": "홍길동", "contact_title": "이사", "initials": "CH",
    "contact_mobile": "010-0000-0000", "contact_email": "name@example.com"})
S["lines"] = [{"id": "L1", "kind": "extra", "key": "extra", "name": "MG10", "commodity": "Magbot Gripper", "qty": 2,
               "unit_price": 1250000, "remark": "", "photo": "product-quotes/1/photos/a.jpg"},
              {"id": "L2", "kind": "extra", "key": "extra", "name": "TCV1", "commodity": "Magbot ATC", "qty": 1,
               "unit_price": 2300000, "remark": "컨트롤박스 포함", "photo": "product-quotes/1/photos/b.jpg"},
              {"id": "L3", "kind": "extra", "key": "extra", "name": "Cable 0.3m", "commodity": "Accessory", "qty": 4,
               "unit_price": 50000, "remark": ""}]
S["next_id"] = 4
S["delivery"] = {"case": 1, "freight": None}
l0, h0 = parts(S)
X0 = qx.build_xlsx(h0, l0, pr.line_photos(l0))
p0 = rb.parse(X0)
check(g_pics(X0) == {"MG10": rb.photo_sig(A), "TCV1": rb.photo_sig(B)}, "읽기: 줄마다 Remark 사진 지문", str(g_pics(X0)))
check(set(p0["pics"]) == {rb.photo_sig(A), rb.photo_sig(B)}, "읽기: 지문 → 사진 바이트")

# ── ② 엑셀에서: L1 그대로, L2 사진 뺌, L3 사진 새로(PNG) ──
C = pic("green", "PNG")
wb = openpyxl.load_workbook(io.BytesIO(X0))
ws = wb.active
drop(ws, row_of(ws, "TCV1"))
im = XLImage(io.BytesIO(C))
im.width, im.height = 60, 45
ws.add_image(im, f"G{row_of(ws, 'Cable')}")
b = io.BytesIO()
wb.save(b)
U = b.getvalue()
pu = rb.parse(U)
_, ch_preview, _ = rb.merge(S, pu)
check("TCV1 사진 삭제" in ch_preview and "Cable 0.3m 사진 추가" in ch_preview and not any("MG10" in x for x in ch_preview),
      "미리보기 바뀐 점: 뺀 사진·새 사진만", str(ch_preview))
n_store = len(STORE)
st1, ch1, wn1 = pr._synced_state(S, qid=1, key="product-quotes/1/up.xlsx", ext=".xlsx", items=[], parsed=pu)
by = {ln["id"]: ln for ln in st1["lines"]}
check(by["L1"]["photo"] == "product-quotes/1/photos/a.jpg" and by["L1"].get("photo_sig") == rb.photo_sig(A),
      "그대로 둔 사진: 다시 저장하지 않고 원래 사진 키", str(by["L1"]))
check("photo" not in by["L2"] and "photo_sig" not in by["L2"], "엑셀에서 뺀 사진: 작성 상태에서도 뺌", str(by["L2"]))
k3 = by["L3"].get("photo") or ""
check(k3.startswith("product-quotes/1/photos/") and len(STORE) == n_store + 1, "새로 넣은 사진: 1장만 저장", k3)
im3 = Image.open(io.BytesIO(STORE.get(k3, b"")))
check(im3.format == "JPEG" and max(im3.size) <= pr.PHOTO_MAX_PX, "저장한 사진은 앱 사진처럼 줄인 JPEG")
check(by["L3"].get("photo_sig") == rb.photo_sig(C), "가져온 사진 지문 = 엑셀 원본 지문")
check(not any("_pic" in ln for ln in st1["lines"]) and wn1 == [], "임시 표시(_pic) 남지 않음·경고 없음", str(wn1))
check(st1.get("base_file", {}).get("state", {}).get("lines") == st1["lines"], "바탕 파일 상태에도 사진 반영")

# 같은 파일을 다시 올림 → 바뀐 점·새 저장 없음
n_store = len(STORE)
st1b, ch1b, _ = pr._synced_state(st1, qid=1, key="product-quotes/1/up.xlsx", ext=".xlsx", items=[], parsed=pu)
check(not any("사진" in x for x in ch1b) and len(STORE) == n_store and st1b["lines"] == st1["lines"],
      "같은 파일 다시 올림: 사진 변화 없음", str(ch1b))

# 엑셀에서 L1 사진을 다른 사진으로 바꿈 → 사진 변경
D = pic("yellow")
wb = openpyxl.load_workbook(io.BytesIO(U))
ws = wb.active
drop(ws, row_of(ws, "MG10"))
im = XLImage(io.BytesIO(D))
ws.add_image(im, f"G{row_of(ws, 'MG10')}")
b = io.BytesIO()
wb.save(b)
st2, ch2, _ = pr._synced_state(st1, qid=1, key="product-quotes/1/up2.xlsx", ext=".xlsx", items=[], parsed=rb.parse(b.getvalue()))
l1 = next(ln for ln in st2["lines"] if ln["id"] == "L1")
check("MG10 사진 변경" in ch2 and l1["photo"] != by["L1"]["photo"] and l1["photo_sig"] == rb.photo_sig(D),
      "엑셀에서 바꾼 사진: 새 사진으로", str(ch2))

# 사진을 바꾼 뒤 앱에서 다시 사진을 올리면 지문은 지운다(quote_line_photo) — 아래 ③에서 같은 상황을 직접 만든다

# ── ③ 올린 파일(U) 위에서 이어 고치기 ──
base = st1["base_file"]["state"]
E, F = pic("purple"), pic("orange")
STORE["product-quotes/1/photos/e.jpg"] = E
STORE["product-quotes/1/photos/f.jpg"] = F
N = copy.deepcopy(base)
for ln in N["lines"]:
    if ln["id"] == "L1":                      # 앱에서 사진 바꿈
        ln["photo"] = "product-quotes/1/photos/e.jpg"
        ln.pop("photo_sig", None)
N["lines"] = [ln for ln in N["lines"] if ln["id"] != "L3"]          # 사진 있던 줄 지움
N["lines"].append({"id": "L4", "kind": "extra", "key": "extra", "name": "MTC", "commodity": "Magbot Tool", "qty": 1,
                   "unit_price": 900000, "remark": "새 줄", "photo": "product-quotes/1/photos/f.jpg"})
N["next_id"] = 5
bl, bh = parts(base)
bh["date"] = ""
nl, nh = parts(N, no="QT-CH20261029-02")
out, notes = qp.patch(U, bl, bh, nl, nh, pr.line_photos(nl))
got = g_pics(out)
check(got == {"MG10": rb.photo_sig(E), "MTC": rb.photo_sig(F)}, "올린 파일에 바꾼 사진·새 줄 사진, 지운 줄 사진은 빠짐", str(got))
wo = openpyxl.load_workbook(io.BytesIO(out)).active
po = rb.parse(out)
in_table = set(range(po["pos"]["head"] + 1, po["pos"]["total"]))
check(sum(1 for i in wo._images if i.anchor._from.col == 6 and i.anchor._from.row + 1 in in_table) == 2,
      "표 안 Remark 칸 그림 2장(남는 그림 없음)", str([(i.anchor._from.row + 1) for i in wo._images]))
check(not any("사진" in x for x in notes), "사진 알림 없음", str(notes))
r_new = row_of(wo, "MTC")
check((wo.row_dimensions[r_new].height or 0) >= 48, "새 줄 행 높이 = 사진 들어갈 만큼",
      str(wo.row_dimensions[r_new].height))
check(wo.cell(r_new, 7).alignment.horizontal == "left", "글 있는 Remark 칸은 왼쪽 정렬(사진 자리 비움)")
check(sum(1 for i in wo._images if i.anchor._from.col != 6) == sum(
    1 for i in openpyxl.load_workbook(io.BytesIO(U)).active._images if i.anchor._from.col != 6), "Remark 밖 그림(도장 등) 그대로")

# 사진을 안 바꾼 고치기 → 그림은 원본 그대로
N2 = copy.deepcopy(base)
next(ln for ln in N2["lines"] if ln["id"] == "L2")["qty"] = 3
nl2, nh2 = parts(N2, no="QT-CH20261029-03")
out2, notes2 = qp.patch(U, bl, bh, nl2, nh2, pr.line_photos(nl2))
check(g_pics(out2) == g_pics(U), "사진 안 바꾼 고치기: 원래 사진 그대로", str(g_pics(out2)))

# 저장소에 사진이 없으면 알림만, 고치기는 계속
N3 = copy.deepcopy(N)
N3["lines"][-1]["photo"] = "product-quotes/1/photos/none.jpg"
nl3, nh3 = parts(N3, no="QT-CH20261029-04")
out3, notes3 = qp.patch(U, bl, bh, nl3, nh3, pr.line_photos(nl3))
check(any("사진을 불러오지 못해" in x for x in notes3) and "MTC" not in g_pics(out3), "없는 사진: 알림 후 사진 없이", str(notes3))

# 발행 파일을 다시 올림 → 작성 상태 그대로(앱 사진은 저장소 바이트로 같다고 봄)
n_store = len(STORE)
st4, ch4, _ = pr._synced_state(N, qid=1, key="product-quotes/1/issued.xlsx", ext=".xlsx", items=[], parsed=rb.parse(out))
check(not any("사진" in x for x in ch4) and len(STORE) == n_store
      and [ln.get("photo") for ln in st4["lines"]] == [ln.get("photo") for ln in N["lines"]], "발행본 다시 올림: 사진 그대로", str(ch4))

print(f"PASS {len(PASS)} / FAIL {len(FAIL)}")
for label, detail in FAIL:
    print("  FAIL", label, "|", detail[:300])
sys.exit(1 if FAIL else 0)
