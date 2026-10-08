"""제안서 작업 1·2단계(자료 읽기·핵심 질문) 검증.

- 기본: 문서 읽기·근거 검증·질문 순서 — 모델은 가짜(chat 주입), DB 불필요.
- --db : 실제 DB·저장소로 run_reading → 답변 → 질문 마치기 (backend 컨테이너에서 실행)
    docker compose run --rm --no-deps -v ./scripts:/scripts backend python /scripts/test_proposal_dialog.py --db
"""
import asyncio
import io
import json
import sys
from pathlib import Path

sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8", errors="replace")
ROOT = Path(__file__).resolve().parents[1]
for cand in (ROOT / "backend", Path("/app")):
    if (cand / "app").is_dir():
        sys.path.insert(0, str(cand))
        break

from app.proposal_project import dialog, extract as ex  # noqa: E402
from app.proposal_project.catalog import CODES  # noqa: E402

PASS: list[str] = []
FAIL: list[tuple[str, str]] = []


def check(ok: bool, label: str, detail: str = "") -> None:
    (PASS.append(label) if ok else FAIL.append((label, detail)))


REQUEST = "고객사 A 약액 보틀 주입 자동화를 실증하려 합니다. 시간당 17병을 처리해야 하며 약액 주입은 210초입니다."
ATTACH = "대상물: 500ml 유리 보틀 2종\n로봇은 협동로봇 1대를 검토 중입니다."
PROJECT = {"intake": {"project": "고객사 A, 실증 목적으로 예산 확보 판단", "volume": "시간당 17병"},
           "request_text": REQUEST}


def fake_chat_factory(answers: dict[str, dict]):
    """문항 묶음마다 해당 문항 답만 돌려주는 가짜 모델."""
    async def chat(messages, **_kw):
        user = messages[-1]["content"]
        asked = [c for c in CODES if f"- {c}:" in user]
        return json.dumps({"items": [dict(code=c, **answers[c]) for c in asked if c in answers]},
                          ensure_ascii=False)
    return chat


ANSWERS = {
    "P02": {"value": "실증 목적 예산 확보 판단", "status": "confirmed", "evidence": "실증 목적으로 예산 확보 판단"},
    "P08": {"value": "시간당 17병", "unit": "병/시간", "status": "confirmed", "evidence": "시간당 17병을 처리해야 하며"},
    "P06": {"value": "500ml 유리 보틀 2종", "status": "confirmed", "evidence": "대상물: 500ml 유리 보틀 2종"},
    "P13": {"value": "협동로봇 1대", "status": "assumed", "evidence": "로봇은 협동로봇 1대를 검토 중입니다."},
    # 날조: 근거 문장이 원문에 없다.
    "P33": {"value": "설치 공간 3m x 2m", "status": "confirmed", "evidence": "설치 공간은 3m x 2m 입니다"},
    # 날조: 근거는 있지만 값의 숫자(25)가 원문에 없다.
    "P41": {"value": "25병 실증부터", "status": "confirmed", "evidence": "시간당 17병을 처리해야 하며"},
}

# ── 1) 문서 읽기 ────────────────────────────────────────────────────────────
import docx  # noqa: E402
import pymupdf  # noqa: E402

buf = io.BytesIO()
d = docx.Document()
d.add_paragraph("설치 폭 1200 mm")
t = d.add_table(rows=1, cols=2)
t.cell(0, 0).text, t.cell(0, 1).text = "무게", "14 kg"
d.save(buf)
got = ex.read_document(buf.getvalue(), "spec.docx", "")
check("설치 폭 1200 mm" in got and "무게 | 14 kg" in got, "docx 문단·표 읽기", got)
pdf = pymupdf.open()
pdf.new_page().insert_text((72, 72), "Throughput 17 bottles per hour")
got = ex.read_document(pdf.tobytes(), "a.pdf", "application/pdf")
check("17 bottles" in (got or ""), "pdf 글자 읽기", str(got))
check(ex.read_document("시간당 17병".encode("cp949"), "req.txt", "text/plain") == "시간당 17병", "txt(cp949) 읽기")
check(ex.read_document(b"\x89PNG\r\n", "x.png", "image/png") is None, "이미지는 글자 근거 아님(None)")
check(ex.read_document(b"not a zip", "broken.docx", "") is None, "깨진 docx 는 None(예외 없음)")

# ── 2) 근거 검증 ────────────────────────────────────────────────────────────
srcs = ex.build_sources(PROJECT, {7: ("사양.txt", ATTACH)})
check([s.kind for s in srcs] == ["form", "form", "request", "attachment"], "출처 순서: 폼 → 요청 → 첨부")
codes = set(CODES)
it, _ = ex.verify({"code": "p08", **ANSWERS["P08"]}, srcs, codes)
check(it is not None and it.source == "request" and it.evidence.startswith("[고객 요청 원문]"),
      "근거 문장이 원문에 있으면 통과·출처 기록", str(it))
it, why = ex.verify({"code": "P33", **ANSWERS["P33"]}, srcs, codes)
check(it is None and "원문에 없음" in why, "근거 문장이 원문에 없으면 버림", why)
it, why = ex.verify({"code": "P41", **ANSWERS["P41"]}, srcs, codes)
check(it is None and "숫자" in why, "값의 숫자가 원문에 없으면 버림", why)
it, _ = ex.verify({"code": "P06", "value": "보틀", "status": "확정", "evidence": " 대상물:  500ml 유리 보틀 2종 "},
                  srcs, codes)
check(it is not None and it.status == "assumed" and it.source == "attachment",
      "공백 차이는 허용, 모르는 상태는 '가정' 으로 낮춤", str(it))
loc_src = [ex.Source("request", "요청", "시간당 17병 수준이며, 주입 1회에 약 210초가 걸려 병목입니다."),
           ex.Source("attachment", "첨부", "운영: 2교대 16시간")]
check(getattr(ex._locate("시간당 17병, 주입 1회 약 210초, 2교대 16시간", loc_src), "label", None) == "요청",
      "여러 곳을 이어 붙인 근거: 조각마다 원문에 있으면 통과(조사 차이 허용)")
check(getattr(ex._locate("시간당 17병 수준이며, 주입 1회에 약 210초가 걸려 병목입니다. 운영: 2교대 16시간", loc_src),
              "label", None) == "요청", "두 출처 문장을 이어 붙인 근거도 문장마다 찾으면 통과")
check(ex._locate("시간당 20병, 주입 1회 약 210초", loc_src) is None, "조각 하나라도 숫자가 다르면 버림")
check(ex._locate("클린룸 등급 Class 10", loc_src) is None, "원문에 없는 조각은 버림")
check(ex.verify({"code": "P99", "value": "x", "evidence": "시간당"}, srcs, codes)[0] is None, "모르는 문항 버림")
check(ex.verify({"code": "P08", "value": "", "evidence": "시간당"}, srcs, codes)[0] is None, "빈 값 버림")

res = asyncio.run(ex.extract(PROJECT, {7: ("사양.txt", ATTACH)}, chat=fake_chat_factory(ANSWERS)))
check(sorted(i.code for i in res.items) == ["P02", "P06", "P08", "P13"], "추출: 근거 있는 값만 남김",
      str([i.code for i in res.items]))
check(sorted(d["code"] for d in res.dropped) == ["P33", "P41"], "추출: 날조 2건 버림", str(res.dropped))
check({i.code: i.status for i in res.items}["P13"] == "assumed", "'검토 중' 은 가정으로 유지")
calls: list[int] = []


async def counting_chat(messages, **kw):
    calls.append(len(messages[-1]["content"]))
    return "{\"items\": []}"

asyncio.run(ex.extract(PROJECT, {}, chat=counting_chat))
check(len(calls) == 12, "폼 묶음마다 1회씩(12회) 나눠 물음", str(calls))
check(asyncio.run(ex.extract({"intake": {}, "request_text": ""}, {}, chat=counting_chat)).items == []
      and len(calls) == 12, "원문이 없으면 모델을 부르지 않음")
check(ex._parse("앞말 {\"items\": [{\"code\": \"P01\"}]} 뒷말") == [{"code": "P01"}], "JSON 앞뒤 군말 허용")
check(ex._parse("엉망") == [], "JSON 아니면 빈 목록")

# ── 3) 질문 순서 ────────────────────────────────────────────────────────────
items = {c: {"status": "empty", "value": "", "evidence": "", "source": ""} for c in CODES}
items["P02"]["status"] = "confirmed"
items["P06"]["status"] = "unknown"      # 필수인데 모름 → 다시 묻는다
items["P09"]["status"] = "unknown"      # 선택인데 모름 → 묻지 않는다
items["P31"]["status"] = "conflict"
qs = [q["code"] for q in dialog.next_questions({"items": items})]
check(qs == ["P01", "P03", "P12", "P06", "P13"],
      "필수 4개(모름 필수 P06 재질문) + 첫 화면 빈 문항 1개 = 5개, 표 3 순서(모름 선택 P09·답한 P02 제외)", str(qs))
allq = [q["code"] for q in dialog.next_questions({"items": items}, limit=100)]
check(len(allq) == 5 and not {"P17", "P19", "P21", "P22", "P23", "P24", "P10", "P44", "P15", "P16", "P31"} & set(allq),
      "한도를 늘려 불러도 5개, 설계·기타 문항은 묻지 않음(AI 가 제안)", str(allq))
items2 = {c: dict(v) for c, v in items.items()}
items2["P08"]["status"] = "conflict"
g2 = [q["code"] for q in dialog.next_questions({"items": items2}, limit=100)]
check("P08" in g2 and "P01" not in g2, "남은 한 자리는 자료 충돌 문항을 먼저", str(g2))
check(dialog.auto_title({"P01": "고객사 A 보틀 라인"}, "안녕하세요.") == "고객사 A 보틀 라인"
      and dialog.auto_title({"P06": "500ml 유리 보틀 2종"}, "안녕하세요.") == "500ml 유리 보틀 2종 자동화 제안"
      and dialog.auto_title({}, "안녕하세요. 반갑습니다\n홍길동입니다.\n보틀 주입 자동화 문의드립니다.")
      == "보틀 주입 자동화 문의드립니다.",
      "자동 제목: 고객·프로젝트 → 대상물 → 인사말이 아닌 첫 줄")
check(dialog.next_questions({"items": items})[0]["group"] == "고객과 목적 및 이번에 결정할 내용",
      "질문에 표 3 묶음 이름", "")
check(dialog.required_open({"items": items}) == ["P03", "P06", "P12", "P13"], "남은 필수 문항")
items3 = {c: dict(v) for c, v in items.items()}
for c in ("P01", "P03", "P12", "P06", "P07", "P08", "P13"):
    items3[c]["status"] = "confirmed"
q3 = [q["code"] for q in dialog.next_questions({"items": items3}, limit=100)]
check(q3 == ["P29", "P30", "P33", "P35", "P43"],
      "필수가 다 차면 첫 화면 빈 문항 5개(표 3 순서), 그 밖(P25 등)은 한도 밖", str(q3))
check("P14" not in q3, "지정·검토 로봇(P13)을 답했으면 P14 는 묻지 않음", str(q3))
items4 = {c: dict(v) for c, v in items3.items()}
for c in ("P29", "P30", "P33"):
    items4[c].update(status="confirmed", source="dialog")
q4 = [q["code"] for q in dialog.next_questions({"items": items4})]
check(q4 == ["P35", "P43"], "이미 물은 3개만큼 한도에서 뺌(남은 2개)", str(q4))
for c in ("P35", "P43"):
    items4[c].update(status="unknown", source="dialog")          # 모름·건너뜀도 물은 것으로 센다
check(dialog.next_questions({"items": items4}) == [], "합쳐서 5개 물었으면 더 묻지 않음(P25 등 비어 있어도)")
items4["P13"].update(status="empty", source="")
check([q["code"] for q in dialog.next_questions({"items": items4})] == ["P13"],
      "컨셉 필수 문항은 한도를 넘어도 채워질 때까지 묻는다")
items5 = {c: dict(v) for c, v in items3.items()}
for c in ("P29", "P30", "P33", "P35", "P43"):
    items5[c]["status"] = "confirmed"
q5 = dialog.next_questions({"items": items5})
check([q["code"] for q in q5] == ["P25", "P27", "P37", "P34", "P38"]
      and q5[-1]["group"] == dialog.QUESTIONS["P38"].group,
      "첫 화면이 다 차면 컨셉에 꼭 필요한 5문항(보관·만재·작업자·공간·연동 신호)", str([(q["code"], q["group"]) for q in q5]))


# ── 4) DB 통합 (--db) ───────────────────────────────────────────────────────
async def db_flow() -> None:
    from sqlalchemy import text

    from app.database import SessionLocal
    from app.migrations import run_migrations
    from app.proposal_project import service
    from app.storage import make_object_key, put_object

    async with SessionLocal() as db:
        await run_migrations(db)
        uid = (await db.execute(text("SELECT id FROM users WHERE username='alex'"))).scalar_one()
        atts = []
        for name, data, mime in (("사양.txt", ATTACH.encode("utf-8"), "text/plain"),
                                 ("scan.pdf", b"%PDF-1.4 broken", "application/pdf"),
                                 ("robot.png", b"\x89PNG\r\n", "image/png")):
            key = make_object_key(uid, name)
            put_object(key=key, data=io.BytesIO(data), length=len(data), mime=mime)
            atts.append(int((await db.execute(
                text("INSERT INTO attachments (user_id, bucket, object_key, original_filename, mime, size_bytes) "
                     "VALUES (:u, 'und-cortex-attachments', :k, :n, :m, :s) RETURNING id"),
                {"u": uid, "k": key, "n": name, "m": mime, "s": len(data)})).scalar_one()))
        await db.commit()
        pid = await service.create_project(
            db, user_id=uid, title="대화 시험", request_text=REQUEST, intake=PROJECT["intake"],
            assets=[service.AssetInput(atts[0], "content"), service.AssetInput(atts[1], "content"),
                    service.AssetInput(atts[2], "appearance", external_ok=True)])
    try:
        await dialog.run_reading(pid, chat=fake_chat_factory(ANSWERS))
        async with SessionLocal() as db:
            p = await service.get_project(db, pid)
            check(p["stage"] == "questioning", "자료 읽기 후 질문 단계", p["stage"])
            filled = {c: (v["status"], v["source"]) for c, v in p["items"].items() if v["status"] != "empty"}
            check(filled == {"P02": ("confirmed", "form"), "P06": ("confirmed", "attachment"),
                             "P08": ("confirmed", "request"), "P13": ("assumed", "attachment")},
                  "문항 4개 채움(출처·상태)", str(filled))
            check([a["readable"] for a in p["assets"]] == [True, False, None],
                  "첨부 읽기 여부: txt 성공 · 깨진 pdf 실패 · 외형 이미지는 대상 아님", str(p["assets"]))
            last = p["messages"][-1]["content"]
            check("4개 문항" in last and "2개는 넣지 않았습니다" in last and "1개는 글자를 읽지 못했습니다" in last,
                  "읽기 결과 안내 메시지", last)
            qs = [q["code"] for q in dialog.next_questions(p)]
            check(qs[:3] == ["P01", "P03", "P09"] and len(qs) == 5, "표 3 순서로 5개 질문", str(qs))

            try:
                await dialog.finish_questions(db, p)
                check(False, "필수 문항이 비면 질문 마치기 거부")
            except service.ProjectError:
                check(True, "필수 문항이 비면 질문 마치기 거부")

            n = await dialog.answer(db, pid, [
                {"code": "P03", "value": "보틀 투입부터 주입 완료 보틀 배출까지"},
                {"code": "P12", "value": ""},                          # 필수 빈칸 → 그대로
                {"code": "P01", "value": ""},                          # 선택 빈칸 → 미확인
                {"code": "P13", "value": "협동로봇 1대, 7축 이상", "unknown": False},
                {"code": "P04", "unknown": True},
            ])
            check(n == 4, "답 4건 기록(필수 빈칸 제외)", str(n))
            p = await service.get_project(db, pid)
            st = {c: (p["items"][c]["status"], p["items"][c]["source"]) for c in ("P03", "P12", "P01", "P13", "P04")}
            check(st == {"P03": ("confirmed", "dialog"), "P12": ("empty", ""), "P01": ("unknown", "dialog"),
                         "P13": ("confirmed", "dialog"), "P04": ("unknown", "dialog")},
                  "답·모름·건너뜀 상태", str(st))
            check(p["messages"][-1]["role"] == "user" and "P03 보틀 투입부터" in p["messages"][-1]["content"],
                  "답변이 진행 기록에 남음")
            try:
                await dialog.answer(db, pid, [{"code": "P77", "value": "x"}])
                check(False, "모르는 문항 답 거부")
            except service.ProjectError:
                await db.rollback()
                check(True, "모르는 문항 답 거부")

        # 다시 읽어도 사용자가 답한 값(P13)은 덮지 않는다.
        await dialog.run_reading(pid, chat=fake_chat_factory(ANSWERS))
        async with SessionLocal() as db:
            p = await service.get_project(db, pid)
            check(p["items"]["P13"]["value"] == "협동로봇 1대, 7축 이상", "다시 읽어도 사용자 답 유지",
                  p["items"]["P13"]["value"])
            await dialog.answer(db, pid, [{"code": "P12", "value": "주입 완료 보틀을 트레이에 적재해 반출"}])
            p = await service.get_project(db, pid)
            await dialog.finish_questions(db, p)
            p = await service.get_project(db, pid)
            check(p["stage"] == "concept", "필수 문항이 차면 컨셉 단계로", p["stage"])

        async def boom(*_a, **_k):
            raise RuntimeError("모델 서버 연결 끊김")
        await dialog.run_reading(pid, chat=boom)
        async with SessionLocal() as db:
            p = await service.get_project(db, pid)
            check(p["stage"] == "intake" and "자료 읽기에 실패" in p["messages"][-1]["content"],
                  "모델 오류 시 처음 단계로 되돌리고 사유 기록", p["stage"])
    finally:
        async with SessionLocal() as db:
            await db.execute(text("DELETE FROM proposal_projects WHERE id = :p"), {"p": pid})
            await db.execute(text("DELETE FROM attachments WHERE id = ANY(:ids)"), {"ids": atts})
            await db.commit()


if "--db" in sys.argv:
    asyncio.run(db_flow())

print(f"PASS {len(PASS)} / FAIL {len(FAIL)}")
for label, detail in FAIL:
    print("  FAIL", label, detail)
sys.exit(1 if FAIL else 0)
