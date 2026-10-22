# -*- coding: utf-8 -*-
"""대여 견적서(사용자 2026-10-22, 방역 로봇 3년 렌탈 견적서 Lu260309 기준) — DB 없이 점검.

    python scripts/test_rental_quote.py

- 수기 작성에서 '제품 대여'를 고르면 대여 문구·중도 해지 주석, 개월 수가 없으면 발행 막기
- 대여 개월 수·시작일 입력, 판매 ↔ 대여 바꾸기(기본 문구만 바뀜), 국내 ↔ 해외(주석 언어)
- 엑셀: Total 줄 금액 칸 = 총액 SUM, 비고 칸 = 월 납입금(총액 ÷ 개월, 천 원 올림), 판매 견적서는 예전 그대로
- 다시 읽기(올린 파일 이어 고치기): 대여 양식을 알아보고 Total 에 더하는 줄을 금액 칸 SUM 에서 읽음
- 영업 건: 기간 끝나는 날·월 납입금
"""
from __future__ import annotations

import datetime as dt
import io
import sys
from datetime import date
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "backend"))
sys.stdout.reconfigure(encoding="utf-8")

import openpyxl  # noqa: E402

from app.company_knowledge import quote_readback as rb  # noqa: E402
from app.company_knowledge import quote_session as qs  # noqa: E402
from app.company_knowledge import quote_xlsx as qx  # noqa: E402
from app.sales_deals import service as svc  # noqa: E402

PASS: list[str] = []
FAIL: list[tuple[str, str]] = []


def check(ok: bool, label: str, detail: str = "") -> None:
    (PASS.append(label) if ok else FAIL.append((label, detail)))


HEAD = {"customer": "상주보건소", "to": "안 현정 담당관님", "contact_name": "전 진웅", "contact_title": "팀장",
        "contact_mobile": "010-0000-0000", "contact_email": "name@example.com", "initials": "LU"}

# 1) 수기 작성 — 목적 '제품 대여'
st = qs.draft_state({}, model="", payload_kg=None, form={**HEAD, "purpose": "rental"}, manual=True)
f = st["form"]
check(f["purpose"] == "rental" and f["comments"] == qs.RENTAL_COMMENTS and f["foot_notes"] == [qs.RENTAL_NOTE],
      "대여 수기 견적: 대여 문구 · 중도 해지 주석", str(f["comments"]))
sale = qs.draft_state({}, model="", payload_kg=None, form=HEAD, manual=True)
check(sale["form"]["purpose"] == "sale" and sale["form"]["comments"] == qs.MANUAL_COMMENTS and not sale["form"]["foot_notes"],
      "판매(기본)는 예전 수기 견적 그대로")
bad = qs.draft_state({}, model="", payload_kg=None, form={**HEAD, "purpose": "lease"}, manual=True)
check(bad["form"]["purpose"] == "sale", "모르는 목적은 판매로")

# 방역 로봇 견적서 품목: 렌탈 36개월 × 425,000 / 설치 2,000,000 / 배송 500,000 / 유지보수 36개월 × 500,000
st, _ = qs.apply_ops(st, [
    {"op": "add", "name": "3년 렌탈(월 납입금)", "qty": 36, "unit": "개월", "unit_price": 425_000},
    {"op": "add", "name": "설치 및 시운전 (국내)", "qty": 1, "unit": "MD", "unit_price": 2_000_000},
    {"op": "add", "name": "배송비", "qty": 1, "unit": "식", "unit_price": 500_000},
    {"op": "add", "name": "유지보수 3년 기준", "qty": 36, "unit": "개월", "unit_price": 500_000},
    {"op": "delivery", "case": 5}], None)
check(any("대여 개월 수" in p for p in qs.problems(st)), "대여 개월 수 없으면 발행 막음", str(qs.problems(st)))

# 2) 개월 수·시작일
st, done = qs.apply_ops(st, [{"op": "set", "field": "rental_months", "value": "36"},
                             {"op": "set", "field": "rental_start", "value": "2026.11.1"}], None)
check(st["form"]["rental_months"] == 36 and st["form"]["rental_start"] == "2026-11-01", "개월 수 36 · 시작일 2026-11-01",
      str((st["form"]["rental_months"], st["form"]["rental_start"])))
st2, _ = qs.apply_ops(st, [{"op": "set", "field": "rental_start", "value": "11월쯤"}], None)
check(st2["form"]["rental_start"] == "2026-11-01", "못 읽는 날짜는 넣지 않음")
check(qs.problems(st) == [], "개월 수까지 채우면 발행 가능", str(qs.problems(st)))
total = sum(ln["qty"] * ln["unit_price"] for ln in st["lines"])
check(total == 35_800_000 and qs.rental_monthly(total, 36) == 995_000, "총액 35,800,000 → 월 995,000원(천 원 올림)",
      f"{total} {qs.rental_monthly(total, 36)}")

# 3) 판매 ↔ 대여
s2, _ = qs.apply_ops(st, [{"op": "set", "field": "purpose", "value": "sale"}], None)
check(s2["form"]["comments"] == qs.MANUAL_COMMENTS and qs.RENTAL_NOTE not in s2["form"]["foot_notes"],
      "판매로 바꾸면 대여 문구·주석 빠짐")
s3, _ = qs.apply_ops(s2, [{"op": "set", "field": "purpose", "value": "rental"}], None)
check(s3["form"]["comments"] == qs.RENTAL_COMMENTS and s3["form"]["foot_notes"].count(qs.RENTAL_NOTE) == 1,
      "다시 대여로 — 주석은 한 번만")
mine, _ = qs.apply_ops(st, [{"op": "set", "field": "comments", "value": "www.magbot.kr\n- 방역로봇 3년 렌탈"},
                            {"op": "set", "field": "purpose", "value": "sale"}], None)
check(mine["form"]["comments"] == "www.magbot.kr\n- 방역로봇 3년 렌탈", "담당자가 고친 Comments 는 목적을 바꿔도 그대로")
en, _ = qs.apply_ops(st, [{"op": "set", "field": "lang", "value": "en"}], None)
check(en["form"]["comments"] == qs.RENTAL_COMMENTS_EN and en["form"]["foot_notes"] == [qs.RENTAL_NOTE_EN],
      "해외로 바꾸면 대여 문구·주석도 영문")

# 4) 엑셀
lines = qs.render_lines(st, "ko", None)
h = qx.quote_header(st["form"], lang="ko", quote_no="QT-LU20261022-01", day=date(2026, 10, 22), pogo=False)
data = qx.build_xlsx(h, lines)
ws = openpyxl.load_workbook(io.BytesIO(data)).active
tot = next(r for r in range(1, ws.max_row + 1) if str(ws.cell(r, 1).value or "").startswith("Total"))
fv, gv = str(ws.cell(tot, 6).value), str(ws.cell(tot, 7).value)
check(fv.startswith("=SUM(") and gv == f"=ROUNDUP(F{tot}/36,-3)" and "월" in ws.cell(tot, 7).number_format,
      "Total 줄: 금액 칸 총액 SUM · 비고 칸 월 납입금 식", f"{fv} | {gv} | {ws.cell(tot, 7).number_format}")
check(any(qs.RENTAL_NOTE.split("\n")[0][2:20] in str(ws.cell(r, 1).value or "") for r in range(tot, ws.max_row + 1)),
      "표 아래 중도 해지 주석")
hs = qx.quote_header(sale["form"] | {"purpose": "sale"}, lang="ko", quote_no="QT-LU20261022-02", day=date(2026, 10, 22), pogo=False)
ws2 = openpyxl.load_workbook(io.BytesIO(qx.build_xlsx(hs, lines))).active
t2 = next(r for r in range(1, ws2.max_row + 1) if str(ws2.cell(r, 1).value or "").startswith("Total"))
check(str(ws2.cell(t2, 7).value).startswith("=SUM(") and hs["rental_months"] is None, "판매 견적서 Total 은 예전 그대로")

# 5) 다시 읽기
p = rb.parse(data)
check(p is not None and p["form"].get("purpose") == "rental" and p["form"].get("rental_months") == 36,
      "다시 읽기: 대여 양식·36개월 알아봄", str(p and p["form"].get("rental_months")))
check(p is not None and not any(r.get("in_total") is False for r in p["rows"]), "다시 읽기: 품목이 미반영 옵션으로 잘못 읽히지 않음")
plain = {**st, "form": {**st["form"], "purpose": "sale", "rental_months": None}}
merged, ch, _ = rb.merge(plain, p)
check(merged["form"]["purpose"] == "rental" and merged["form"]["rental_months"] == 36 and any("대여 개월 수" in c for c in ch),
      "올린 대여 견적서로 작성 상태도 대여로", str(ch))

# 6) 영업 건 대여 조건
r = svc._rental({"rental": {"months": 36, "start": "2026-11-01"}}, 35_800_000)
check(r == {"biz": "rental", "rental": {"months": 36, "start": "2026-11-01", "end": "2029-10-31", "monthly": 995_000}},
      "영업 건: 2026-11-01 + 36개월 → 2029-10-31, 월 995,000", str(r))
r = svc._rental({"rental": {"months": 3, "start": "2026-11-30"}}, 1_000_000)
check(r["rental"]["end"] == "2027-02-27" and r["rental"]["monthly"] == 334_000, "월말 시작도 끝나는 날 계산", str(r))
check(svc._rental({}, 1_000) == {} and svc._rental({"rental": {"months": 12}}, None)["rental"]["start"] is None,
      "판매 건은 그대로 · 시작일 없으면 비워 둠")
check(svc._add_months(dt.date(2026, 1, 31), 1) == dt.date(2026, 2, 28), "개월 더하기: 1/31 + 1개월 = 2/28")

# 7) 시리얼 칸 — '36 개월' 줄은 제품 1대, 품목 Service 는 빠짐
from app.sales_deals import units  # noqa: E402
tl = units.target_lines([{"name": "방역 로봇 3년 렌탈", "qty": 36, "unit": "개월", "unit_price": 425_000},
                         {"name": "유지보수 3년", "commodity": "Service", "qty": 36, "unit": "개월", "unit_price": 500_000},
                         {"name": "TCV1 Tool Plate", "qty": 3, "unit_price": 400_000}])
check([(x["name"], x["qty"]) for x in tl] == [("방역 로봇 3년 렌탈", 1), ("TCV1 Tool Plate", 3)],
      "시리얼 칸: 개월 줄은 1개, Service 줄은 빠짐, 일반 품목은 수량만큼", str(tl))

# 8) 단가 = 1대당 월 납입금(사용자 2026-10-22) — TCC1 Master 1대 월 200,000 · Tool Plate 2개 월 40,000 · 설치(Service) 1식 500,000, 12개월
t = qs.draft_state({}, model="", payload_kg=None, form={**HEAD, "purpose": "rental", "rental_months": 12}, manual=True)
t, _ = qs.apply_ops(t, [
    {"op": "add", "name": "TCC1 Master T.C (Robot Side)", "commodity": "Magbot ATC", "qty": 1, "unit_price": 200_000},
    {"op": "add", "name": "TCC1 Tool Plate (Tool Side)", "commodity": "Magbot ATC", "qty": 2, "unit_price": 40_000},
    {"op": "add", "name": "설치 및 교육", "commodity": "Service", "qty": 1, "unit": "식", "unit_price": 500_000},
    {"op": "delivery", "case": 5}], None)
check(qs.subtotal(t) == 3_860_000, "합계 = 1×200,000×12 + 2×40,000×12 + 설치 500,000(한 번) = 3,860,000", str(qs.subtotal(t)))
rl = qs.render_lines(t, "ko", None)
items = [r for r in rl if not r.get("row")]
check([r["amount"] for r in items] == [2_400_000, 960_000, 500_000] and [r.get("months") for r in items] == [12, 12, None],
      "줄 금액: 제품 줄은 ×12개월, Service 줄은 한 번", str([(r["amount"], r.get("months")) for r in items]))
h2 = qx.quote_header(t["form"], lang="ko", quote_no="QT-LU20261022-03", day=date(2026, 10, 22), pogo=False)
w2 = openpyxl.load_workbook(io.BytesIO(qx.build_xlsx(h2, rl))).active
rows2 = [r for r in range(17, w2.max_row + 1) if w2.cell(r, 4).value in (1, 2) and isinstance(w2.cell(r, 6).value, str)]
fs = [w2.cell(r, 6).value for r in rows2]
check(fs[:2] == [f"=D{rows2[0]}*E{rows2[0]}*12", f"=D{rows2[1]}*E{rows2[1]}*12"] and fs[2] == f"=D{rows2[2]}*E{rows2[2]}"
      and "월 납입금" in str(w2["E16"].value), "엑셀: 제품 줄 금액 =D×E×12, Service 줄 =D×E, 단가 칸 '(월 납입금)'", str((fs, w2["E16"].value)))
check(qs.rental_monthly(qs.subtotal(t), 12) == 322_000, "월 납입금 = 3,860,000 ÷ 12 → 322,000(천 원 올림)")
si = svc.snapshot_items(rl)
check([x["unit_price"] for x in si] == [2_400_000, 40_000 * 12, 500_000] and sum(x["unit_price"] * x["qty"] for x in si) == 3_860_000,
      "영업 건 품목: 단가 × 개월(1대당 계약 금액), 합계 3,860,000", str(si))
sale_t = {**t, "form": {**t["form"], "purpose": "sale"}}
check(qs.subtotal(sale_t) == 200_000 + 80_000 + 500_000, "판매 견적서면 개월을 곱하지 않음")

# 9) 한 번만 내는 줄(사용자 2026-10-22) — 유지보수 36개월 · 단위 '식' · 배송비/운반비, 제품 이름의 '물류·운반'은 월 납입금 그대로
R36 = {**HEAD, "purpose": "rental", "rental_months": 36}
ROBOT = {"op": "add", "name": "방역 로봇", "commodity": "Magbot", "qty": 1, "unit_price": 900_000}


def rental_total(ops: list[dict]) -> int:
    s = qs.draft_state({}, model="", payload_kg=None, form=R36, manual=True)
    s, _ = qs.apply_ops(s, ops, None)
    return qs.subtotal(s)


check(rental_total([ROBOT, {"op": "add", "name": "유지보수 3년", "commodity": "Maintenance", "qty": 36, "unit": "개월",
                            "unit_price": 500_000}]) == 900_000 * 36 + 18_000_000, "유지보수 36개월 × 50만 = 1,800만(개월 다시 안 곱함)")
check(rental_total([ROBOT, {"op": "add", "name": "유지보수 3년", "commodity": "Maintenance", "qty": 1, "unit": "식",
                            "unit_price": 18_000_000}]) == 900_000 * 36 + 18_000_000, "유지보수 1식 1,800만 — 단위 '식'은 한 번")
check(rental_total([ROBOT, {"op": "add", "name": "배송비", "qty": 1, "unit_price": 150_000}]) == 900_000 * 36 + 150_000,
      "배송비 — 물류비처럼 한 번")
check(qs.month_factor(R36, {"name": "운반비", "unit": "ea"}) == 1 and qs.month_factor(R36, {"name": "택배비"}) == 1,
      "운반비·택배비 줄(올린 파일 등)도 한 번")
check(qs.month_factor(R36, {"name": "운반 로봇 AMR", "commodity": "물류 로봇"}) == 36, "'운반 로봇'·'물류 로봇' 제품은 월 납입금 × 개월")

# 10) 대여 견적서 품목 + < OPTION > — 품목 묶음에도 Sub total + 회색 줄(사용자 2026-10-22). 판매는 예전대로 없음
OPT = [ROBOT, {"op": "add", "name": "추가 Tool Plate", "commodity": "Magbot ATC", "qty": 1, "unit_price": 40_000, "option": True},
       {"op": "delivery", "case": 5}]
for purpose, want in (("rental", ["subtotal", "gap", "section", "option", "gap"]), ("sale", ["gap", "section", "option", "gap"])):
    s = qs.draft_state({}, model="", payload_kg=None, form={**R36, "purpose": purpose}, manual=True)
    s, _ = qs.apply_ops(s, OPT, None)
    rows = [r.get("row") for r in qs.render_lines(s, "ko", None)]
    check(rows[1:] == want, f"{'대여' if purpose == 'rental' else '판매'} 품목 + OPTION → 품목 뒤 {' · '.join(map(str, want[:2]))}", str(rows))

for label in PASS:
    print("PASS", label)
for label, detail in FAIL:
    print("FAIL", label, "—", detail)
print(f"\n{len(PASS)} PASS / {len(FAIL)} FAIL")
sys.exit(1 if FAIL else 0)
