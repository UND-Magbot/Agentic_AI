"""제안서 프로젝트 → 고객용 PPT(project_deck) 검증 — 모델·DB 불필요(가짜 주입).

프로젝트 값은 가이드(UND 영업 AI 제안서 작성 가이드)의 삼성웰스토리 예시 답변을 그대로 쓴다(guide_items.json).
합격 기준은 가이드 표 9 "결과물과 에이전트 검수 사례" 중 PPT 쪽 항목.
"""
import asyncio
import io
import json
import re
import sys
from pathlib import Path

sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8", errors="replace")
ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "backend"))

from PIL import Image  # noqa: E402
from pptx import Presentation  # noqa: E402
from pptx.enum.shapes import MSO_SHAPE_TYPE  # noqa: E402

from app import proposal_llm  # noqa: E402
from app.proposal import project_deck as pd  # noqa: E402
from app.proposal_project.catalog import guide_notes  # noqa: E402

PASS: list[str] = []
FAIL: list[tuple[str, str]] = []


def check(ok: bool, label: str, detail: str = "") -> None:
    (PASS.append(label) if ok else FAIL.append((label, detail)))


# ── 가이드 예시로 만든 프로젝트 ─────────────────────────────────────────────
_TAG_STATUS = (("미확인", "unknown"), ("확인 필요", "unknown"), ("추가 결정", "unknown"), ("검토", "assumed"),
               ("확인", "confirmed"), ("고객", "confirmed"), ("채택", "adopted"), ("지정", "adopted"),
               ("현재", "adopted"), ("제안", "assumed"), ("운영", "adopted"), ("기본", "adopted"), ("승인", "adopted"))


def status_of(example: str) -> str:
    tag = re.match(r"\[([^\]]+)\]", example)
    t = tag.group(1) if tag else ""
    return next((s for k, s in _TAG_STATUS if k in t), "assumed")


items = {c: {"value": d["example"], "unit": "", "status": status_of(d["example"]), "nature": "", "source": "guide",
             "evidence": "가이드 예시", "checked_at": None, "pages": [], "version": 1}
         for c, d in guide_notes().items()}
check(items["P08"]["status"] == "confirmed" and "2,400" in items["P08"]["value"], "예시 P08 = 확인됨 2,400개")


def png(color) -> bytes:
    b = io.BytesIO()
    Image.new("RGB", (1600, 900), color).save(b, "PNG")
    return b.getvalue()


ALTS = [{"id": "A", "name": "휴머노이드형", "item_codes": ["P14"]},
        {"id": "B", "name": "AMR 양팔형", "item_codes": ["P15"]},
        {"id": "C", "name": "상부 레일 델타형", "item_codes": ["P16"]}]
IMAGES = [
    {"id": 1, "attachment_id": 101, "alt_id": None, "version": 1, "approved": True,
     "labels": [{"text": "승강 적재부 A~H", "x": 0.7, "y": 0.6}, {"text": "외측 카세트 교체", "x": 0.95, "y": 0.2}]},
    {"id": 2, "attachment_id": 102, "alt_id": "A", "version": 1, "approved": True, "labels": []},
    {"id": 3, "attachment_id": 103, "alt_id": "B", "version": 1, "approved": True, "labels": []},
    {"id": 4, "attachment_id": 104, "alt_id": "B", "version": 2, "approved": True,
     "labels": [{"text": "DOBOT 팔 2대", "x": 0.4, "y": 0.3}, {"text": "범위 밖", "x": 1.4, "y": 0.3}]},
    {"id": 5, "attachment_id": 105, "alt_id": "B", "version": 3, "approved": False, "labels": []},   # 미승인 최신본
    # C 는 승인 이미지 없음
]
QUOTE = [
    {"group": "공통", "item": "전후 비전", "qty": 2, "unit": "대", "basis": "통합 구성 개념 수량"},
    {"group": "공통", "item": "승강 적재부", "qty": 8, "unit": "식", "basis": "통합 구성 개념 수량"},
    {"group": "공통", "item": "이송·변환·버퍼", "qty": 1, "unit": "식", "basis": "통합 구성 개념 수량"},
    {"group": "공통", "item": "미취출 감지 센서", "qty": 1, "unit": "세트", "basis": "통합 구성 개념 수량"},
    {"group": "통합·실증", "item": "통합 제어부", "qty": 1, "unit": "식", "basis": "통합 구성 개념 수량"},
    {"group": "대안", "alt_id": "A", "item": "휴머노이드 로봇", "qty": 1, "unit": "대", "basis": "사용자 지정"},
    {"group": "대안", "alt_id": "B", "item": "DOBOT 10kg급 팔", "qty": 2, "unit": "대", "basis": "사용자 지정"},
    {"group": "대안", "alt_id": "B", "item": "AMR", "qty": 1, "unit": "대", "basis": "사용자 지정"},
    {"group": "대안", "alt_id": "B", "item": "양팔 EOAT", "qty": 2, "unit": "식", "basis": "개념 수량"},
    {"group": "대안", "alt_id": "C", "item": "델타 로봇 + 상부 레일", "qty": 1, "unit": "식", "basis": "사용자 지정"},
    {"group": "대안", "alt_id": "C", "item": "델타 EOAT", "qty": 1, "unit": "식", "basis": "개념 수량"},
    {"group": "옵션", "item": "툴체인저·툴스탠드", "qty": 1, "unit": "식", "basis": ""},
    {"group": "옵션", "item": "교대 턴테이블", "qty": 1, "unit": "식", "basis": ""},
    {"group": "현장 적용", "item": "현장 적용비(설치·운송·시운전)", "qty": 1, "unit": "식", "basis": ""},
]
project = {"id": 1, "title": "삼성웰스토리 식기세척 후단 자동화 실증", "request_text": "300식/끼, 약 2,400개를 2시간 이내",
           "items": items, "alternatives": ALTS, "images": IMAGES, "pages": [], "quote_lines": QUOTE,
           "output": {"max_pages": 10}}

# ── 1) 기본 페이지 구성 ─────────────────────────────────────────────────────
pages = pd.build_default_pages(project)
check([p["type"] for p in pages] == ["cover", "overview", "flow", "common_concept", "alternative", "alternative",
                                     "alternative", "equipment", "poc", "quote"], "대안 3개 → 가이드 표 7과 같은 10쪽",
      str([p["type"] for p in pages]))
check([p["page_no"] for p in pages] == list(range(1, 11)), "쪽 번호 1~10")
alt_b = pages[5]
check(alt_b["title"] == "AMR 양팔형" and alt_b["item_codes"][0] == "P15" and alt_b["asset_ids"] == [4],
      "대안 B: 고유 문항 P15 + 승인된 최신 이미지(v2, 미승인 v3 제외)", str(alt_b))
check(pages[6]["asset_ids"] == [] and pages[3]["asset_ids"] == [1], "승인 이미지 없는 대안 C 는 비움, 공통은 공통 이미지")
five = dict(project, alternatives=ALTS + [{"id": "D", "name": "겐트리형"}, {"id": "E", "name": "6축형"}])
p5 = pd.build_default_pages(five)
check(len(p5) == 10 and "poc" not in [p["type"] for p in p5] and "common_concept" not in [p["type"] for p in p5]
      and sum(p["type"] == "alternative" for p in p5) == 5, "대안 5개: 실증·공통 컨셉을 빼서 10쪽 이내",
      str([p["type"] for p in p5]))

# ── 2) 견적·주요 항목(표 9: 대안 합산 금지, 가격 없으면 별도 협의, 수량 일치) ─────
rows, totals = pd.quote_rows(project)
check(all(r[3] == pd.PRICE_TBD for r in rows), "근거 단가 없으면 전부 '별도 협의'")
check([t[0] for t in totals] == ["기본 구성 = 공통 + A안 · 휴머노이드형", "기본 구성 = 공통 + B안 · AMR 양팔형",
                                 "기본 구성 = 공통 + C안 · 상부 레일 델타형"] and all(t[1] == pd.PRICE_TBD for t in totals),
      "합계는 대안별 '공통 + 1종'만(세 대안 합산 없음)", str(totals))
eq = pd.equipment_rows(project)
qty_q = {r[1]: r[2] for r in rows}
check(all(qty_q[r[1]] == r[2] for r in eq) and len(eq) == len(QUOTE) - 1, "주요 항목과 견적의 장비 수량 일치(현장비만 제외)")
check(qty_q["승강 적재부"] == "8식" and qty_q["DOBOT 10kg급 팔"] == "2대", "수량·단위 표기")
priced = [dict(l, unit_price=1_000_000) if l["group"] in ("공통", "통합·실증") or l.get("alt_id") == "A"
          else dict(l) for l in QUOTE]
priced[0] = dict(priced[0], unit_price=5_000_000)
_, tot2 = pd.quote_rows(dict(project, quote_lines=priced))
common_sum = 5_000_000 * 2 + 1_000_000 * (8 + 1 + 1 + 1)
check(tot2[0][1] == f"{common_sum + 1_000_000:,}원", "A안 합계 = 공통 + A안 단가×수량", str(tot2))
check(tot2[1][1] == pd.PRICE_TBD, "단가 근거가 빠진 대안은 합계도 별도 협의(부분 합산 금지)")
nobasis = [dict(l, unit_price=1000, basis="") for l in QUOTE]
check(all(r[3] == pd.PRICE_TBD for r in pd.quote_rows(dict(project, quote_lines=nobasis))[0]),
      "단가가 있어도 근거가 없으면 금액을 쓰지 않음")

# ── 3) 문구 작성(가짜 모델) ─────────────────────────────────────────────────
calls: list[str] = []


def reply_for(user: str) -> dict:
    if "(flow)" in user:
        return {"headline": "세척 후 취출부터 종류별 적재까지 한 흐름으로 자동화합니다",
                "cards": [{"heading": "흐름", "bullets": ["소형 약 20%는 별도 분기"]}],
                "steps": [{"label": "세척기 배출", "note": "핑거 타입 배출부", "kind": "main"},
                          {"label": "플랫 변환·버퍼", "note": "속도 조정", "kind": "main"},
                          {"label": "세척 후 비전", "note": "위치·자세 확인", "kind": "main"},
                          {"label": "로봇 취출", "kind": "main"},
                          {"label": "종류별 승강 적재", "note": "A~H 8개 위치", "kind": "main"},
                          {"label": "소형 분기", "note": "Sorter", "kind": "branch", "from": 1},
                          {"label": "미취출 감지·정지", "note": "끝단 전 센서", "kind": "exception", "from": 3},
                          {"label": "버퍼 수용", "kind": "exception", "from": 3}]}
    if "(overview)" in user:
        if "[보완]" not in user:           # 첫 시도: 확인된 수치를 빠뜨림 + 근거 없는 수치·등급
            return {"headline": "식기 후단 자동화 가능성을 실증합니다",
                    "cards": [{"heading": "목표", "bullets": ["처리 효율 99.5% 달성", "IP69K 등급 로봇 적용"]}]}
        return {"headline": "식기 약 2,400개를 2시간 이내 처리하는 구성을 실증합니다",
                "cards": [{"heading": "고객 요구", "bullets": ["300식/끼, 약 3,000개/끼", "IP69K 등급 로봇 적용"]},
                          {"heading": "범위", "bullets": ["세척기 배출 이후 ~ 카세트 반출"]}]}
    if "(poc)" in user:
        return {"oops": True}              # 형식 오류 → 문항 값 대체
    return {"headline": "요지", "cards": [{"heading": "구성", "bullets": ["공통 모듈 재사용"]}]}


async def fake_chat(messages, **kw):
    user = messages[-1]["content"]
    calls.append(user)
    if "표지 문구" in messages[0]["content"]:
        return json.dumps({"title": "식기세척 후단 자동화 실증 제안", "subtitle": "세 가지 로봇 컨셉 비교",
                           "customer": "삼성웰스토리"}, ensure_ascii=False)
    return json.dumps(reply_for(user), ensure_ascii=False)


proposal_llm.chat = fake_chat
report = pd.DeckReport()
text = asyncio.run(pd.compose(project, pages, report))
check(text.customer == "삼성웰스토리" and text.title.startswith("식기세척"), "표지 문구")
ov = text.pages[2]
allt = pd._page_text_all(ov)
check("2,400개" in allt and sum("(overview)" in c for c in calls) == 2, "확인된 수치가 빠지면 한 번 다시 쓰게 함", allt)
check("IP69K" not in allt and "IP69K" in report.masked, "문항에 없는 등급(IP69K)은 가림", str(report.masked))
check(not any("99.5" in m for m in report.masked) or "99.5" not in allt, "첫 시도의 근거 없는 수치는 최종본에 없음")
check(any("5쪽" in w or "9쪽" in w for w in report.warnings) and text.pages[9].cards
      and text.pages[9].cards[0].heading.startswith("P"), "형식 오류 페이지는 문항 값으로 대체 + 경고", str(report.warnings))
flow = text.pages[3]
check([s.kind for s in flow.steps].count("main") == 5 and flow.steps[5].src == 1 and flow.steps[6].src == 3,
      "공정 단계: 정상 5 + 분기·예외(갈라지는 단계 번호)")
bad_flow = pd._parse_page({"steps": [{"label": "세척기 배출", "kind": "main"}, {"label": "플랫 변환·버퍼", "kind": "main"},
                                     {"label": "로봇 취출", "kind": "main"}, {"label": "미취출 감지·정지", "kind": "main"},
                                     {"label": "종류별 적재", "kind": "main"}]}, True)
check([(s_.label, s_.kind, s_.src) for s_ in bad_flow.steps if s_.kind != "main"] == [("미취출 감지·정지", "exception", 2)]
      and [s_.label for s_ in bad_flow.steps if s_.kind == "main"] == ["세척기 배출", "플랫 변환·버퍼", "로봇 취출", "종류별 적재"],
      "모델이 정지 단계를 정상 흐름에 넣어도 예외로 내림(버퍼는 정상 흐름 유지)",
      str([(s_.label, s_.kind, s_.src) for s_ in bad_flow.steps]))
p35 = next(u for u in calls if "(poc)" in u)
check("P35" in p35 and "[미확인]" not in p35.split("P35", 1)[1][:6] or "미확인" in p35, "문항 상태를 모델에 함께 전달")
check(not any("P31 [" in u and "비어 있음" in u for u in calls), "비어 있음·충돌 문항은 전달하지 않음")

# 성능 단정 → 검토(실측: 문항은 "클린룸 Class 100" 환경 사실뿐인데 "Class 100 클린룸 환경 대응 가능")
for src_, want in (("Class 100 클린룸 환경 대응 가능", "Class 100 클린룸 환경 대응 검토"),
                   ("클린룸 대응이 가능합니다", "클린룸 대응을 검토합니다"),
                   ("정밀도 보장", "정밀도 검토"),
                   ("스카라 로봇 기반 생산성 극대화", "스카라 로봇 기반 생산성 향상"),
                   ("처리량 1,200개 확보", "처리량 1,200개 확보 검토"),
                   ("공간 확보를 위한 배치", "공간 확보를 위한 배치"),          # 성능 표현 없는 확보는 그대로
                   ("자동화 가능성 검증", "자동화 가능성 검증"),              # "가능성"은 단정 아님
                   ("시간당 1,200개 처리 목표 달성", "시간당 1,200개 처리 목표 달성")):   # 이미 목표·검토 표현
    got, _ = pd.soften(src_)
    check(got == want, f"성능 단정 낮춤: {src_}", got)


async def cover_bad_then_good(messages, **kw):
    calls.append("cover")
    n = calls.count("cover")
    if n == 1:
        return json.dumps({"title": "식기세척 후단 자동화", "subtitle": "양산 투자 제안", "customer": "삼성웰스토리"},
                          ensure_ascii=False)
    return json.dumps({"title": "식기세척 후단 자동화 실증 제안", "subtitle": "세 가지 로봇 컨셉 비교",
                       "customer": "삼성웰스토리"}, ensure_ascii=False)


async def cover_always_bad(messages, **kw):
    return json.dumps({"title": "식기세척 후단 자동화", "subtitle": "양산 투자 제안", "customer": "삼성웰스토리"},
                      ensure_ascii=False)


p02_project = dict(project, items=dict(items, P02=dict(items["P02"], status="confirmed",
                                                        value="실증 결과를 보고 양산 투자를 결정합니다")))
calls.clear()
proposal_llm.chat = cover_bad_then_good
dt, r1 = pd.DeckText(), pd.DeckReport()
asyncio.run(pd._compose_cover(p02_project, pages[0], dt, r1))
check(calls.count("cover") == 2 and "실증" in dt.title and not r1.warnings, "표지가 P02(실증)와 어긋나면 한 번 다시 씀",
      f"{dt.title} / {dt.subtitle}")
proposal_llm.chat = cover_always_bad
dt2, r2 = pd.DeckText(), pd.DeckReport()
asyncio.run(pd._compose_cover(p02_project, pages[0], dt2, r2))
check(dt2.subtitle.startswith("자동화 실증 제안") and "양산 투자" not in dt2.subtitle and r2.warnings,
      "다시 써도 어긋나면 부제를 실증 제안으로 고치고 경고", f"{dt2.subtitle} {r2.warnings}")
proposal_llm.chat = fake_chat

# ── 4) 조판 ─────────────────────────────────────────────────────────────────
imgs = {1: png((200, 210, 220)), 2: png((180, 190, 200)), 4: png((170, 200, 210)), 5: png((0, 0, 0))}
data, rep = pd.render(project, pages, text, imgs, report=report)
prs = Presentation(io.BytesIO(data))


def slide_texts(s):
    return ["".join(r.text for p in sh.text_frame.paragraphs for r in p.runs) for sh in s.shapes if sh.has_text_frame]


check(len(prs.slides) == 10 and rep.pages == 10, "10쪽 이내(구성과 같은 쪽수)", str(len(prs.slides)))
bad = [sh for s in prs.slides for sh in s.shapes
       if sh.shape_type != MSO_SHAPE_TYPE.LINE and (sh.width <= 0 or sh.height <= 0)]
check(not bad, "크기 0 이하 도형 없음(연결선 제외)", str([b.name for b in bad]))
check(rep.overflow == [], "넘친 글상자·표 없음", str(rep.overflow))
cover = slide_texts(prs.slides[0])
check("삼성웰스토리" in cover and "PREPARED FOR" not in cover, "표지 고객사(제출본 식 — 영문 머리표 없음)", str(cover)[:200])
fs = prs.slides[2]
def big_pics(s):
    return [sh for s_ in [s] for sh in s_.shapes if sh.shape_type == MSO_SHAPE_TYPE.PICTURE and sh.width > 914400 * 2]


check(not big_pics(fs), "공정도는 그림이 아닌 도형(편집 가능)")
check(sum(sh.shape_type == MSO_SHAPE_TYPE.LINE for sh in fs.shapes) >= 4 + 3 + 3, "단계 화살표 4 + 분기 연결선 3 + 범례 3")
ft = slide_texts(fs)
check(all(any(k in x for x in ft) for k in ("미취출 감지·정지", "소형 분기", "세척기 배출", "핑거 타입 배출부")),
      "공정도에 정상·분기·예외 단계(설명은 상자 안)", str(ft)[:300])
common = prs.slides[3]
check(len(big_pics(common)) == 1, "공통 컨셉 이미지")
ct = slide_texts(common)
check("승강 적재부 A~H" in ct and "외측 카세트 교체" in ct, "이미지 위 라벨(편집 가능한 글상자)")
bt = slide_texts(prs.slides[5])
check("DOBOT 팔 2대" in bt and "범위 밖" not in bt, "좌표가 그림 밖인 라벨은 버림")
pics_b = big_pics(prs.slides[5])
check(len(pics_b) == 1 and pics_b[0].image.blob == imgs[4], "대안 B 는 승인된 v2 이미지", "")
check("승인된 컨셉 이미지가 없습니다" in slide_texts(prs.slides[6]) and any("7쪽" in w for w in rep.warnings),
      "승인 이미지 없는 대안 C: 자리 표시 + 경고")
qt = slide_texts(prs.slides[9])
qtab = [c.text for sh in prs.slides[9].shapes if sh.has_table for r in sh.table.rows for c in r.cells]
check("금액" in qtab and "별도 협의" in qtab and any("공통 + B안 · AMR 양팔형" in t for t in qt),
      "견적 표 + 대안별 기본 구성(택 1)", str(qt)[:300])
check(not any("A안" in t and "B안" in t for t in qt), "한 줄에 두 대안을 더한 합계 없음")
check(pd.QUOTE_HEADLINE in qt, "견적 쪽 핵심 메시지는 고정 문구(ROI 등 덧붙임 방지)")
etab = [c.text for sh in prs.slides[7].shapes if sh.has_table for r in sh.table.rows for c in r.cells]
check("승강 적재부" in etab and "8식" in etab and "현장 적용비(설치·운송·시운전)" not in etab, "주요 항목 표(현장비 제외)")
check(any("근거 단가가 없는 항목" in t for t in qt) and any("합산하지 않습니다" in t for t in qt), "견적 주석: 별도 협의·합산 금지")

# ── 4-1) 공정 컨셉 1개(기본) — 비교 표현 없이 ──────────────────────────────
single = dict(project, alternatives=[{"id": "A", "name": "협동로봇 셀", "item_codes": ["P14"]}],
              quote_lines=[l for l in QUOTE if l.get("alt_id") in (None, "A")])
check(not pd.comparing(single) and pd.comparing(project), "비교안 판단: 1개면 아님, 3개면 비교")
sp = pd.build_default_pages(single)
check([p["type"] for p in sp].count("alternative") == 1 and len(sp) == 8, "컨셉 1개 → 8쪽", str([p["type"] for p in sp]))
srows, stot = pd.quote_rows(single)
check(stot == [["기본 구성 = 공통 + 공정 컨셉", pd.PRICE_TBD]], "견적 합계 한 줄: 공통 + 공정 컨셉", str(stot))
check(all(r[4] != "선택 1종" for r in srows) and any(r[0] == "공정 컨셉 · 협동로봇 셀" for r in srows),
      "컨셉 행은 '포함', 구분은 '공정 컨셉'", str([r[0] for r in srows]))
st = pd.DeckText(title="t", pages={p["page_no"]: pd.PageText("요지", [pd.Card("구성", ["a"])],
                                                           [pd.FlowStep("투입"), pd.FlowStep("취출"), pd.FlowStep("적재")])
                                  for p in sp})
sdata, srep = pd.render(single, sp, st, {2: imgs[2]})
sprs = Presentation(io.BytesIO(sdata))
alltext = [x for s_ in sprs.slides for x in slide_texts(s_)]
tabtext = [c.text for s_ in sprs.slides for sh in s_.shapes if sh.has_table for r in sh.table.rows for c in r.cells]
check(not any(x.startswith("CONCEPT") for x in alltext), "컨셉 1개면 CONCEPT A 같은 비교 표시 없음")
check(not any("대안" in x or "택 1" in x or "합산하지" in x for x in alltext + tabtext),
      "대안·택 1·합산 문구 없음", str([x for x in alltext + tabtext if "대안" in x or "택 1" in x or "합산" in x])[:300])
check(pd.QUOTE_HEADLINE_SINGLE in alltext and srep.overflow == [], "견적 핵심 메시지(단일) + 넘침 없음", str(srep.overflow))

# ── 5) generate_pptx: 승인 구성 없으면 기본안 + 안내 ─────────────────────────
data2, rep2 = asyncio.run(pd.generate_pptx(project, images=imgs, products=[]))
check(len(Presentation(io.BytesIO(data2)).slides) == 10 and any("기본 구성" in w for w in rep2.warnings),
      "승인 구성이 없으면 기본 구성 + 안내")
custom = [dict(p) for p in pages if p["type"] != "poc"]
for n, p in enumerate(custom, 1):
    p["page_no"] = n
data3, rep3 = asyncio.run(pd.generate_pptx(dict(project, pages=custom), images=imgs, products=[]))
check(len(Presentation(io.BytesIO(data3)).slides) == 9 and not any("기본 구성" in w for w in rep3.warnings),
      "승인된 구성(9쪽)을 그대로 따름")
over = pages + [dict(pages[1], page_no=11)]
_, rep4 = pd.render(project, over, text, imgs)
check(any("상한 10쪽 초과" in w for w in rep4.warnings), "쪽수 상한 초과 경고")

# ── 6) 제출본 배치(2026-09-30): 핵심 수치·묶음 견적·비교표·공통 모듈 ─────────────
kp = pd._parse_page({"headline": "h", "kpis": [{"label": "처리 대상", "value": "2,400", "unit": "개"},
                                               {"label": "정확도", "value": "99.9", "unit": "%"},
                                               {"label": "글자만", "value": "많음", "unit": ""}]}, False)
check(len(kp.kpis) == 2, "핵심 수치는 숫자가 든 값만", str(kp.kpis))
rk = pd.DeckReport()
nums, src = pd._allowed(project)
pd._mask_page(kp, nums, src, rk)
check([k["value"] for k in kp.kpis] == ["2,400"], "근거 없는 핵심 수치(99.9%)는 칸째 뺌", str(kp.kpis))
kt = pd._parse_page({"kpis": [{"label": "품목", "value": "1", "unit": "종"}, {"label": "처리량", "value": "2,400", "unit": "개"}]}, False)
pd._mask_page(kt, nums, src, pd.DeckReport())
check([k["label"] for k in kt.kpis] == ["처리량"], "'1종' 같은 개수 1은 핵심 수치에서 뺌(QA 2회차)", str(kt.kpis))
check(pd._facts({"items": {"P06": {"status": "confirmed", "value": "브래킷 1종", "unit": ""}}}, ["P06"]) == {},
      "개수 1은 빠진 확인값으로 치지 않음(핵심 수치 규칙과 충돌 해소, QA 3회차)")
check("confirm" in pd._EDIT_SYSTEM and "확인 자료" in pd._EDIT_SYSTEM, "프롬프트 수정: 칸 이름 설명(확인 자료 → confirm)")
groups = pd.quote_groups(project)
check([g[0] for g in groups][:2] == ["공통 설비", "통합·실증"] and any(g[0] == "B. AMR 양팔형" for g in groups),
      "견적 묶음 행: 공통 묶음 + 컨셉별 한 줄", str([g[0] for g in groups]))
brow = next(g for g in groups if g[0] == "B. AMR 양팔형")
check("DOBOT 10kg급 팔 2대" in brow[1] and brow[2] == "택 1안", "컨셉 줄: 품목·수량 묶음 + 택 1안", str(brow))
check(not any("현장 적용비" in g[1] for g in groups), "현장 적용비는 표 밖 주석")
mx = pd.equipment_matrix(project)
check(mx and mx[0] == ["구분", "A. 휴머노이드형", "B. AMR 양팔형", "C. 상부 레일 델타형"]
      and [r[0] for r in mx[1:]][:2] == ["로봇 구성", "그리퍼·툴"], "주요 항목 비교표(행: 로봇·그리퍼, 열: 컨셉)", str(mx))
nocommon = dict(project, images=[i for i in IMAGES if i["alt_id"]])
cp = [dict(p, asset_ids=[]) if p["type"] == "common_concept" else p for p in pages]
cdata, crep = pd.render(nocommon, cp, text, imgs)
ctext = slide_texts(Presentation(io.BytesIO(cdata)).slides[3])
check("승인된 컨셉 이미지가 없습니다" not in ctext and not any("4쪽" in w for w in crep.warnings),
      "공통 컨셉 이미지가 없으면 빈 자리 대신 모듈 칸")

# ── 7) 완성본 프롬프트로 수정(apply_edit) ─────────────────────────────────────────
state = pd.deck_state(pages, text)
p2, t2 = pd.deck_from_state(state)
check(len(p2) == len(pages) and pd._page_to_dict(t2.pages[3]) == pd._page_to_dict(text.pages[3]),
      "문구 저장·복원(JSON 왕복)")
res = pd.apply_edit(project, p2, t2, {
    "cover": {"title": "식기세척 후단 자동화 실증 제안서"},
    "pages": [{"page_no": 2, "headline": "실증으로 확인할 3가지 — 인식·파지·적재", "title": "과제 개요"},
              {"page_no": 9, "remove": True}, {"page_no": 1, "remove": True},
              {"page_no": 5, "headline": "처리 속도 5,000개/시간을 검토합니다"}],
    "order": [1, 2, 4, 3, 5, 6, 7, 8, 10], "note": "개요 요지를 바꿨습니다"}, "2쪽 요지를 바꿔 줘")
check(len(res.pages) == 9 and res.pages[0]["type"] == "cover", "쪽 삭제(9쪽) + 표지는 못 뺌", str([p["type"] for p in res.pages]))
check(res.pages[1]["title"] == "과제 개요" and res.text.pages[2].headline.startswith("실증으로 확인할 3가지"),
      "쪽 제목·요지 고침")
check(res.pages[2]["type"] == "common_concept" and res.pages[3]["type"] == "flow", "쪽 순서 바꿈(3↔4)",
      str([p["type"] for p in res.pages]))
check("5,000개" not in res.text.pages[5].headline and res.report.masked, "고칠 때도 근거 없는 수치는 가림",
      res.text.pages[5].headline)
check(res.text.title == "식기세척 후단 자동화 실증 제안서" and "표지" in res.changed and "쪽 순서" in res.changed,
      "표지·순서 변경 기록", str(res.changed))
check([p["page_no"] for p in res.pages] == list(range(1, 10)), "쪽 번호 다시 매김")
edata, erep = pd.render(project, res.pages, res.text, imgs)
check(len(Presentation(io.BytesIO(edata)).slides) == 9 and erep.overflow == [], "고친 문구로 다시 조판", str(erep.overflow))
t7 = pd.deck_from_state(state)[1]
t7.pages[3] = pd._parse_page({"headline": "h", "cards": [{"heading": "운영 연계", "bullets": ["가" * 40, "나" * 40, "다" * 40]}],
                              "steps": [{"label": f"단계{i}", "note": "설명", "kind": "main"} for i in range(7)]
                              + [{"label": "미취출 감지", "kind": "exception", "from": 5},
                                 {"label": "버퍼 수용", "kind": "exception", "from": 6}]}, True)
d7, r7 = pd.render(project, pages, t7, imgs)
bad7 = [sh.name for s in Presentation(io.BytesIO(d7)).slides for sh in s.shapes
        if sh.shape_type != MSO_SHAPE_TYPE.LINE and (sh.width <= 0 or sh.height <= 0)]
check(not bad7, "7단계(두 줄) 공정도 + 예외 + 요점: 크기 0 이하 도형 없음(파일 손상 방지)", str(bad7))
rq = pd.requested_pages("2쪽 요지를 '실증으로 확인할 3가지'가 드러나게 바꾸고, 실증 쪽 확인 자료에 배출 영상도 넣어 줘.", pages)
check(rq == {2, 9}, "요청 속 쪽 알아보기: 쪽 번호 + '실증 쪽'(따옴표 속 '실증으로'는 아님)", str(rq))
check(pd.requested_pages("견적 페이지 주석을 줄여 줘, 표지 부제도", pages) == {1, 10}, "쪽 이름: 견적 페이지·표지")


async def _edit_twice():
    calls = []

    async def fake(messages, **_kw):
        calls.append(messages[-1]["content"])
        if "[보완]" in messages[-1]["content"]:
            return json.dumps({"pages": [{"page_no": 9, "confirm": "배출 영상과 도면을 받아 확인합니다"}], "note": "9쪽 보완"},
                              ensure_ascii=False)
        return json.dumps({"pages": [{"page_no": 2, "headline": "실증으로 확인할 3가지"}], "note": "2쪽"}, ensure_ascii=False)
    orig = pd.proposal_llm.chat
    pd.proposal_llm.chat = fake
    try:
        return await pd.edit_deck(project, state, "2쪽 요지를 바꾸고, 실증 쪽 확인 자료에 배출 영상도 넣어 줘.", images=imgs,
                                  products=[]), calls
    finally:
        pd.proposal_llm.chat = orig
(_b, _r, st2, er), ecalls = asyncio.run(_edit_twice())
check(len(ecalls) == 2 and er.changed == ["2쪽", "9쪽"] and "영상" in st2["text"]["9"]["confirm"],
      "요청한 쪽을 일부만 고치면 빠진 쪽만 한 번 더 요청(QA 4회차)", str(er.changed))
none = pd.apply_edit(project, *pd.deck_from_state(state), {"pages": [], "note": "견적 표는 구성 단계에서"}, "견적 바꿔")
check(none.changed == [] and "구성" in none.note, "바꿀 게 없으면 변경 없음 + 안내")

print(f"PASS {len(PASS)} / FAIL {len(FAIL)}")
for label, detail in FAIL:
    print(f"  ✗ {label}: {detail}")
sys.exit(1 if FAIL else 0)
