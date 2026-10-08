"""방금 생성한 _test_expense_fresh.xlsx 검증."""
from __future__ import annotations

import io
import sys
from pathlib import Path

sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8", errors="replace")
sys.stderr = io.TextIOWrapper(sys.stderr.buffer, encoding="utf-8", errors="replace")

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parent.parent
PATH = ROOT / "_test_expense_fresh.xlsx"

wb = load_workbook(PATH)

# 1) 내역 시트 — H/I/J/K/L 열 (개인카드)
print("=== expense_내역 (개인카드 H~L, rows 5~10) ===")
ws = wb["expense_내역"]
for r in range(5, 11):
    cells = []
    for c_letter in ("H", "I", "J", "K", "L"):
        cell = ws[f"{c_letter}{r}"]
        v = cell.value
        if v is not None:
            cells.append(
                f"{c_letter}{r}={v!r}({type(v).__name__})[{cell.number_format}]"
            )
    if cells:
        print("  " + " | ".join(cells))

# 2) 영수증 시트 — 캡션 + 이미지 anchor
print("\n=== expense_개인카드영수증 첨부 ===")
wr = wb["expense_개인카드영수증 첨부"]
for r in [4, 26, 48, 70]:
    cells = []
    for c in range(1, 18):
        v = wr.cell(row=r, column=c).value
        if v is not None and str(v).strip():
            cells.append(f"{get_column_letter(c)}{r}={v!r}")
    print(f"  R{r}: " + (" | ".join(cells) if cells else "(empty)"))

print(f"\n  IMAGES: {len(wr._images)}")
for i, im in enumerate(wr._images):
    a = im.anchor
    fr = a._from
    to = a.to
    print(
        f"    [{i}] "
        f"from=({get_column_letter(fr.col+1)}{fr.row+1} +{fr.colOff},{fr.rowOff}) "
        f"to=({get_column_letter(to.col+1)}{to.row+1} +{to.colOff},{to.rowOff})"
    )
