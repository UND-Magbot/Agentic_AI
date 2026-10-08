"""meeting_builder 단위 검증 — LLM/STT/인프라 불필요.

검증 항목:
1. 템플릿 38개 시트가 보존되고 새 주차 시트 1개가 append (총 39개).
2. 새 시트명이 명명 규칙('YYYY년 MM월DD일')을 따른다.
3. 제목(B2)/회의 주요 내용(D5)/지시사항(D13) 이 정확히 기입된다.
4. 새 시트의 핵심 병합 영역(B2:N3, D5:N12, D13:N16)이 보존된다.
5. 새 시트가 활성 시트로 설정된다.
6. 파일명 규칙.
"""
import io
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "backend"))

from openpyxl import load_workbook  # noqa: E402

from app.meeting_builder import (  # noqa: E402
    MeetingReport,
    build_filename,
    build_meeting_xlsx,
)

_MAIN = (
    "1. 독일 하노버 전시회 출장 결과\n"
    "- 시멘스 등 주요 기업 방문, 신규 파트너 미팅 활발\n\n"
    "2. 시장 진입 전략\n"
    "- B2G 영역 선제 공략 필요"
)
_DIRECTIVES = (
    "1. B2G 우선 시장 진입 전략 추진\n"
    "- Lab, 학교 등 연구 기관 협력 전략 구체화\n\n"
    "2. 제품 구조 재정비"
)


def main() -> int:
    failures: list[str] = []

    # 템플릿 원본 시트 수 측정.
    tpl = ROOT / "backend" / "app" / "templates" / "meeting_template.xlsx"
    base_count = len(load_workbook(tpl).sheetnames)

    report = MeetingReport(
        meeting_date="2026-05-19",
        main_content=_MAIN,
        directives=_DIRECTIVES,
        author="배재병",
    )
    data = build_meeting_xlsx(report)
    print(f"생성된 xlsx: {len(data)} bytes")

    wb = load_workbook(io.BytesIO(data))

    # 1) 시트 수 = 원본 + 1
    if len(wb.sheetnames) != base_count + 1:
        failures.append(
            f"시트 수: {len(wb.sheetnames)} (기대 {base_count + 1})"
        )

    # 2) 새 시트명
    expected_name = "2026년 05월19일"
    if expected_name not in wb.sheetnames:
        failures.append(f"새 시트명 '{expected_name}' 없음 — 실제: {wb.sheetnames[-1]!r}")
        ws = wb.worksheets[-1]
    else:
        ws = wb[expected_name]

    # 3) 셀 내용
    title = ws["B2"].value
    if title != "주간회의보고 (2026-05-19)":
        failures.append(f"제목 B2: {title!r}")
    if ws["D5"].value != _MAIN.strip():
        failures.append(f"회의 주요 내용 D5 불일치: {ws['D5'].value!r}")
    if ws["D13"].value != _DIRECTIVES.strip():
        failures.append(f"지시사항 D13 불일치: {ws['D13'].value!r}")

    # 4) 병합 영역 보존
    merged = {str(r) for r in ws.merged_cells.ranges}
    for need in ("B2:N3", "D5:N12", "D13:N16"):
        if need not in merged:
            failures.append(f"병합 영역 누락: {need} (실제: {sorted(merged)})")

    # 5) 활성 시트
    if wb.active.title != ws.title:
        failures.append(f"활성 시트: {wb.active.title!r} (기대 {ws.title!r})")

    # 6) 파일명
    fn = build_filename(report)
    if fn != "주간 회의록 보고_배재병_20260519.xlsx":
        failures.append(f"파일명: {fn!r}")
    print(f"파일명: {fn}")

    # 출력물 보존 — 육안 확인용.
    out = ROOT / "_test_meeting_builder.xlsx"
    out.write_bytes(data)
    print(f"저장: {out}")

    if failures:
        print(f"\nFAIL ({len(failures)}건):")
        for f in failures:
            print(f"  - {f}")
        return 1
    print("\nPASS — 6개 검증 항목 모두 통과")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
