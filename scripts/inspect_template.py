"""템플릿 + 기존 v7 결과물 검증 — UTF-8 강제, 셀 데이터 형식, 이미지 anchor."""
from __future__ import annotations

import io
import sys
from pathlib import Path

# Windows 콘솔 cp949 → utf-8 강제.
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8", errors="replace")
sys.stderr = io.TextIOWrapper(sys.stderr.buffer, encoding="utf-8", errors="replace")

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter


def inspect_book(path: Path, name: str) -> None:
    print(f"\n=== {name}: {path} ===")
    wb = load_workbook(path)
    print(f"  sheets: {wb.sheetnames}")
    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        print(f"\n  -- sheet: {sheet_name}")
        print(f"     dims: {ws.dimensions}, max_row={ws.max_row}, max_col={ws.max_column}")
        col_widths = {}
        for c_letter, dim in ws.column_dimensions.items():
            if dim.width:
                col_widths[c_letter] = round(dim.width, 2)
        if col_widths:
            print(f"     col widths: {col_widths}")
        if hasattr(ws, "_images") and ws._images:
            print(f"     IMAGES: {len(ws._images)}")
            for i, im in enumerate(ws._images):
                anchor = im.anchor
                kind = type(anchor).__name__
                msg = f"        [{i}] type={kind}"
                if hasattr(anchor, "_from") and anchor._from:
                    fr = anchor._from
                    msg += (
                        f"\n           from: col={fr.col}({get_column_letter(fr.col+1)}) "
                        f"row={fr.row+1} colOff={fr.colOff} rowOff={fr.rowOff}"
                    )
                if hasattr(anchor, "to") and anchor.to:
                    to = anchor.to
                    msg += (
                        f"\n           to:   col={to.col}({get_column_letter(to.col+1)}) "
                        f"row={to.row+1} colOff={to.colOff} rowOff={to.rowOff}"
                    )
                print(msg)


def dump_data_cells(path: Path, sheet_name: str, rng: list[str]) -> None:
    print(f"\n=== Cell values: {path.name} :: {sheet_name} ===")
    wb = load_workbook(path)
    ws = wb[sheet_name]
    for cell in rng:
        c = ws[cell]
        print(
            f"  {cell}: value={c.value!r} type={type(c.value).__name__} "
            f"num_fmt={c.number_format!r}"
        )


def dump_row_range(path: Path, sheet_name: str, rows: list[int], max_col: int = 17) -> None:
    print(f"\n=== Row dump: {path.name} :: {sheet_name} (rows={rows}) ===")
    wb = load_workbook(path)
    ws = wb[sheet_name]
    for r in rows:
        row_cells = []
        for c in range(1, max_col + 1):
            v = ws.cell(row=r, column=c).value
            if v is not None and str(v).strip():
                row_cells.append(f"{get_column_letter(c)}{r}={v!r}")
        if row_cells:
            print(f"  R{r}: " + " | ".join(row_cells))


if __name__ == "__main__":
    root = Path(__file__).resolve().parent.parent
    template = root / "backend" / "app" / "templates" / "expense_template.xlsx"
    v7 = root / "expense_v7_full.xlsx"

    inspect_book(template, "TEMPLATE")
    if v7.exists():
        inspect_book(v7, "V7 RESULT")
        dump_data_cells(
            v7,
            "expense_내역",
            [
                "I5", "I6", "I7", "I8",
                "B5", "B6", "B7", "B8",
                "K5", "K6", "K7", "K8",
                "D5", "D6", "D7", "D8",
            ],
        )
        # 영수증 시트 - row 4(슬롯1 캡션), 25(슬롯1 하단), 26(슬롯2 캡션), 27 등.
        dump_row_range(
            v7,
            "expense_개인카드영수증 첨부",
            [3, 4, 5, 24, 25, 26, 27, 28, 47, 48, 49, 50, 69, 70, 71],
            max_col=17,
        )
        dump_row_range(
            template,
            "expense_개인카드영수증 첨부",
            [3, 4, 5, 24, 25, 26, 27, 28, 47, 48, 49, 50, 69, 70, 71],
            max_col=17,
        )
