"""견적서 작성 대화(company_knowledge/quote_session.py) 점검 — 초안·질문 순서·말 읽기·숫자 검증·직접 수정·발행 검사·엑셀 줄.

    docker compose run --rm --no-deps -e PYTHONPATH=/app -v ./backend/app:/app/app -v ./scripts:/scripts -v ./docs:/docs backend \\
        python /scripts/test_quote_session.py
"""
from __future__ import annotations

import io
import sys
from datetime import date

import openpyxl

from app.company_knowledge import product_prices as pp
from app.company_knowledge import quote_session as qs
from app.company_knowledge import quote_xlsx as qx

PASS: list[str] = []
FAIL: list[tuple[str, str]] = []


def check(ok: bool, label: str, detail: str = "") -> None:
    (PASS.append(label) if ok else FAIL.append((label, detail)))


rows = pp.parse_book()


def acc(model: str = "TCV1") -> dict:
    items = [
        {"name": f"{model} MASTER T.C", "code": None, "qty": 1, "price_model": model, "price_item": "MASTER T.C", "optional": False},
        {"name": "IB (Interface Bracket)", "code": None, "qty": 1, "price_model": model, "price_item": "IB", "optional": True},
        {"name": f"{model} T.P", "code": None, "qty": 2, "price_model": model, "price_item": "T.P", "optional": False},
        {"name": "Pogo Pin Male (PPM)", "code": "PPM", "qty": 2, "price_model": model, "price_item": "PPM/PPF", "optional": False},
        {"name": "Pogo Pin Female (PPF)", "code": "PPF", "qty": 2, "price_model": model, "price_item": "PPM/PPF", "optional": False},
        {"name": "Cable 0.3m", "code": None, "qty": 4, "price_model": model, "price_item": "Cable 0.3m", "optional": False},
    ]
    a = {"applicable": True, "items": items}
    pp.attach_prices(items, rows)
    return a


# 1) 초안 — 추천 구성품 그대로, 체크 안 한 IB 는 넣을 수 있는 옵션으로
st = qs.draft_state(acc(), model="TCV1", payload_kg=10, form={"customer": "", "initials": "VT"})
check([ln["key"] for ln in st["lines"]] == ["MASTER T.C", "T.P", "PPM", "PPF", "Cable 0.3m"] and [o["key"] for o in st["options"]] == ["IB"]
      and qs.subtotal(st) == 2_300_000 + 800_000 + 500_000 + 500_000 + 200_000 and st["form"]["delivery"] == "통상 3 ~ 4주"
      and st["form"]["initials"] == "VT",
      "초안: 추천 품목·단가표 가격, IB 는 옵션 목록, 조건 기본값, 미리 채운 담당자 이니셜", str(qs.subtotal(st)))

# 2) 질문 순서 — 고객사 → 납품 방식(4가지) → 물류비 / 설치 인원·일수 → 1인 1일 단가 → 추가 항목 → 담당자
check(qs.next_question(st)["id"] == "customer", "첫 질문: 고객사")
st, done = qs.apply_ops(st, [{"op": "set", "field": "customer", "value": "가나테크"}, {"op": "set", "field": "to", "value": "김철수 대표님"}],
                        "가나테크, 김철수 대표님")
q = qs.next_question(st)
check(st["form"]["customer"] == "가나테크" and q["id"] == "delivery" and len(q["choices"]) == 5 and len(done) == 2
      and q["choices"][-1] == "5. 별도 협의",
      "고객사 반영 → 다음은 납품 방식(선택 버튼 5개, 5. 별도 협의)", str(q))

# 납품 방식 4가지 — 품목 줄이 방식에 맞춰진다(영업팀 기준: 1 그대로 / 2 물류비 / 3 인건비 / 4 물류비+인건비)
base = qs.subtotal(st)
check([qs.delivery_case(x) for x in ("1", "2번", "3. UND 인원 설치 납품", "4. 화물 발송 + UND 인원 설치", "택배로 보내요", "화물로 직납",
                                     "저희 인원이 설치하면서 납품", "화물로 보내고 설치하러 가요", "글쎄요")] == [1, 2, 3, 4, 1, 2, 3, 4, None],
      "납품 방식 읽기: 번호·버튼 글·말 → 1~4, 모호하면 묻기")
check([qs.delivery_case(x) for x in ("화물은 안 써요, 저희가 직접 가서 설치해요", "택배 말고 화물로", "설치 없이 화물로만 보내요",
                                     "택배로 보내고 설치는 안 해요")] == [3, 2, 2, 1],
      "부정 표현: '화물은 안 써요'·'택배 말고'·'설치 없이'·'설치는 안 해요'를 반대로 읽지 않음")
check([qs.delivery_case(x) for x in ("5", "5. 별도 협의", "납품은 별도 협의할게요", "추후 협의")] == [5, 5, 5, 5],
      "5 별도 협의: 번호·버튼 글·말")
c5, _ = qs.apply_ops(st, qs.quick_ops(st, "5. 별도 협의"), "5. 별도 협의")
check(qs.subtotal(c5) == base and not any(ln["kind"] in ("freight", "work") for ln in c5["lines"])
      and qs.next_question(c5)["id"] == "extra" and not any("납품 방식" in x for x in qs.problems(c5)),
      "5 별도 협의: 품목 그대로(물류비·인건비 없음), 다음 질문으로")
c1, _ = qs.apply_ops(st, qs.quick_ops(st, "1. 택배 발송"), "1. 택배 발송")
check(qs.subtotal(c1) == base and not any(ln["kind"] in ("freight", "work") for ln in c1["lines"])
      and qs.next_question(c1)["id"] == "extra", "1 택배: 품목 그대로, 바로 추가 항목 질문")
PLACE = "직납 및 배차 비용 별도\n(화물 비용 별도)"
check(st["form"]["place"] == PLACE and all(c["form"]["place"] == PLACE for c in (c1,)),
      "납품장소 = '직납 및 배차 비용 별도 (화물 비용 별도)' — 초안부터, 방식과 관계없이")
c2, _ = qs.apply_ops(st, qs.quick_ops(st, "2"), "2")
check(qs.next_question(c2)["id"] == "freight" and "물류비 금액 필요" in qs.problems(c2), "2 화물 직납: 물류비 금액을 묻는다")
c2, _ = qs.apply_ops(c2, qs.quick_ops(c2, "15만원이요"), "15만원이요")
check(qs.subtotal(c2) == base + 150_000 and any(ln["kind"] == "freight" and ln["unit_price"] == 150_000 for ln in c2["lines"])
      and c2["form"]["place"] == PLACE and qs.next_question(c2)["id"] == "extra",
      "2 화물 직납: 물류비 150,000원 줄, 납품장소 문구는 그대로")
c3, _ = qs.apply_ops(st, qs.quick_ops(st, "3번 2명 3일"), "3번 2명 3일")
check(qs.next_question(c3)["id"] == "extra" and not any(ln["kind"] == "freight" for ln in c3["lines"])
      and any(ln["kind"] == "work" and ln["qty"] == 6 and ln["unit_price"] == qs.WORK_RATE for ln in c3["lines"])
      and qs.subtotal(c3) == base + 6 * 800_000,
      "3 설치 납품: '3번 2명 3일' → 인건비 6 인·일 × 회사 기준 80만원(1인 8H), 단가는 묻지 않음")
check(any(ln["key"] == "work" and "1인 8H 기준" in ln["remark"] for ln in qs.render_lines(c3, "ko", None)), "견적서 비고에 '1인 8H 기준'")
c3, _ = qs.apply_ops(c3, [{"op": "work", "rate": 500000}], "인건비 단가 50만원으로 해 줘")
check(qs.subtotal(c3) == base + 3_000_000 and qs.next_question(c3)["id"] == "extra"
      and not any("8H" in ln["remark"] for ln in qs.render_lines(c3, "ko", None) if ln["key"] == "work"),
      "말하면 단가 바꿀 수 있음: 2명×3일×50만원 = 3,000,000원(기준 아님 → 비고에 8H 안 붙음)")
c4, _ = qs.apply_ops(st, qs.quick_ops(st, "화물로 보내고 설치하러 가요"), "화물로 보내고 설치하러 가요")
for m in ("10만원", "1명 2일"):
    c4, _ = qs.apply_ops(c4, qs.quick_ops(c4, m), m)
check(qs.subtotal(c4) == base + 100_000 + 1_600_000 and {ln["kind"] for ln in c4["lines"]} >= {"freight", "work"}
      and qs.next_question(c4)["id"] == "extra", "4 화물 + 설치: 물류비 100,000 + 인건비 1명×2일×80만원(회사 기준)")
# 대화로 방식 바꾸기 — 설치 빼기(4→2), 물류비 추가(1→2), 인건비 줄 삭제(3→1)
c4b, d4 = qs.apply_ops(c4, [{"op": "work", "needed": False}], "설치는 빼 주세요")
check(c4b["delivery"]["case"] == 2 and not any(ln["kind"] == "work" for ln in c4b["lines"])
      and any(ln["kind"] == "freight" for ln in c4b["lines"]), "'설치는 빼 주세요' → 4에서 2(화물 직납)로", str(d4))
c1b, _ = qs.apply_ops(c1, [{"op": "add", "name": "물류비", "qty": 1, "unit_price": 80000}], "물류비 8만원 추가")
check(c1b["delivery"]["case"] == 2 and qs.subtotal(c1b) == base + 80_000 and not any(ln["kind"] == "extra" for ln in c1b["lines"]),
      "택배(1)에 '물류비 8만원 추가' → 화물 직납(2)으로, 물류비 줄 하나")
wl = next(ln for ln in c3["lines"] if ln["kind"] == "work")
c3b, _ = qs.apply_ops(c3, [{"op": "remove", "line": wl["id"]}], "인건비 빼 줘")
check(c3b["delivery"]["case"] == 1 and qs.subtotal(c3b) == base, "인건비 줄 삭제 → 3에서 1(택배)로")

st = c4
st, _ = qs.apply_ops(st, [{"op": "add", "name": "교육비", "qty": 1, "unit_price": 100000}], "교육비 10만원 추가해 줘")
check(any(ln["name"] == "교육비" and ln["kind"] == "extra" for ln in st["lines"]) and qs.next_question(st)["id"] == "contact",
      "대화 '교육비 10만원 추가' → 추가 항목 줄, 추가 항목 질문 끝, 다음은 담당자")
st2, done2 = qs.apply_ops(st, [{"op": "add", "name": "출장비", "qty": 1, "unit_price": 300000}], "출장비도 넣어 줘")
check(not any(ln["name"] == "출장비" for ln in st2["lines"]) and not done2, "말에 없는 금액(30만원)은 AI 가 지어내도 넣지 않음")

early = qs.draft_state(acc(), model="TCV1", payload_kg=10, form={"customer": "가나테크"})       # 아직 납품 방식 질문 단계
guard, gd = qs.apply_ops(early, [{"op": "set", "field": "contact_name", "value": "김철수"}], "김철수 대표님께 보낼 거예요")
check(not gd and not guard["form"]["contact_name"], "고객 쪽 받는 분을 AI 가 담당자로 넣어도 막음(담당자 질문 전·'담당' 말 없음)")

# 3) 품목 수정·삭제·옵션
tp = next(ln for ln in st["lines"] if ln["key"] == "T.P")
st, _ = qs.apply_ops(st, [{"op": "update", "line": tp["id"], "qty": 3}, {"op": "option", "name": "IB", "include": True}],
                     "T.P 3개로 바꾸고 IB 넣어 줘")
check(next(ln for ln in st["lines"] if ln["key"] == "T.P")["qty"] == 3 and any(ln["key"] == "IB" for ln in st["lines"]) and not st["options"],
      "대화로 T.P 3개·IB 포함")
ppf = next(ln for ln in st["lines"] if ln["key"] == "PPF")
st, _ = qs.apply_ops(st, [{"op": "remove", "line": ppf["id"]}], "PPF 빼 줘")
check(not any(ln["key"] == "PPF" for ln in st["lines"]), "대화로 품목 삭제")
st, _ = qs.apply_ops(st, [{"op": "set", "field": "contact_name", "value": "홍길동"}, {"op": "set", "field": "contact_title", "value": "이사"},
                          {"op": "set", "field": "contact_mobile", "value": "010-0000-0000"},
                          {"op": "set", "field": "contact_email", "value": "name@example.com"}], "홍길동 이사 010-0000-0000 name@example.com")
check(qs.next_question(st) is None and qs.problems(st) == [], "담당자까지 채우면 질문 끝, 발행 가능", str(qs.problems(st)))

# 4) 직접 수정(왼쪽 화면) — 바뀐 것만 변경으로
lines_in = [{"id": ln["id"], "name": ln["name"], "qty": ln["qty"], "unit_price": ln["unit_price"], "remark": ln.get("remark", "")}
            for ln in st["lines"] if ln["kind"] not in ("work", "freight")]
lines_in[0]["unit_price"] = 2_100_000                                   # 마스터 단가 직접 조정
lines_in.append({"name": "시운전 교육", "qty": 1, "unit_price": 200000, "remark": "1회"})
ops = qs.edit_ops(st, {"form": {"payment": "발주시 100%"}, "lines": lines_in, "delivery": {"case": 4, "freight": 120000},
                      "work": {"people": 2, "days": 4, "rate": 500000}})
st, done = qs.apply_ops(st, ops, None)
check(st["form"]["payment"] == "발주시 100%" and st["lines"][0]["unit_price"] == 2_100_000
      and any(ln["name"] == "시운전 교육" for ln in st["lines"]) and next(ln for ln in st["lines"] if ln["kind"] == "work")["qty"] == 8
      and next(ln for ln in st["lines"] if ln["kind"] == "freight")["unit_price"] == 120_000,
      "직접 수정: 결제조건·마스터 단가·새 항목·물류비 120,000·설치 2명 4일 반영", str(done))

# 5) 발행 검사(C13)
bad = qs.draft_state(acc("TCK100"), model="TCK100", payload_kg=None, form={"customer": "x"})
check(any("가격 미정" in p for p in qs.problems(bad)) and any("이니셜" in p for p in qs.problems(bad))
      and any("납품 방식" in p for p in qs.problems(bad)),
      "단가표 0원·이니셜 없음·납품 방식 미선택 → 발행 막음", str(qs.problems(bad)))

# 6) 엑셀 줄 — 본체 → 옵션 → 물류비 → 설치 인건비 → 추가 항목, 언어별 문구, 해외는 달러
rl = qs.render_lines(st, "ko", None)
kinds = [ln["key"] for ln in rl]
check(kinds[:2] == ["MASTER T.C", "T.P"] and kinds[-4:-2] == ["freight", "work"] and rl[-3]["unit"] == "인·일"
      and rl[-3]["remark"] == "2명 × 4일" and rl[-4]["desc"] == "물류비" and rl[-3]["desc"] == "설치 인건비"
      and {rl[-2]["desc"], rl[-1]["desc"]} == {"교육비", "시운전 교육"} and rl[0]["desc"] == "TCV1\n- MASTER T.C (Payload 10Kg)",
      "엑셀 줄 순서(물류비 → 설치 인건비 → 추가 항목)·'인·일'·비고", str(kinds))
re_ = qs.render_lines(st, "en", 1380)
check(re_[-3]["desc"] == "Installation labor" and re_[-4]["desc"] == "Freight" and re_[-3]["unit"] == "man-days"
      and re_[0]["unit_price"] == round(2_100_000 / 1380, 2), "영문: 물류비·인건비 문구·단위(8 man-days 복수형), 달러 환산")
check(re_[0]["desc"] == "TCV1\n- Master Tool Changer (Payload: 10 kg)" and re_[1]["desc"] == "TCV1\n- Tool Plate (Payload: 10 kg)"
      and rl[0]["desc"] == "TCV1\n- MASTER T.C (Payload 10Kg)",
      "영문 본체 문구 = 해외 기준 견적서 표기(Master Tool Changer·Tool Plate·Payload: 10 kg), 국내는 그대로", str(re_[0]["desc"]))
check([qs.unit_label(u, "en", n) for u, n in (("식", 1), ("식", 2), ("세트", 2), ("개", 3), ("개월", 1), ("개월", 12), ("set", 2), ("box", 2))]
      == ["lot", "lots", "sets", "ea", "month", "months", "sets", "box"] and qs.unit_label("식", "ko", 2) == "식",
      "영문 단위 단·복수(1 lot / 2 lots, 2 sets — 해외 기준 견적서), ea·모르는 단위는 그대로, 국내는 그대로")
h = qx.quote_header(st["form"] | {"initials": "VT"}, lang="ko", quote_no="UND-VT20261006-01", day=date(2026, 10, 6), pogo=True)
ws = openpyxl.load_workbook(io.BytesIO(qx.build_xlsx(h, rl))).active
vals = {c.coordinate: c.value for row in ws.iter_rows() for c in row if c.value is not None}
work_row = next(k for k, v in vals.items() if v == "설치 인건비")
place_row = next(k for k, v in vals.items() if v == "납품장소")
check(ws[f"D{work_row[1:]}"].number_format == '0" 인·일"' and ws[f"D{work_row[1:]}"].value == 8
      and vals.get(f"G{place_row[1:]}") == "직납 및 배차 비용 별도\n(화물 비용 별도)",
      "엑셀: 인건비 수량 '8 인·일', 납품장소 '직납 및 배차 비용 별도 (화물 비용 별도)'")
en_st, _ = qs.apply_ops(st, [{"op": "set", "field": "lang", "value": "en"}], None)
check(en_st["form"]["place"] == "UND Korea Factory", "영문 납품장소 문구(해외 기준 견적서, 2026-10-14)")

# v1.2 액세서리 — 수량 미정 줄은 빼지 않고 발행을 막고, 조건부 수량은 담당자 확인 전 발행 불가(PGR24)
ca = acc()
for it in ca["items"]:
    if it["price_item"] in ("PPM/PPF", "Cable 0.3m"):
        it.update(conditional=True, basis="조건부 근거")
ca["items"].append({"name": "Pneumatic Male (PMM)", "code": "PMM", "qty": None, "price_model": "TCV1", "price_item": "PMM/PMF", "optional": False})
pp.attach_prices(ca["items"], rows)
cs = qs.draft_state(ca, model="TCV1", payload_kg=10)
pr = qs.problems(cs)
check(any(ln["key"] == "PMM" and ln["qty"] is None for ln in cs["lines"]) and any("PMM" in x and "수량 미정" in x for x in pr)
      and sum("조건부 수량" in x for x in pr) == 3,
      "v1.2: 수량 미정(PMM) 줄은 남겨 발행을 막고, 조건부 PPM·PPF·케이블은 확인 전 발행 불가", str(pr))
ppm_id = next(ln["id"] for ln in cs["lines"] if ln["key"] == "PPM")
ppf_id = next(ln["id"] for ln in cs["lines"] if ln["key"] == "PPF")
cs2, done2 = qs.apply_ops(cs, qs.edit_ops(cs, {"lines": [{**{k: ln.get(k) for k in ("id", "name", "qty", "unit_price", "remark")},
                                                       **({"confirmed": True} if ln["id"] == ppm_id else {})}
                                                      for ln in cs["lines"] if ln["kind"] not in ("work", "freight")]}), None)
cs3, _ = qs.apply_ops(cs2, [{"op": "update", "line": ppf_id, "qty": 3}], "PPF 3개로")
pr3 = qs.problems(cs3)
check(not any("Pogo Pin Male" in x and "조건부" in x for x in pr3) and not any("Pogo Pin Female" in x and "조건부" in x for x in pr3)
      and any("조건부" in x for x in pr3) and any("확인" in d for d in done2),
      "[수량 확인]·수량 수정 → 그 줄만 확인됨, 남은 조건부 줄(케이블)은 계속 막음", str(pr3))

# 입금 계좌(사용자 2026-10-08) — 기본은 회사 계좌, 대화·직접 수정으로 바꿈, 말에 없는 계좌번호는 안 받음, 언어 바꾸면 그 언어 기본 계좌
bk = qs.draft_state(acc(), model="TCV1", payload_kg=10, form={"customer": "x"})
check(bk["form"]["bank"] == pp.QUOTE_ISSUER["bank"] and bk["form"]["account_no"] == pp.QUOTE_ISSUER["account_no"]
      and bk["form"]["account_holder"] == pp.QUOTE_ISSUER["account_holder"], "입금 계좌 기본값 = 회사 계좌", str({k: bk["form"][k] for k in qs.BANK_FIELDS}))
msg = "입금 계좌를 신한은행 110-123-456789 로 바꿔 줘"
b2, done = qs.apply_ops(bk, [{"op": "set", "field": "bank", "value": "신한은행"}, {"op": "set", "field": "account_no", "value": "110-123-456789"}], msg)
check(b2["form"]["bank"] == "신한은행" and b2["form"]["account_no"] == "110-123-456789" and any("입금 계좌번호" in x for x in done),
      "대화로 입금 은행·계좌번호 변경", str(done))
b3, _ = qs.apply_ops(bk, [{"op": "set", "field": "account_no", "value": "999-999-9999"}], "계좌 바꾸고 싶어요")
check(b3["form"]["account_no"] == pp.QUOTE_ISSUER["account_no"], "말에 없는 계좌번호는 넣지 않음")
b4, _ = qs.apply_ops(b2, [{"op": "set", "field": "lang", "value": "en"}], None)
check(b4["form"]["account_no"] == pp.QUOTE_ISSUER_EN["account_no"] and b4["form"]["bank"] == pp.QUOTE_ISSUER_EN["bank"],
      "해외로 바꾸면 영문 기본 계좌로", str({k: b4["form"][k] for k in qs.BANK_FIELDS}))
hf = {**b2["form"], "customer": "x", "contact_name": "a", "contact_mobile": "1", "contact_email": "e"}
hd = qx.quote_header(hf, lang="ko", quote_no="VT20261008", day=date(2026, 10, 8), pogo=False)
hd0 = qx.quote_header({**hf, "bank": "", "account_no": "", "account_holder": ""}, lang="ko", quote_no="VT20261008", day=date(2026, 10, 8), pogo=False)
check(hd["issuer"]["account_no"] == "110-123-456789" and hd["issuer"]["bank"] == "신한은행"
      and hd0["issuer"]["account_no"] == pp.QUOTE_ISSUER["account_no"], "견적서 머리: 바꾼 계좌, 비우면 회사 기본 계좌")
check("account_no" in qs.QUOTE_SYSTEM and "결제조건(payment)과 다르다" in qs.QUOTE_SYSTEM, "AI 지시: 입금 계좌 ≠ 결제조건, 거절·지어내기 금지")

# ── 해외 견적: 환율은 선택, 달러 단가 직접 입력(사용자 2026-10-13 — 해외 기준 견적서처럼 $ 로만) ──
CT = {"customer": "ABC Robotics", "initials": "AL", "contact_name": "Alex Jeon", "contact_mobile": "+82-10-0000-0000",
      "contact_email": "a@example.com", "lang": "en"}
us = qs.draft_state(acc(), model="TCV1", payload_kg=10, form=CT)
us, _ = qs.apply_ops(us, [{"op": "delivery", "case": 5}, {"op": "no_extra"}], None)
pr0 = qs.problems(us)
check(any("달러 단가 없음" in x for x in pr0) and not any("환율 필요" == x[-5:] for x in pr0)
      and (qs.next_question(us) or {}).get("id") == "fx",
      "해외 + 환율 없음 + 달러 단가 없음 → '달러 단가 없음'(환율 자체는 필수 아님), AI 는 달러 단가나 환율을 물음", str(pr0))
pay = {"lines": [{"id": ln["id"], "name": ln["name"], "qty": ln["qty"], "unit_price": ln["unit_price"], "remark": ln.get("remark", ""),
                  "usd": 100.0 + i} for i, ln in enumerate(us["lines"])]}
us2, _ = qs.apply_ops(us, qs.edit_ops(us, pay), None)
rl = qs.render_lines(us2, "en", None)
check(not qs.problems(us2) and qs.next_question(us2) is None and rl[0]["unit_price"] == 100.0
      and qs.subtotal_usd(us2) == round(sum(ln["qty"] * (100.0 + i) for i, ln in enumerate(us["lines"])), 2),
      "환율 없이 달러 단가만 적으면 발행 가능 — 견적서 단가 = 적은 달러, 달러 합계", str(qs.problems(us2)))
fxs, _ = qs.apply_ops(us, [{"op": "set", "field": "fx_rate", "value": 1400}], None)
one = {"lines": [{"id": ln["id"], "name": ln["name"], "qty": ln["qty"], "unit_price": ln["unit_price"], "remark": ln.get("remark", ""),
                  "usd": 999.0 if i == 0 else None} for i, ln in enumerate(fxs["lines"])]}
fxs2, _ = qs.apply_ops(fxs, qs.edit_ops(fxs, one), None)
r2 = qs.render_lines(fxs2, "en", 1400)
check(not qs.problems(fxs2) and r2[0]["unit_price"] == 999.0 and r2[1]["unit_price"] == round(fxs2["lines"][1]["unit_price"] / 1400, 2),
      "환율을 넣으면 비운 줄은 원화 ÷ 환율, 달러를 적은 줄은 그 값", str([x["unit_price"] for x in r2[:2]]))
clr = {"lines": [{**x, "usd": None} for x in one["lines"]]}
fxs3, _ = qs.apply_ops(fxs2, qs.edit_ops(fxs2, clr), None)
check(fxs3["lines"][0].get("usd") is None and qs.render_lines(fxs3, "en", 1400)[0]["unit_price"] == round(fxs3["lines"][0]["unit_price"] / 1400, 2),
      "달러 단가를 지우면 다시 환율 환산")
ad, _ = qs.apply_ops(us2, qs.edit_ops(us2, {"lines": [*pay["lines"], {"id": None, "name": "Training", "qty": 1, "unit_price": None, "usd": 500}]}), None)
check(ad["lines"][-1]["name"] == "Training" and ad["lines"][-1]["usd"] == 500 and ad["lines"][-1]["unit_price"] is None
      and not qs.problems(ad) and qs.render_lines(ad, "en", None)[-1]["unit_price"] == 500.0,
      "해외 견적에 달러 단가만 적은 새 항목 추가(원화 단가 없음)")
check(any("Training" in x for x in qs.problems({**ad, "form": {**ad["form"], "lang": "ko"}})),
      "같은 견적을 국내로 바꾸면 원화 단가 없는 줄은 '단가 없음'")
fr, _ = qs.apply_ops(us2, [{"op": "delivery", "case": 2, "freight": 150000}], None)
check(any("환율 필요" in x for x in qs.problems(fr)),
      "해외 + 물류비(원화) 줄 + 환율 없음 → 환율 필요(물류비를 달러로 바꿀 수 없음)", str(qs.problems(fr)))

# ── 대화로 달러 단가('단가 1000달러', 사용자 2026-10-13) — AI 가 낸 변경을 코드가 검증 ──
check(qs.says_usd("마스터 단가 1000달러로") and qs.says_usd("$1,200") and qs.says_usd("USD 900") and qs.says_usd("1500불")
      and qs.says_usd("1만 달러") and not qs.says_usd("단가 150만원") and not qs.says_usd("1,500,000원") and not qs.says_usd("2개로"),
      "달러 표시('달러'·'$'·'USD'·'불') 알아봄, 원화 표시가 있으면 달러로 안 봄")
L1 = us["lines"][0]["id"]
c1, d1 = qs.apply_ops(us, [{"op": "update", "line": L1, "usd": 1000}], "마스터 단가 1000달러로 해 줘")
check(c1["lines"][0]["usd"] == 1000 and any("$1,000.00" in x for x in d1), "대화: '1000달러' → 그 줄 달러 단가 $1,000", str(d1))
c2, _ = qs.apply_ops(us, [{"op": "update", "line": L1, "unit_price": 1200}], "마스터 $1,200")
check(c2["lines"][0]["usd"] == 1200 and c2["lines"][0]["unit_price"] == us["lines"][0]["unit_price"],
      "AI 가 달러 금액을 원화 칸(unit_price)에 넣어도 말에 달러가 있으면 달러 단가로 옮김(원화 단가표 가격은 그대로)")
c3, _ = qs.apply_ops(us, [{"op": "update", "line": L1, "usd": 1500}], "마스터 1000달러")
check(not c3["lines"][0].get("usd"), "말에 없는 달러 금액(1500)은 넣지 않음 — 숫자는 말한 것만")
c4, _ = qs.apply_ops(us, [{"op": "add", "name": "Training", "qty": 1, "unit_price": 500}], "교육비 500달러 추가")
check(c4["lines"][-1]["name"] == "Training" and c4["lines"][-1]["usd"] == 500 and c4["lines"][-1]["unit_price"] is None,
      "대화: 새 항목을 달러로 추가 → 달러 단가만(원화 없음)", str(c4["lines"][-1]))
c5, _ = qs.apply_ops(us, [{"op": "update", "line": L1, "unit_price": 1500000}], "마스터 단가 150만원")
check(c5["lines"][0]["unit_price"] == 1500000 and not c5["lines"][0].get("usd"), "원화로 말하면 지금처럼 원화 단가")
check(qs.quick_ops(us, "1000달러") is None and qs.quick_ops(us, "1380") == [{"op": "set", "field": "fx_rate", "value": 1380.0}],
      "환율을 묻는 중에 '1000달러'라고 하면 환율로 읽지 않음(숫자만 말하면 환율)")
check("달러 합계" in qs.describe(us2) and "$" in qs.describe(us2) and "달러 합계" not in qs.describe({**us2, "form": {**us2["form"], "lang": "ko"}}),
      "AI 에게 주는 견적 상태에 해외 견적이면 환율·달러 합계도")

c6, d6 = qs.apply_ops(us, [{"op": "add", "name": "TCC1", "qty": 2, "unit_price": None},
                            {"op": "update", "line": "L99", "usd": 1000}], "TCC1 2세트, 단가 1000달러로 추가해 줘")
check(c6["lines"][-1]["name"] == "TCC1" and c6["lines"][-1]["usd"] == 1000 and c6["lines"][-1]["qty"] == 2,
      "AI 가 '단가 없는 추가 + 없는 줄 단가 변경'으로 나눠 내도 한 항목으로 합쳐 넣음(실제 사내 AI 응답 꼴)", str(d6))

# ── 묶음(옵션·소제목·사양·문구)·NEGO(사용자 2026-10-13, 영업팀 요청 14·9번) ──
MF = {"customer": "상주보건소", "initials": "AL", "contact_name": "전진웅", "contact_mobile": "010-0000-0000", "contact_email": "a@example.com"}
mb = qs.draft_state({}, model="", payload_kg=None, manual=True, form=MF)
mb, _ = qs.apply_ops(mb, [{"op": "delivery", "case": 5}, {"op": "no_extra"}], None)
check([b["id"] for b in qs.blocks(mb)] == ["B0"] and qs.blocks(mb)[0]["type"] == "items", "예전·새 견적 모두 기본 묶음(B0) 하나로 시작")
pay = {"lines": [{"id": None, "name": "UND-DR1", "commodity": "방역 로봇", "qty": 1, "unit_price": 15_300_000, "remark": "", "block": "B0"},
                 {"id": None, "name": "- 교육 포함", "commodity": "설치 및 시운전 (국내)", "qty": 1, "unit_price": 2_000_000, "remark": "", "block": "new-a"},
                 {"id": None, "name": "배송비", "commodity": "", "qty": 1, "unit_price": 500_000, "remark": "", "block": "new-a"},
                 {"id": None, "name": "CS-01", "commodity": "충전 스테이션", "qty": 1, "unit_price": 1_200_000, "remark": "", "block": "new-o"},
                 {"id": None, "name": "16L", "commodity": "예비 물탱크", "qty": 1, "unit_price": 425_000, "remark": "", "block": "new-o"}],
       "blocks": [{"id": "B0", "type": "items", "title": "", "subtotal": False},
                  {"id": "new-s", "type": "spec", "title": "방역 로봇 주요 사양",
                   "rows": [{"label": "무게", "value": "55kg"}, {"label": "방역특징", "value": "초음파 미립자 분무\n1000 ㎡ / 15분"}, {"label": "", "value": ""}]},
                  {"id": "new-a", "type": "items", "title": "설치 및 시운전", "subtotal": True},
                  {"id": "new-o", "type": "option", "title": ""},
                  {"id": "new-t", "type": "text", "text": "렌탈 조건은 별도 계약서를 따릅니다."}],
       "nego": {"mode": "amount", "value": 800_000, "label": "특별 할인"},
       "form": {"total_note": "월 995,000원", "foot_notes": "본 견적은 1개월 유효\n* 중도 해지 시 위약금"}}
mb2, done_b = qs.apply_ops(mb, qs.edit_ops(mb, pay), None)
bl = qs.blocks(mb2)
check([b["type"] for b in bl] == ["items", "spec", "items", "option", "text"] and [b["id"] for b in bl] == ["B0", "B1", "B2", "B3", "B4"]
      and len(bl[1]["rows"]) == 2 and bl[2]["subtotal"] is True,
      "화면의 새 묶음(임시 id) → B1~B4, 빈 사양 줄은 뺌", str([(b["id"], b["type"]) for b in bl]))
check([ln.get("block") for ln in mb2["lines"]] == [None, "B2", "B2", "B3", "B3"], "줄이 새 묶음 id 를 따라감(기본 묶음은 표시 없음)",
      str([ln.get("block") for ln in mb2["lines"]]))
check(qs.subtotal(mb2) == 15_300_000 + 2_500_000 - 800_000 and not qs.problems(mb2),
      "합계 = 품목 묶음만(옵션 1,625,000 제외) − NEGO 800,000 = 17,000,000", str((qs.subtotal(mb2), qs.problems(mb2))))
rl = qs.render_lines(mb2, "ko", None)
check([r.get("row") for r in rl] == [None, "subtotal", "gap", "section", "spec", "spec", "gap", "section", None, None, "subtotal", "gap",
                                     "section", "option", "option", "gap", "text", "gap", "nego", "gap"],
      "견적서 줄 순서: 품목 → 사양(소제목) → 구분·소제목·품목·소계 → 구분·< OPTION >·옵션 → 문구 → NEGO", str([r.get("row") for r in rl]))
check([r["amount"] for r in rl if r.get("row") == "subtotal"] == [15_300_000, 2_500_000]
      and next(r for r in rl if r.get("row") == "section" and "OPTION" in r["text"])["text"] == "< OPTION >" and rl[-2]["amount"] == -800_000 and rl[-2]["desc"] == "NEGO (특별 할인)" and rl[-1]["row"] == "gap"
      and qx.total(rl) == 17_000_000, "소계 2,500,000, 옵션 제목 기본 '< OPTION >', NEGO -800,000, 견적 합계 17,000,000", str(qx.total(rl)))
mp, _ = qs.apply_ops(mb2, qs.edit_ops(mb2, {"nego": {"mode": "percent", "value": 10, "label": ""}}), None)
check(qs.subtotal(mp) == round(17_800_000 * 0.9) and qs.render_lines(mp, "ko", None)[-2]["desc"] == "NEGO (10%)",
      "NEGO % = 옵션 뺀 합계(17,800,000) 기준 10% → 견적서엔 금액, 이름 'NEGO (10%)'", str(qs.subtotal(mp)))
mn, dn_ = qs.apply_ops(mb2, qs.edit_ops(mb2, {"nego": None}), None)
check(not mn.get("nego") and qs.subtotal(mn) == 17_800_000 and "NEGO 삭제" in dn_, "NEGO 비우면 삭제")
big, _ = qs.apply_ops(mb2, [{"op": "nego", "amount": 20_000_000}], None)
check(any("NEGO 할인 금액" in x for x in qs.problems(big)), "NEGO 가 합계보다 크면 발행 막음")
cl, _ = qs.apply_ops(mb2, qs.edit_ops(mb2, {"form": {"total_note": ""}}), None)
check(cl["form"]["total_note"] == "" and mb2["form"]["foot_notes"] == ["본 견적은 1개월 유효", "* 중도 해지 시 위약금"],
      "Total 아래 문구는 지울 수 있음, 예전 글(줄바꿈)로 온 주석은 줄마다 한 칸", str(mb2["form"]["foot_notes"]))
fn, _ = qs.apply_ops(mb2, qs.edit_ops(mb2, {"form": {"foot_notes": ["본 견적은 1개월 유효", "", "* 중도 해지 시, 잔여 개월 수 × 월 납입료\n위약금을 일시불로 납부"]}}), None)
check(fn["form"]["foot_notes"] == ["본 견적은 1개월 유효", "* 중도 해지 시, 잔여 개월 수 × 월 납입료\n위약금을 일시불로 납부"],
      "주석 칸 목록 — 빈 칸은 빼고, 칸 안 줄바꿈은 그대로", str(fn["form"]["foot_notes"]))
fe, _ = qs.apply_ops(fn, qs.edit_ops(fn, {"form": {"foot_notes": []}}), None)
check(fe["form"]["foot_notes"] == [], "주석을 다 빼면 빈 목록")
rm = {"blocks": [b for b in qs.blocks(mb2) if b["id"] != "B2"],
      "lines": [{**{k: ln.get(k) for k in ("id", "name", "qty", "unit_price", "remark", "usd")}, "block": ln.get("block")} for ln in mb2["lines"]
                if ln.get("block") != "B2"]}
mr, _ = qs.apply_ops(mb2, qs.edit_ops(mb2, rm), None)
check([b["id"] for b in qs.blocks(mr)] == ["B0", "B1", "B3", "B4"] and len(mr["lines"]) == 3, "묶음을 지우면 그 줄도 같이 빠짐(화면이 함께 보냄)")
# 대화: 옵션·소제목·NEGO
c1, d1 = qs.apply_ops(mb, [{"op": "add", "name": "CS-01", "commodity": "충전 스테이션", "qty": 1, "unit_price": 1_200_000, "option": True}],
                      "충전 스테이션 120만원 옵션으로 넣어 줘")
check(qs.blocks(c1)[-1]["type"] == "option" and c1["lines"][-1]["block"] == qs.blocks(c1)[-1]["id"] and "Total 미반영" in d1[-1]
      and qs.subtotal(c1) == 0, "대화: '옵션으로' → 옵션 묶음을 만들어 넣고 기본은 Total 미반영", str(d1))
c2, d2 = qs.apply_ops(mb, [{"op": "add", "name": "설치 1MD", "qty": 1, "unit_price": 2_000_000, "group": "설치 및 시운전"},
                           {"op": "add", "name": "배송비", "qty": 1, "unit_price": 500_000, "group": "설치 및 시운전"}],
                      "설치 및 시운전 묶음 만들고 소계 보여 줘. 설치 1MD 200만원, 배송비 50만원")
g = [b for b in qs.blocks(c2) if b.get("title") == "설치 및 시운전"]
check(len(g) == 1 and g[0]["subtotal"] and [ln.get("block") for ln in c2["lines"]] == [g[0]["id"]] * 2 and "묶음 추가: 설치 및 시운전 (소계 표시)" in d2,
      "대화: 소제목 묶음을 한 번만 만들고 두 줄을 넣음(소계 표시)", str(d2))
c3, d3 = qs.apply_ops(c2, [{"op": "nego", "amount": 300_000}], "30만원 네고해 줘")
c4, _ = qs.apply_ops(c2, [{"op": "nego", "amount": 900_000}], "30만원 네고해 줘")
check(qs.subtotal(c3) == 2_200_000 and not c4.get("nego"), "대화: '30만원 네고' → NEGO -300,000, 말에 없는 금액은 안 받음", str(d3))
c5, _ = qs.apply_ops(c3, [{"op": "nego", "remove": True}], "네고 빼 줘")
check(not c5.get("nego"), "대화: 네고 빼기")
c6, _ = qs.apply_ops(c2, [{"op": "update", "line": c2["lines"][1]["id"], "option": True}], "배송비는 옵션으로 돌려 줘")
check(qs.subtotal(c6) == 2_000_000, "대화: 기존 줄을 옵션 묶음으로 옮기면 합계에서 빠짐")
fr2, _ = qs.apply_ops(mb, [{"op": "add", "name": "화물 운송", "qty": 1, "unit_price": 150_000, "option": True}], "화물 운송 15만원 옵션으로")
check(fr2["lines"][-1]["name"] == "화물 운송" and not (fr2.get("delivery") or {}).get("freight"),
      "옵션 묶음에 넣는 '화물' 항목은 납품 방식 물류비로 바꾸지 않음")
uen, _ = qs.apply_ops(us2, [{"op": "nego", "percent": 10}], "10% 할인")
check(qs.subtotal_usd(uen) == round(qs.subtotal_usd(us2) * 0.9, 2) and qs.render_lines(uen, "en", None)[-2]["amount"] == -round(qs.subtotal_usd(us2) * 0.1, 2),
      "해외(달러) 견적 NEGO % 는 달러 합계 기준", str(qs.subtotal_usd(uen)))

check({425000.0, 120000000.0, 35000.0} <= qs.said_numbers("물탱크 42만 5천원, 장비 1억 2천만원, 소모품 3만5천원")
      and 425000.0 in qs.said_numbers("42만5천"), "'42만 5천원'·'1억 2천만원'·'3만5천원'처럼 이어 말한 금액도 읽음",
      str(sorted(qs.said_numbers("물탱크 42만 5천원, 장비 1억 2천만원, 소모품 3만5천원"))))
c7, _ = qs.apply_ops(c3, [{"op": "nego", "percent": 5}], "네고는 5%로")
check(c7["nego"] == {"mode": "percent", "value": 5.0, "label": ""} or c7["nego"]["mode"] == "percent", "네고 % 로 바꾸기")
lab0, _ = qs.apply_ops(mb2, [{"op": "nego", "percent": 5}], "네고 5%로 바꿔 줘")
check(lab0["nego"]["label"] == "특별 할인", "대화로 네고 금액·%만 바꾸면 이름(특별 할인)은 그대로", str(lab0["nego"]))

g2, _ = qs.apply_ops(mb, [{"op": "group", "title": "유지보수", "subtotal": True},
                          {"op": "add", "name": "유지보수 3년", "qty": 36, "unit_price": 500_000}], "유지보수 묶음 만들고 소계, 유지보수 3년 36개월 월 50만원")
gb = [b for b in qs.blocks(g2) if b.get("title") == "유지보수"]
check(len(gb) == 1 and g2["lines"][-1]["block"] == gb[0]["id"], "AI 가 묶음 만들기와 항목 추가를 따로 내도 그 말에서 만든 묶음에 넣음(실제 사내 AI 응답 꼴)")

u1, _ = qs.apply_ops(mb, [{"op": "add", "name": "유지보수 3년", "qty": 36, "unit": "개월", "unit_price": 500_000}], "유지보수 36개월 월 50만원")
u2, _ = qs.apply_ops(u1, qs.edit_ops(u1, {"lines": [{"id": u1["lines"][-1]["id"], "name": "유지보수 3년", "qty": 36, "unit_price": 500_000,
                                                      "remark": "", "unit": "개\"월"}]}), None)
check(u1["lines"][-1]["unit"] == "개월" and qs.render_lines(u1, "ko", None)[-1]["unit"] == "개월" and u2["lines"][-1]["unit"] == "개월",
      "수량 단위(대화 '36개월' → 개월, 화면 수정) — 견적서 수량 칸 '36 개월', 따옴표는 뺌", str(u2["lines"][-1].get("unit")))

# ── 사용자 지정 묶음 + 서술형 줄 + 물류비·인건비 넣을 묶음(사용자 2026-10-13) ──
sg = qs.draft_state({}, model="", payload_kg=None, manual=True, form=MF)
sg, _ = qs.apply_ops(sg, [{"op": "no_extra"}], None)
pg = {"blocks": [{"id": "B0", "type": "items", "title": "", "subtotal": False},
                 {"id": "new-g", "type": "items", "title": "설치 및 시운전", "subtotal": True}],
      "lines": [{"id": None, "name": "UND-DR1", "commodity": "방역 로봇", "qty": 1, "unit_price": 15_300_000, "remark": "", "block": "B0"},
                {"id": None, "name": "55kg", "commodity": "무게", "qty": None, "unit_price": None, "remark": "", "block": "B0", "kind": "note"},
                {"id": None, "name": "- 교육 포함", "commodity": "설치 및 시운전 (국내)", "qty": 1, "unit": "MD", "unit_price": 2_000_000,
                 "remark": "", "block": "new-g"}],
      "delivery": {"case": 4, "freight": 500_000, "block": "new-g"}, "work": {"people": 1, "days": 1, "rate": 500_000}}
sg2, dg = qs.apply_ops(sg, qs.edit_ops(sg, pg), None)
note = next(ln for ln in sg2["lines"] if ln["kind"] == "note")
check(note["commodity"] == "무게" and note["name"] == "55kg" and note["qty"] is None and not qs.problems(sg2),
      "서술형 줄(항목명 / 내용, 가격 없음)은 수량·단가 검사에서 빠짐", str(qs.problems(sg2)))
check(sg2["delivery"]["block"] == "B1" and any("물류비·인건비 넣을 묶음: 설치 및 시운전" in x for x in dg),
      "납품 방식 칸에서 고른 묶음(새 묶음 임시 id → B1)", str(sg2["delivery"]))
rg = qs.render_lines(sg2, "ko", None)
rws = [(r.get("row"), r.get("label") or r.get("text") or r.get("commodity")) for r in rg]
check(rws[:2] == [(None, "방역 로봇"), ("spec", "무게")] and ("section", "설치 및 시운전") in rws
      and [r.get("commodity") for r in rg if not r.get("row")][1:] == ["설치 및 시운전 (국내)", "Service", "Service"],
      "견적서: 서술형 줄은 사양 줄로, 물류비·인건비 줄은 고른 '설치 및 시운전' 묶음 안(소계 포함)", str(rws))
subt = [r for r in rg if r.get("row") == "subtotal"][-1]
check(subt["amount"] == 2_000_000 + 500_000 + 500_000 and qs.subtotal(sg2) == 15_300_000 + 3_000_000 and qx.total(rg) == 18_300_000,
      "소계 = 설치 + 물류비 + 인건비, 합계는 서술형 줄과 무관", str((subt["amount"], qs.subtotal(sg2))))
sg3, _ = qs.apply_ops(sg2, qs.edit_ops(sg2, {"blocks": [{"id": "B0", "type": "items", "title": "", "subtotal": False}],
                                              "lines": [{"id": ln["id"], "name": ln["name"], "qty": ln["qty"], "unit_price": ln["unit_price"],
                                                         "remark": "", "block": "B0"} for ln in sg2["lines"]
                                                        if ln["kind"] not in ("work", "freight") and ln.get("block") != "B1"]}), None)
check(qs.auto_block(sg3) == "B0" and all(r.get("row") != "section" for r in qs.render_lines(sg3, "ko", None)),
      "고른 묶음을 지우면 물류비·인건비는 다시 기본 품목으로")

# ── 국내↔해외 전환 — 기본 문구만 바뀌고 직접 고친 문구는 그대로(사용자 2026-10-13) ──
lk = qs.draft_state({}, model="", payload_kg=None, manual=True, form={"customer": "X"})
le, _ = qs.apply_ops(lk, qs.edit_ops(lk, {"form": {**lk["form"], "lang": "en"}}), None)
check(le["form"]["comments"] == qs.MANUAL_COMMENTS_EN and le["form"]["delivery"] == "4 ~ 6 weeks after payment"
      and le["form"]["bank"] == "KOOKMIN BANK", "해외로: 조건 문구·계좌가 영문 기본으로(수기 견적 Comments 는 Product Supply)", str(le["form"]["comments"]))
lb, _ = qs.apply_ops(le, qs.edit_ops(le, {"form": {**le["form"], "lang": "ko"}}), None)
check(lb["form"]["comments"] == qs.MANUAL_COMMENTS and lb["form"]["delivery"] == "통상 3 ~ 4주" and lb["form"]["place"] == lk["form"]["place"]
      and lb["form"]["bank"] == "국민은행" and lb["form"]["payment"] == lk["form"]["payment"],
      "다시 국내로: 수기 견적 Comments('제품 공급 건')·납기·납품장소·계좌가 국내 기본으로 돌아옴", str({k: lb["form"][k] for k in ("comments", "delivery")}))
ak = qs.draft_state(acc(), model="TCV1", payload_kg=10, form={"customer": "x"})
ae, _ = qs.apply_ops(ak, qs.edit_ops(ak, {"form": {**ak["form"], "lang": "en"}}), None)
ab, _ = qs.apply_ops(ae, qs.edit_ops(ae, {"form": {**ae["form"], "lang": "ko"}}), None)
check(ab["form"]["comments"] == ak["form"]["comments"], "툴체인저 견적은 툴체인저 Comments 로 돌아옴")

# 납품 방식 칸에서 '별도 묶음을 만들어 넣기' — 항목 없이 물류비·인건비만 든 새 묶음(사용자 2026-10-13)
na = qs.draft_state({}, model="", payload_kg=None, manual=True, form=MF)
na, _ = qs.apply_ops(na, [{"op": "no_extra"}, {"op": "add", "name": "UND-DR1", "qty": 1, "unit_price": 15_300_000}], None)
na2, dna = qs.apply_ops(na, qs.edit_ops(na, {
    "blocks": [{"id": "B0", "type": "items", "title": "", "subtotal": False}, {"id": "new-77", "type": "items", "title": "설치 및 시운전", "subtotal": True}],
    "delivery": {"case": 4, "freight": 500_000, "block": "new-77"}, "work": {"people": 1, "days": 2, "rate": 1_000_000},
    "lines": [{"id": ln["id"], "name": ln["name"], "qty": ln["qty"], "unit_price": ln["unit_price"], "remark": "", "block": "B0"} for ln in na["lines"]]}), None)
rn = qs.render_lines(na2, "ko", None)
check(qs.auto_block(na2) == "B1" and [r.get("row") for r in rn] == [None, "subtotal", "gap", "section", None, None, "subtotal", "gap"]
      and rn[-2]["amount"] == 500_000 + 2_000_000 and not qs.problems(na2),
      "새 묶음(임시 id)을 만들며 물류비·인건비를 넣음 → '설치 및 시운전' 소제목 아래 물류비·인건비 + 소계", str([r.get("row") for r in rn]))

# 단락 규칙(사용자 2026-10-13): 품목 묶음 2개 이상이면 품목 묶음마다 Sub total(OPTION 은 없음), 단락이 2개 이상이면 단락마다 끝에 회색 줄
one = qs.render_lines(na, "ko", None)
check(all(not r.get("row") for r in one), "기본 묶음 하나뿐인 견적은 예전처럼 Sub total·회색 줄 없음", str([r.get("row") for r in one]))
lg, _ = qs.apply_ops(na, [{"op": "add", "name": "CHARGING DOCK", "qty": 1, "unit_price": 6_131_580, "option": True}], "충전독 6131580원 옵션으로")
# OPTION 은 Total 반영 줄이 있을 때만 Sub total 을 셀 묶음(사용자 2026-10-23) — 회색 줄은 단락마다 그대로
check([r.get("row") for r in qs.render_lines(lg, "ko", None)] == [None, "gap", "section", "option", "gap"],
      "기본 + OPTION(미반영): 둘 다 Sub total 없이, 단락마다 회색 줄", str([r.get("row") for r in qs.render_lines(lg, "ko", None)]))
lg_in, _ = qs.apply_ops(na, [{"op": "add", "name": "CHARGING DOCK", "qty": 1, "unit_price": 6_131_580, "option": True, "in_total": True}], "충전독 6131580원 옵션으로, Total 반영")
check([r.get("row") for r in qs.render_lines(lg_in, "ko", None)] == [None, "subtotal", "gap", "section", "option", "subtotal", "gap"],
      "기본 + OPTION(반영): 기본 품목·OPTION 둘 다 Sub total", str([r.get("row") for r in qs.render_lines(lg_in, "ko", None)]))
three = {**lg, "blocks": [*qs.blocks(lg), {"id": "B9", "type": "items", "title": "설치 및 시운전"}],
         "lines": [*lg["lines"], {"id": "x9", "kind": "extra", "key": "extra", "name": "설치", "commodity": "Service", "qty": 1,
                                  "unit_price": 500_000, "remark": "", "block": "B9"}]}
rows3 = [r.get("row") for r in qs.render_lines(three, "ko", None)]
check(rows3.count("subtotal") == 2 and rows3[rows3.index("option") + 1] != "subtotal",
      "기본 + 사용자 묶음 + OPTION(미반영): 기본·사용자 묶음에만 Sub total, OPTION 은 없음", str(rows3))
for ln in three["lines"]:
    if ln.get("kind") == "option" or ln.get("name") == "CHARGING DOCK":
        ln["in_total"] = True
rows3b = [r.get("row") for r in qs.render_lines(three, "ko", None)]
check(rows3b.count("subtotal") == 3, "기본 + 사용자 묶음 + OPTION(반영): 셋 다 Sub total", str(rows3b))

ng, _ = qs.apply_ops(na, [{"op": "nego", "amount": 300_000}], "30만원 네고")
check([r.get("row") for r in qs.render_lines(ng, "ko", None)] == [None, "nego", "gap"],
      "묶음이 하나여도 NEGO 가 있으면 Total 바로 위에 회색 줄", str([r.get("row") for r in qs.render_lines(ng, "ko", None)]))

print(f"PASS {len(PASS)} / FAIL {len(FAIL)}")
for label, detail in FAIL:
    print(f"  ✗ {label}: {detail}")
sys.exit(1 if FAIL else 0)
