"""백엔드가 실제로 만든 attachment 의 xlsx 를 MinIO 에서 직접 받아 무결성 검증.

HTTP 다운로드는 별도 auth/proxy 가 얽혀 있으므로 검증 단계에서는 backend storage
모듈을 직접 호출(같은 코드 경로) — MinIO 스트림의 내용물 자체를 확인.
"""
from __future__ import annotations

import asyncio
import datetime as _dt
import io
import os
import sys
from pathlib import Path

sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8", errors="replace")
sys.stderr = io.TextIOWrapper(sys.stderr.buffer, encoding="utf-8", errors="replace")

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT / "backend"))

# .env 로드 + MinIO 엔드포인트를 외부에서 본 호스트:포트로 강제.
# (compose 내부에선 'minio:9000' 이지만 호스트에서 실행 시엔 'localhost:9000')
# 반드시 backend.app.config import 전에 설정해야 한다.
with open(ROOT / ".env", "rb") as f:
    for line in f.read().decode("utf-8", errors="replace").splitlines():
        if "=" in line and not line.startswith("#"):
            k, v = line.split("=", 1)
            os.environ.setdefault(k.strip(), v.strip())
os.environ["MINIO_ENDPOINT"] = "localhost:9000"

from openpyxl import load_workbook
from sqlalchemy import select

from app.database import SessionLocal
from app.models import Attachment
from app.storage import get_object_stream

ATTACHMENT_ID = int(sys.argv[1]) if len(sys.argv) > 1 else 26


async def main():
    print(f"=== Verifying attachment id={ATTACHMENT_ID} ===\n")

    async with SessionLocal() as db:
        res = await db.execute(
            select(Attachment).where(Attachment.id == ATTACHMENT_ID)
        )
        att = res.scalar_one_or_none()
        if not att:
            print(f"FAIL: attachment id={ATTACHMENT_ID} not found in DB")
            sys.exit(1)
        print(f"  DB row: id={att.id} user_id={att.user_id}")
        print(f"          filename={att.original_filename!r}")
        print(f"          mime={att.mime}")
        print(f"          size={att.size_bytes:,}B")
        print(f"          object_key={att.object_key}")

    assert "spreadsheetml.sheet" in att.mime, f"MIME mismatch: {att.mime}"
    assert "유엔디_expense" in att.original_filename, \
        f"filename pattern unexpected: {att.original_filename}"
    print(f"PASS DB metadata 정상")

    # MinIO 에서 직접 fetch
    response = get_object_stream(att.object_key)
    try:
        chunks = []
        for chunk in response.stream(amt=64 * 1024):
            chunks.append(chunk)
        body = b"".join(chunks)
    finally:
        response.close()
        response.release_conn()
    print(f"PASS MinIO fetch OK, body size={len(body):,}B")
    assert len(body) == att.size_bytes, f"size mismatch: db={att.size_bytes} got={len(body)}"
    print(f"PASS body size == DB.size_bytes")

    # openpyxl 검증
    wb = load_workbook(io.BytesIO(body))
    print(f"PASS openpyxl load — sheets={wb.sheetnames}")
    expected_sheets = ["작성방법", "expense_내역", "expense_개인카드영수증 첨부"]
    for sn in expected_sheets:
        assert sn in wb.sheetnames, f"누락 시트: {sn}"
    print(f"PASS 3개 시트 모두 존재")

    # 내역 시트
    ws = wb["expense_내역"]
    print(f"\n[expense_내역]")
    print(f"  L1 (작성자) = {ws['L1'].value!r}")
    author = ws["L1"].value
    assert author == "배재병", f"author mismatch: {author}"
    print(f"PASS author=배재병")

    expected_rows = [
        ("해당없음", _dt.date(2026, 2, 7),
         "현대 글로비스 공장 출입 위한 안전화 구매", 72000, "워크업 대구 반야월점"),
        ("여비교통비", _dt.date(2026, 2, 9),
         "평택지제역 -> 동대구역 SRT 승차권 구매", 29500, "주식회사 에스알"),
        ("여비교통비", _dt.date(2026, 2, 10),
         "숙소 근처 -> 평택지제역 콜택시 이용", 30000, "미래 대리우전"),
    ]
    for r, expected in enumerate(expected_rows, start=5):
        cat = ws[f"H{r}"].value
        date = ws[f"I{r}"].value
        purpose = ws[f"J{r}"].value
        amount = ws[f"K{r}"].value
        vendor = ws[f"L{r}"].value
        print(f"  Row{r}: H={cat!r} I={date!r}({type(date).__name__}) "
              f"J={purpose!r} K={amount} L={vendor!r}")
        assert cat == expected[0], f"row{r}.cat: want {expected[0]} got {cat}"
        assert isinstance(date, (_dt.date, _dt.datetime)), \
            f"row{r}.date 타입 잘못: {type(date).__name__} — v3 fix 깨짐"
        if isinstance(date, _dt.datetime):
            date = date.date()
        assert date == expected[1], f"row{r}.date: want {expected[1]} got {date}"
        assert expected[2] in str(purpose), \
            f"row{r}.purpose: want contains {expected[2]} got {purpose}"
        assert int(amount) == expected[3], f"row{r}.amount: want {expected[3]} got {amount}"
        assert vendor == expected[4], f"row{r}.vendor: want {expected[4]} got {vendor}"
    print(f"PASS 3개 개인카드 라인 모두 정확")

    # 법인카드 영역(A-F) 비어있어야 — 사용자 'X' 명시
    for r in range(5, 8):
        for c in ("A", "B", "C", "D", "E", "F"):
            v = ws[f"{c}{r}"].value
            assert v is None, f"법인카드 {c}{r}={v!r} — 비어있어야"
    print(f"PASS 법인카드 영역 비어있음")

    # 라인 수 한도 — 8행 이하 비어있어야 (3건만)
    for r in range(8, 11):
        for c in ("H", "I", "J", "K", "L"):
            v = ws[f"{c}{r}"].value
            assert v is None, f"row{r}.{c}={v!r} — 라인 3건 초과"
    print(f"PASS 8행 이하 비어있음 — 라인 정확히 3건")

    # 영수증 시트
    wr = wb["expense_개인카드영수증 첨부"]
    print(f"\n[영수증 시트]")
    images = list(getattr(wr, "_images", []))
    print(f"  images={len(images)}")
    assert len(images) == 3, f"이미지 3장이어야: {len(images)}"
    print(f"PASS 이미지 3장")

    slots = [(0, "C4", "D4", "E4"), (1, "G4", "H4", "I4"), (2, "K4", "L4", "M4")]
    for idx, dc, pc, ac in slots:
        d = wr[dc].value
        p = wr[pc].value
        a = wr[ac].value
        print(f"  slot[{idx}]: {dc}={d!r} {pc}={p!r} {ac}={a!r}")
        assert d and p and a, f"slot{idx} 캡션 누락"
    print(f"PASS 슬롯 0/1/2 캡션 채워짐")

    # 미사용 슬롯 캡션 0 (v4 fix)
    for name, *cells in [
        ("slot3", "O4", "P4", "Q4"),
        ("slot4", "C26", "D26", "E26"),
        ("slot7", "O26", "P26", "Q26"),
        ("slot15", "O70", "P70", "Q70"),
    ]:
        for cell in cells:
            v = wr[cell].value
            assert v is None, f"{name}.{cell}={v!r} — 비어있어야 (v4 fix)"
    print(f"PASS 미사용 슬롯 캡션 0 (v4 fix 회귀 X)")

    # 이미지 anchor 인셋 (v4 fix)
    for i, im in enumerate(images):
        a = im.anchor
        fr = a._from
        to = a.to
        print(f"  image[{i}] from=col{fr.col}row{fr.row}+{fr.colOff},{fr.rowOff} "
              f"to=col{to.col}row{to.row}+{to.colOff},{to.rowOff}")
        assert fr.colOff > 0 and fr.rowOff > 0, f"image[{i}] 인셋 누락"
    print(f"PASS 이미지 인셋 적용")

    print("\n" + "=" * 60)
    print("✓ ALL VERIFICATION PASSED")
    print("=" * 60)
    print(f"  - DB metadata, MinIO fetch, body 크기 정합")
    print(f"  - 시트 구조 무결 (3 sheets)")
    print(f"  - 작성자/3개 라인/날짜(datetime)/금액/공급자 모두 정확")
    print(f"  - 법인카드 영역 빈 칸 (사용자 'X' 반영)")
    print(f"  - 영수증 3장, 인셋 적용")
    print(f"  - 미사용 슬롯 잔재 0")


if __name__ == "__main__":
    asyncio.run(main())
