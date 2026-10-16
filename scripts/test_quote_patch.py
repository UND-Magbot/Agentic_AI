"""올린 견적서 파일 위에서 이어서 고치기(company_knowledge/quote_patch.py) 점검 — 사용자 2026-10-15.

    MSYS_NO_PATHCONV=1 docker compose run --rm --no-deps -e PYTHONPATH=/app -v ./backend/app:/app/app -v ./scripts:/scripts \\
        -v ./docs:/docs backend python /scripts/test_quote_patch.py

① 회사 양식 견적서를 사람이 엑셀에서 꾸민 파일(빨간 굵은 글씨·채우기·직접 병합·행 높이·그림·텍스트 상자)
② 그 파일로 작성 상태를 맞춘 뒤(quote_readback) 이어서 고침(수량·비고·줄 추가·삭제·Comments·주석·NEGO)
③ patch 결과 — 값은 '이어서 고친 상태로 새로 만든 견적서'와 같고, 사람이 한 서식은 그대로인지
"""
from __future__ import annotations

import asyncio
import copy
import io
import re
import sys
import zipfile
from datetime import date

import openpyxl
import pymupdf
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Font, PatternFill

from app.company_knowledge import quote_patch as qp
from app.company_knowledge import quote_readback as rb
from app.company_knowledge import quote_session as qs
from app.company_knowledge import quote_xlsx as qx

PASS: list[str] = []
FAIL: list[tuple[str, str]] = []
DAY = date(2026, 10, 15)


def check(ok: bool, label: str, detail: str = "") -> None:
    (PASS.append(label) if ok else FAIL.append((label, detail)))


def parts(state: dict, no: str = "QT-LU20261015-01") -> tuple[list, dict]:
    f = state["form"]
    lang = f.get("lang") or "ko"
    lines = qs.render_lines(state, lang, f.get("fx_rate"))
    pogo = any(ln.get("key") in ("PPM", "PPF") for ln in lines if not ln.get("row"))
    return lines, qx.quote_header(f, lang=lang, quote_no=no, day=DAY, pogo=pogo)


def edit(state: dict, payload: dict) -> dict:
    new, _ = qs.apply_ops(copy.deepcopy(state), qs.edit_ops(state, payload), None)
    return new


def row_of(ws, text: str) -> int:
    return next(r for r in range(1, ws.max_row + 1) if str(ws.cell(r, 2).value or "").startswith(text))


def add_textbox(data: bytes, row0: int) -> bytes:
    """그림판 XML 에 텍스트 상자 하나(엑셀에서 넣은 도형 흉내) — row0 = 0부터 센 행."""
    zin = zipfile.ZipFile(io.BytesIO(data))
    draw = next(n for n in zin.namelist() if re.match(r"xl/drawings/drawing\d+\.xml$", n))
    xml = zin.read(draw).decode()
    shape = (f'<xdr:twoCellAnchor xmlns:xdr="http://schemas.openxmlformats.org/drawingml/2006/spreadsheetDrawing" '
             f'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"><xdr:from><xdr:col>0</xdr:col><xdr:colOff>0</xdr:colOff><xdr:row>{row0}</xdr:row><xdr:rowOff>0</xdr:rowOff>'
             f'</xdr:from><xdr:to><xdr:col>2</xdr:col><xdr:colOff>0</xdr:colOff><xdr:row>{row0 + 1}</xdr:row><xdr:rowOff>0</xdr:rowOff></xdr:to>'
             '<xdr:sp macro="" textlink=""><xdr:nvSpPr><xdr:cNvPr id="99" name="TextBox 99"/><xdr:cNvSpPr txBox="1"/></xdr:nvSpPr>'
             '<xdr:spPr><a:xfrm><a:off x="0" y="0"/><a:ext cx="1000000" cy="300000"/></a:xfrm><a:prstGeom prst="rect"><a:avLst/></a:prstGeom>'
             '</xdr:spPr><xdr:txBody><a:bodyPr/><a:p><a:r><a:t>도형 글</a:t></a:r></a:p></xdr:txBody></xdr:sp><xdr:clientData/></xdr:twoCellAnchor>')
    xml = re.sub(r"(</(?:\w+:)?wsDr>)", lambda m: shape + m.group(1), xml)
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as z:
        for n in zin.namelist():
            z.writestr(n, xml if n == draw else zin.read(n))
    return out.getvalue()


def shape_row(data: bytes) -> int | None:
    z = zipfile.ZipFile(io.BytesIO(data))
    for n in z.namelist():
        if re.match(r"xl/drawings/drawing\d+\.xml$", n):
            m = re.search(r"TextBox 99", (x := z.read(n).decode()))
            if m:
                a = x.rfind("twoCellAnchor", 0, m.start())
                return int(re.search(r"<(?:\w+:)?row>(\d+)</(?:\w+:)?row>", x[a:]).group(1))
    return None


# ── ① 작성 상태 S 와 사람이 꾸민 파일 U ──
S = qs.draft_state({}, model="", payload_kg=None, manual=True, form={
    "customer": "가나테크", "to": "김철수 대표님", "contact_name": "홍길동", "contact_title": "이사", "initials": "LU",
    "contact_mobile": "010-0000-0000", "contact_email": "name@example.com", "foot_notes": ["납기는 발주 후 협의"]})
S["lines"] = [{"id": "L1", "kind": "extra", "key": "extra", "name": "MG10\n- Magnetic Gripper", "commodity": "Magbot Gripper",
               "qty": 2, "unit_price": 1250000, "remark": ""},
              {"id": "L2", "kind": "extra", "key": "extra", "name": "TCV1", "commodity": "Magbot ATC", "qty": 1,
               "unit_price": 2300000, "remark": "컨트롤박스 포함"},
              {"id": "L3", "kind": "extra", "key": "extra", "name": "Cable 0.3m", "commodity": "Accessory", "qty": 4,
               "unit_price": 50000, "remark": ""}]
S["next_id"] = 4
S["delivery"] = {"case": 1, "freight": None}
lines0, head0 = parts(S)
wb = openpyxl.load_workbook(io.BytesIO(qx.build_xlsx(head0, lines0)))
ws = wb.active
r_tcv = row_of(ws, "TCV1")
ws.cell(r_tcv, 7).font = Font(name="맑은 고딕", size=11, bold=True, color="FFFF0000")      # 빨간 굵은 비고
ws.cell(r_tcv, 1).fill = PatternFill("solid", fgColor="FFCCE5FF")                            # 품목 칸 채우기
ws.row_dimensions[r_tcv].height = 55
bottom = ws.max_row + 2
ws.merge_cells(f"A{bottom}:C{bottom + 1}")                                                   # 맨 아래 직접 병합
ws.cell(bottom, 1).value = "직접 병합한 안내"
img = XLImage("/app/app/company_knowledge/data/quote/seal.png")
img.width = img.height = 40
ws.add_image(img, f"E{bottom}")                                                              # 직접 넣은 그림
buf = io.BytesIO()
wb.save(buf)
U = add_textbox(buf.getvalue(), bottom - 1)
n_img = len(openpyxl.load_workbook(io.BytesIO(U)).active._images)

# ② 그 파일에 맞춘 작성 상태(base) → 이어서 고침(N)
base, ch, wn = rb.merge(S, rb.parse(U))
check(ch == [] and wn == [], "꾸민 파일(값은 그대로)로 맞춘 상태 = 원래 상태", f"{ch} {wn}")
N = edit(base, {"form": {"comments": "www.magbot.kr\n- 이어서 고친 공급 건", "foot_notes": ["납기는 발주 후 협의", "운송 보험 별도"]},
                "lines": [{"id": "L1", "name": "MG10\n- Magnetic Gripper", "commodity": "Magbot Gripper", "qty": 3, "unit_price": 1250000,
                           "remark": "수량 늘림"},
                          {"id": "L2", "name": "TCV1", "commodity": "Magbot ATC", "qty": 1, "unit_price": 2300000, "remark": "컨트롤박스 포함"},
                          {"name": "출장 교육", "commodity": "Service", "qty": 1, "unit_price": 300000, "remark": "1일"}],
                "nego": {"mode": "amount", "value": 100000, "label": "특별 할인"}})
bl, bh = parts(base)
nl, nh = parts(N, no="QT-LU20261015-02")
bh["date"] = ""
out, notes = qp.patch(U, bl, bh, nl, nh)

# ③ 값 — 새로 만든 견적서와 같은가
pa, pb = rb.parse(out), rb.parse(qx.build_xlsx(nh, nl))
check(pa["rows"] == pb["rows"], "줄 값 = 이어서 고친 상태로 새로 만든 견적서", str([(x, y) for x, y in zip(pa["rows"], pb["rows"]) if x != y][:2]))
check(pa["form"] == pb["form"] and pa["quote_no"] == "QT-LU20261015-02", "머리·조건·주석·발행번호 같음",
      str({k: (pa["form"].get(k), pb["form"].get(k)) for k in pb["form"] if pa["form"].get(k) != pb["form"].get(k)}))

# ③ 서식 — 사람이 한 것 그대로
wo = openpyxl.load_workbook(io.BytesIO(out))
w2 = wo.active
r2 = row_of(w2, "TCV1")
f = w2.cell(r2, 7).font
check(f.b and f.color is not None and f.color.rgb == "FFFF0000", "빨간 굵은 비고 글씨 유지", str(f.color and f.color.rgb))
check(w2.cell(r2, 1).fill.fgColor.rgb == "FFCCE5FF" and w2.row_dimensions[r2].height == 55, "채우기·행 높이 유지")
merged = [str(m) for m in w2.merged_cells.ranges if w2.cell(m.min_row, 1).value == "직접 병합한 안내"]
shift = w2.max_row - ws.max_row
check(merged == [f"A{bottom + shift}:C{bottom + 1 + shift}"], "맨 아래 직접 병합이 줄 수 변화만큼 이동", f"{merged} shift={shift}")
check(len(w2._images) == n_img, "그림 수 그대로(로고·서명·직인·직접 넣은 그림)", f"{len(w2._images)} vs {n_img}")
check(shape_row(out) == bottom - 1 + shift, "텍스트 상자 되살림·같이 이동", f"{shape_row(out)} vs {bottom - 1 + shift} ({notes})")
check(not any(str(c.value or "").startswith("Cable 0.3m") for row in w2.iter_rows() for c in row), "지운 줄은 행째 빠짐")
tot = next(r for r in range(1, w2.max_row + 1) if str(w2.cell(r, 1).value or "").startswith("Total"))
check(str(w2.cell(tot, 7).value).startswith("=SUM("), "합계 수식 다시 씀", str(w2.cell(tot, 7).value))
r_new = row_of(w2, "출장 교육")
fc = w2.cell(r_new, 7).font.color
check((fc is None or fc.rgb != "FFFF0000") and w2.cell(r_new, 1).fill.fgColor.rgb != "FFCCE5FF",
      "한 줄만 강조한 서식(빨간 비고·채우기)은 새 줄로 안 퍼짐", f"{fc and fc.rgb} {w2.cell(r_new, 1).fill.fgColor.rgb}")

# 표 전체 글꼴·색을 바꾼 파일(궁서체·기울임·노란 채우기, 사용자 2026-10-16) — 새 줄도 같은 모양
wb4 = openpyxl.load_workbook(io.BytesIO(U))
w4 = wb4.active
for r in (row_of(w4, "MG10"), row_of(w4, "TCV1"), row_of(w4, "Cable 0.3m")):
    for c in range(1, 8):
        w4.cell(r, c).font = Font(name="궁서", size=10, italic=True)
        w4.cell(r, c).fill = PatternFill("solid", fgColor="FFFFF2CC")
b4 = io.BytesIO()
wb4.save(b4)
out4, _ = qp.patch(b4.getvalue(), bl, bh, nl, nh)
w5 = openpyxl.load_workbook(io.BytesIO(out4)).active
r5, ref5 = row_of(w5, "출장 교육"), openpyxl.load_workbook(io.BytesIO(qx.build_xlsx(nh, nl))).active
rr5 = row_of(ref5, "출장 교육")
cells5 = [w5.cell(r5, c) for c in range(1, 8) if type(w5.cell(r5, c)).__name__ != "MergedCell"]
check(all(x.font.name == "궁서" and x.font.i and x.fill.fgColor.rgb == "FFFFF2CC" for x in cells5),
      "새로 넣은 줄도 올린 파일 품목 줄 모양(궁서·기울임·노란 채우기)", str([(x.font.name, x.font.i, x.fill.fgColor.rgb) for x in cells5]))
check(all(w5.cell(r5, c).number_format == ref5.cell(rr5, c).number_format for c in (4, 5, 6)),
      "새 줄 수량·단가·금액 숫자 형식은 회사 양식 그대로", str([w5.cell(r5, c).number_format for c in (4, 5, 6)]))
check(rb.parse(out4)["rows"] == pb["rows"], "모양 바꾼 파일에서도 줄 값 같음")

pdf = asyncio.run(qx.to_pdf(out))
txt = "".join(p.get_text() for p in pymupdf.open(stream=pdf, filetype="pdf"))
want = 3 * 1250000 + 2300000 + 300000 - 100000
check(f"{want:,}" in txt and "이어서 고친 공급 건" in txt and "운송 보험 별도" in txt, "PDF 합계·고친 글", f"합계 {want:,} 있음={f'{want:,}' in txt}")

# 다른 프로그램(LibreOffice — 엑셀처럼 'xdr:' 접두어로 저장)이 저장한 파일 위에서도
import subprocess  # noqa: E402
import tempfile  # noqa: E402
from pathlib import Path  # noqa: E402

with tempfile.TemporaryDirectory() as t:
    (Path(t) / "u.xlsx").write_bytes(U)
    subprocess.run(["soffice", f"-env:UserInstallation=file://{t}/p", "--headless", "--convert-to", "xlsx:Calc MS Excel 2007 XML",
                    "--outdir", f"{t}/o", f"{t}/u.xlsx"], check=True, capture_output=True, timeout=120)
    LO = (Path(t) / "o" / "u.xlsx").read_bytes()
lo_draw = next(zipfile.ZipFile(io.BytesIO(LO)).read(n).decode() for n in zipfile.ZipFile(io.BytesIO(LO)).namelist()
               if re.match(r"xl/drawings/drawing\d+\.xml$", n))
if "<xdr:sp" in lo_draw:
    base_lo, _, _ = rb.merge(S, rb.parse(LO))
    bl2, bh2 = parts(base_lo)
    bh2["date"] = ""
    out2, notes2 = qp.patch(LO, bl2, bh2, nl, nh)
    w3 = openpyxl.load_workbook(io.BytesIO(out2)).active
    check(rb.parse(out2)["rows"] == pb["rows"] and shape_row(out2) is not None and "도형 글" in zipfile.ZipFile(io.BytesIO(out2)).read(
        qp._drawing_of_sheet(zipfile.ZipFile(io.BytesIO(out2)))).decode(),
        "LibreOffice 저장 파일(xdr: 접두어): 값 반영·텍스트 상자 유지", f"{notes2}")
    check(asyncio.run(qx.to_pdf(out2))[:4] == b"%PDF", "고친 파일이 PDF 로 바뀜(파일 손상 없음)")
else:
    check(False, "LibreOffice 가 텍스트 상자를 저장하지 않아 xdr: 접두어 경로를 못 봄", lo_draw[:200])

# 표 구조를 엑셀에서 바꾼 파일 → PatchError(조용히 망가뜨리지 않음)
wb3 = openpyxl.load_workbook(io.BytesIO(U))
wb3.active.delete_rows(row_of(wb3.active, "Cable 0.3m"))
b3 = io.BytesIO()
wb3.save(b3)
try:
    qp.patch(b3.getvalue(), bl, bh, nl, nh)
    check(False, "표 구조가 바뀐 파일은 PatchError")
except qp.PatchError as e:
    check("맞지 않" in str(e), "표 구조가 바뀐 파일은 PatchError", str(e))

# 같은 내용이면 그대로(값 칸 안 건드림)
same, _ = qp.patch(U, bl, bh, bl, {**bh, "date": DAY.isoformat()})
check(rb.parse(same)["rows"] == rb.parse(U)["rows"], "바뀐 게 없으면 표 그대로")


def _img(pdf: bytes):
    pm = pymupdf.open(stream=pdf, filetype="pdf")[0].get_pixmap(dpi=80)
    return pm.samples, (pm.width, pm.height)


# 실제 고객 파일(사용자 2026-10-15, docs/examples — git 제외라 없으면 건너뜀) — CC 이름을 흰 글자로 숨김 + 결제조건 행은 엑셀 자동 높이
REAL = next(Path("/docs/examples/quotes").glob("*QT-CH20261015-02.xlsx"), None)
if REAL is None:
    print("SKIP 실제 고객 파일 검사 — docs/examples/quotes 에 파일 없음")
else:
    PCM = REAL.read_bytes()
    pp_ = rb.parse(PCM)
    check(pp_["form"].get("cc") == "" and bool(pp_["form"].get("to")), "흰 글자로 숨긴 CC 는 읽지 않음", str(pp_["form"]))
    base_p, _, _ = rb.merge(S, pp_)
    blp, bhp = parts(base_p, no=pp_["quote_no"])
    bhp["date"] = ""
    outp, _ = qp.patch(PCM, blp, bhp, blp, {**bhp, "date": DAY.isoformat()})
    a9 = openpyxl.load_workbook(io.BytesIO(outp), rich_text=True).active["A9"].value
    check(type(a9).__name__ == "CellRichText" and rb._s(a9).endswith("CC :"), "고친 파일에서도 CC 숨김 글자 유지", repr(a9))
    check(_img(asyncio.run(qx.to_pdf(outp))) == _img(asyncio.run(qx.to_pdf(PCM))), "미리보기 PDF = 올린 파일 PDF(숨긴 글자·자동 행 높이)")

print(f"PASS {len(PASS)} / FAIL {len(FAIL)}")
for label, detail in FAIL:
    print(f"  ✗ {label}: {detail}")
sys.exit(1 if FAIL else 0)
