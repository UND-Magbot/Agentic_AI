"""Excel 의 m/d aaa 포맷 + 컬럼 폭 시뮬레이션으로 최종 표시 확인."""
from __future__ import annotations

import datetime as _dt
import io
import sys
from pathlib import Path

sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8", errors="replace")

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parent.parent
PATH = ROOT / "_test_artifacts" / "user_scenario.xlsx"

WEEKDAYS_KO = ["월", "화", "수", "목", "금", "토", "일"]


def render_md_aaa(d) -> str:
    if isinstance(d, _dt.datetime):
        d = d.date()
    if isinstance(d, _dt.date):
        return f"{d.month}/{d.day} {WEEKDAYS_KO[d.weekday()]}"
    return str(d) if d is not None else ""


def render_amount(v) -> str:
    if isinstance(v, (int, float)):
        return f"{int(v):,}"
    return str(v) if v is not None else ""


wb = load_workbook(PATH)
ws = wb["expense_내역"]
print("=== expense_내역 (개인카드 영역, 우측) 시뮬레이션 ===")
print(f"{'계정과목':<10} {'사용일자':<10} {'지출사유':<40} {'합계':<10} {'공급자':<20}")
print("-" * 100)
for r in range(5, 11):
    cat = ws[f"H{r}"].value
    dt = ws[f"I{r}"].value
    purpose = ws[f"J{r}"].value
    amt = ws[f"K{r}"].value
    ven = ws[f"L{r}"].value
    if any(v is not None for v in [cat, dt, purpose, amt, ven]):
        print(
            f"{cat or '':<10} {render_md_aaa(dt):<10} "
            f"{(purpose or '')[:40]:<40} {render_amount(amt):<10} {ven or '':<20}"
        )

wr = wb["expense_개인카드영수증 첨부"]
print("\n=== 영수증 시트 캡션 (슬롯 0..15) ===")
slots = [
    (0, "C4", "D4", "E4"), (1, "G4", "H4", "I4"), (2, "K4", "L4", "M4"), (3, "O4", "P4", "Q4"),
    (4, "C26", "D26", "E26"), (5, "G26", "H26", "I26"), (6, "K26", "L26", "M26"), (7, "O26", "P26", "Q26"),
    (8, "C48", "D48", "E48"), (9, "G48", "H48", "I48"), (10, "K48", "L48", "M48"), (11, "O48", "P48", "Q48"),
    (12, "C70", "D70", "E70"), (13, "G70", "H70", "I70"), (14, "K70", "L70", "M70"), (15, "O70", "P70", "Q70"),
]
for slot_idx, dc, pc, ac in slots:
    d = wr[dc].value
    p = wr[pc].value
    a = wr[ac].value
    if d or p or a:
        print(f"  slot[{slot_idx:02d}] {dc}={d!r}  {pc}={p!r}  {ac}={a!r}")
    else:
        print(f"  slot[{slot_idx:02d}] (empty - cleared)")
print(f"\n  이미지 {len(wr._images)} 장:")
for i, im in enumerate(wr._images):
    a = im.anchor
    print(
        f"    [{i}] from=col{a._from.col}row{a._from.row} (+{a._from.colOff},{a._from.rowOff}) "
        f"to=col{a.to.col}row{a.to.row} (+{a.to.colOff},{a.to.rowOff})"
    )
