# -*- coding: utf-8 -*-
"""견적서 부가세 별도 / 포함(사용자 2026-10-28 — 공공기관 견적은 부가세 포함, 10%) 점검.

    MSYS_NO_PATHCONV=1 docker compose run --rm -T --no-deps -e PYTHONPATH=/app -e PYTHONIOENCODING=utf-8 \\
        -v ./backend/app:/app/app -v ./scripts:/scripts -v ./docs:/docs backend python /scripts/test_quote_vat.py

포함이면 견적서 단가부터 × 1.1(원 단위 반올림), 표 위 'VAT 포함', Total 옆 '(부가세 포함)'. 작성 상태와 영업 건은 늘 공급가액.
화면 수정 → 견적서 줄·엑셀 → 되읽기(별도↔포함) → 올린 파일 위에서 이어 고치기 → 영업 건 품목까지 본다.
"""
from __future__ import annotations

import copy
import io
import sys
from datetime import date

import openpyxl

from app.company_knowledge import quote_patch as qp
from app.company_knowledge import quote_readback as rb
from app.company_knowledge import quote_session as qs
from app.company_knowledge import quote_xlsx as qx
from app.sales_deals import service as svc

PASS: list[str] = []
FAIL: list[tuple[str, str]] = []
DAY = date(2026, 10, 28)


def check(ok: bool, label: str, detail: str = "") -> None:
    (PASS.append(label) if ok else FAIL.append((label, detail)))


def parts(state: dict, no: str = "QT-TS20261028-01") -> tuple[list, dict]:
    lang = state["form"].get("lang") or "ko"
    return qs.render_lines(state, lang, state["form"].get("fx_rate")), \
        qx.quote_header(state["form"], lang=lang, quote_no=no, day=DAY, pogo=False)


def build(state: dict) -> bytes:
    lines, header = parts(state)
    return qx.build_xlsx(header, lines)


def edit(state: dict, payload: dict) -> dict:
    new, _ = qs.apply_ops(copy.deepcopy(state), qs.edit_ops(state, payload), None)
    return new


def ws_of(data: bytes):
    return openpyxl.load_workbook(io.BytesIO(data)).active


def total_row(ws) -> int:
    return next(r for r in range(17, ws.max_row + 1) if str(ws.cell(r, 1).value or "").startswith("Total"))


def row_of(ws, text: str) -> int:
    return next(r for r in range(1, ws.max_row + 1) if text in str(ws.cell(r, 2).value or ""))


S = qs.draft_state({"applicable": True, "items": []}, model="", payload_kg=None, form={
    "customer": "[테스트] 공공기관", "to": "김철수 주무관님", "contact_name": "홍길동", "contact_title": "이사", "initials": "TS",
    "contact_mobile": "010-0000-0000", "contact_email": "name@example.com"})
S.update(manual=True, asked=["extra"], delivery={"case": 1, "freight": None})
S["lines"] = [{"id": "L1", "kind": "extra", "key": "extra", "name": "TCV1 MASTER T.C", "commodity": "Magbot ATC", "qty": 1,
               "unit_price": 2_300_000, "remark": "", "price_status": "set"},
              {"id": "L2", "kind": "extra", "key": "extra", "name": "TCV1 Tool Plate", "commodity": "Magbot ATC", "qty": 2,
               "unit_price": 400_000, "remark": "", "price_status": "set"},
              {"id": "L3", "kind": "extra", "key": "extra", "name": "끝전 확인 부품", "commodity": "Part", "qty": 3,
               "unit_price": 12_345, "remark": "", "price_status": "set"}]
S["next_id"] = 4
S["nego"] = {"mode": "amount", "value": 100_000, "label": ""}
SUPPLY = 2_300_000 + 2 * 400_000 + 3 * 12_345 - 100_000

# 1. 기본 = 부가세 별도 — 지금과 똑같다
lines0, head0 = parts(S)
ws0 = ws_of(build(S))
check(not qs.vat_included(S) and qx.total(lines0) == SUPPLY, "기본은 VAT 별도 — 합계 = 공급가액", str(qx.total(lines0)))
check(ws0["G15"].value == "VAT 별도" and str(ws0.cell(total_row(ws0), 1).value).startswith("Total (부가세 별도)"),
      "별도 — 표 위 'VAT 별도', Total (부가세 별도)", str((ws0["G15"].value, ws0.cell(total_row(ws0), 1).value)))

# 2. 화면에서 [VAT 포함] — 잘못된 값은 무시
check(edit(S, {"form": {"vat": "보통"}})["form"].get("vat") in (None, "", "excluded"), "잘못된 부가세 값은 무시")
SI = edit(S, {"form": {"vat": "included"}})
check(SI["form"]["vat"] == "included" and qs.vat_included(SI), "[VAT 포함] 저장")
check([ln["unit_price"] for ln in SI["lines"]] == [2_300_000, 400_000, 12_345], "작성 상태 단가는 그대로 공급가액")
lines1, head1 = parts(SI)
items = [r for r in lines1 if not r.get("row")]
check([r["unit_price"] for r in items] == [2_530_000, 440_000, 13_580], "견적서 단가 = 공급가액 × 1.1(원 단위 반올림)",
      str([r["unit_price"] for r in items]))
check([r["supply_unit"] for r in items] == [2_300_000, 400_000, 12_345], "줄마다 공급가액(supply_unit) 보관")
nego = next(r for r in lines1 if r.get("row") == "nego")
check(nego["amount"] == -100_000 and nego["supply_unit"] == -90_909,
      "NEGO 는 견적서엔 적은 금액 그대로, 영업 건용 공급가액은 ÷ 1.1(사용자 2026-10-29)", str(nego))
check(qx.total(lines1) == 2_530_000 + 880_000 + 3 * 13_580 - 100_000, "견적서 Total = × 1.1 단가의 합 − NEGO", str(qx.total(lines1)))
SUPPLY_I = SUPPLY + 100_000 - 90_909                  # 부가세 포함 견적서의 공급가액 — NEGO 는 ÷ 1.1
check("VAT 포함 · 공급가액" in qs.total_text(SI) and f"{SUPPLY_I:,}원" in qs.total_text(SI), "대화 합계 글 — 포함 금액·공급가액",
      qs.total_text(SI))

# 3. 엑셀 — 'VAT 포함', Total (부가세 포함), 단가 칸 × 1.1
ws1 = ws_of(build(SI))
r1 = row_of(ws1, "TCV1 MASTER")
check(ws1["G15"].value == "VAT 포함" and str(ws1.cell(total_row(ws1), 1).value).startswith("Total (부가세 포함)"),
      "포함 — 표 위 'VAT 포함', Total (부가세 포함)", str((ws1["G15"].value, ws1.cell(total_row(ws1), 1).value)))
check(ws1.cell(r1, 5).value == 2_530_000, "엑셀 단가 칸 = 2,530,000", str(ws1.cell(r1, 5).value))
check(head0["labels"]["total"] == "Total (부가세 별도)", "양식 기본 문구는 바뀌지 않음(다른 견적서에 안 번짐)")

# 4. 영업 건 품목은 공급가액 — 명세서가 부가세를 또 붙이지 않게
deal_items = svc.snapshot_items(lines1)
deal_sum = sum(it["unit_price"] * it["qty"] for it in deal_items)
check(deal_sum == SUPPLY_I == qs.subtotal(SI), "영업 건 품목 합 = 공급가액(NEGO ÷ 1.1)", str(deal_items))
# 거래명세서(공급가액 + 부가세 10%) 합계 = 견적서 Total — 단가 줄마다 원 단위 반올림이라 몇 원 차이까지
check(abs(round(deal_sum * 1.1) - qx.total(lines1)) <= 3, "명세서 합계(공급가액 × 1.1) ≈ 견적서 Total",
      f"{round(deal_sum * 1.1)} vs {qx.total(lines1)}")

# 5. 되읽기 — 포함 견적서 파일을 별도 상태에 올리면 포함으로 바뀌고 단가는 공급가액으로
M, ch, wn = rb.merge(S, rb.parse(build(SI)))
check(M["form"].get("vat") == "included" and [ln["unit_price"] for ln in M["lines"]] == [2_300_000, 400_000, 12_345]
      and (M.get("nego") or {}).get("value") == 100_000, "포함 파일 되읽기 → VAT 포함 · 단가는 공급가액 · NEGO 는 적힌 금액 그대로", f"{ch} {wn} {M.get('nego')}")
check(any("부가세" in c for c in ch), "되읽기 바뀐 점에 부가세", str(ch))
M2, ch2, wn2 = rb.merge(SI, rb.parse(build(SI)))
check(ch2 == [] and wn2 == [], "고치지 않은 포함 파일 되읽기 → 바뀐 점 없음", f"{ch2} {wn2}")
M3, ch3, _ = rb.merge(SI, rb.parse(build(S)))
check(M3["form"].get("vat") == "excluded" and [ln["unit_price"] for ln in M3["lines"]] == [2_300_000, 400_000, 12_345],
      "별도 파일을 포함 상태에 올리면 별도로", str(ch3))

# 6. 올린 파일 위에서 이어 고치기 — 별도 → 포함으로 바꾸면 표 위 문구·Total·단가가 따라감
bl, bh = parts(S)
nl, nh = parts(SI, no="QT-TS20261028-02")
bh["date"] = ""
out, _ = qp.patch(build(S), bl, bh, nl, nh)
wp = ws_of(out)
check(wp["G15"].value == "VAT 포함" and str(wp.cell(total_row(wp), 1).value).startswith("Total (부가세 포함)")
      and wp.cell(row_of(wp, "TCV1 MASTER"), 5).value == 2_530_000,
      "이어 고치기 — 'VAT 포함'·Total 문구·단가 반영", str((wp["G15"].value, wp.cell(total_row(wp), 1).value)))

# 7. 해외(달러) 견적서는 부가세 없음 — 포함으로 둬도 무시
SE = copy.deepcopy(SI)
SE["form"].update(lang="en", fx_rate=1400)
check(not qs.vat_included(SE), "해외 견적서는 부가세 포함 무시")

# 8. 대여 견적서 — 월 납입금 단가도 × 1.1, 영업 건엔 공급가액 × 개월
SR = edit(edit(S, {"form": {"purpose": "rental"}}), {"form": {"rental_months": 12, "vat": "included"}})
lr = [r for r in qs.render_lines(SR, "ko", None) if not r.get("row")]
check(lr[0]["unit_price"] == 2_530_000 and lr[0].get("months") in (None, 12), "대여 — 단가 × 1.1", str(lr[0]))
it = svc.snapshot_items(lr)
check(it[0]["unit_price"] == 2_300_000 * (lr[0].get("months") or 1), "대여 영업 건 = 공급가액 × 개월", str(it[0]))

for label, detail in FAIL:
    print(f"FAIL {label} :: {detail}")
print(f"PASS {len(PASS)} / FAIL {len(FAIL)}")
sys.exit(1 if FAIL else 0)
