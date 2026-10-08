"""확정 결과 저장·재활용(confirmed) 검증.

- 기본: 수정본 글자 추출(.drawio/.pptx/.docx)·확정 마커 형식 — 서버 불필요.
- --db : 실제 DB·임베딩으로 초안 → 확정 → 유사 요청 검색까지 (backend 컨테이너에서 실행).
    docker compose run --rm --no-deps -v ./scripts:/scripts -v ./docs:/docs backend \
        python /scripts/test_confirmed.py --db
"""
import asyncio
import io
import json
import re
import sys
from pathlib import Path

sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8", errors="replace")
ROOT = Path(__file__).resolve().parents[1]
for cand in (ROOT / "backend", Path("/app")):
    if (cand / "app").is_dir():
        sys.path.insert(0, str(cand))
        break

from app import confirmed  # noqa: E402
from app.diagram import drawio  # noqa: E402
from app.diagram.spec import ConceptMapSpec  # noqa: E402

PASS: list[str] = []
FAIL: list[tuple[str, str]] = []


def check(ok: bool, label: str, detail: str = "") -> None:
    (PASS.append(label) if ok else FAIL.append((label, detail)))


SPEC_PATH = next(p for p in (ROOT / "docs/concept_drawio/chem_spec.json",
                             Path("/docs/concept_drawio/chem_spec.json")) if p.is_file())
SPEC = ConceptMapSpec.model_validate(json.loads(SPEC_PATH.read_text(encoding="utf-8")))

# ── 1) 확정 마커 ↔ frontend message.tsx _CONFIRM_RE ─────────────────────────
FRONT_RE = re.compile(r"\[이 결과를 확정\]\(/api/proposal-records/(\d+)/confirm\)")
m = FRONT_RE.fullmatch(confirmed.confirm_link(42))
check(bool(m) and m.group(1) == "42", "확정 마커 형식 = 프론트 정규식")

# ── 2) 수정본 글자 추출 ─────────────────────────────────────────────────────
xml = drawio.build_drawio(SPEC).replace("투입 존", "투입 컨베이어(담당자 수정)")
txt = confirmed.extract_revised_text("수정본.drawio", xml.encode("utf-8"))
check("투입 컨베이어(담당자 수정)" in txt, "drawio 수정 라벨 추출")
check("<b>" not in txt and "<br>" not in txt, "HTML 태그 제거")
check(txt.index("병/액 주입 자동화 라인 개념도") < txt.index("투입 컨베이어"), "위→아래 순서")

import docx  # noqa: E402

d = docx.Document()
d.add_paragraph("1. 핵심 메시지")
d.add_paragraph("담당자가 고친 문장")
t = d.add_table(rows=1, cols=2)
t.rows[0].cells[0].text, t.rows[0].cells[1].text = "항목", "내용"
buf = io.BytesIO()
d.save(buf)
dt = confirmed.extract_revised_text("본문.docx", buf.getvalue())
check("담당자가 고친 문장" in dt and "항목 | 내용" in dt, "docx 문단·표 추출")

from pptx import Presentation  # noqa: E402

prs = Presentation()
for title, body in [("1. 핵심 메시지", "담당자가 고친 슬라이드 문장"), ("2. 현황", "두 번째 장")]:
    sl = prs.slides.add_slide(prs.slide_layouts[1])
    sl.shapes.title.text = title
    sl.placeholders[1].text = body
buf = io.BytesIO()
prs.save(buf)
pt = confirmed.extract_revised_text("제안서본문_수정.pptx", buf.getvalue())
check("담당자가 고친 슬라이드 문장" in pt and pt.index("[슬라이드 1]") < pt.index("[슬라이드 2]"),
      "pptx 슬라이드 순서대로 추출", pt)

for name, data, label in [("a.pdf", b"%PDF", "허용 안 된 형식"), ("a.pptx", b"not zip", "깨진 pptx"),
                          ("a.drawio", b"<mxfile><diagram>eJzLSM3JyQcABiwCFQ==</diagram></mxfile>",
                           "압축 drawio 안내"),
                          ("a.drawio", b"not xml", "깨진 drawio")]:
    try:
        confirmed.extract_revised_text(name, data)
        check(False, f"{label} → 오류")
    except confirmed.RecordError:
        check(True, f"{label} → 오류")


# ── 3) DB·임베딩 통합 (--db) ────────────────────────────────────────────────
async def db_flow() -> None:
    from sqlalchemy import select, text

    from app.database import SessionLocal
    from app.migrations import run_migrations
    from app.models import User

    chem = ("한화로봇 14kg 약액 주입 공정. 로봇 투입 -> 비전 센터링 -> 너트러너로 뚜껑 풀기 -> "
            "액주입 -> 불량 시 병 치우기 -> 뚜껑 잠그기. 약액 주입 210초.")
    similar = "협동로봇으로 병 뚜껑 풀고 약액 충전 후 다시 캡 체결하는 라인, 비전으로 병 위치 확인"
    unrelated = "주유소 4족 보행 로봇 야간 순찰 및 화재 감지 제안"

    async with SessionLocal() as db:
        await run_migrations(db)
        uids = (await db.execute(select(User.id).order_by(User.id).limit(2))).scalars().all()
        owner = uids[0]
        other = uids[1] if len(uids) > 1 else owner + 10_000
        rid = await confirmed.create_draft(
            db, user_id=owner, kind="concept_map", title="테스트 약액 라인", request_text=chem,
            content=SPEC.model_dump(), result_attachment_ids=[])
        rec = await confirmed.get_record(db, rid)
        check(rec.status == "draft", "초안 저장")
        check(not await confirmed.find_confirmed(db, chem, kind="concept_map", min_score=0.99),
              "확정 전에는 검색 안 됨")

        try:
            await confirmed.confirm(db, record_id=rid, user_id=other)
            check(False, "남의 기록 확정 거부")
        except confirmed.RecordError:
            check(True, "남의 기록 확정 거부")

        rec = await confirmed.confirm(db, record_id=rid, user_id=owner)
        check(rec.status == "confirmed", "확정")
        hits = await confirmed.find_confirmed(db, similar, kind="concept_map", top_k=5)
        scores = {r.id: s for r, s in hits}
        check(rid in scores, "비슷한 요청 → 확정본 검색", str(scores))
        far = await confirmed.find_confirmed(db, unrelated, kind="concept_map", top_k=5)
        check(rid not in {r.id for r, _ in far}, "다른 공정 요청 → 검색 안 됨",
              str([(r.id, round(s, 3)) for r, s in far]))
        block = await confirmed.get_confirmed_examples(db, similar, kind="concept_map")
        check("옮겨 쓰지 않는다" in block and "테스트 약액 라인" not in block[:200],
              "예시 블록에 복사 금지 안내")
        check(not await confirmed.find_confirmed(db, similar, kind="proposal_body"),
              "종류(kind)가 다르면 섞이지 않음")
        print(f"  유사도: 비슷한 요청 {scores.get(rid, 0):.3f}")

        # 정리 — 테스트 행만 지운다.
        await db.execute(text("DELETE FROM proposal_records WHERE id = :id"), {"id": rid})
        await db.commit()


if "--db" in sys.argv:
    asyncio.run(db_flow())

print(f"PASS {len(PASS)} / FAIL {len(FAIL)}")
for label, d in FAIL:
    print("  FAIL", label, d)
sys.exit(1 if FAIL else 0)
