"""주간 회의록 보고 — 풀 E2E 검증.

실제 회의 녹음(m4a)을 백엔드에 업로드하고 '주간 회의 보고 작성해줘' 요청을 보내
음성 인식 → 요약 → xlsx 생성 → 다운로드까지 전 파이프라인을 검증한다.

전제: docker compose 스택(backend/minio/embed/vllm) 가동 중 (http://localhost:8002).

검증 항목:
1. 로그인 → 토큰 발급
2. 오디오 첨부 업로드 (audio MIME 화이트리스트 통과)
3. /v1/generate — 회의 의도 감지 → fast-path → 응답에 성공 메시지 + 다운로드 링크
4. 결과 xlsx 다운로드 → 새 주차 시트 + 제목/회의 주요 내용/지시사항 채워짐
"""
import io
import os
import re
import sys
import time
from pathlib import Path

import httpx
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
BACKEND = "http://localhost:8002"
AUDIO = ROOT / "알파시티1로31길.m4a"

# 시드 superadmin 계정 (infra/postgres/init/02_users.sql).
LOGIN = {"username": "superadmin", "password": os.environ["SEED_SUPERADMIN_PASSWORD"]}


def main() -> int:
    fails: list[str] = []

    if not AUDIO.exists():
        print(f"FAIL: 녹음 파일 없음 — {AUDIO}")
        return 1

    # 백엔드 헬스 체크.
    try:
        h = httpx.get(f"{BACKEND}/health", timeout=10)
        print(f"백엔드 health: {h.status_code} {h.text[:200]}")
    except Exception as e:
        print(f"FAIL: 백엔드에 연결할 수 없습니다 ({BACKEND}) — {e}")
        return 1

    # ── 1) 로그인 ──────────────────────────────────────────────────────
    r = httpx.post(f"{BACKEND}/v1/auth/login", json=LOGIN, timeout=30)
    if r.status_code != 200:
        print(f"FAIL: 로그인 실패 {r.status_code} — {r.text[:300]}")
        return 1
    token = r.json()["access_token"]
    headers = {"Authorization": f"Bearer {token}"}
    print("1) 로그인 OK")

    # ── 2) 오디오 업로드 ────────────────────────────────────────────────
    with open(AUDIO, "rb") as f:
        r = httpx.post(
            f"{BACKEND}/v1/attachments",
            files={"file": (AUDIO.name, f, "audio/mp4")},
            headers=headers,
            timeout=120,
        )
    if r.status_code != 201:
        print(f"FAIL: 오디오 업로드 실패 {r.status_code} — {r.text[:300]}")
        return 1
    att = r.json()
    att_id = att["id"]
    print(f"2) 업로드 OK — attachment id={att_id} mime={att['mime']} size={att['size_bytes']}")

    # ── 3) /v1/generate — 회의 보고 요청 ───────────────────────────────
    # 프론트(router.ts composeRequestContextBlock)가 만드는 컨텍스트 블록을 재현.
    system_prompt = (
        "당신은 한국어 전용 사내 AI 어시스턴트입니다.\n\n"
        "## 현재 요청 컨텍스트\n"
        "- 현재 날짜: 2026-05-19 (YYYY-MM-DD).\n"
        "- 현재 사용자(작성자 후보): 배재병.\n"
        "- 첨부된 파일(업로드 순, 총 1개):\n"
        f"  1) id={att_id} ({AUDIO.name}, audio/mp4)\n"
        f"- 도구가 attachment_ids 인자를 받으면 위 1개 ID 를 전달: [{att_id}].\n"
    )
    body = {
        "systemPrompt": system_prompt,
        "messages": [{"role": "user", "content": "주간 회의 보고 작성해줘"}],
    }
    print("3) 회의 보고 생성 요청 — 스트리밍 (STT+요약, 수 분 소요)...")
    t0 = time.time()
    parts: list[str] = []
    # 청크 간 최대 간격 측정 — STT 침묵 구간에 하트비트가 흐르는지 검증.
    # undici(프론트 fetch) 기본 body 타임아웃 5분 — 그 전에 청크가 흘러야 한다.
    last = time.time()
    max_gap = 0.0
    with httpx.stream(
        "POST", f"{BACKEND}/v1/generate", json=body, timeout=1200
    ) as resp:
        if resp.status_code != 200:
            print(f"FAIL: /v1/generate {resp.status_code}")
            return 1
        for chunk in resp.iter_text():
            now = time.time()
            max_gap = max(max_gap, now - last)
            last = now
            parts.append(chunk)
    full = "".join(parts)
    elapsed = time.time() - t0
    print(f"   응답 수신 ({elapsed:.0f}초, 최대 청크 간격 {max_gap:.0f}초):")
    if max_gap > 120:
        fails.append(
            f"청크 간격 {max_gap:.0f}초 > 120초 — 하트비트 미작동, "
            "프론트 fetch 5분 타임아웃 위험"
        )
    print("   " + full.replace("\n", "\n   "))

    if "✓ 주간 회의록 보고" not in full:
        fails.append("응답에 성공 메시지('✓ 주간 회의록 보고') 누락")

    m = re.search(r"\[([^\]]+\.xlsx)\]\((/api/attachments/(\d+)/download)\)", full)
    if not m:
        fails.append("응답에 다운로드 링크([파일명](url)) 누락")
        _report(fails)
        return 1 if fails else 0

    result_filename = m.group(1)
    result_att_id = int(m.group(3))
    print(f"4) 결과 파일: {result_filename} (attachment id={result_att_id})")

    # ── 4) 결과 xlsx 다운로드 + 검증 ───────────────────────────────────
    r = httpx.get(
        f"{BACKEND}/v1/attachments/{result_att_id}/download",
        headers=headers,
        timeout=60,
    )
    if r.status_code != 200:
        fails.append(f"결과 xlsx 다운로드 실패 {r.status_code}")
        _report(fails)
        return 1

    xlsx = r.content
    out_path = ROOT / "_test_meeting_e2e.xlsx"
    out_path.write_bytes(xlsx)
    print(f"   결과 xlsx 저장: {out_path} ({len(xlsx)} bytes)")

    try:
        wb = load_workbook(io.BytesIO(xlsx))
    except Exception as e:
        fails.append(f"결과 xlsx 가 openpyxl 로 열리지 않음: {e}")
        _report(fails)
        return 1

    ws = wb.active
    title = str(ws["B2"].value or "")
    d5 = str(ws["D5"].value or "")
    d13 = str(ws["D13"].value or "")
    print(f"   시트 수: {len(wb.sheetnames)}, 활성 시트: {wb.active.title!r}")
    print(f"   B2(제목): {title}")
    print(f"   D5(회의 주요 내용) {len(d5)}자: {d5[:300]}")
    print(f"   D13(지시사항) {len(d13)}자: {d13[:300]}")

    if "주간회의보고" not in title:
        fails.append(f"제목(B2) 비정상: {title!r}")
    if len(d5) < 20:
        fails.append(f"회의 주요 내용(D5) 이 비었거나 너무 짧음 ({len(d5)}자)")
    if len(d13) < 5:
        fails.append(f"지시사항(D13) 이 비었음 ({len(d13)}자)")
    if len(wb.sheetnames) < 2:
        fails.append("시트가 1개뿐 — 새 주차 시트 append 실패")

    _report(fails)
    return 1 if fails else 0


def _report(fails: list[str]) -> None:
    if fails:
        print(f"\nFAIL ({len(fails)}건):")
        for f in fails:
            print(f"  - {f}")
    else:
        print("\nPASS - 회의 보고 E2E 전 파이프라인 통과")


if __name__ == "__main__":
    raise SystemExit(main())
