"""해외 견적서 해외 기준 견적서 표기·한글 칸 검사·AI 번역(사용자 2026-10-14) — DB 없이 순수 로직만.

    MSYS_NO_PATHCONV=1 docker compose run --rm --no-deps -e PYTHONPATH=/app -v ./backend/app:/app/app -v ./scripts:/scripts backend \\
        python /scripts/test_quote_translate.py
"""
from __future__ import annotations

import asyncio
import json
import sys
from datetime import date

from app.company_knowledge import product_prices as pp
from app.company_knowledge import quote_session as qs
from app.company_knowledge import quote_translate as qt
from app.company_knowledge import quote_xlsx as qx

PASS: list[str] = []
FAIL: list[tuple[str, str]] = []


def check(ok: bool, label: str, detail: str = "") -> None:
    (PASS.append(label) if ok else FAIL.append((label, detail)))


head = {"customer": "ABC Robotics Co., Ltd.", "initials": "AL", "contact_name": "전알렉스", "contact_title": "이사",
        "contact_mobile": "010-1234-5678", "contact_email": "alex@example.com",
        "contact_name_en": "Alex Jeon", "contact_honorific": "Mr."}
st = qs.draft_state({}, model="", payload_kg=None, form=head, manual=True)
st, _ = qs.apply_ops(st, [{"op": "add", "name": "TCC1\n- Payload: 5 kg", "commodity": "Magbot ATC", "qty": 2, "usd": 1000, "remark": "커플러 포함"},
                          {"op": "add", "name": "설치 교육", "commodity": "교육", "qty": 1, "usd": 500, "unit": "식"},
                          {"op": "add", "name": "보관함", "commodity": "Accessory", "qty": 3, "usd": 20, "unit": "박스"},
                          {"op": "delivery", "case": 5}], None)
check(qt.cells(st) == [], "국내 견적서면 한글 칸 검사 없음")

# 해외로 — 해외 기본 문구·비고
st, _ = qs.apply_ops(st, [{"op": "set", "field": "lang", "value": "en"}], None)
f = st["form"]
check(f["delivery"] == "4 ~ 6 weeks after payment" and f["place"] == "UND Korea Factory"
      and f["payment"].startswith("100% advance payment") and f["note"] == "Ex-work UND Korea Factory",
      "해외로 바꾸면 해외 기본 문구(납기·장소·결제·비고 Ex-work)", json.dumps({k: f[k] for k in ("delivery", "place", "payment", "note")}))
check(pp.QUOTE_LABELS["en"]["spec"] == "Description" and pp.QUOTE_LABELS["en"]["total"] == "Total Amount",
      "영문 항목명 Commodity / Description · Total Amount")
han = qt.cells(st)
paths = {c["path"]: c for c in han}
lid = {ln["name"].split("\n")[0]: ln["id"] for ln in st["lines"]}
check(f"line.{lid['TCC1']}.remark" in paths and f"line.{lid['설치 교육']}.name" in paths and f"line.{lid['설치 교육']}.commodity" in paths,
      "한글 칸: 비고·규격·품목", str(list(paths)))
check(f"line.{lid['설치 교육']}.unit" not in paths and f"line.{lid['보관함']}.unit" in paths,
      "단위 '식'은 자동으로 lot(검사 안 함), 표에 없는 한글 단위 '박스'는 한글 칸", str(list(paths)))
check(not any(p.startswith("form.contact") for p in paths), "영문 이름이 있으면 담당자 칸은 통과")

# 담당자·번호 — 'Mr. Alex Jeon', 82-
hd = qx.quote_header(st["form"], lang="en", quote_no="QT-AL20261014-01", day=date(2026, 10, 14), pogo=False)
check(hd["contact"]["name"] == "Alex Jeon" and hd["contact"]["title"] == "Mr." and hd["contact"]["mobile"] == "82-10-1234-5678",
      "해외 담당자 = Mr. Alex Jeon, 휴대폰 82-10-1234-5678", str(hd["contact"]))
hk = qx.quote_header({**st["form"], "lang": "ko"}, lang="ko", quote_no="X", day=date(2026, 10, 14), pogo=False)
check(hk["contact"]["name"] == "전알렉스" and hk["contact"]["title"] == "이사" and hk["contact"]["mobile"] == "010-1234-5678",
      "국내 견적서는 그대로(한글 이름·직함·010)", str(hk["contact"]))
check(qx.intl_mobile("+82-10-1") == "+82-10-1" and qx.intl_mobile("82-10-1") == "82-10-1" and qx.intl_mobile("") == "",
      "이미 국가번호면 그대로")
check(qs.clean_honorific("mr") == "Mr." and qs.clean_honorific("Ms") == "Ms." and qs.clean_honorific("Dr.") == "",
      "호칭은 Mr./Ms. 만")
lines = qs.render_lines(st, "en", None)
units = {r["desc"].split("\n")[0]: r.get("unit") for r in lines if not r.get("row")}
check(units.get("설치 교육") == "lot" and units.get("보관함") == "박스", "해외 견적서 단위 식 → lot", str(units))

# 영문 이름이 비면 한글 이름이 담당자 칸 한글로 — 이름도 AI 가 옮기되 '확인 필요'(name=True)
s2 = {**st, "form": {**st["form"], "contact_name_en": "", "to": "황 과장"}}
c2 = {c["path"]: c for c in qt.cells(s2)}
check(c2.get("form.contact_name_en", {}).get("name") is True and c2.get("form.to", {}).get("name") is True
      and c2.get(f"line.{lid['TCC1']}.remark", {}).get("name") is False,
      "이름(담당자·받는 분)은 이름 확인 표시", str(c2))
check("담당자 영문 이름 비어 있음" not in qs.problems(s2), "영문 이름이 비어도 한글 이름이 있으면 '비어 있음'은 아님(한글 칸으로 막힘)")


# AI 번역 — 칸에는 한글 원문, 영어는 state.en_text 에(사용자 2026-10-14). 가짜 AI: 정상 / 한글 남김 / 모델명 빠짐
calls: list[int] = []


async def fake_chat(messages, **kw):
    items = json.loads(messages[1]["content"])["items"]
    calls.append(len(items))
    out = []
    for x in items:
        t = x["text"]
        en = {"커플러 포함": "Couplers included", "설치 교육": "Installation training", "교육": "Training",
              "박스": "box", "보관함": "Storage case"}.get(t, t)
        out.append({"id": x["id"], "en": en})
    return json.dumps({"items": out})

check(len(qt.untranslated(st)) == len(han) and all(c["en"] is None for c in qt.overview(st)), "처음엔 모든 한글 칸이 번역 전")
new, done, failed = asyncio.run(qt.ensure(st, fake_chat))
ov = {c["path"]: c for c in qt.overview(new)}
check(ov[f"line.{lid['TCC1']}.remark"]["en"] == "Couplers included" and ov[f"line.{lid['설치 교육']}.name"]["by"] == "ai",
      "AI 번역이 번역 표에 저장(by=ai)", str(ov))
check(new["lines"][0]["remark"] == "커플러 포함" and st.get("en_text") is None, "칸의 한글 원문은 그대로(원래 상태도 그대로)")
check(qt.untranslated(new) == [] and not failed, "번역 뒤 번역 전 칸 없음", str(qt.untranslated(new)))
loc = qt.localized(new)
check(loc["lines"][0]["remark"] == "Couplers included" and qt.cells(loc) == [], "견적서에 찍을 상태는 영어(한글 칸 없음)", str(qt.cells(loc)))
n0 = len(calls)
again, d2, _ = asyncio.run(qt.ensure(new, fake_chat))
check(again is new and not d2 and len(calls) == n0, "번역이 다 있으면 AI 를 다시 부르지 않음")

# 원문을 고치면 그 칸만 다시 번역
ed = json.loads(json.dumps(new))
next(x for x in ed["lines"] if x["id"] == lid["설치 교육"])["name"] = "교육"
check([c["path"] for c in qt.untranslated(ed)] == [f"line.{lid['설치 교육']}.name"], "원문을 고친 칸만 번역 전으로", str(qt.untranslated(ed)))
ed2, d3, _ = asyncio.run(qt.ensure(ed, fake_chat))
check(calls[-1] == 1 and d3[0]["en"] == "Training", "고친 칸 하나만 AI 에 보냄", str(calls))


async def bad_chat(messages, **kw):
    items = json.loads(messages[1]["content"])["items"]
    return json.dumps({"items": [{"id": x["id"], "en": "교육" if i == 0 else "Payload kg"} for i, x in enumerate(items)]})

s3 = {**st, "lines": [*st["lines"], {**st["lines"][0], "id": "L99", "remark": "TCV1 30kg 포함"}]}
n3, _, f3 = asyncio.run(qt.ensure(s3, bad_chat))
why = {x["path"]: x["why"] for x in f3}
check(why.get(f"line.{lid['TCC1']}.remark") == "한글이 남음" and why.get("line.L99.remark", "").startswith("모델명·숫자 빠짐"),
      "AI 답 검사: 한글이 남거나 모델명·숫자(TCV1·30)가 빠지면 안 씀", str(why))
check(f"line.{lid['TCC1']}.remark" in {c["path"] for c in qt.untranslated(n3)}, "못 바꾼 칸은 번역 전으로 남아 발행을 막음")

# 담당자가 영어 고치기 · [한글 그대로]
p_rm = f"line.{lid['TCC1']}.remark"
u = qt.set_en(new, p_rm, en="Coupler included (1 set)")
check(next(c for c in qt.overview(u) if c["path"] == p_rm)["by"] == "user", "담당자가 고친 영어는 by=user")
u2 = qt.set_en(u, p_rm, en="")
check(next(c for c in qt.overview(u2) if c["path"] == p_rm)["en"] is None, "영어를 비우면 다음에 AI 가 다시 번역")
for bad, why_ in (("커플러", "한글"), ("x" * 101, "칸 길이")):
    try:
        qt.set_en(new, p_rm, en=bad)
        check(False, f"영어 고치기 검사 — {why_} 거절")
    except ValueError:
        check(True, f"영어 고치기 검사 — {why_} 거절")
k = qt.set_en(n3, p_rm, keep=True)
check(p_rm not in {c["path"] for c in qt.untranslated(k)} and qt.localized(k)["lines"][0]["remark"] == "커플러 포함",
      "[한글 그대로] 칸은 발행을 막지 않고 한글로 찍힘")
check(p_rm not in qt.set_en(k, p_rm, keep=False).get("keep_ko", []), "[한글 그대로] 끄기")
# 칸을 직접 영어로 고치면 그 칸의 번역은 지움
dd = json.loads(json.dumps(new))
dd["lines"][0]["remark"] = "Couplers incl."
dd2, _, _ = asyncio.run(qt.ensure({**dd, "lines": [*dd["lines"], {**dd["lines"][0], "id": "L98", "remark": "보관함"}]}, fake_chat))
check(p_rm not in dd2["en_text"], "한글이 아닌 칸의 예전 번역은 정리", str(list(dd2["en_text"])))

# 엑셀 — 영문 양식에 담당자 줄·전화번호
xl = qx.build_xlsx(qx.quote_header(loc["form"], lang="en", quote_no="QT-AL20261014-01", day=date(2026, 10, 14), pogo=False),
                   qs.render_lines(loc, "en", None))
import io

import openpyxl

ws = openpyxl.load_workbook(io.BytesIO(xl)).active
check("Tel : 82-10-1234-5678" in ws["A4"].value and "Mr. Alex Jeon (alex@example.com)" in ws["A4"].value,
      "엑셀 Quotation by: Tel : 82-… / Mr. Alex Jeon (메일)", repr(ws["A4"].value))
check(ws["B16"].value == "Description" and ws["A16"].value == "Commodity", "엑셀 표 머리 Commodity / Description")

# 국내로 되돌리면 비고 기본(Ex-work)은 지움, 담당자가 쓴 비고는 유지
k1, _ = qs.apply_ops(st, [{"op": "set", "field": "lang", "value": "ko"}], None)
k2, _ = qs.apply_ops({**st, "form": {**st["form"], "note": "Custom remark"}}, [{"op": "set", "field": "lang", "value": "ko"}], None)
check(k1["form"]["note"] == "" and k2["form"]["note"] == "Custom remark", "국내로: 기본 비고만 지우고 직접 쓴 비고는 그대로")
check(qt._ok("- 페이로드 5kg", "- Payload: 5 kg", 100) is None and qt._ok("- 페이로드 5kg", "- Payload: 6 kg", 100) is not None,
      "AI 답 검사: '5kg' → '5 kg'(해외 기준 견적서 띄어쓰기)는 통과, 숫자가 바뀌면 거절")

print(f"PASS {len(PASS)} / FAIL {len(FAIL)}")
for label, detail in FAIL:
    print(f"  ✗ {label}: {detail}")
sys.exit(1 if FAIL else 0)
